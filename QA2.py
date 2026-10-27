import sys
import os
import re
import shutil
import time
import json
from io import BytesIO
from dataclasses import dataclass, replace
from datetime import datetime, time as datetime_time
from functools import lru_cache
from urllib.parse import unquote, urlparse, urljoin, urlunparse

import checklist_parameters

from PyQt6.QtWidgets import (
    QApplication,
    QCompleter,
    QButtonGroup,
    QCheckBox,
    QComboBox,
    QFrame,
    QGraphicsOpacityEffect,
    QLabel,
    QStyle,
    QAbstractItemView,
    QHBoxLayout,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QTabWidget,
    QTreeWidget,
    QTreeWidgetItem,
    QFileDialog,
    QVBoxLayout,
    QWidget,
)
from PyQt6.QtGui import QFont, QColor, QBrush, QTextOption
from PyQt6.QtCore import QEasingCurve, QPropertyAnimation, QTimer, Qt, QStringListModel

from PyQt6.uic import loadUi

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import pymupdf as fitz
except ImportError:
    fitz = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

if pytesseract is not None and os.name == "nt":
    tesseract_path = shutil.which("tesseract")
    if not tesseract_path:
        candidates = (
            os.path.join(
                os.environ.get("LOCALAPPDATA", ""),
                "Programs", "Tesseract-OCR", "tesseract.exe",
            ),
            os.path.join(
                os.environ.get("ProgramFiles", r"C:\Program Files"),
                "Tesseract-OCR", "tesseract.exe",
            ),
        )
        tesseract_path = next(
            (candidate for candidate in candidates if os.path.isfile(candidate)),
            None,
        )
    if tesseract_path:
        pytesseract.pytesseract.tesseract_cmd = tesseract_path

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None
    load_workbook = None

try:
    import win32com.client as win32
except ImportError:
    win32 = None


# Structured BOM exports mark section breaks with "Structure N: <title>" and
# render the top-level (bold) assembly rows starting with a lone "b" token,
# e.g. "b 43C1696G50: ASSY.SEISMIC STRUCTURE, 36"W X ...". Nested rows start
# with "a"/"d" instead and are intentionally skipped.
STRUCTURE_HEADER_RE = re.compile(r"Structure\s+(\d+)\s*:\s*(.+)")

# Structure extracted when the scope checkbox is ticked.
SCOPED_STRUCTURE = "2"
ASSEMBLY_ROW_RE = re.compile(r"^b\s+([^\s:]{2,25}):\s*(.+)$")

# Local copy of the Index2 workbook, kept alongside this script.
INDEX2_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Index2.xlsx")
METADATA_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Metadata.xlsx")
SEMANTIC_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "semantic_aliases.example.json")
_CHECKLIST_NAME = "3. MCOE_Checklist.xlsx"
_CHECKLIST_CANDIDATES = (
    os.path.join(os.path.dirname(os.path.abspath(__file__)), _CHECKLIST_NAME),
    os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "Docs", _CHECKLIST_NAME),
)
CHECKLIST_PATH = next((path for path in _CHECKLIST_CANDIDATES if os.path.exists(path)), _CHECKLIST_CANDIDATES[0])

CHECKLIST_LIVE_URLS = (
    "https://eaton.sharepoint.com/sites/Team_ETO-SystemEnhancements/Shared%20Documents/"
    "Elpaso%20Queue/MCOE%20Audit/3.%20MCOE_Checklist.xlsx",
    "https://modernitar-ear-dhs.etn.com/sites/Team_ETO-SystemEnhancements/Shared%20Documents/"
    "Elpaso%20Queue/MCOE%20Audit/3.%20MCOE_Checklist.xlsx?d=wb51086bc36af488da61f26aec15460d2",
)
CHECKLIST_OPEN_TIMEOUT = 60
CODE_SHEET_NAME = "Code Wise Parameters"
COMMON_SHEET_NAMES = ("Comman parameters", "Common parameters")
PARAMETER_COLUMN_COUNT = 5
VALIDATION_COLORS = {"PASS": "#188038", "FAIL": "#d93025", "UNRESOLVED": "#b06000", "EXCLUDED": "#68788b"}

# Live share link for Index2.xlsx. Raw COM Workbooks.Open() can't resolve this
# (it creates a blank placeholder), so it's launched via the "ms-excel:ofe|u|"
# protocol instead - the same mechanism a browser's "Open in Desktop App" uses,
# which correctly loads the real live/co-authored document.
INDEX2_LIVE_URL = (
    "https://eaton.sharepoint.com/:x:/r/sites/Team_PnSDesignglobal/"
    "PnS%20Design%20Handbook/Spreadsheets/PRLC_switchboard/Index2.xlsx"
    "?d=w00845da0e710471caf4e54c165388371&csf=1&web=1&e=IQrUxh"
)

# QuickPrint PDFs state e.g. "Enclosure Type: Type 3R (nonwalk-in) Flat Roof".
# The captured value picks which sheet to open in Index2's linked workbook.
ENCLOSURE_TYPE_RE = re.compile(r"Enclosure\s+Type\s*:\s*Type\s*([0-9A-Za-z]+)", re.IGNORECASE)

ENCLOSURE_SHEET_MAP = {
    "3R": "BASICSTR-NEMA 3R",
    "1": "BASICSTR-NEMA 1",
}

# Splits e.g. "43C1696G50" -> base "43C1696", suffix "G50".
STYLE_SUFFIX_RE = re.compile(r"^(.+?)([A-Za-z]+\d+)$")


def link_file_name(link_target):
    """File name of an Index2 hyperlink, whether it is a full URL or relative."""

    path = urlparse(link_target).path if not os.path.isabs(link_target) else link_target

    return unquote(path).replace("\\", "/").rsplit("/", 1)[-1]


def excel_document_url(folder, target):
    """Turn Index2's browser-style :x:/r links into Excel document URLs."""
    if os.path.isabs(target):
        return target
    url = urljoin(folder.rstrip("/") + "/", target)
    parsed = urlparse(url)
    path = re.sub(r"^/:[a-z]:/r/", "/", parsed.path, flags=re.I)
    if re.search(r"\.xls[xmb]?$", path, re.I):
        return urlunparse(parsed._replace(path=path, query="", fragment=""))
    return url


def snapshot_cell_value(value):
    """Convert COM dates to serializable, timezone-free Excel wall-clock values."""
    if isinstance(value, datetime):
        return datetime(value.year, value.month, value.day, value.hour, value.minute, value.second, value.microsecond)
    if isinstance(value, datetime_time):
        return datetime_time(value.hour, value.minute, value.second, value.microsecond)
    return value


def normalize_excel_rows(values):
    """Convert Excel COM's scalar, one-dimensional, or two-dimensional values to rows."""

    if values is None:
        return ()

    if isinstance(values, (tuple, list)):
        if not values:
            return ()
        if all(isinstance(row, (tuple, list)) for row in values):
            rows = [tuple(row) for row in values]
        else:
            rows = [tuple(values)]
    else:
        rows = [(values,)]

    column_count = max(map(len, rows), default=0)
    return tuple(tuple(snapshot_cell_value(value) for value in row) + (None,) * (column_count - len(row)) for row in rows)


# Checklist labels that the app can fill from the BOM/Index2/QuickPrint itself.
ASSEMBLY_LABELS = {"assembly no.", "assembly number", "assembly drawing"}
WIDTH_LABELS = {"width (wide)", "width", "structure width"}
DEPTH_LABELS = {"depth", "structure depth"}
ENCLOSURE_LABELS = {"nema 3r/1r", "nema 1/3r", "type 1/3r", "type 1&3r"}


def normalize_code(value):
    """Case/space/punctuation/extension-insensitive key, so 'GND BOND' matches 'GNDBOND'."""

    return re.sub(r"[^A-Z0-9]", "", re.sub(r"\.XLS[XM]?$", "", str(value or "").strip().upper()))


def split_style_number(style_number):

    match = STYLE_SUFFIX_RE.match(style_number)

    if not match or not (
        any(character.isalpha() for character in match.group(1))
        and any(character.isdigit() for character in match.group(1))
        or match.group(1).isdigit() and len(match.group(1)) >= 6
    ):
        return style_number, ""

    return match.group(1), match.group(2)


def build_index2_map(index2_path):
    """Maps a base assembly number (e.g. '43C1696') to the linked workbook's
    hyperlink target, read from Index2's comma-separated family lists."""

    if load_workbook is None or not os.path.exists(index2_path):
        return {}

    try:
        wb = load_workbook(index2_path, data_only=True)
    except Exception:
        return {}

    ws = wb["INDEX2"] if "INDEX2" in wb.sheetnames else wb.active

    mapping = {}

    for (cell,) in ws.iter_rows(min_col=11, max_col=11):
        if not isinstance(cell.value, str):
            continue
        link_cell = ws.cell(row=cell.row, column=5)
        if not link_cell.hyperlink or not link_cell.hyperlink.target:
            continue
        for token in cell.value.split(","):
            token = token.strip().upper()
            if token:
                mapping.setdefault(token, link_cell.hyperlink.target)

    wb.close()

    return mapping


def build_index2_map_from_com(ws):
    """Same mapping as build_index2_map(), but read from a live win32com
    Worksheet - used when there's no (or a stale) local Index2.xlsx copy."""

    used = ws.UsedRange
    values = normalize_excel_rows(used.Value)
    start_row = used.Row
    index_column = 11 - used.Column

    if not values or not 0 <= index_column < used.Columns.Count:
        return {}

    row_to_link = {}
    for hl in ws.Hyperlinks:
        try:
            if hl.Range.Column == 5:
                row_to_link[hl.Range.Row] = hl.Address
        except Exception:
            continue

    mapping = {}

    for offset, row_vals in enumerate(values):

        link_target = row_to_link.get(start_row + offset)

        if not link_target:
            continue

        value = row_vals[index_column]
        if isinstance(value, str):
            for token in value.split(","):
                token = token.strip().upper()
                if token:
                    mapping.setdefault(token, link_target)

    return mapping


@dataclass(frozen=True)
class Index2RootEntry:
    root: str
    row: int
    family: str
    link: str


def read_index2_root_catalog(ws):
    """Inventory column K independently of usable hyperlinks; retain duplicate source rows."""
    used = ws.UsedRange
    values = normalize_excel_rows(used.Value)
    root_column, family_column = 11 - used.Column, 5 - used.Column
    if not values or not 0 <= root_column < used.Columns.Count:
        raise ValueError("Index2 column K is not available in the worksheet data.")
    links = {
        hyperlink.Range.Row: str(hyperlink.Address or "")
        for hyperlink in ws.Hyperlinks if hyperlink.Range.Column == 5
    }
    entries = []
    for offset, row in enumerate(values):
        value = row[root_column]
        if not isinstance(value, str):
            continue
        family = str(row[family_column] or "").strip() if 0 <= family_column < len(row) else ""
        roots = set()
        for token in value.split(","):
            token = token.strip().upper()
            base, suffix = split_style_number(token)
            root = base if re.fullmatch(r"[GH]\d+", suffix) else token
            if re.fullmatch(r"[A-Z0-9-]{6,}", root) and any(c.isdigit() for c in root):
                roots.add(root)
        row_number = used.Row + offset
        entries.extend(Index2RootEntry(root, row_number, family, links.get(row_number, "")) for root in sorted(roots))
    return tuple(entries)


def _format_dim(value):
    """Excel returns whole numbers as floats via COM (36.0) - trim that."""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def com_grid_lookup(ws, base_number, suffix):
    """COM equivalent of the width/depth grid search: column B is the
    per-row width, the row directly above the assembly's merged row-band
    holds the per-column depth. `ws` is a live win32com Worksheet object."""

    found = ws.Cells.Find(
        What=base_number, LookAt=1, LookIn=-4163,
        SearchOrder=1, SearchDirection=1, MatchCase=False, SearchFormat=False,
    )

    if found is None:
        return "", ""

    merge_area = found.MergeArea
    min_row = merge_area.Row
    max_row = min_row + merge_area.Rows.Count - 1
    header_row = min_row - 1
    target_col = found.Column

    depth_headers = {}
    col = target_col + 1
    blank_run = 0

    while blank_run < 3 and (col - target_col) < 100:
        value = ws.Cells(header_row, col).Value
        if value is None or str(value).strip() == "":
            blank_run += 1
        else:
            blank_run = 0
            depth_headers[col] = _format_dim(value)
        col += 1

    suffix_norm = suffix.strip().upper()

    for r in range(min_row, max_row + 1):

        width_value = ws.Cells(r, 2).Value

        for col, depth_value in depth_headers.items():

            cell_value = ws.Cells(r, col).Value

            if cell_value is not None and str(cell_value).strip().upper() == suffix_norm:
                width = _format_dim(width_value) if width_value is not None else ""
                return width, depth_value

    return "", ""


def snapshot_grid_lookup(sheet, base_number, suffix):
    """Same BASICSTR grid lookup as the COM reader, using the captured cells."""
    found = next((
        cell for row in sheet.iter_rows() for cell in row
        if cell.value is not None and str(cell.value).strip().upper() == base_number.strip().upper()
    ), None)
    if found is None:
        return "", ""
    top = bottom = found.row
    for area in sheet.merged_cells.ranges:
        if found.coordinate in area:
            top, bottom = area.min_row, area.max_row
            break
    if top <= 1:
        raise ValueError(f"{sheet.title}: {base_number} has no depth header row.")
    depths = {}
    blank_run = 0
    column = found.column + 1
    while blank_run < 3 and column - found.column < 100:
        value = sheet.cell(top - 1, column).value
        if value is None or not str(value).strip():
            blank_run += 1
        else:
            blank_run = 0
            depths[column] = _format_dim(value)
        column += 1
    for row in range(top, bottom + 1):
        for column, depth in depths.items():
            value = sheet.cell(row, column).value
            if value is not None and str(value).strip().upper() == suffix.strip().upper():
                width = sheet.cell(row, 2).value
                return _format_dim(width) if width is not None else "", depth
    return "", ""


@dataclass(frozen=True)
class DimensionResult:
    width: str = ""
    depth: str = ""
    message: str = ""
    source: str = ""


@dataclass(frozen=True)
class ParameterRule:
    parameter_name: str
    keywords: tuple[str, ...]


@dataclass(frozen=True)
class ConfigurationRecord:
    assembly_code: str
    category: str
    parameters: tuple[ParameterRule, ...]
    manually_mapped: bool = False
    checklist_row: int = 0


@dataclass(frozen=True)
class CodeMapping:
    name: str
    code: str
    parameters: tuple[str, ...]
    aliases: tuple[str, ...] = ()
    manually_mapped: bool | None = None
    checklist_row: int = 0


def _clean_cells(row):

    return ["" if cell is None else str(cell).strip() for cell in row]


def parse_common_parameters(rows):
    """Names from the 'Comman parameters' sheet: first column, header row skipped, no repeats."""

    names = []

    for row in list(rows or ())[1:]:
        name = _clean_cells(row[:1])[0] if row else ""
        if name and name not in names:
            names.append(name)

    return tuple(names)


def parse_code_mappings(rows):
    """'Code Wise Parameters' rows (header first): Name, Code, Parameter 1..N ->
    {normalized code: CodeMapping}. At least PARAMETER_COLUMN_COUNT parameters; more when a
    row defines them."""

    body = [_clean_cells(row) for row in list(rows or ())[1:]]
    count = max(
        [PARAMETER_COLUMN_COUNT]
        + [max((i for i, cell in enumerate(cells[2:], start=1) if cell), default=0) for cells in body]
    )
    mappings = {}

    for cells in body:
        cells = (cells + [""] * (2 + count))[:2 + count]
        key = normalize_code(cells[1])
        if not key or key == "CODE" or key in mappings:
            continue
        parameters = tuple(cells[2:])
        mappings[key] = CodeMapping(cells[0], cells[1], parameters, parameters)

    return mappings


def read_local_sheet_values(path, sheet_name):

    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed.")

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        return list(workbook[sheet_name].iter_rows(values_only=True))
    finally:
        workbook.close()


@dataclass(frozen=True)
class PDFExtraction:
    pages: tuple[str, ...] = ()
    ocr_pages: tuple[int, ...] = ()
    issues: tuple[str, ...] = ()

    @property
    def text(self):
        return "\n".join(self.pages)


def extract_pdf_pages(path):
    """Extract selectable PDF text and use OCR for pages with little/no text."""

    if pdfplumber is None:
        return PDFExtraction(issues=("pdfplumber is not installed.",))

    pages = []
    ocr_pages = []
    issues = []
    rendered_pdf = None

    try:
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                if len(text.strip()) < 20:
                    if fitz is None or pytesseract is None or Image is None:
                        missing = []
                        if fitz is None:
                            missing.append("PyMuPDF")
                        if pytesseract is None:
                            missing.append("pytesseract")
                        if Image is None:
                            missing.append("Pillow")
                        issues.append(
                            f"Page {page_number}: OCR unavailable (missing {', '.join(missing)}). "
                            "Install Python packages with 'pip install pymupdf pytesseract pillow' "
                            "and install the Tesseract OCR executable on PATH."
                        )
                    else:
                        try:
                            if rendered_pdf is None:
                                rendered_pdf = fitz.open(path)
                            pixmap = rendered_pdf.load_page(page_number - 1).get_pixmap(
                                dpi=300, alpha=False,
                            )
                            with Image.open(BytesIO(pixmap.tobytes("png"))) as image:
                                ocr_text = pytesseract.image_to_string(image)
                            if ocr_text.strip():
                                text = "\n".join(part for part in (text, ocr_text) if part.strip())
                                ocr_pages.append(page_number)
                            else:
                                issues.append(f"Page {page_number}: OCR returned no readable text.")
                        except Exception as exc:
                            issues.append(
                                f"Page {page_number}: OCR failed ({type(exc).__name__}: {exc})."
                            )
                pages.append(text)
    except Exception as exc:
        issues.append(f"PDF read failed ({type(exc).__name__}: {exc}).")
    finally:
        if rendered_pdf is not None:
            rendered_pdf.close()

    return PDFExtraction(tuple(pages), tuple(ocr_pages), tuple(issues))


# ----------------------------------------------------------------------
# Excel "Current Extraction Coverage" parameters, extracted with the ECOE audit's
# regex approach. Patterns marked (ECOE) are copied unchanged from
# ECOE_AUDIT._extract_parameter_values / _unit_device_values.
# ----------------------------------------------------------------------

EXCEL_PARAMETERS = tuple(checklist_parameters.COMMON_PARAMETERS)

_ECOE_BUS_MATERIAL = [  # (ECOE) "Bus Material"
    r"(?:BUS\s+MATERIAL|MATERIAL)[^\n,;:=-]{0,8}[:=\-]?\s*((?:TIN|SILVER)\s+PLATED\s+(?:COPPER|ALUMINUM|ALUMINIUM)|COPPER|ALUMINUM|ALUMINIUM)",
    r"\b((?:SILVER|TIN)\s+PLATED\s+(?:COPPER|ALUMINUM|ALUMINIUM)|COPPER|ALUMINUM|ALUMINIUM)\s+BUS\b",
    r"\b(SILVER\s+PLATED\s+CU)\b",
    r"\b\d{3,5}\s+AMP\s+(SILVER\s+PLATED\s+(?:COPPER|CU))\b",
    r"\b\d{3,5}\s+(SILVER\s+PLATED\s+(?:COPPER|CU))\b",
]
_ECOE_BUS_AMPS = [  # (ECOE) "Bus Amps"
    r"DEVICE\s+SIZE\s*\(?\s*FRAME\s+AMPS?\s*\)?[^\n,;]{0,24}?\b(\d{3,5})\b",
    r"\bCHASSIS\s+STRUCTURE\s+(\d{3,5})\b",
    r"(?:MAIN\s+BUS|BUS\s+AMP(?:S|ACITY)?)\s*[:=]?\s*(\d{3,5})\s*A?\b",
    r"\b(\d{3,5})\s*(?:AMP|A)\s+(?:SILVER|TIN|COPPER|ALUMINUM)",
    r"(?:BUS|BUS\s+AMP(?:S|ACITY)?)[^\n,;]{0,24}?(\d{3,5})\s*A\b",
]
_ECOE_VOLTAGE = [  # (ECOE) "Service voltage"
    r"(?:SERVICE|SYS\.?\s*VOLT(?:AGE)?)\s*[:=]?\s*(\d{2,4}\s*(?:Y\s*)?/?\s*\d{2,4}\s*V[^\n,;]{0,24}|\d{2,4}\s*V[^\n,;]{0,24})",
    r"(\d{2,4}\s*/\s*\d{2,4}\s*V\s+(?:1|3)\s+PHASE\s+(?:3|4)\s+WIRE)",
    r"(\d{2,4}\s*/\s*\d{2,4}\s*V\s*(?:1|3)\s*-?\s*PHASE)",
    r"(\d{2,4}\s*/\s*\d{2,4}\s*V\s*(?:1|3)\s*-?\s*PHASE\s*(?:3|4)\s*-?\s*WIRE)",
    r"(\d{2,4}\s*(?:Y\s*)?/?\s*\d{2,4}\s*V)",
]
_ECOE_DEPTH = [  # (ECOE) "Depth (Inner)"
    r"DEPTH\s*\(?\s*INNER\s*\)?\s*[-:]?\s*IN\.?[^\n,;]{0,12}?(\d{1,3}(?:\.\d+)?)",
    r"\bFRONT\s+(\d{1,3}(?:\.\d{2})?)\s*\[\s*\d{3,4}(?:\.\d+)?\s*\]",
    r"\bFRONT\s+REAR\s+DEPTH\s+STRUCTURE\b[^\n]{0,180}?\b0\.00\s*\[\s*0\s*\]\s+(\d{1,3}(?:\.\d{2})?)\s*\[\s*\d{3,4}(?:\.\d+)?\s*\]",
]
_ECOE_HEIGHT = [  # (ECOE) "Height"
    r"HEIGHT\s*[-:]?\s*INCHES?[^\n,;]{0,12}?(\d{1,3}(?:\.\d+)?)",
    r"(?:WEIGHT,?\s*APPROX\.|SECTION\s+NO\.?)[^\n]{0,300}?\b(\d{2,3}\.\d{2})\s*\[\s*\d{3,4}(?:\.\d+)?\s*\]",
]
_ECOE_WIDTH = [  # (ECOE) "Width"
    r"W\s*D?TH\s*[-:]?\s*INCHES?[^\n,;]{0,12}?(\d{1,3}(?:\.\d+)?)",
    r"SECTION\s+WIDTH\s+IN\s*\[\s*MM\s*\][^\n]{0,300}?\b(\d{2,3}\.\d{2})\s*\[\s*\d{3,4}(?:\.\d+)?\s*\]",
]
# Structured BOM descriptions give size as e.g. '36"W X 24D'.
_BOM_WIDTH = r"\b(\d{2,3})\s*\"?\s*W\b"
_BOM_DEPTH = r"\b(\d{2,3})\s*\"?\s*D\b"

EXCEL_PARAMETER_PATTERNS = {
    "Material": _ECOE_BUS_MATERIAL,
    "Structure Depth": _ECOE_DEPTH + [_BOM_DEPTH],
    "Assembly No.": [r"\bASSEMBLY\s+(?:NO\.?|NUMBER)\s*[:=]\s*([A-Z0-9-]{4,25})"],
    "Width": _ECOE_WIDTH + [_BOM_WIDTH],
    "Amperage": _ECOE_BUS_AMPS,
    "Depth": _ECOE_DEPTH + [_BOM_DEPTH],
    "Bus Rating": [r"\bBUS\s+RATING\s*[:=]?\s*(\d{3,5})\s*A?\b"] + _ECOE_BUS_AMPS,
    "Height": _ECOE_HEIGHT,
    "Voltage": _ECOE_VOLTAGE,
    "3/1 Phase": [r"\b(1|3)\s*-?\s*PHASE\b"],
    "3/4 Wire": [r"\b(3|4)\s*-?\s*WIRE\b"],
    "Seismic/Non-Seismic": [
        r"\b(?:SEISMIC(?:\s*/\s*NON[- ]?SEISMIC)?|STRUCTURE\s+RATING)\s*[:=]\s*(NON[- ]?SEISMIC|SEISMIC|YES|NO)\b",
        r"^(NON[- ]?SEISMIC|SEISMIC)$",
    ],
    "Heat Test/Density": [
        r"\bHEAT\s+TEST(?:\s*/\s*DENSITY)?\s*[:=\-]\s*(\d+(?:\.\d+)?\s*(?:DEG(?:REES)?\s*)?C?\b|[A-Z]{3,12}\b)",
        r"\bDENSITY\s*[:=\-]\s*(\d+(?:\.\d+)?\s*%?|[A-Z]{3,12}\b)",
        r"\b(HEAT\s+TEST)\b(?!\s*[:=/\-])",
    ],
    "NEMA 1/3R": [  # (ECOE) "Enclosure" type group, plus BOM 'NEMA/3R'
        r"\b(TYPE\s*1|TYPE\s*3R)\b",
        r"\bNEMA\s*/?\s*(3R|1)\b",
    ],
    "Flat Roof/Sloped Roof": [r"\b(FLAT|SLOPED)\s+ROOF\b"],  # (ECOE) "Enclosure" roof group
    "Poles": [  # (ECOE) "Poles" and the _unit_device_values pole pattern
        r"\b(?:FEEDER\s+BREAKER|MAIN\s+BREAKER|BREAKER)[^\n,;]{0,100}?\b(\d{1,2})P\b",
        r"\b(\d+)\s*P\b",
    ],
    "Part Name": [r"\bPART\s+NAME\s*[:=]\s*([A-Z0-9][A-Z0-9 &()./\-]{1,60}?)(?=\s{2,}|\s[A-Z ]+[:=]|[,;]|$)"],
    "Utility Name": [
        r"\bUTILITY\s+NAME\s*[:=]\s*([A-Z0-9][A-Z0-9 &./\-]{1,60}?)(?=\s{2,}|\s[A-Z ]+[:=]|[,;]|$)",
        r"\b(?:PULL\s+SECTION|UTILITY\s+METERING)\s*-\s*([A-Z][A-Z0-9&. ]{2,40}?)\s*(?:\(|,|;|$)",
    ],
    "Utility Amperage": [  # (ECOE) Frame Amps pull-section/metering pattern
        r"\b(\d+)\s*A\s+(?:UTILITY\s+METERING|(?:UNDER\s*GROUND\s+)?PULL\s+SECTION)\b",
        r"\bUTILITY\s+AMP(?:S|ERAGE)?\s*[:=]\s*(\d{2,5})\b",
    ],
    "With/Without Lugs": [r"\b(WITHOUT\s+LUGS?|WITH\s+LUGS?|W/O\s+LUGS?)\b"],
    "Bus Position": [
        r"\bBUS\s+(?:POSITION|LOCATION)\s*[:=\-]?\s*(AUXILIARY\s*\(?TOP\)?(?:\s+AND\s+STANDARD\s*\(?BOTTOM\)?)?|STANDARD\s*\(?BOTTOM\)?)",
        r"\bBUS\s+(?:POSITION|LOCATION)\s*[:=\-]?\s*(TOP|BOTTOM|FRONT|REAR|LEFT|RIGHT)\b",
    ],
    "Bus Capacity": [r"\bBUS\s+CAPACITY\s*[:=]?\s*(\d{3,5})\s*A?\b"] + _ECOE_BUS_AMPS,
    "Chassis Rating (Amperage)": [
        r"\bCHASSIS\s+STRUCTURE\s+(\d{3,5})\b",  # (ECOE) "Bus Amps"
        r"\bCHASSIS\s+RATING[^\n,;]{0,20}?(\d{3,5})\s*A?\b",
    ],
    "Number of Lugs": [  # (ECOE) lug_pattern count group
        r"\b(?:NEUTRAL\s+TERMINALS?|TERMINALS?)(?:,?\s*MECHANICAL)?\s*,?\s*\((\d+)\)",
        r"\bNUMBER\s+OF\s+LUGS\s*[:=]?\s*(\d+)\b",
    ],
    "Number of Drillings": [
        r"\b(\d+)\s*DRILLINGS?\b",
        r"\((\d+)\)\s*(?:EUSERC\s+)?PRESS\s*BOLTS?\b",
    ],
    "Service/Non-Service Enterance": [r"\b(NON[- ]?SERVICE\s+ENTR[AE]NCE|SERVICE\s+ENTR[AE]NCE)\b"],
    "Section Width": [_ECOE_WIDTH[1], r"\bSECTION\s+WIDTH\s*[:=]?\s*(\d{2,3}(?:\.\d+)?)"],
    "EUSERC Drill Number": [r"\bDRILL(?:ING)?\s*(?:NUMBER|NO\.?|#)\s*[:=]?\s*([A-Z0-9-]{1,10})\b"],
    "Width (Wide)": _ECOE_WIDTH + [_BOM_WIDTH],
}


def ecoe_normalize(text):
    """(ECOE) normalize(): collapse whitespace and upper-case."""

    return re.sub(r"\s+", " ", str(text)).strip().upper()


def excel_variable_name(parameter):
    """'Width (Wide)' -> 'Width_Wide'; '3/1 Phase' -> '_3_1_Phase' (a name cannot start with a digit)."""

    name = re.sub(r"\W+", "_", parameter.strip()).strip("_")

    return f"_{name}" if name[:1].isdigit() else name


def extract_excel_parameter_values(source_text):
    """{Excel parameter: [values]} from one source, using the ECOE value clean-up."""

    text = ecoe_normalize(source_text)
    values = {}

    for parameter, patterns in EXCEL_PARAMETER_PATTERNS.items():
        flattened = []
        for pattern in patterns:
            for match in re.findall(pattern, text, re.I):
                if isinstance(match, tuple):
                    match = " ".join(part for part in match if part)
                value = ecoe_normalize(match).replace("ALUMINIUM", "ALUMINUM")
                if parameter == "Material":
                    value = re.sub(r"\bCU\b", "COPPER", value)
                elif parameter == "Seismic/Non-Seismic":
                    value = {"YES": "SEISMIC", "NO": "NON-SEISMIC", "NON SEISMIC": "NON-SEISMIC"}.get(value, value)
                value = re.sub(r"\s+(?:MINIMUM|INTERRUPT|HEAT TEST|GROUND BUS BOLTED TO).*$", "", value)
                if value and value not in flattened:
                    flattened.append(value)
        # (ECOE) drop values that are strict substrings of another value
        values[parameter] = [
            item for item in flattened
            if not any(item != other and item in other for other in flattened)
        ]

    return values


class ExcelParameterVariables:
    """One attribute per Excel parameter, named after it (see excel_variable_name):
    the first value found, or '' when not found. `values` keeps every value found,
    keyed by the exact Excel parameter name."""

    def __init__(self, values, names=EXCEL_PARAMETERS):
        self.values = {name: list(values.get(name, ())) for name in names}

        for name, found in self.values.items():
            setattr(self, excel_variable_name(name), found[0] if found else "")

    @classmethod
    def from_sources(cls, source_texts, assembly_numbers=(), names=EXCEL_PARAMETERS):
        merged = {}

        for text in source_texts:
            for name, found in extract_excel_parameter_values(text).items():
                target = merged.setdefault(name, [])
                target.extend(value for value in found if value not in target)

        for number in assembly_numbers:
            if number not in merged.setdefault("Assembly No.", []):
                merged["Assembly No."].append(number)

        return cls(merged, names)

    def resolve(self, query):
        """Exact Excel parameter name for a typed name or variable name, else None."""

        key = normalize_code(query)

        for name in self.values:
            if key in (normalize_code(name), normalize_code(excel_variable_name(name))):
                return name

        return None

    def describe(self, name):
        found = self.values[name]

        if not found:
            return f"{name} = not found"

        extra = f"   (all: {', '.join(found)})" if len(found) > 1 else ""

        return f"{name} = {found[0]}{extra}"


class ConfigurationRepository:
    """Loads checklist fields and their Database_Index2 label variants."""

    def __init__(self, checklist_path=None):
        self.checklist_path = checklist_path if checklist_path is not None else (
            METADATA_PATH if os.path.exists(METADATA_PATH) else CHECKLIST_PATH
        )
        self._records = self.load_configuration()
        self._parameter_rules = self._collect_parameter_rules()

    @staticmethod
    def _normalized(value):
        return re.sub(r"\s+", " ", str(value).strip()).casefold() if value is not None else ""

    def load_configuration(self):
        if load_workbook is None or not os.path.exists(self.checklist_path):
            return {}

        workbook = load_workbook(self.checklist_path, data_only=True, read_only=True)
        try:
            if "MCOE Checklist" not in workbook.sheetnames:
                return {}

            rows = tuple(workbook["MCOE Checklist"].iter_rows(values_only=True))
            codes = {normalize_code(row[1]) for row in rows[1:] if len(row) > 1 and row[1]}
            database = {}
            if "Database_Index2" in workbook.sheetnames:
                for row in workbook["Database_Index2"].iter_rows(values_only=True):
                    for offset in (0, 1):
                        if offset < len(row) and normalize_code(row[offset]) in codes:
                            database.setdefault(normalize_code(row[offset]), row[offset + 1:])
                            break
            records = {}

            for row_number, row in enumerate(rows[1:], start=2):
                code_value = row[1] if len(row) > 1 else None
                if not isinstance(code_value, str) or not code_value.strip():
                    continue

                parameters = []
                aliases = database.get(normalize_code(code_value), ())
                for index, name in enumerate(row[2:13]):
                    alias = aliases[index] if index < len(aliases) else None
                    if not isinstance(name, str) or not name.strip():
                        name = alias
                    if not isinstance(name, str) or not name.strip():
                        continue

                    keywords = []
                    for candidate in (name, alias):
                        if isinstance(candidate, str) and candidate.strip():
                            if self._normalized(candidate) not in {self._normalized(item) for item in keywords}:
                                keywords.append(candidate.strip())
                    parameters.append(ParameterRule(name.strip(), tuple(keywords)))

                code = code_value.strip().upper()
                category = row[0]
                flag = row[13] if len(row) > 13 else None
                records[code] = ConfigurationRecord(
                    code,
                    str(category).strip() if category is not None else "",
                    tuple(parameters),
                    flag is True or isinstance(flag, str) and flag.strip().upper() == "TRUE",
                    row_number,
                )

            return records
        finally:
            workbook.close()

    def _collect_parameter_rules(self):
        rules = {}
        for record in self._records.values():
            for rule in record.parameters:
                key = self._normalized(rule.parameter_name)
                existing = rules.get(key)
                if existing is None:
                    rules[key] = rule
                    continue
                keywords = list(existing.keywords)
                known = {self._normalized(keyword) for keyword in keywords}
                for keyword in rule.keywords:
                    if self._normalized(keyword) not in known:
                        keywords.append(keyword)
                        known.add(self._normalized(keyword))
                rules[key] = ParameterRule(existing.parameter_name, tuple(keywords))
        return tuple(rules.values())

    def get_configuration(self, assembly_code):
        key = str(assembly_code).strip().upper()
        return self._records.get(key) or next((
            record for code, record in self._records.items() if normalize_code(code) == normalize_code(key)
        ), None)

    def code_mappings(self):
        """Checklist C:M fields and Database_Index2 aliases, with N kept as metadata."""
        return {
            normalize_code(record.assembly_code): CodeMapping(
                record.category, record.assembly_code,
                tuple(rule.parameter_name for rule in record.parameters),
                tuple(next((label for label in rule.keywords if label != rule.parameter_name), rule.parameter_name)
                      for rule in record.parameters),
                record.manually_mapped, record.checklist_row,
            )
            for record in self._records.values()
        }

    def reload(self):
        self._records = self.load_configuration()
        self._parameter_rules = self._collect_parameter_rules()

    def use_code_mappings(self, mappings):
        """Replaces the checklist with Sheet1 rows: each parameter name is its own keyword."""

        self._records = {
            key: ConfigurationRecord(
                mapping.code, mapping.name,
                tuple(
                    ParameterRule(name, tuple(dict.fromkeys((name, alias or name))))
                    for name, alias in zip(mapping.parameters, mapping.aliases or mapping.parameters) if name
                ),
                mapping.manually_mapped is True,
                mapping.checklist_row,
            )
            for key, mapping in mappings.items() if any(mapping.parameters)
        }
        self._parameter_rules = self._collect_parameter_rules()

    def extract_parameters(self, text, source):
        repository = ParameterRepository()
        for rule in self._parameter_rules:
            labels = sorted(rule.keywords, key=len, reverse=True)
            alternatives = "|".join(re.escape(label) for label in labels)
            pattern = re.compile(rf"(?im)^\s*(?:{alternatives})\s*[:=]\s*(\S(?:.*\S)?)\s*$")
            for match in pattern.finditer(text or ""):
                value = match.group(1)
                if ClassifierWorkbookParser.parameter_key(rule.parameter_name) == "PHASE":
                    # ABC feeder labels describe phase order, not the system's phase count.
                    phase = re.fullmatch(r"(1|3|SINGLE|THREE)(?:\s*-?\s*(?:PHASE|PH))?", value, re.I)
                    if phase:
                        value = "1" if phase[1].upper() in {"1", "SINGLE"} else "3"
                    elif re.match(r"[ABC](?:\s*[-/,]?\s*[ABC]){0,2}(?=\s|$)", value, re.I):
                        repository.add_parameter("Phase Order", value, source)
                        continue
                repository.add_parameter(rule.parameter_name, value, source)
        return repository


class ParameterRepository:
    """Retains extracted values with provenance, including conflicting sources."""

    def __init__(self):
        self.parameters = {}

    def add_parameter(self, name, value, source):
        if name is None or value is None or not str(value).strip():
            return
        key = ConfigurationRepository._normalized(name)
        entry = (str(value).strip(), str(source))
        self.parameters.setdefault(key, []).append(entry)

    def get_parameter(self, name):
        entries = self.parameters.get(ConfigurationRepository._normalized(name), [])
        values = {value for value, _source in entries}
        return next(iter(values)) if len(values) == 1 else None

    def get_entries(self, name):
        return tuple(self.parameters.get(ConfigurationRepository._normalized(name), ()))

    def has_conflict(self, name):
        entries = self.get_entries(name)
        return len({value for value, _source in entries}) > 1

    def merge_sources(self, *repositories):
        for repository in repositories:
            for key, entries in repository.parameters.items():
                target = self.parameters.setdefault(key, [])
                for entry in entries:
                    if entry not in target:
                        target.append(entry)

    def clear(self):
        self.parameters.clear()

    @property
    def value_count(self):
        return sum(len(entries) for entries in self.parameters.values())

    @property
    def conflict_count(self):
        return sum(
            len({value for value, _source in entries}) > 1
            for entries in self.parameters.values()
        )


def section_bus_requirements(source_text, source):
    requirements = {}
    for line in source_text.replace("\\P", "\n").splitlines():
        text = ecoe_normalize(line)
        sections = re.search(
            r"\b(?:SECTIONS?|STRUCTURES?)\s+(\d+(?:\s*(?:,|&|AND|-|TO)\s*\d+)*)\b",
            text, re.I,
        )
        ratings = re.findall(r"\b(DENSITY|HEAT\s+TEST(?:ED)?)\s+(?:RATED\s+)?BUS(?:ES)?\b", text, re.I)
        if not sections or not ratings:
            continue
        numbers = []
        for member in re.split(r"\s*(?:,|&|AND)\s*", sections.group(1), flags=re.I):
            bounds = re.split(r"\s*(?:-|TO)\s*", member, flags=re.I)
            first, last = int(bounds[0]), int(bounds[-1])
            if 1 <= first <= last <= 999:
                numbers.extend(range(first, last + 1))
        for number in numbers:
            repository = requirements.setdefault(number, ParameterRepository())
            for rating in ratings:
                value = "DENSITY" if rating.upper() == "DENSITY" else "HEAT TEST"
                repository.add_parameter("Heat Test/Density", value, f"{source}: {line.strip()}")
    return requirements


def is_bus_description(description):
    return bool(
        re.search(r"\b(?:X-?BUS|CROSS\s*BUS|BUS(?:ES)?|SPLICES?)\b", description, re.I)
        or re.search(r"\bUGPS\b.*\d.*\b(?:CU|CUAG|CUSN|AL)\b", description, re.I)
    )


def assembly_bus_requirements(requirements, section_number, description):
    if not is_bus_description(description):
        return None
    return requirements.get(int(section_number)) if section_number is not None else None


def extract_table_selectors(text, source):
    repository = ParameterRepository()
    for match in re.finditer(r"(?im)^\s*Bus\s+(?:Amps|Amperage)\s*:\s*(\d+)\b", text):
        repository.add_parameter("Bus Amperage", match[1], source)
    for name in ("NG Drawout", "Structure Material", "Test Plug", "Metering Configuration", "Table Row", "Door Vents"):
        pattern = rf"(?im)^\s*{re.escape(name)}\s*[:=]\s*([^\r\n]+)"
        for match in re.finditer(pattern, text):
            repository.add_parameter(name, match.group(1).strip(), source)
    return repository


class DWGFile:
    """Read native DWG text and section blocks without closing existing documents.

    Known Config door blocks supply hinge side only when their width and top-left
    insertion point identify one unrotated, unscaled SECTION_NUMBER structure.
    """

    def __init__(self, file_path):
        self.file_path = file_path
        self.section_parameters = {}

    @staticmethod
    def section_block_parameters(sections, doors, source):
        result = {}
        for number, x, y, width, height, depth in sections:
            repository = result.setdefault(str(number), ParameterRepository())
            for name, value in (("Width (Wide)", width), ("Height", height), ("Depth", depth)):
                if value:
                    repository.add_parameter(name, _format_dim(value), f"{source}: section {number} block")
        for name, x, y, width, side in doors:
            owners = {
                str(number) for number, sx, sy, sw, sh, _depth in sections
                if sh and abs(sx - x) < 0.01 and abs(sy + sh - y) < 0.01 and abs(sw - width) < 0.01
            }
            if len(owners) == 1:
                result[next(iter(owners))].add_parameter("Hinge Side", side, f"{source}: {name}")
        return result

    @staticmethod
    def _entity_text(entity):
        object_name = str(getattr(entity, "ObjectName", ""))
        if object_name in {"AcDbText", "AcDbMText", "AcDbAttribute"}:
            return str(getattr(entity, "TextString", "") or "")

        if object_name == "AcDbBlockReference":
            attributes = []
            try:
                for attribute in entity.GetAttributes():
                    tag = str(getattr(attribute, "TagString", "") or "").strip()
                    value = str(getattr(attribute, "TextString", "") or "").strip()
                    if tag and value:
                        attributes.append(f"{tag}: {value}")
            except Exception:
                return ""
            return "\n".join(attributes)

        return ""

    def read_text(self):
        if win32 is None:
            raise RuntimeError("pywin32 (win32com) is not installed.")

        try:
            app = win32.GetActiveObject("AutoCAD.Application")
        except Exception:
            app = win32.Dispatch("AutoCAD.Application")

        path = os.path.abspath(self.file_path)
        document = next((
            candidate for candidate in app.Documents
            if os.path.normcase(str(candidate.FullName)) == os.path.normcase(path)
        ), None)
        opened = document is None
        if opened:
            document = app.Documents.Open(path, True)
        try:
            text_chunks = []
            sections, doors = [], []
            for space_name in ("ModelSpace", "PaperSpace"):
                space = getattr(document, space_name)
                for entity in space:
                    text = self._entity_text(entity)
                    if text:
                        text_chunks.append(text)
                    if space_name != "ModelSpace" or entity.ObjectName != "AcDbBlockReference":
                        continue
                    name = str(getattr(entity, "Name", ""))
                    door = re.fullmatch(r"Config\d+_(\d+)_(Rt|Lft)", name, re.I)
                    number = re.search(r"(?im)^SECTION_NUMBER:\s*(\d+)\s*$", text)
                    if not (door or number):
                        continue
                    if abs(float(entity.Rotation)) > 0.001 or any(
                        abs(float(getattr(entity, axis)) - 1) > 0.001
                        for axis in ("XScaleFactor", "YScaleFactor")
                    ):
                        continue
                    x, y, _z = entity.InsertionPoint
                    if door:
                        doors.append((name, x, y, float(door[1]), "RIGHT" if door[2].upper() == "RT" else "LEFT"))
                    elif entity.IsDynamicBlock:
                        properties = {str(p.PropertyName).lower(): p.Value for p in entity.GetDynamicBlockProperties()}
                        width, height = properties.get("width"), properties.get("height")
                        if isinstance(width, (float, int)) and isinstance(height, (float, int)):
                            depth = properties.get("depth")
                            sections.append((number[1], x, y, width, height, depth if isinstance(depth, (int, float)) else None))
            self.section_parameters = self.section_block_parameters(sections, doors, os.path.basename(self.file_path))
            return "\n".join(text_chunks)
        finally:
            if opened:
                document.Close(False)

    def extract_parameters(self, configuration_repository, drawing_text=None):
        if drawing_text is None:
            drawing_text = self.read_text()

        return configuration_repository.extract_parameters(
            drawing_text, os.path.basename(self.file_path),
        )


def table_dimension_lookup(values, family, base, suffix):
    """Read only labelled layouts verified in the linked family workbooks."""

    def normalized(value):
        return str(value).strip().upper() if value is not None else ""

    base_pattern = re.compile(rf"(?<![A-Z0-9]){re.escape(base.upper())}(?![A-Z0-9])")
    matches = []
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            if base_pattern.search(normalized(value)):
                matches.append((row_index, column_index))

    if not matches:
        return None

    results = []
    for row_index, column_index in matches:
        row = values[row_index]
        if family == "REARCOVER":
            if column_index + 1 >= len(row) or normalized(row[column_index + 1]) != suffix.upper():
                continue
            width_columns = {
                header_column
                for header in values[:row_index]
                for header_column, value in enumerate(header)
                if normalized(value) == "STRUCTURE WIDTH"
            }
            if len(width_columns) == 1:
                width_value = row[next(iter(width_columns))]
                if width_value is not None:
                    results.append(DimensionResult(
                        width=_format_dim(width_value),
                        message="Width found; depth not specified in rear-cover table.",
                    ))
        elif family == "WPFLAT":
            for header_index in range(row_index + 1, min(row_index + 5, len(values))):
                header = [normalized(value) for value in values[header_index]]
                if "DEPTH" not in header or "GROUP" not in header:
                    continue
                depth_column = header.index("DEPTH")
                group_column = header.index("GROUP")
                for data_row in values[header_index + 1:header_index + 3]:
                    if normalized(data_row[group_column]) == suffix.upper() and data_row[depth_column] is not None:
                        results.append(DimensionResult(
                            depth=_format_dim(data_row[depth_column]),
                            message="Depth range from workbook; width not specified.",
                        ))
        elif family == "BUSSPLIC":
            headers = {
                normalized(value)
                for header in values[row_index + 1:row_index + 6]
                for value in header[column_index:column_index + 8]
            }
            if {"OFFSET", "BARS PER PHASE", "GROUPS", "AMP RATING"} <= headers:
                for data_row in values[row_index + 6:row_index + 7]:
                    if any(normalized(value) == suffix.upper() for value in data_row[column_index:column_index + 8]):
                        results.append(DimensionResult(
                            message="Group found; table lists offset, bus size and amp ratings, not Width/Depth.",
                        ))

    if not results:
        return DimensionResult(message="Base found; group or dimension layout not supported.")
    if len({(result.width, result.depth, result.message) for result in results}) != 1:
        return DimensionResult(message="Conflicting table matches; dimensions unresolved.")
    return results[0]


GROUP_HEADERS = {"GROUP", "GROUPS"}
HEADER_LOOKBACK = 2
CONTEXT_LOOKBACK = 6
FULL_SHEET_CELL_LIMIT = 40000
NUMERIC_VALUE_RE = re.compile(r"^([0-9]*\.?[0-9]+)\s*[A-Z\"'%]*$")


def _cell_text(value):

    return "" if value is None else _format_dim(value).upper()


def _cell_at(row, index):

    return row[index] if index < len(row) else None


def values_match(cell, wanted):
    """Equal ignoring case/punctuation, or numerically equal when units differ (2000A vs 2000)."""

    left, right = _cell_text(cell), _cell_text(wanted)

    if not left or not right:
        return False

    if normalize_code(left) == normalize_code(right):
        return True

    left_number, right_number = NUMERIC_VALUE_RE.match(left), NUMERIC_VALUE_RE.match(right)

    return bool(left_number and right_number and float(left_number.group(1)) == float(right_number.group(1)))


@dataclass(frozen=True)
class TableCandidate:
    sheet_name: str
    header_labels: tuple[frozenset[str], ...]
    group_column: int
    rows: tuple[tuple, ...]
    context: str
    keyword_score: int = 0
    classifier_type: str = ""
    parameter_names: tuple[str, ...] = ()

    def column_for(self, parameter_name):
        key = normalize_code(parameter_name)

        for index, labels in enumerate(self.header_labels):
            if key and key in labels:
                return index

        return None

    def rows_for_base(self, base):
        """Rows naming the base assembly, else all rows if the base only appears in the title area."""

        pattern = re.compile(rf"(?<![A-Z0-9]){re.escape(base.upper())}(?![A-Z0-9])")
        own = tuple(row for row in self.rows if any(pattern.search(_cell_text(cell)) for cell in row))

        if own:
            return own

        return self.rows if pattern.search(self.context) else ()


class WorkbookParser:
    """Finds tables that carry a single 'Group'/'Groups' header column."""

    @staticmethod
    def extract_tables(values, sheet_name):

        tables = []

        for row_index, row in enumerate(values):

            group_columns = [column for column, value in enumerate(row) if _cell_text(value) in GROUP_HEADERS]

            if len(group_columns) != 1:
                continue

            labels = [set() for _ in row]
            header_index = row_index

            while header_index >= max(0, row_index - HEADER_LOOKBACK):
                header = values[header_index]
                if not any(_cell_text(value) for value in header):
                    break
                for column, value in enumerate(header):
                    if column < len(labels) and _cell_text(value):
                        labels[column].add(normalize_code(value))
                header_index -= 1

            data = []
            for data_row in values[row_index + 1:]:
                if not any(_cell_text(value) for value in data_row):
                    break
                if any(_cell_text(value) in GROUP_HEADERS for value in data_row):
                    break
                data.append(tuple(data_row))

            context_rows = values[max(0, row_index - CONTEXT_LOOKBACK):row_index]
            context = " ".join(_cell_text(value) for context_row in context_rows for value in context_row)

            tables.append(TableCandidate(
                sheet_name,
                tuple(frozenset(label_set) for label_set in labels),
                group_columns[0],
                tuple(data),
                context,
            ))

        return tables


class ClassifierWorkbookParser:
    """Adapt the generalized parser's configurations to the audit table interface."""

    @staticmethod
    @lru_cache(maxsize=1)
    def semantics():
        with open(SEMANTIC_PATH, encoding="utf-8") as stream:
            return json.load(stream)

    @staticmethod
    @lru_cache(maxsize=512)
    def parameter_key(name):
        key = normalize_code(name)
        for parent, info in ClassifierWorkbookParser.semantics()["parameters"].items():
            if key in {normalize_code(label) for label in (parent, *info.get("aliases", ()))}:
                key = normalize_code(parent)
                break
        return {
            "WIDTHWIDE": "STRUCTUREWIDTH",
            "HEATTESTDENSITY": "TEST",
            "31PHASE": "PHASE",
            "31PHASE31": "PHASE",
            "34WIRE": "WIRES",
            "34WIRE34": "WIRES",
            "WIRE": "WIRES",
        }.get(key, key)

    @staticmethod
    @lru_cache(maxsize=512)
    def parameter_labels(name):
        key = ClassifierWorkbookParser.parameter_key(name)
        labels = {normalize_code(name), key}
        for parent, info in ClassifierWorkbookParser.semantics()["parameters"].items():
            if ClassifierWorkbookParser.parameter_key(parent) == key:
                labels.update(normalize_code(label) for label in (parent, *info.get("aliases", ())))
        for alias in ("Width (Wide)", "Heat Test/Density", "3/1 Phase", "3/1 Phase (3/1)"):
            if ClassifierWorkbookParser.parameter_key(alias) == key:
                labels.add(normalize_code(alias))
        return frozenset(labels)

    @staticmethod
    def extract_tables(workbook, file_name, base, parser=None):
        from Classifier import parse_assembly

        tables = []
        parsed_sheets = parser.parse(base) if parser is not None else parse_assembly(
            file_name, base, METADATA_PATH, SEMANTIC_PATH, workbook=workbook,
        )
        for parsed in parsed_sheets:
            configurations = parsed["result"]["Configurations"]
            for configuration in configurations:
                group = configuration.get("Group No.", configuration.get("Group"))
                if not group:
                    raise ValueError(f"{file_name} / {parsed['sheet_name']}: configuration has no group.")
                parameters = tuple(
                    name for name in configuration
                    if name not in ("Group", "Group No.")
                    and ClassifierWorkbookParser.parameter_key(name) != "ASSEMBLYNUMBER"
                )
                names = ("Assembly Number", *parameters, "Group")
                tables.append(TableCandidate(
                    parsed["sheet_name"],
                    tuple(ClassifierWorkbookParser.parameter_labels(name) for name in names),
                    len(names) - 1,
                    ((base, *(configuration[name] for name in parameters), group),),
                    base,
                    classifier_type=parsed["parser_type"],
                    parameter_names=parameters,
                ))
        return tuple(tables)

    @staticmethod
    def values_match(derived, wanted):
        from Classifier import AnchoredTableParser

        keyword_mapping = {
            normalize_code(key): value
            for key, value in ClassifierWorkbookParser.semantics().get("keyword_mapping", {}).items()
        }

        derived = _cell_text(keyword_mapping.get(normalize_code(derived), derived))
        wanted = _cell_text(keyword_mapping.get(normalize_code(wanted), wanted))
        derived = _cell_text(AnchoredTableParser.TESTS.get(derived, derived))
        wanted = _cell_text(AnchoredTableParser.TESTS.get(wanted, wanted))
        if derived == "ALL":
            return bool(wanted)
        minimum = re.fullmatch(r'(\d+(?:\.\d+)?)\s*[DW"]?\s*MIN\.?', derived)
        number = NUMERIC_VALUE_RE.fullmatch(wanted)
        if minimum:
            return bool(number and float(number[1]) >= float(minimum[1]))
        bounds = re.fullmatch(r"(\d+)\s*-\s*(\d+)\s*A?", derived)
        wanted_bounds = re.fullmatch(r"(\d+)\s*-\s*(\d+)\s*A?", wanted)
        if bounds:
            if wanted_bounds:
                return float(bounds[1]) <= float(wanted_bounds[1]) <= float(wanted_bounds[2]) <= float(bounds[2])
            return bool(number and float(bounds[1]) <= float(number[1]) <= float(bounds[2]))
        if wanted_bounds:
            return False
        return _grid_value_matches(derived, wanted)

    @staticmethod
    def consistent_value(values):
        values = list(dict.fromkeys(values))
        if not values or any(
            not ClassifierWorkbookParser.values_match(left, right)
            for index, left in enumerate(values) for right in values[index + 1:]
        ):
            return ""
        specific = [
            value for value in values
            if MATERIAL_ALIASES.get(_cell_text(value), _cell_text(value)) in {"CUAG", "CUSN", "CUTN", "TIN PLATED ALUMINUM"}
        ]
        return specific[0] if specific else values[0]

    @staticmethod
    def identify_group(tables, base, wanted):
        """Require distinguishing inputs; never choose using the BOM's group."""
        amperages = [(name, value) for name, value in wanted.items()
                     if ClassifierWorkbookParser.parameter_key(name) == "AMPERAGE" and value]
        if len(amperages) == 1:
            name, value = amperages[0]
            bounds = re.fullmatch(r"(\d+)\s*-\s*(\d+)\s*A?", str(value))
            columns = [next((table.column_for(parameter) for parameter in table.parameter_names
                             if ClassifierWorkbookParser.parameter_key(parameter) == "AMPERAGE"), None)
                       for table in tables]
            if bounds and columns and all(column is not None for column in columns):
                ratings = {str(row[column]) for table, column in zip(tables, columns)
                           for row in table.rows_for_base(base)}
                if ratings and all(re.fullmatch(r"\d+", rating) for rating in ratings):
                    low, high = map(int, bounds.groups())
                    if str(low) not in ratings or str(high) not in ratings or low > high:
                        return "", f"Classifier: amperage range {value} is not bounded by listed ratings."
                    results = [
                        ClassifierWorkbookParser.identify_group(tables, base, {**wanted, name: rating})
                        for rating in sorted(ratings, key=int) if low <= int(rating) <= high
                    ]
                    groups = {group for group, _message in results}
                    if len(groups) == 1 and "" not in groups:
                        return results[0][0], f"{results[0][1]} Same group for all listed ratings in {value}A."
                    reason = next((message for group, message in results if not group),
                                  "Classifier: different groups apply within the amperage range.")
                    return "", reason
        inputs = {}
        for name, value in wanted.items():
            if value:
                inputs.setdefault(ClassifierWorkbookParser.parameter_key(name), []).append(value)
        if "DOORTYPE" in inputs and "ASSEMBLYTYPE" not in inputs:
            inputs["ASSEMBLYTYPE"] = ["DOOR"]
        groups = set()
        missing = set()
        sources = set()
        names = {name for table in tables for name in table.parameter_names}
        invariants = {
            name for name in names
            if not name.startswith("Unparsed ")
            and all(table.column_for(name) is not None for table in tables)
            and len({str(row[table.column_for(name)]) for table in tables for row in table.rows_for_base(base)}) == 1
        }
        for table in tables:
            required = {name: inputs.get(ClassifierWorkbookParser.parameter_key(name), ()) for name in table.parameter_names}
            for row in table.rows_for_base(base):
                if any(
                    values and not all(ClassifierWorkbookParser.values_match(
                        row[table.column_for(name)], value,
                    ) for value in values)
                    for name, values in required.items()
                ):
                    continue
                absent = {name for name, values in required.items() if not values and name not in invariants}
                if not required:
                    absent.add("unparsed table selection criteria")
                missing.update(absent)
                if not absent:
                    groups.add(_cell_text(row[table.group_column]))
                    sources.add(f"{table.sheet_name} ({table.classifier_type})")
        if len(groups) == 1 and not missing:
            return next(iter(groups)), f"Classifier group found in {', '.join(sorted(sources))}."
        if len(groups) > 1:
            return "", f"Classifier: multiple groups match: {', '.join(sorted(groups))}."
        if missing:
            return "", f"Classifier: missing job inputs: {', '.join(sorted(missing))}."
        return "", "Classifier: no configuration matches the job inputs."


class TableSelector:

    @staticmethod
    def select_best(tables, base, parameter_names):
        """Tables for the base assembly with the most checklist parameters as columns."""

        scored = []

        for table in tables:
            if not table.rows_for_base(base):
                continue
            score = sum(table.column_for(name) is not None for name in parameter_names)
            if score:
                scored.append(replace(table, keyword_score=score))

        if not scored:
            return ()

        best = max(table.keyword_score for table in scored)

        return tuple(table for table in scored if table.keyword_score == best)


class GroupDeterminationEngine:

    @staticmethod
    def identify_group(tables, base, parameter_values):
        """Returns (group, message): the one group whose row matches every parameter value."""

        candidates = TableSelector.select_best(tables, base, list(parameter_values))

        if not candidates:
            return "", "No table with a Group column matches the assembly and checklist parameters."

        groups = {}

        for table in candidates:
            columns = {
                name: table.column_for(name)
                for name in parameter_values if table.column_for(name) is not None
            }
            for row in table.rows_for_base(base):
                if all(values_match(_cell_at(row, column), parameter_values[name]) for name, column in columns.items()):
                    group = _cell_text(_cell_at(row, table.group_column))
                    if group:
                        groups.setdefault(group, set()).add(table.sheet_name)

        if not groups:
            return "", "No table row matches the extracted parameter values."

        if len(groups) > 1:
            return "", f"Multiple groups match: {', '.join(sorted(groups))}."

        group = next(iter(groups))

        return group, f"Group found in {', '.join(sorted(groups[group]))}."


class TableValueExtractor:

    @staticmethod
    def extract(tables, base, group, labels):
        """Reads each parameter from the family-workbook row of this base + group.

        `labels` maps a parameter name to the header labels to look for. A value is
        returned only when every matching row agrees on it."""

        group_key = normalize_code(group)
        found = {}

        if not group_key:
            return found

        for name, candidates in labels.items():
            values = set()
            for table in tables:
                rows = [
                    row for row in table.rows_for_base(base)
                    if normalize_code(_cell_text(_cell_at(row, table.group_column))) == group_key
                ]
                for label in candidates:
                    column = table.column_for(label)
                    if column is None:
                        continue
                    for row in rows:
                        value = _cell_at(row, column)
                        if value is not None and _cell_text(value):
                            values.add(_format_dim(value))
            if len(values) == 1:
                found[name] = next(iter(values))

        return found


GROUP_CODE_RE = re.compile(r"^(?:[A-Z]{1,3}\d{1,3}|N/A)$")
DRAWING_TOKEN_RE = re.compile(r"\b(?:\d{1,2}[A-Z]\d{4,5}|\d[A-Z]\d{5})\b")  # 42C5495, 2C11909, 1A85142
GRID_MATERIAL_RE = re.compile(r"\b(ALUM(?:INUM)?|CUAG|CUSN|CUTN|CU|AL|COPPER)\b")
GRID_HEAT_DENSITY_RE = re.compile(r"\b(HEAT TEST(?:ED)?|HT|DENSITY|DEN)\b")
SHEET_WIDTH_RE = re.compile(r"(\d{2,3})\s*\"?\s*W(?:IDE)?\b")
GRID_DEPTH_RE = re.compile(r"\b(\d{2,3}\s*D(?:\s*MIN)?)\b")
GRID_AMPERAGE_RE = re.compile(r"\b(\d{2,5}(?:\s*-\s*\d{2,5})?)\s*A\b")


@dataclass(frozen=True)
class GridMatch:
    sheet_name: str
    headers: tuple[str, ...]   # column headings above the group cell, nearest first
    section: str               # title row above the table, e.g. 'THREE PHASE'
    row_labels: tuple[str, ...]
    title_text: str            # all title rows of the sheet
    group: str = ""
    axis_values: tuple[tuple[str, str], ...] = ()
    required_parameters: tuple[str, ...] = ()

    @property
    def sheet_width(self):
        for text in (self.sheet_name.upper(), self.title_text):
            match = SHEET_WIDTH_RE.search(text)
            if match:
                return match.group(1)
        return ""


def _base_row_above(texts, row_index, base_pattern):
    """Row of the assembly drawing for a table cell: its own row, else the nearest row above,
    unless a different drawing number is met first (that row belongs to another assembly)."""

    for index in range(row_index, max(-1, row_index - 25), -1):
        row_texts = texts[index]
        if any(base_pattern.search(text) for text in row_texts):
            return index
        if any(DRAWING_TOKEN_RE.search(text) for text in row_texts):
            return None

    return None


def find_dimension_grid_matches(values, sheet_name, base, group):
    matches = []
    recognized = False
    for header_index, header in enumerate(values):
        for drawing_column in range(1, len(header) - 1):
            if tuple(normalize_code(cell) for cell in header[drawing_column - 1:drawing_column + 2]) != (
                "STRUCTUREWIDTH", "ASSEMBLYDRAWING", "DEPTH",
            ):
                continue
            recognized = True
            if header_index + 1 >= len(values):
                continue
            depth_row = values[header_index + 1]
            depths = {}
            for column in range(drawing_column + 1, len(depth_row)):
                value = _cell_text(depth_row[column])
                if not re.fullmatch(r"\d+(?:\.\d+)?", value):
                    break
                depths[column] = _format_dim(float(value))
            current_base = ""
            section = _cell_text(_cell_at(values[header_index - 1], drawing_column - 1)) if header_index else ""
            for row in values[header_index + 2:]:
                width = _cell_text(_cell_at(row, drawing_column - 1))
                if not re.fullmatch(r"\d+(?:\.\d+)?", width):
                    break
                drawing = _cell_text(_cell_at(row, drawing_column))
                if drawing:
                    current_base = drawing
                if current_base != base.upper():
                    continue
                for column, depth in depths.items():
                    actual = _cell_text(_cell_at(row, column))
                    if not GROUP_CODE_RE.fullmatch(actual) or actual == "N/A" or (group and actual != group.upper()):
                        continue
                    axes = [("Width (Wide)", _format_dim(float(width))), ("Depth", depth)]
                    height = re.search(r'(\d+)\s*"?\s*HIGH\b', section)
                    if height:
                        axes.append(("Height", height.group(1)))
                    enclosure = re.search(r"NEMA\s*(3R|1)\b", sheet_name.upper())
                    if enclosure:
                        axes.append(("NEMA 1/3R", enclosure.group(1)))
                    matches.append(GridMatch(
                        sheet_name, (), section, (), "", actual,
                        tuple(axes), tuple(name for name, _value in axes),
                    ))
    return matches if recognized else None


def find_assembly_column_matches(values, sheet_name, base, group):
    matches = []
    style_pattern = re.compile(rf"{re.escape(base.upper())}(G\d+)")
    for header_index, header in enumerate(values):
        width_columns = [column for column, value in enumerate(header) if _cell_text(value) == "WIDTH"]
        if not width_columns or header_index < 2:
            continue
        section = next((
            _cell_text(next(value for value in row if _cell_text(value)))
            for row in reversed(values[:header_index - 1])
            if sum(bool(_cell_text(value)) for value in row) == 1
        ), "")
        for column, label in enumerate(header):
            if _cell_text(label) != "ASSEMBLY":
                continue
            width_column = next((index for index in reversed(width_columns) if index < column), None)
            if width_column is None:
                continue
            channel = _cell_text(_cell_at(values[header_index - 1], column)).rstrip(" *")
            height_text = _cell_text(_cell_at(values[header_index - 2], width_column))
            height = re.search(r'(\d+)\s*"?\s*HIGH\b', height_text)
            for row in values[header_index + 1:]:
                width = _cell_text(_cell_at(row, width_column))
                if not re.fullmatch(r"\d+(?:\.\d+)?", width):
                    break
                style = style_pattern.fullmatch(_cell_text(_cell_at(row, column)))
                if not style or (group and style.group(1) != group.upper()):
                    continue
                axes = [("Width (Wide)", _format_dim(float(width)))]
                if height:
                    axes.append(("Height", height.group(1)))
                if channel:
                    axes.append(("Channel Type", channel))
                matches.append(GridMatch(sheet_name, (height_text, channel), section, (), "", style.group(1), tuple(axes)))
    return matches


def end_cap_position(text):
    text = _cell_text(text)
    if not re.search(r"\bEND\s*CAPS?\b", text):
        return ""
    positions = {
        value for value, pattern in (
            ("LEFT", r"\b(?:LEFT|LH)\b"),
            ("RIGHT", r"\b(?:RIGHT|RH)\b"),
            ("CENTER", r"\b(?:INTERMEDIATE|CENTER|CENTRE|CTR)\b"),
        ) if re.search(pattern, text)
    }
    return next(iter(positions)) if len(positions) == 1 else ""


def find_smart_table_matches(values, sheet_name, base, group):
    """Keep each INPUT/OUTPUT table, row enclosure type and output position scoped."""
    matches = []
    recognized = False
    aliases = {
        "STRUCTUREWIDTH": "Width (Wide)", "STRUCTUREDEPTH": "Depth",
        "STRUCTUREHEIGHT": "Height", "SEISMIC": "Seismic/Non-Seismic",
        "STRUCTUREMATERIAL": "Structure Material", "NGDRAWOUT": "NG Drawout",
    }
    style_pattern = re.compile(rf"{re.escape(base.upper())}(G\d+|H\d+)")
    for marker_index, marker in enumerate(values[:-1]):
        inputs = next((column for column, cell in enumerate(marker) if _cell_text(cell) == "INPUT"), None)
        outputs = next((column for column, cell in enumerate(marker) if _cell_text(cell) == "OUTPUT"), None)
        if inputs is None or outputs is None or inputs >= outputs:
            continue
        end = next((column for column in range(outputs + 1, len(marker)) if _cell_text(marker[column])), len(marker))
        header = values[marker_index + 1]
        for row in values[marker_index + 2:]:
            if not any(_cell_text(cell) for cell in row) or (
                any(_cell_text(cell) == "INPUT" for cell in row)
                and any(_cell_text(cell) == "OUTPUT" for cell in row)
            ):
                break
            axes = []
            for column in range(inputs, outputs):
                name = _cell_text(_cell_at(header, column))
                value = _cell_text(_cell_at(row, column))
                if name and value and value != "N/A":
                    parameter = aliases.get(normalize_code(name), name.title())
                    if normalize_code(name) == "MATERIAL" and re.search(r"\b(?:GALVANIZED|STEEL)\b", value):
                        parameter = "Structure Material"
                    axes.append((parameter, value))
            enclosure_text = next((value for name, value in axes if name == "Structure Type"), sheet_name.upper())
            enclosure = re.search(r"TYPE\s*(3R|1)\b", enclosure_text)
            mixed_enclosure = re.search(r"TYPE\s*1\s*(?:&|AND|/|OR)\s*(?:TYPE\s*)?3R\b", enclosure_text)
            if enclosure and not mixed_enclosure:
                axes.append(("NEMA 1/3R", enclosure.group(1)))
            for column in range(outputs, end):
                style = style_pattern.fullmatch(_cell_text(_cell_at(row, column)))
                if style:
                    recognized = True
                    if not group or style.group(1) == group.upper():
                        output_label = _cell_text(_cell_at(header, column))
                        position = end_cap_position(output_label)
                        output_axes = axes + ([("End Cap Position", position)] if position else [])
                        matches.append(GridMatch(
                            sheet_name, (), output_label, (), "",
                            style.group(1), tuple(output_axes), tuple(name for name, _value in output_axes),
                        ))
    return matches if recognized else None


def find_labeled_table_matches(values, sheet_name, base, group):
    matches = []
    recognized = False
    base_pattern = re.compile(rf"(?<![A-Z0-9]){re.escape(base.upper())}(?![A-Z0-9])")
    for header_index, header in enumerate(values):
        labels = [_cell_text(value) for value in header]
        depth_column = next((column for column, label in enumerate(labels) if label == "DEPTH"), None)
        group_column = next((column for column, label in enumerate(labels) if label == "GROUP"), None)
        description_column = next((column for column, label in enumerate(labels) if label == "DESCRIPTION"), None)
        widths = {column: hit.group(1) for column, label in enumerate(labels) if (hit := SHEET_WIDTH_RE.fullmatch(label))}
        if not ((depth_column is not None and group_column is not None) or (description_column is not None and widths)):
            continue
        context = values[max(0, header_index - 4):header_index]
        if not any(base_pattern.search(_cell_text(value)) for row in context for value in row):
            continue
        recognized = True
        section = next((
            _cell_text(value) for row in reversed(context) for value in row
            if _cell_text(value) and not base_pattern.search(_cell_text(value))
        ), "")
        for row in values[header_index + 1:]:
            if depth_column is not None and group_column is not None:
                depth = re.fullmatch(r"(\d+(?:\.\d+)?)D", _cell_text(_cell_at(row, depth_column)))
                if not depth:
                    break
                actual = _cell_text(_cell_at(row, group_column))
                if GROUP_CODE_RE.fullmatch(actual) and actual != "N/A" and (not group or actual == group.upper()):
                    matches.append(GridMatch(sheet_name, (), section, (), "", actual, (("Depth", depth.group(1)),), ("Depth",)))
                continue
            label = _cell_text(_cell_at(row, description_column))
            if label in ("DESCRIPTION", "WITH TEST PLUG", "W/O TEST PLUG", "BLANK DOORS", "L-SHAPED MOUNTING PAN") or DRAWING_TOKEN_RE.fullmatch(label):
                break
            if not label:
                continue
            for column, width in widths.items():
                actual = _cell_text(_cell_at(row, column))
                if not GROUP_CODE_RE.fullmatch(actual) or actual == "N/A" or (group and actual != group.upper()):
                    continue
                axes = [("Width (Wide)", width), ("Table Row", label)]
                if "TEST PLUG" in section:
                    axes.append(("Test Plug", "WITHOUT" if section.startswith("W/O") else "WITH"))
                matches.append(GridMatch(sheet_name, (), section, (), "", actual, tuple(axes), tuple(name for name, _value in axes)))
    return matches if recognized else None


def find_grid_group_matches(values, sheet_name, base, group):
    """Grid layouts (e.g. WCUTO22XMAIN, UGPORD2) put the group code in a cell; the column
    headings above it and the row labels beside it say material, density, amperage, width..."""

    smart_matches = find_smart_table_matches(values, sheet_name, base, group)
    if smart_matches is not None:
        return smart_matches

    labeled_matches = find_labeled_table_matches(values, sheet_name, base, group)
    if labeled_matches is not None:
        return labeled_matches

    dimension_matches = find_dimension_grid_matches(values, sheet_name, base, group)
    if dimension_matches is not None:
        return dimension_matches

    assembly_matches = find_assembly_column_matches(values, sheet_name, base, group)

    base_pattern = re.compile(rf"(?<![A-Z0-9]){re.escape(base.upper())}(?![A-Z0-9])")
    group_text = (group or "").strip().upper()
    texts = [[_cell_text(value) for value in row] for row in values]
    titles = " ".join(
        next(t for t in row_texts if t)
        for row_texts in texts if sum(bool(t) for t in row_texts) == 1
    )
    matches = list(assembly_matches)

    for row_index, row_texts in enumerate(texts):
        for column, text in enumerate(row_texts):
            is_group = text == group_text if group_text else bool(GROUP_CODE_RE.match(text)) and text != "N/A"
            if not is_group or _base_row_above(texts, row_index, base_pattern) is None:
                continue

            headers, section = [], ""
            for above in range(row_index - 1, -1, -1):
                above_texts = texts[above]
                if sum(bool(t) for t in above_texts) == 1:
                    section = next(t for t in above_texts if t)
                    if re.search(r"(?:PHASE|NEUTRAL) CROSS BUS ASSEMBLY", section):
                        for parent in range(above - 1, max(-1, above - 50), -1):
                            parent_texts = texts[parent]
                            if not any(parent_texts):
                                break
                            if "ASSY DWG" in parent_texts and "AMPS" in parent_texts:
                                header_rows = []
                                for header_row in texts[parent:above]:
                                    if any(GROUP_CODE_RE.fullmatch(value) for value in header_row):
                                        break
                                    header_rows.append(header_row)
                                for header_row in reversed(header_rows):
                                    cell = next((header_row[index] for index in range(min(column, len(header_row) - 1), -1, -1) if header_row[index]), "")
                                    if cell:
                                        headers.append(cell)
                                break
                    break
                # Merged headings keep their text only in the leftmost cell.
                cell = next((above_texts[i] for i in range(min(column, len(above_texts) - 1), -1, -1) if above_texts[i]), "")
                if cell and not GROUP_CODE_RE.match(cell):
                    headers.append(cell)

            matches.append(GridMatch(
                sheet_name, tuple(headers), section,
                tuple(t for index, t in enumerate(row_texts) if t and index < column), titles, text,
                required_parameters=("Width (Wide)",) if normalize_code(section).startswith("CROSSBUSMOUNTING") else (),
            ))

    return matches


def grid_parameter_values(match):
    """Excel parameter name -> value, read from one grid match (column headings first,
    then the row labels, then the sheet name/title)."""

    found = {}
    sources = (*match.headers, *match.row_labels)

    def first(pattern, texts):
        for text in texts:
            hit = pattern.search(text)
            if hit:
                return hit
        return None

    material = first(GRID_MATERIAL_RE, sources)
    if material:
        found["Material"] = material.group(1)

    heat_density = first(GRID_HEAT_DENSITY_RE, sources)
    if heat_density:
        found["Heat Test/Density"] = "DENSITY" if heat_density.group(1) in ("DENSITY", "DEN") else "HEAT TEST"

    phase = re.search(r"\b(THREE|SINGLE|3|1)\s*-?\s*(?:PHASE|PH)\b", " ".join((match.section, *match.headers)))
    if phase:
        found["3/1 Phase"] = "1" if phase.group(1) in ("SINGLE", "1") else "3"

    amperage = first(GRID_AMPERAGE_RE, match.row_labels)
    if amperage:
        found["Amperage"] = re.sub(r"\s+", "", amperage.group(1))

    width = first(SHEET_WIDTH_RE, sources)
    found["Width (Wide)"] = width.group(1) if width else match.sheet_width
    if not found["Width (Wide)"]:
        del found["Width (Wide)"]

    depth = first(GRID_DEPTH_RE, sources)
    if depth:
        found["Depth"] = re.sub(r"\s+", " ", depth.group(1))

    drillings = first(re.compile(r"(\d+)\s*DRILLINGS?"), sources)
    if drillings:
        found["Number of Drillings"] = drillings.group(1)

    found.update(match.axis_values)
    return found


def merge_grid_matches(matches, structure_width=""):
    """Parameter values all matches agree on. Several matches on different sheets are
    narrowed to the sheet whose width equals the structure's width first; amperages that
    differ (e.g. 2500A and 3000A rows sharing a group) are listed together."""

    all_values = [grid_parameter_values(match) for match in matches]

    if len({match.sheet_name for match in matches}) > 1 and structure_width:
        narrowed = [values for values in all_values if values.get("Width (Wide)") == str(structure_width)]
        all_values = narrowed or all_values

    merged = {}
    for key in {key for values in all_values for key in values}:
        found = sorted({values[key] for values in all_values if key in values})
        if len(found) == 1:
            merged[key] = found[0]
        elif key == "Amperage":
            merged[key] = "/".join(found)

    return merged


def description_parameter_values(description):
    """Parameters readable from an assembly description, e.g. '800A CUAG DEN 3PH'."""

    from Classifier import AnchoredTableParser

    text = (description or "").upper()
    found = grid_parameter_values(GridMatch("", (text,), "", (text,), ""))
    found.update(AnchoredTableParser.attributes(text))
    position = end_cap_position(text)
    if position:
        found["End Cap Position"] = position
    door_widths = re.search(r"\b(?:DBL|DOUBLE)[ ,]*(?:DR|DOORS?)?[ ,]*(\d{2})\s*-\s*(\d{2})\b", text)
    if door_widths:
        found["Structure Width - Left (Or Single) Door"] = door_widths[1]
        found["Structure Width - Right Door"] = door_widths[2]
    if re.search(r"\b\d+\s*A\s+MAX\b", text):
        found.pop("Amperage", None)

    # Preserve the range rather than treating its upper endpoint as the job rating.
    short = re.search(r"\b(\d{2})\s*-\s*(\d{4})\b", text)
    full = re.search(r"\b(\d{3,4})\s*-\s*(\d{3,4})\s*A?\b", text)
    if full:
        found["Amperage"] = f"{full.group(1)}-{full.group(2)}"
    elif short:
        found["Amperage"] = f"{short.group(1)}00-{short.group(2)}"

    height = re.search(r'(\d+)\s*"?\s*(?:HIGH|HEIGHT|H)\b', text)
    if height:
        found["Height"] = height.group(1)
    seismic = re.search(r"\b(NON[- ]?SEISMIC|SEISMIC)\b", text)
    if seismic:
        found["Seismic/Non-Seismic"] = "NON-SEISMIC" if seismic.group(1).startswith("NON") else "SEISMIC"

    return found


MATERIAL_ALIASES = {
    "COPPER": "CU", "ALUMINUM": "AL", "ALUM": "AL",
    "SILVER PLATED COPPER": "CUAG", "TIN PLATED COPPER": "CUSN",
}
COPPER_MATERIALS = {"CU", "CUAG", "CUSN", "CUTN"}
ALUMINUM_MATERIALS = {"AL", "TIN PLATED ALUMINUM"}


def _grid_value_matches(derived, wanted):
    """Sheet value vs job value: '2500/3000' lists, '2500-3000' ranges (either side),
    '30D MIN', and generic CU/AL headings for their supported plated materials."""

    def clean(text):
        text = re.sub(r"\s*MIN$", "", _cell_text(text).strip())
        text = re.sub(r"\s+ONLY$", "", text)
        return MATERIAL_ALIASES.get(text, text)

    wanted = clean(wanted)
    wanted_number = NUMERIC_VALUE_RE.match(wanted)
    wanted_bounds = re.fullmatch(r"(\d+)\s*-\s*(\d+)", wanted)

    if re.fullmatch(r"SEISMIC \((?:FREESTANDING OR WALL SECURED|WALL SECURED OR FREESTANDING)\)", clean(derived)):
        return wanted in {"SEISMIC", "SEISMIC (FREESTANDING)", "SEISMIC (WALL SECURED)"}
    alternatives = re.split(r"\s+OR\s+(?![^()]*\))", str(derived), flags=re.I)
    if len(alternatives) > 1 and not any(re.search(r"\d", part) for part in alternatives):
        return any(_grid_value_matches(part, wanted) for part in alternatives)

    for part in str(derived).split("/"):
        part = clean(part)
        if values_match(part, wanted):
            return True
        if part in COPPER_MATERIALS and wanted in COPPER_MATERIALS and "CU" in (part, wanted):
            return True
        if part in ALUMINUM_MATERIALS and wanted in ALUMINUM_MATERIALS and "AL" in (part, wanted):
            return True
        bounds = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        number = NUMERIC_VALUE_RE.match(part)
        if bounds and wanted_number and int(bounds.group(1)) <= float(wanted_number.group(1)) <= int(bounds.group(2)):
            return True
        if wanted_bounds and number and int(wanted_bounds.group(1)) <= float(number.group(1)) <= int(wanted_bounds.group(2)):
            return True

    return False


def table_description_tokens(text):
    text = str(text or "").upper()
    text = re.sub(r"\b(?:X[ -]?BUS|CROSS[ -]?BUS|CRS\s+BUS)\b", "CROSSBUS", text)
    text = re.sub(r"\b(?:MTG|MNTG|MOUNT)\b", "MOUNTING", text)
    text = re.sub(r"\bAUX\b", "AUXILIARY", text)
    ignored = {"ASSY", "ASSEMBLY", "SWBC", "SWBD", "TABLE", "TABLES", "FOR", "THE", "AND", "OF", "WIDE", "WIDTH", "DWG", "DRAWING"}
    return {word for word in re.findall(r"\b[A-Z]+\b", text) if word not in ignored}


def select_grid_tables(matches, common_parameters, description=""):
    identities = {(match.sheet_name, match.section) for match in matches}
    if len(identities) <= 1:
        return matches, ""

    description_tokens = table_description_tokens(description)
    table_tokens = {
        identity: table_description_tokens(identity[1] or identity[0])
        for identity in identities
    }
    if "CROSSBUS" in description_tokens:
        relevant = {
            identity: tokens for identity, tokens in table_tokens.items()
            if "CROSSBUS" in tokens
            and ("MOUNTING" in tokens) == ("MOUNTING" in description_tokens)
        }
        if relevant:
            table_tokens = relevant
            matches = [match for match in matches if (match.sheet_name, match.section) in relevant]
        elif any("CROSSBUS" in tokens for tokens in table_tokens.values()):
            return [], "Table selection unresolved: the assembly description does not match a cross-bus table type."

    shared_tokens = set.intersection(*table_tokens.values())
    scores = {
        identity: len((tokens - shared_tokens) & description_tokens)
        for identity, tokens in table_tokens.items()
    }
    best_score = max(scores.values())
    if best_score:
        selected = {identity for identity, score in scores.items() if score == best_score}
        matches = [match for match in matches if (match.sheet_name, match.section) in selected]
    if len({(match.sheet_name, match.section) for match in matches}) == 1:
        return matches, ""

    prefix = "CROSSBUSMOUNTING"
    mounting_tables = [match for match in matches if normalize_code(match.section).startswith(prefix)]
    if not mounting_tables:
        remaining_tokens = {frozenset(table_description_tokens(match.section or match.sheet_name)) for match in matches}
        if description_tokens and len(remaining_tokens) > 1 and any("CROSSBUS" in tokens for tokens in remaining_tokens):
            return [], "Table selection unresolved: the assembly description does not uniquely identify a cross-bus table."
        return matches, ""

    for name in ("Bus Position", "Part Name"):
        values = common_parameters.get(name, ())
        if isinstance(values, str):
            values = (values,)
        selectors = {normalize_code(value) for value in values if value}
        if not selectors:
            continue
        if len(selectors) != 1:
            return [], f"Table selection unresolved: conflicting saved {name} values."
        selector = next(iter(selectors))
        selected = [
            match for match in mounting_tables
            if selector in (normalize_code(match.section), normalize_code(match.section)[len(prefix):])
        ]
        if not selected:
            return [], f"Table selection unresolved: saved {name} does not identify a mounting table."
        return selected, ""

    return [], "Table selection unresolved: save Bus Position or Part Name before matching a mounting group."


def workbook_migration(sheets):
    for _sheet_name, rows in sheets:
        for row in rows[:8]:
            for value in row:
                text = _cell_text(value)
                target = re.search(r"MAINTAINED ON SPREADSHEET\s*:\s*(.+?\.XLSX?)\b", text)
                if target:
                    return target.group(1), text
                if "OBSOLET" in text or "OBSOLTED" in text:
                    return "", text
    return None


def find_grid_group(sheets, base, wanted, common_parameters=None, description=""):
    """(group, message) for grid-layout family sheets: the one group whose table agrees with
    every job value. Fixed root applicability fields need no extra input, but dimensions
    and varying selectors do. Supplied values must agree even with fixed fields."""

    migration = workbook_migration(sheets)
    if migration is not None:
        return "", f"Workbook retired: {migration[1]}. Resolve the current workbook through Index2."
    matches = [
        match for sheet_name, values in sheets
        for match in find_grid_group_matches(values, sheet_name, base, "")
    ]
    wanted = dict(wanted)
    selector_text = description.upper()
    for name in ("Test Plug", "Table Row", "Metering Configuration", "Door Vents"):
        saved = (common_parameters or {}).get(name, ())
        if isinstance(saved, str):
            saved = (saved,)
        if len(saved) > 1:
            return "", f"Table selection unresolved: conflicting saved {name} values."
        if len(saved) == 1:
            wanted[name] = saved[0]
    if re.search(r"\b(?:W/O|WITHOUT|NO)\s+TEST\s+PLUG\b", selector_text):
        wanted["Test Plug"] = "WITHOUT"
    elif re.search(r"\bWITH\s+TEST\s+PLUG\b", selector_text):
        wanted["Test Plug"] = "WITH"
    row_text = normalize_code(wanted.get("Table Row") or wanted.get("Metering Configuration") or wanted.get("Door Vents") or description)
    row_options = {
        value for match in matches for name, value in match.axis_values
        if name == "Table Row" and normalize_code(value) in row_text
    }
    if row_options:
        longest = max(len(normalize_code(value)) for value in row_options)
        rows = [value for value in row_options if len(normalize_code(value)) == longest]
        if len(rows) == 1:
            wanted["Table Row"] = rows[0]
    matches, selection_error = select_grid_tables(matches, common_parameters or {}, description)
    if selection_error:
        return "", selection_error

    by_table = {}
    for match in matches:
        by_table.setdefault((match.group, match.sheet_name, match.section), []).append(match)

    if not by_table:
        return "", "Base not found in grid-layout sheets."

    wanted = {normalize_code(name): value for name, value in wanted.items() if value}
    groups = {}
    missing = set()
    configurations = [
        {normalize_code(name): value for name, value in grid_parameter_values(match).items()}
        for table_matches in by_table.values() for match in table_matches
    ]
    required = {normalize_code(name) for table_matches in by_table.values()
                for match in table_matches for name in match.required_parameters}
    invariants = {
        name for name in required
        if ClassifierWorkbookParser.parameter_key(name) not in {
            "STRUCTUREWIDTH", "STRUCTUREDEPTH", "STRUCTUREHEIGHT", "HEIGHT",
        }
        and all(name in configuration for configuration in configurations)
        and len({configuration[name] for configuration in configurations}) == 1
    }

    for (group, sheet_name, section), matches in by_table.items():
        for match in matches:
            derived = {normalize_code(name): value for name, value in grid_parameter_values(match).items()}
            compared = [name for name in wanted if name in derived]
            if any(not _grid_value_matches(derived[name], wanted[name]) for name in compared):
                continue
            absent = [name for name in match.required_parameters
                      if not wanted.get(normalize_code(name)) and normalize_code(name) not in invariants]
            if absent:
                missing.update(absent)
                continue
            if compared and all(_grid_value_matches(derived[name], wanted[name]) for name in compared):
                evidence = f"{sheet_name} / {section}" if section else sheet_name
                if match.required_parameters:
                    inputs = ", ".join(f"{name}={wanted[normalize_code(name)]}" for name in match.required_parameters
                                       if normalize_code(name) in wanted)
                    evidence += f" [inputs: {inputs}]"
                groups.setdefault(group, set()).add(evidence)

    if missing:
        return "", f"Missing required table inputs: {', '.join(sorted(missing))}."
    if not groups:
        return "", "No grid group matches the job parameter values."
    if len(groups) > 1:
        return "", f"Multiple groups match: {', '.join(sorted(groups))}."

    group = next(iter(groups))

    return group, f"Group found in {', '.join(sorted(groups[group]))}."


def structure_dimensions_of(items):
    """Read dimensions from the structure assembly, not generic drawing dimension tables."""

    for _style, description, _w, _d in items:
        if "STRUCTURE" in description.upper():
            text = description.upper()
            pair = re.search(r'\b(\d{2,3})\s*X\s*(\d{2,3})\s+STRUCTURE\s+ASSY\b', text)
            if pair:
                return pair.group(1), pair.group(2)
            width = SHEET_WIDTH_RE.search(text)
            depth = re.search(_BOM_DEPTH, text)
            if width or depth:
                return width.group(1) if width else "", depth.group(1) if depth else ""

    return "", ""


def extract_job_dimensions(text, source):
    """Keep explicitly labeled job dimensions separate from section-specific values."""

    global_values = ParameterRepository()
    sections = {}
    current = global_values
    header = re.compile(r"^\s*(?:SECTION|STRUCTURE)\s+(?:NO\.?\s*)?(\d+)\b", re.I)
    dimension = re.compile(
        r'\b(?P<name>STRUCTURE\s+DEPTH|SECTION\s+DEPTH|DEPTH(?:\s*\(INNER\))?|'
        r'STRUCTURE\s+WIDTH|SECTION\s+WIDTH|WIDTH(?:\s*\(WIDE\))?)'
        r'\s*(?:[:=]|-(?!\d))\s*(?P<value>\d{1,3}(?:\.\d+)?)'
        r'(?:\s*(?:INCHES|INCH|IN\.?|["WD]))?(?=\s|[,;]|$)',
        re.I,
    )
    for line in (text or "").splitlines():
        match = header.match(line)
        if match:
            current = sections.setdefault(match.group(1), ParameterRepository())
        for match in dimension.finditer(line):
            name = "Depth" if "DEPTH" in match.group("name").upper() else "Width (Wide)"
            current.add_parameter(name, match.group("value"), source)
        for name, patterns in (
            ("Depth", _ECOE_DEPTH),
            ("Width (Wide)", _ECOE_WIDTH),
        ):
            for pattern in patterns:
                for match in re.finditer(pattern, line, re.I):
                    current.add_parameter(name, match.group(1), source)
    return global_values, sections


def dimension_entries(repository, name):
    aliases = DEPTH_LABELS if ConfigurationRepository._normalized(name) in DEPTH_LABELS else WIDTH_LABELS | {"section width"}
    return tuple(
        entry for alias in sorted(aliases)
        for entry in repository.get_entries(alias)
    )


def dimension_value(entries):
    values = set()
    for value, _source in entries:
        match = NUMERIC_VALUE_RE.fullmatch(_cell_text(value))
        if not match:
            return ""
        values.add(_format_dim(float(match.group(1))))
    return next(iter(values)) if len(values) == 1 else ""


def structure_width_of(items):
    return structure_dimensions_of(items)[0]


@dataclass(frozen=True)
class ValidationResult:
    expected_group: str = ""
    actual_group: str = ""
    status: str = "UNRESOLVED"
    message: str = ""


class ValidationEngine:

    @staticmethod
    def compare(expected_group, actual_group, message=""):

        if not normalize_code(expected_group):
            return ValidationResult("", actual_group, "UNRESOLVED", "Assembly number has no group suffix.")

        if not actual_group:
            return ValidationResult(expected_group, "", "UNRESOLVED", message)

        status = "PASS" if normalize_code(expected_group) == normalize_code(actual_group) else "FAIL"

        return ValidationResult(expected_group, actual_group, status, message)


class ExcelComSession:
    """Reads live family workbooks (e.g. BASICSTR.xlsx) automatically:
    raw COM Workbooks.Open() can't resolve SharePoint 'doc2.aspx' share links
    (it creates a blank placeholder instead of loading the real document), so
    Index2 is launched via the "ms-excel:ofe|u|" protocol handler instead (the
    same mechanism a browser's "Open in Desktop App" button uses). Once truly
    loaded, Index2's own `.Path` resolves to its real SharePoint folder, which
    is then combined with the target filename and opened directly - no
    manual steps and no relying on Hyperlink.Follow(). Already-open
    workbooks are reused when possible.

    Assembly-specific reads use Classifier.py with the adjacent Metadata.xlsx
    and semantic_aliases.example.json. Live values and merged ranges are copied
    in memory; no local workbook export or parser output files are required.
    Family workbooks opened by this session are closed after capture; Index2
    and workbooks already open in the user's Excel are never closed.
    Unrecognized layouts retain the existing readers. Parser failures are
    surfaced to the audit as unresolved rather than silently falling back.
    Family COM operations run in excel_capture.py with a hard deadline. A timed
    out Excel session is not reused for further uncached families in that run.
    """

    INDEX2_LOAD_TIMEOUT = 150
    FAMILY_CAPTURE_TIMEOUT = 45

    def __init__(self, status_callback=None):
        self._app = None
        self._index2_wb = None
        self._index2_map = None
        self._index2_catalog = None
        self._family_workbooks = {}
        self._group_tables = {}
        self._sheet_values = {}
        self._source_names = {}
        self._classifier_snapshots = {}
        self._classifier_tables = {}
        self._classifier_parsers = {}
        self._family_read_errors = {}
        self._index2_folder = None
        self._replacement_links = None
        self._excel_blocked_message = ""
        self._assembly_targets = {}
        self._captured_roots = {}
        self.status_callback = status_callback

    def _report(self, text):
        if self.status_callback:
            self.status_callback(text)

    def _get_app(self):

        if win32 is None:
            raise RuntimeError("pywin32 (win32com) is not installed.")

        if self._app is not None:
            try:
                self._app.Workbooks.Count  # noqa: B018 - liveness probe
                return self._app
            except Exception:
                self._app = None

        try:
            self._app = win32.GetActiveObject("Excel.Application")
        except Exception:
            self._app = None

        return self._app

    def _find_open_workbook(self, app, target_name):

        target_name = target_name.strip().upper()
        target_stem = os.path.splitext(target_name)[0]

        for wb in app.Workbooks:
            name = wb.Name.strip().upper()
            if name == target_name or os.path.splitext(name)[0] == target_stem:
                return wb

        return None

    def _wait_for_real_index2(self, app, before_names, timeout):
        """Polls until Index2.xlsx shows up and isn't the transient 'doc2'
        placeholder Excel briefly shows while resolving the share link."""

        deadline = time.time() + timeout
        waited = 0
        relaunched = False

        while time.time() < deadline:

            self._report(f"Waiting for Excel to load Index2.xlsx... ({waited}s)")

            try:
                app = self._get_app() or app
                for wb in app.Workbooks:
                    if os.path.splitext(wb.Name)[0].upper() == "INDEX2":
                        return wb
            except Exception:
                pass

            if not relaunched and waited == 20 and self._get_app() is None:
                # os.startfile's shell handoff can occasionally be dropped;
                # firing it again is harmless if the first one did land.
                self._report("Still nothing yet - relaunching Excel...")
                os.startfile(f"ms-excel:ofe|u|{INDEX2_LIVE_URL}")
                relaunched = True

            time.sleep(1)
            waited += 1

        raise RuntimeError(
            "Timed out waiting for Excel to load Index2.xlsx. If Excel is still opening it "
            "(e.g. a sign-in prompt), finish that there, then click Start Automation again."
        )

    def _get_index2_workbook(self):

        app = self._get_app()

        if app is not None:
            existing = self._find_open_workbook(app, "Index2.xlsx")
            if existing is not None:
                self._index2_wb = existing
                return existing

        if self._index2_wb is not None:
            try:
                self._index2_wb.Name  # still alive?
                return self._index2_wb
            except Exception:
                self._index2_wb = None

        before = [wb.Name for wb in app.Workbooks] if app is not None else []

        self._report("Launching Excel to open Index2 (live)...")
        os.startfile(f"ms-excel:ofe|u|{INDEX2_LIVE_URL}")

        wb = self._wait_for_real_index2(app, before, self.INDEX2_LOAD_TIMEOUT)
        self._index2_wb = wb

        return wb

    def get_index2_map(self):
        """Base assembly number -> family workbook filename, built once from
        the live Index2 worksheet (no local file needed)."""

        if self._index2_map is not None:
            return self._index2_map

        wb = self._get_index2_workbook()
        self._report("Reading Index2's assembly index (live)...")
        self._index2_map = build_index2_map_from_com(wb.Sheets(1))
        self._index2_folder = wb.Path
        self._cache_replacement_links(wb)

        return self._index2_map

    def get_index2_catalog(self, refresh=False):
        """Read only Index2, never family workbooks, for the coverage inventory."""
        if refresh or self._index2_catalog is None:
            workbook = self._get_index2_workbook()
            self._index2_catalog = read_index2_root_catalog(workbook.Sheets(1))
        return self._index2_catalog

    def _cache_replacement_links(self, workbook):
        links = {}
        for hyperlink in workbook.Worksheets(1).Hyperlinks:
            if hyperlink.Range.Column != 5:
                continue
            address = str(hyperlink.Address or "")
            if not address:
                continue
            for name in (str(hyperlink.Range.Value or ""), link_file_name(address)):
                links.setdefault(normalize_code(name), set()).add(address)
        self._replacement_links = links

    def _replacement_link(self, target, notice, previous_link=""):
        if self._replacement_links is None:
            self._cache_replacement_links(self._get_index2_workbook())
        candidates = self._replacement_links.get(normalize_code(link_file_name(target)), set())
        roots = self._assembly_targets.get(previous_link.upper(), ())
        if not candidates and roots:
            catalog = self.get_index2_catalog()
            alternatives = [{
                entry.link for entry in catalog if entry.root == root and entry.link
                and link_file_name(entry.link).casefold() != link_file_name(previous_link).casefold()
            } for root in roots]
            candidates = set.intersection(*alternatives)
            if len(candidates) == 1:
                self._report(f"Retired {target}: Index2 lists one common replacement for {', '.join(sorted(roots))}.")
        if len(candidates) != 1:
            raise RuntimeError(f"Replacement {target} has no unique Index2 link. {notice}")
        replacement = next(iter(candidates))
        if replacement.startswith("../"):
            from urllib.parse import urljoin
            if self._index2_folder is None:
                self._index2_folder = self._get_index2_workbook().Path
            replacement = urljoin(str(self._index2_folder).rstrip("/") + "/", replacement)
        return replacement

    def get_family_workbook(self, link_target):
        """Reuse or open the live workbook linked by Index2."""

        app = self._get_app()
        is_absolute = "://" in link_target or os.path.isabs(link_target)
        file_name = link_file_name(link_target)

        if app is not None:
            existing = self._find_open_workbook(app, file_name)
            if existing is not None:
                return existing

        cached = self._family_workbooks.get(link_target.upper())
        if cached is not None:
            try:
                cached.Name
                return cached
            except Exception:
                self._family_workbooks.pop(link_target.upper(), None)

        self._report(f"Opening {file_name}...")

        if is_absolute:
            url = excel_document_url("", link_target)
            if app is None:
                raise RuntimeError("Excel is not running. Open Index2 before starting the audit.")
        else:
            index2_wb = self._get_index2_workbook()
            app = self._get_app()
            folder = index2_wb.Path
            if not folder:
                raise RuntimeError("Index2's live folder could not be resolved.")
            url = excel_document_url(folder, link_target)
        target_wb = app.Workbooks.Open(
            url, ReadOnly=True, UpdateLinks=0, Notify=False,
            IgnoreReadOnlyRecommended=True, AddToMru=False,
        )
        self._family_workbooks[link_target.upper()] = target_wb
        return target_wb

    def _excel_call(self, operation, callback):
        for attempt in range(4):
            try:
                return callback()
            except Exception as exc:
                details = getattr(exc, "excepinfo", None)
                codes = (getattr(exc, "hresult", None), details[-1] if details else None)
                if not any(isinstance(code, int) and code & 0xFFFFFFFF == 0x800AC472 for code in codes) or attempt == 3:
                    raise
                self._report(f"Excel is busy during {operation}; retry {attempt + 1}/3.")
                time.sleep(0.5 * (attempt + 1))

    def _close_owned_family(self, link_target):
        key = link_target.upper()
        workbook = self._family_workbooks.get(key)
        if workbook is None:
            return
        if os.path.splitext(workbook.Name)[0].upper() == "INDEX2":
            return
        self._report(f"Closing {workbook.Name}; captured data remains cached.")
        try:
            self._excel_call(f"closing {workbook.Name}", lambda: workbook.Close(False))
        except Exception as exc:
            self._report(f"Could not close {link_file_name(link_target)}: {exc}")
            raise
        self._family_workbooks.pop(key, None)

    def read_sheets_values(self, url, sheet_names, open_if_missing=True):
        """Read requested sheets, optionally opening the workbook when it is not already open."""

        if open_if_missing:
            self._get_index2_workbook()
        app = self._get_app()
        if app is None:
            raise RuntimeError("Excel is not running.")

        file_name = link_file_name(url)
        existing = self._find_open_workbook(app, file_name)
        workbook = existing

        if workbook is None:
            if not open_if_missing:
                raise RuntimeError(f"Open {file_name} in Excel to use the live checklist.")
            # Direct Workbooks.Open() is rejected for this site and pops a modal error;
            # the ms-excel protocol uses Excel's own sign-in, like Index2 does.
            workbook = self._open_with_sign_in(url, file_name, CHECKLIST_OPEN_TIMEOUT)

        sheets = {}

        try:
            for sheet_name in sheet_names:
                try:
                    values = normalize_excel_rows(workbook.Worksheets(sheet_name).UsedRange.Value)
                except Exception:
                    continue
                if values:
                    sheets[sheet_name] = values
        finally:
            if existing is None:
                workbook.Close(False)

        return sheets

    def _open_with_sign_in(self, url, file_name, timeout=None):
        """Direct Open() fails until Office is signed in to the site; the
        ms-excel protocol makes Excel show its sign-in prompt, then this waits."""

        os.startfile(f"ms-excel:ofe|u|{url}")

        deadline = time.time() + (timeout or self.INDEX2_LOAD_TIMEOUT)
        waited = 0

        while time.time() < deadline:
            self._report(f"Sign in to SharePoint in the Excel window if prompted... ({waited}s)")
            try:
                workbook = self._find_open_workbook(self._get_app(), file_name)
                if workbook is not None:
                    return workbook
            except Exception:
                pass
            time.sleep(1)
            waited += 1

        raise RuntimeError(
            f"Timed out waiting for {file_name}. Sign in to SharePoint in Excel, then click Start Automation again."
        )

    def get_family_worksheet(self, link_target, sheet_name):
        return self.get_family_workbook(link_target).Worksheets(sheet_name)

    def lookup_dimensions(self, link_target, sheet_name, base, suffix):
        self.get_group_tables(link_target, assembly_base=base)
        key = link_target.upper()
        file_name = self._source_names[key]
        family = os.path.splitext(file_name)[0].upper()
        if family == "BASICSTR":
            if not sheet_name:
                return DimensionResult(message="BASICSTR requires an enclosure type from QuickPrint.")
            sheet = self._classifier_snapshots[key][sheet_name]
            width, depth = snapshot_grid_lookup(sheet, base, suffix)
            return DimensionResult(
                width, depth,
                "Width and depth found." if width and depth else "Base/group not found in enclosure grid.",
                f"{file_name} / {sheet.title}",
            )

        results = []
        for name, values in self.get_sheet_values(link_target):
            result = table_dimension_lookup(values, family, base, suffix)
            if result is not None:
                results.append(DimensionResult(
                    result.width, result.depth, result.message,
                    f"{file_name} / {name}",
                ))

        if not results:
            return DimensionResult(message="Base not found on active worksheets.", source=file_name)
        if len({(result.width, result.depth, result.message) for result in results}) != 1:
            return DimensionResult(
                message="Different worksheet results; selection unresolved.",
                source="; ".join(result.source for result in results),
            )
        result = results[0]
        return DimensionResult(
            result.width, result.depth, result.message,
            "; ".join(result.source for result in results),
        )

    def get_group_tables(self, link_target, _visited=(), assembly_base=None):
        key = link_target.upper()
        if assembly_base:
            base, _group = split_style_number(assembly_base.strip().upper())
            assembly_base = base
            targets = self._assembly_targets.setdefault(key, set())
            targets.add(base)
        targets = self._assembly_targets.get(key, set())
        captured = self._captured_roots.get(key)
        if captured and not targets.issubset(captured):
            self._group_tables.pop(key, None)
            self._sheet_values.pop(key, None)
            self._classifier_parsers.pop(key, None)
            snapshot = self._classifier_snapshots.pop(key, None)
            if snapshot is not None:
                snapshot.close()
            self._classifier_tables = {
                cache_key: tables for cache_key, tables in self._classifier_tables.items() if cache_key[0] != key
            }
        error_key = (key, assembly_base.upper() if assembly_base else None)
        error = self._family_read_errors.get((key, None)) or self._family_read_errors.get(error_key)
        if error:
            raise RuntimeError(error)
        try:
            return self._read_group_tables(link_target, _visited, assembly_base)
        except Exception as exc:
            if key not in self._group_tables:
                error_key = (key, None)
            message = f"{link_file_name(link_target)}: {type(exc).__name__}: {exc}"
            self._family_read_errors[error_key] = message
            self._report(f"Family read failed (cached for this run): {message}")
            raise

    def _capture_family(self, link_target):
        from excel_capture import capture_family

        if self._excel_blocked_message:
            raise RuntimeError(self._excel_blocked_message)
        if "://" in link_target or os.path.isabs(link_target):
            url = excel_document_url("", link_target)
        else:
            if self._index2_folder is None:
                self.get_index2_map()
            if not self._index2_folder:
                raise RuntimeError("Index2's live folder could not be resolved.")
            url = excel_document_url(self._index2_folder, link_target)
        self._report(f"Opening {link_file_name(link_target)} (limit {self.FAMILY_CAPTURE_TIMEOUT}s)...")
        try:
            name, sheets, snapshot_bytes = capture_family(
                url, self.FAMILY_CAPTURE_TIMEOUT, self._report,
                assemblies=tuple(sorted(self._assembly_targets.get(link_target.upper(), ()))),
            )
        except TimeoutError as exc:
            self._excel_blocked_message = (
                f"Excel did not finish reading {link_file_name(link_target)}. "
                "Further uncached family reads were skipped because Excel may still be blocked. "
                "Resolve any Excel dialogs and restart the audit. Index2 has been left open."
            )
            self._report(str(exc))
            raise
        snapshot = load_workbook(BytesIO(snapshot_bytes), data_only=True)
        return name, sheets, snapshot

    def _capture_live_family(self, link_target, assemblies=()):
        """Worker-only COM capture; the parent enforces its deadline."""
        workbook = self.get_family_workbook(link_target)
        try:
            name = workbook.Name
            roots = tuple(dict.fromkeys(split_style_number(str(value).strip().upper())[0] for value in assemblies))
            if roots:
                self._report(f"Finding root assemblies in {name}: {', '.join(roots)}")
                snapshot = self._capture_root_tables(workbook, roots)
                return name, tuple((sheet.title, tuple(sheet.values)) for sheet in snapshot.worksheets), snapshot
            self._report(f"Reading active worksheets in {name}")
            sheets = []
            for sheet in workbook.Worksheets:
                if sheet.Name.upper().startswith("OBSOLETE"):
                    continue
                self._report(f"Reading {name} / {sheet.Name}")
                values = normalize_excel_rows(sheet.UsedRange.Value)
                if values:
                    sheets.append((sheet.Name, values))
            self._report(f"Capturing merged cells in {name}")
            snapshot = self._snapshot_workbook(workbook, dict(sheets))
            return name, tuple(sheets), snapshot
        finally:
            self._close_owned_family(link_target)

    def _capture_root_tables(self, workbook, roots):
        """Find roots natively; preserve small sheets, restrict inflated sheets to root regions."""
        snapshot = Workbook()
        snapshot.remove(snapshot.active)
        for sheet in workbook.Worksheets:
            if sheet.Name.upper().startswith("OBSOLETE"):
                continue
            bounds = set()
            for root in roots:
                pattern = re.compile(rf"(?<![A-Z0-9]){re.escape(root)}(?:[GH]\d+)?(?![A-Z0-9])", re.IGNORECASE)
                previous = None
                seen = set()
                while True:
                    options = dict(
                        What=root, LookAt=2, LookIn=-4163, SearchOrder=1,
                        SearchDirection=1, MatchCase=False, SearchFormat=False,
                    )
                    if previous is not None:
                        options["After"] = previous
                    found = self._excel_call(f"finding {root} in {sheet.Name}", lambda: sheet.Cells.Find(**options))
                    if found is None or (found.Row, found.Column) in seen:
                        break
                    seen.add((found.Row, found.Column))
                    previous = found
                    if not pattern.search(str(found.Value or "")):
                        continue
                    self._report(f"Found {root} in {workbook.Name} / {sheet.Name}, row {found.Row}")
                    region = found.CurrentRegion
                    top, left = region.Row, region.Column
                    bottom = top + region.Rows.Count - 1
                    right = left + region.Columns.Count - 1
                    if region.Rows.Count <= found.MergeArea.Rows.Count and bottom < sheet.Rows.Count:
                        last_column = get_column_letter(sheet.Columns.Count)
                        search = sheet.Range(
                            f"A{bottom + 1}:{last_column}{min(sheet.Rows.Count, bottom + 8)}"
                        )
                        following = search.Find(
                            What="*", LookIn=-4163, LookAt=2, SearchOrder=1,
                            SearchDirection=1, MatchCase=False, SearchFormat=False,
                        )
                        if following is not None:
                            body = following.CurrentRegion
                            bottom = max(bottom, body.Row + body.Rows.Count - 1)
                            left = min(left, body.Column)
                            right = max(right, body.Column + body.Columns.Count - 1)
                    bounds.add((max(1, top - CONTEXT_LOOKBACK), left, bottom, right))
            if bounds and self._copy_small_sheet(sheet, snapshot):
                continue
            blocks = []
            for top, left, bottom, right in sorted(bounds):
                if blocks and top <= blocks[-1][2] and left <= blocks[-1][3] and right >= blocks[-1][1]:
                    a, b, c, d = blocks.pop()
                    top, left, bottom, right = min(top, a), min(left, b), max(bottom, c), max(right, d)
                blocks.append((top, left, bottom, right))
            if not blocks:
                continue
            target = snapshot.create_sheet(sheet.Name)
            cursor = 1
            for top, left, bottom, right in blocks:
                while True:
                    region = sheet.Range(f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}")
                    rows = normalize_excel_rows(region.Value)
                    merged = self._merged_ranges(sheet, region, rows)
                    expanded = (
                        min([top] + [a for a, _b, _c, _d in merged]),
                        min([left] + [b for _a, b, _c, _d in merged]),
                        max([bottom] + [c for _a, _b, c, _d in merged]),
                        max([right] + [d for _a, _b, _c, d in merged]),
                    )
                    if expanded == (top, left, bottom, right):
                        break
                    top, left, bottom, right = expanded
                for offset, row in enumerate(rows):
                    for column, value in enumerate(row, start=left):
                        if value is not None:
                            target.cell(cursor + offset, column, snapshot_cell_value(value))
                for a, b, c, d in merged:
                    target.merge_cells(
                        start_row=cursor + a - top, start_column=b,
                        end_row=cursor + c - top, end_column=d,
                    )
                cursor += len(rows) + 2
        if not snapshot.worksheets:
            notices = []
            for sheet in workbook.Worksheets:
                if not sheet.Name.upper().startswith("OBSOLETE"):
                    values = normalize_excel_rows(sheet.Range("A1:Z40").Value)
                    notices.append((sheet.Name, values))
            if workbook_migration(notices) is None:
                snapshot.close()
                raise LookupError(f"Root assemblies {', '.join(roots)} were not found in {workbook.Name}.")
            for name, rows in notices:
                target = snapshot.create_sheet(name)
                for row in rows:
                    target.append(row)
        return snapshot

    def prepare_assembly_targets(self, sections, index_map):
        for _num, _title, items in sections:
            for style, _description, _width, _depth in items:
                base, _group = split_style_number(style)
                link = index_map.get(base.upper())
                if link:
                    self._assembly_targets.setdefault(link.upper(), set()).add(base.upper())

    def _read_group_tables(self, link_target, _visited=(), assembly_base=None):
        """Group-bearing tables from every active sheet of the family workbook, cached."""

        key = link_target.upper()
        if key in _visited:
            raise RuntimeError("Replacement workbook cycle detected.")

        if key not in self._group_tables:
            name, sheets, snapshot = self._capture_family(link_target)
            self._source_names[key] = name
            self._classifier_snapshots[key] = snapshot
            self._captured_roots[key] = frozenset(self._assembly_targets.get(key, ()))
            tables = [
                table for sheet_name, values in sheets
                for table in WorkbookParser.extract_tables(values, sheet_name)
            ]
            migration = workbook_migration(sheets)
            if not any(any(_cell_text(value) for row in rows for value in row) for _name, rows in sheets):
                raise RuntimeError(f"{self._source_names[key]} has no readable active table data. Cannot determine a group.")
            if migration is not None:
                target, notice = migration
                if not target:
                    raise RuntimeError(f"Workbook retired: {notice}. Current smart workbook must be identified in Index2.")
                replacement = self._replacement_link(target, notice, link_target)
                self._assembly_targets.setdefault(replacement.upper(), set()).update(
                    self._assembly_targets.get(key, ())
                )
                replacement_tables = self.get_group_tables(replacement, (*_visited, key))
                self._group_tables[key] = replacement_tables
                self._sheet_values[key] = tuple(
                    (f"{link_file_name(replacement)} / {sheet_name}", rows)
                    for sheet_name, rows in self._sheet_values[replacement.upper()]
                )
                self._classifier_snapshots[key].close()
                self._classifier_snapshots[key] = self._classifier_snapshots[replacement.upper()]
                self._source_names[key] = self._source_names[replacement.upper()]
            else:
                self._group_tables[key] = tuple(tables)
                self._sheet_values[key] = tuple(sheets)

        if assembly_base:
            cache_key = (key, assembly_base.upper())
            if cache_key not in self._classifier_tables:
                if key not in self._classifier_parsers:
                    from Classifier import WorkbookAssemblyParser
                    self._classifier_parsers[key] = WorkbookAssemblyParser(
                        self._source_names[key], METADATA_PATH, SEMANTIC_PATH, self._classifier_snapshots[key],
                    )
                self._report(f"Classifying {assembly_base} in {self._source_names[key]}...")
                self._classifier_tables[cache_key] = ClassifierWorkbookParser.extract_tables(
                    self._classifier_snapshots[key], self._source_names[key], assembly_base,
                    parser=self._classifier_parsers[key],
                )
            classified = self._classifier_tables[cache_key]
            if classified:
                return classified
            self._report(f"Classifier found no configurations for {assembly_base}; using existing table readers.")
        return self._group_tables[key]

    def _copy_small_sheet(self, sheet, snapshot):
        """Keep a complete family sheet in place: neighbouring tables and titles decide ownership.

        Packing only the regions around a root cut table titles and merged headers away and let
        another root's table (same columns, adjacent rows) be mistaken for this root's table.
        """
        used = sheet.UsedRange
        if used.Rows.Count * used.Columns.Count > FULL_SHEET_CELL_LIMIT:
            return False
        rows = normalize_excel_rows(used.Value)
        target = snapshot.create_sheet(sheet.Name)
        for row_index, row in enumerate(rows, start=used.Row):
            for column_index, value in enumerate(row, start=used.Column):
                if value is not None:
                    target.cell(row_index, column_index, snapshot_cell_value(value))
        for top, left, bottom, right in self._merged_ranges(sheet, used, rows):
            target.merge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
        return True

    @staticmethod
    def _snapshot_workbook(workbook, sheet_rows=None):
        """Copy live values and merge geometry, including UsedRange's row/column offset."""
        if Workbook is None:
            raise RuntimeError("openpyxl is not installed.")
        snapshot = Workbook()
        snapshot.remove(snapshot.active)
        for sheet in workbook.Worksheets:
            if sheet.Name.upper().startswith("OBSOLETE"):
                continue
            target = snapshot.create_sheet(sheet.Name)
            used = sheet.UsedRange
            rows = sheet_rows.get(sheet.Name, ()) if sheet_rows is not None else normalize_excel_rows(used.Value)
            for row_index, row in enumerate(rows, start=used.Row):
                for column_index, value in enumerate(row, start=used.Column):
                    if value is None:
                        continue
                    target.cell(row_index, column_index, snapshot_cell_value(value))
            for top, left, bottom, right in ExcelComSession._merged_ranges(sheet, used, rows):
                target.merge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
        return snapshot

    @staticmethod
    def _merged_ranges(sheet, used, rows):
        """Skip unmerged blocks in one COM call instead of probing every cell."""
        if used.MergeCells is False:
            return ()
        populated = [(r, c) for r, row in enumerate(rows) for c, value in enumerate(row) if value is not None]
        if not populated:
            return ()
        pending = [(
            used.Row + min(r for r, _c in populated),
            used.Column + min(c for _r, c in populated),
            used.Row + max(r for r, _c in populated),
            used.Column + max(c for _r, c in populated),
        )]
        merged = set()
        while pending:
            top, left, bottom, right = pending.pop()
            if any(a <= top and b <= left and c >= bottom and d >= right for a, b, c, d in merged):
                continue
            address = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
            state = sheet.Range(address).MergeCells
            if state is False:
                continue
            if state is True:
                area = sheet.Cells(top, left).MergeArea
                bounds = (area.Row, area.Column, area.Row + area.Rows.Count - 1, area.Column + area.Columns.Count - 1)
                merged.add(bounds)
                if bounds[0] <= top and bounds[1] <= left and bounds[2] >= bottom and bounds[3] >= right:
                    continue
            if top == bottom and left == right:
                raise RuntimeError(f"Could not resolve merged cell {sheet.Name}!{address}.")
            if bottom - top >= right - left:
                middle = (top + bottom) // 2
                pending.extend(((top, left, middle, right), (middle + 1, left, bottom, right)))
            else:
                middle = (left + right) // 2
                pending.extend(((top, left, bottom, middle), (top, middle + 1, bottom, right)))
        return merged

    def get_sheet_values(self, link_target):
        """[(sheet name, values)] of the family workbook's active sheets."""

        self.get_group_tables(link_target)

        return self._sheet_values.get(link_target.upper(), ())

    def close(self):

        for link_target in tuple(self._family_workbooks):
            self._close_owned_family(link_target)
        for snapshot in self._classifier_snapshots.values():
            snapshot.close()
        self._classifier_snapshots.clear()
        self._classifier_tables.clear()
        self._classifier_parsers.clear()
        self._family_read_errors.clear()
        self._index2_folder = None
        self._replacement_links = None
        self._excel_blocked_message = ""
        self._assembly_targets.clear()
        self._captured_roots.clear()
        self._source_names.clear()
        self._group_tables.clear()
        self._sheet_values.clear()
        self._app = None
        self._index2_wb = None
        self._index2_map = None
        self._index2_catalog = None
        self._family_workbooks.clear()


def extract_enclosure_type(pdf_text):
    """Returns (enclosure_value, sheet_name) or (None, None) if not found."""

    match = ENCLOSURE_TYPE_RE.search(pdf_text)

    if not match:
        return None, None

    value = match.group(1).strip().upper()
    sheet_name = ENCLOSURE_SHEET_MAP.get(value, f"BASICSTR-NEMA {value}")

    return value, sheet_name


def parse_structured_bom(pdf_text):
    """Group top-level (bold) assembly style numbers section-wise by Structure."""

    sections = {}
    order = []
    current_num = None

    for raw_line in pdf_text.splitlines():

        line = raw_line.strip()

        if not line:
            continue

        header_match = STRUCTURE_HEADER_RE.search(line)

        if header_match:

            num = header_match.group(1)
            title = re.split(r"\(cont\.\)|SXJ\d|\s{2,}", header_match.group(2))[0].strip()

            if num not in sections:
                sections[num] = {"title": title or f"Structure {num}", "items": []}
                order.append(num)

            current_num = num
            continue

        if current_num is None:
            continue

        row_match = ASSEMBLY_ROW_RE.match(line)

        if not row_match:
            continue

        style_number = row_match.group(1).strip()
        description = re.split(r"\s{2,}", row_match.group(2))[0].strip()

        sections[current_num]["items"].append((style_number, description, "", ""))

    return [(num, sections[num]["title"], sections[num]["items"]) for num in order]


class AuditReviewDemo(QMainWindow):

    def __init__(self):
        super().__init__()

        ui_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "QA2.ui")
        loadUi(ui_path, self)

        self.setStyleSheet(self._build_stylesheet())

        self.checkbox_items = {}
        self.bom_pdf_path = None
        self.enclosure_type = None
        self.enclosure_sheet = None

        self._index2_map = {}
        self._excel_com = ExcelComSession(status_callback=self._set_status)
        self._configuration_repository = ConfigurationRepository()
        self._parameter_repository = ParameterRepository()
        self._section_bus_requirements = {}
        self.parameter_variables = ExcelParameterVariables({})
        self._common_parameters = EXCEL_PARAMETERS
        self._dimension_cache = {}
        self._dimension_results = {}
        self._assembly_families = {}
        self._assembly_links = {}
        self._validation_results = {}
        self._validation_checks = {}
        self._coverage_expected_checks = {}
        self._dimension_lookup_error = None
        self._code_mappings = {}
        self._pending_folder = None
        self._pdf_read_issues = []
        self._pdf_ocr_page_count = 0
        self._coverage_catalog = None
        self._coverage_indexed_roots = set()
        self._coverage_linked_roots = set()
        self._coverage_job_roots = {}
        self._coverage_excluded_styles = set()

        self.browse_pdf_btn.clicked.connect(self.browse_job_folder)
        self.start_btn.clicked.connect(self.start_automation)

        self.bom_list.setAlternatingRowColors(True)
        self._parameter_columns = PARAMETER_COLUMN_COUNT
        self._configure_columns(PARAMETER_COLUMN_COUNT)
        self.bom_list.itemChanged.connect(self.update_progress)
        self._build_testing_tab()
        self._build_audit_workspace()

    def _build_audit_workspace(self):
        self.setMinimumSize(960, 640)
        self.resize(1440, 900)
        self.setFont(QFont("Aptos", 10))
        old_central = self.takeCentralWidget()
        central = QWidget()
        root = QHBoxLayout(central)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)
        sidebar = QFrame()
        sidebar.setObjectName("sidebar")
        sidebar.setFixedWidth(208)
        navigation = QVBoxLayout(sidebar)
        navigation.setContentsMargins(18, 22, 18, 18)
        navigation.setSpacing(10)
        brand = QLabel("MCOE\nAUDIT")
        brand.setObjectName("brand")
        navigation.addWidget(brand)
        navigation.addSpacing(12)
        self.navigation_group = QButtonGroup(self)
        for index, (title, icon) in enumerate((
            ("Checklist", QStyle.StandardPixmap.SP_FileDialogDetailedView),
            ("Testing", QStyle.StandardPixmap.SP_FileDialogContentsView),
        )):
            button = QPushButton(title)
            button.setCheckable(True)
            button.setChecked(index == 0)
            button.setIcon(self.style().standardIcon(icon))
            self.navigation_group.addButton(button, index)
            navigation.addWidget(button)
        self.navigation_group.idClicked.connect(self.main_tabs.setCurrentIndex)
        self.main_tabs.currentChanged.connect(lambda index: self.navigation_group.button(index).setChecked(True))
        navigation.addSpacing(12)
        log_header = QHBoxLayout()
        log_title = QLabel("ACTIVITY LOG")
        log_title.setObjectName("sidebarCaption")
        log_header.addWidget(log_title, 1)
        self.clear_log_button = QPushButton()
        self.clear_log_button.setObjectName("clearLog")
        self.clear_log_button.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_TrashIcon))
        self.clear_log_button.setToolTip("Clear activity log")
        self.clear_log_button.setAccessibleName("Clear activity log")
        self.clear_log_button.setFixedSize(26, 26)
        log_header.addWidget(self.clear_log_button)
        navigation.addLayout(log_header)
        self.automation_activity = QLabel("Idle")
        self.automation_activity.setObjectName("activityStatus")
        self.automation_activity.setTextFormat(Qt.TextFormat.PlainText)
        self.automation_activity.setWordWrap(True)
        self.automation_activity.setMinimumWidth(0)
        self.automation_activity.setMaximumHeight(90)
        navigation.addWidget(self.automation_activity)
        self.automation_log = QPlainTextEdit()
        self.automation_log.setObjectName("automationLog")
        self.automation_log.setAccessibleName("Automation activity log")
        self.automation_log.setReadOnly(True)
        self.automation_log.setPlaceholderText("Live updates will appear here\nwhen automation starts.")
        self.automation_log.setMaximumBlockCount(1000)
        self.automation_log.setMinimumSize(0, 100)
        self.automation_log.setMaximumHeight(220)
        self.automation_log.setWordWrapMode(QTextOption.WrapMode.WrapAnywhere)
        self.automation_log.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.clear_log_button.clicked.connect(self.automation_log.clear)
        navigation.addWidget(self.automation_log, 1)
        scope_title = QLabel("AUDIT SCOPE")
        scope_title.setObjectName("sidebarCaption")
        navigation.addWidget(scope_title)
        self.test_scope_checkbox.setText("Structure 2 only")
        navigation.addWidget(self.test_scope_checkbox)
        self.enclosure_label.setWordWrap(True)
        navigation.addWidget(self.enclosure_label)
        navigation.addStretch()
        root.addWidget(sidebar)

        workspace = QWidget()
        content = QVBoxLayout(workspace)
        content.setContentsMargins(24, 24, 24, 16)
        content.setSpacing(14)
        heading_row = QHBoxLayout()
        heading_text = QVBoxLayout()
        heading_text.setSpacing(2)
        eyebrow = QLabel("ENGINEERING / QUALITY ASSURANCE")
        eyebrow.setObjectName("eyebrow")
        heading_text.addWidget(eyebrow)
        heading = QLabel("Job Review")
        heading.setObjectName("pageTitle")
        heading_text.addWidget(heading)
        heading_row.addLayout(heading_text)
        heading_row.addStretch()
        self.browse_pdf_btn.setText("Job Folder")
        self.browse_pdf_btn.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_DirOpenIcon))
        self.start_btn.setText("Start Automation")
        self.start_btn.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_MediaPlay))
        heading_row.addWidget(self.browse_pdf_btn)
        heading_row.addWidget(self.start_btn)
        content.addLayout(heading_row)
        self.bom_pdf_label.setWordWrap(True)
        self.bom_pdf_label.setObjectName("jobLabel")
        self.bom_pdf_label.setMinimumWidth(0)
        content.addWidget(self.bom_pdf_label)

        summary = QFrame()
        summary.setObjectName("summary")
        metric_rows = QVBoxLayout(summary)
        metric_rows.setContentsMargins(0, 0, 0, 0)
        self.summary_labels = {}
        for definitions in (
            (("all_assemblies", "ALL BOLD ASSEMBLIES"), ("gh_assemblies", "G/H ASSEMBLIES"),
             ("present", "G/H IN INDEX2"), ("missing", "G/H NOT IN INDEX2"),
             ("coverage_rate", "INDEX2 PRESENCE RATE")),
            (("validation_total", "G/H VALIDATION TOTAL"), ("PASS", "PASS"), ("FAIL", "FAIL"),
             ("UNRESOLVED", "UNRESOLVED"), ("validation_rate", "VALIDATION PASS RATE")),
        ):
            metrics = QHBoxLayout()
            metrics.setSpacing(10)
            for key, title in definitions:
                card = QFrame()
                card.setObjectName("card" + key)
                card.setProperty("metricCard", True)
                metric = QVBoxLayout(card)
                metric.setContentsMargins(14, 10, 14, 10)
                metric.setSpacing(3)
                number = QLabel("0")
                number.setObjectName("metric" + key)
                number.setProperty("metric", True)
                caption = QLabel(title)
                caption.setObjectName("metricCaption")
                caption.setWordWrap(True)
                metric.addWidget(number)
                metric.addWidget(caption)
                metrics.addWidget(card, 1)
                self.summary_labels[key] = number
            metric_rows.addLayout(metrics)
        content.addWidget(summary)

        self.centerGroup.setTitle("")
        self.centerLayout.setContentsMargins(0, 0, 0, 0)
        toolbar = QHBoxLayout()
        self.audit_search = QLineEdit()
        self.audit_search.setPlaceholderText("Search assemblies, descriptions, codes...")
        self.audit_search.setClearButtonEnabled(True)
        self.audit_search.setAccessibleName("Search audit assemblies")
        self.audit_search.setMinimumWidth(180)
        self.status_filter = QComboBox()
        self.status_filter.addItems(["All statuses", "PASS", "FAIL", "UNRESOLVED", "EXCLUDED"])
        self.status_filter.setAccessibleName("Filter overall assembly validation result")
        self.status_filter.setToolTip(
            "Filters the overall result for each unique G/H assembly present in Index2. "
            "All its structure rows are shown as evidence: any FAIL wins; all checks must PASS; otherwise UNRESOLVED."
        )
        self.status_filter.setMinimumWidth(138)
        self.audit_coverage_filter = QComboBox()
        self.audit_coverage_filter.addItems([
            "All Index2 coverage", "Present in Index2", "Present with link", "Not linked",
            "Listed without link", "Not present in Index2", "Not checked", "Excluded (no G/H suffix)",
        ])
        self.audit_coverage_filter.setAccessibleName("Filter audit Index2 root links")
        self.audit_suffix_filter = QComboBox()
        self.audit_suffix_filter.addItems(["All bold assemblies", "G/H assemblies", "Without G/H suffix"])
        self.audit_suffix_filter.setAccessibleName("Filter bold assemblies by G/H suffix")
        self.parameter_toggle = QCheckBox("Parameter columns")
        toolbar.addWidget(self.audit_search, 1)
        toolbar.addWidget(self.status_filter)
        for title, icon, action in (
            ("Expand structures", QStyle.StandardPixmap.SP_ArrowDown, self.bom_list.expandAll),
            ("Collapse structures", QStyle.StandardPixmap.SP_ArrowUp, self.bom_list.collapseAll),
        ):
            button = QPushButton()
            button.setIcon(self.style().standardIcon(icon))
            button.setToolTip(title)
            button.setAccessibleName(title)
            button.setFixedSize(34, 34)
            button.clicked.connect(action)
            toolbar.addWidget(button)
        self.centerLayout.insertLayout(0, toolbar)
        coverage_filters = QHBoxLayout()
        coverage_filters.addWidget(self.audit_suffix_filter)
        coverage_filters.addWidget(self.audit_coverage_filter)
        coverage_filters.addWidget(self.parameter_toggle)
        coverage_filters.addStretch()
        self.centerLayout.insertLayout(1, coverage_filters)
        self.visible_rows = QLabel("0 assemblies")
        self.visible_rows.setObjectName("metricCaption")
        self.centerLayout.insertWidget(2, self.visible_rows)
        self.audit_coverage_summary = QLabel()
        self.audit_coverage_summary.setWordWrap(True)
        self.audit_coverage_summary.setObjectName("metricCaption")
        self.centerLayout.insertWidget(3, self.audit_coverage_summary)
        self.bom_list.setUniformRowHeights(True)
        self.bom_list.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.bom_list.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.bom_list.setVerticalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.bom_list.setIndentation(18)
        self.bom_list.header().setStretchLastSection(False)
        self.bom_list.header().setSectionsMovable(True)
        self.row_details = QPlainTextEdit()
        self.row_details.setReadOnly(True)
        self.row_details.setPlaceholderText("Assembly details")
        self.row_details.setMinimumHeight(60)
        self.row_details.setMaximumHeight(118)
        self.centerLayout.addWidget(self.row_details)
        self.bom_list.currentItemChanged.connect(self._show_row_details)
        self.audit_search.textChanged.connect(self._apply_audit_filters)
        self.status_filter.currentIndexChanged.connect(self._apply_audit_filters)
        self.audit_coverage_filter.currentIndexChanged.connect(self._apply_audit_filters)
        self.audit_suffix_filter.currentIndexChanged.connect(self._apply_audit_filters)
        self.parameter_toggle.toggled.connect(self._toggle_parameter_columns)
        self.main_tabs.tabBar().hide()
        self._tab_opacity = QGraphicsOpacityEffect(self.main_tabs)
        self.main_tabs.setGraphicsEffect(self._tab_opacity)
        self._tab_animation = QPropertyAnimation(self._tab_opacity, b"opacity", self)
        self._tab_animation.setDuration(220)
        self._tab_animation.setEasingCurve(QEasingCurve.Type.OutCubic)
        self.main_tabs.currentChanged.connect(self._animate_workspace)
        self._workspace_revealed = False
        content.addWidget(self.main_tabs, 1)
        run_status = QHBoxLayout()
        self.run_indicator = QLabel("Ready - select a job folder to begin")
        self.run_indicator.setObjectName("runIndicator")
        self.run_indicator.setTextFormat(Qt.TextFormat.PlainText)
        self.run_indicator.setWordWrap(True)
        self.run_indicator.setAccessibleName("Automation run state")
        run_status.addWidget(self.run_indicator, 1)
        self.run_busy_bar = QProgressBar()
        self.run_busy_bar.setAccessibleName("Automation in progress")
        self.run_busy_bar.setRange(0, 0)
        self.run_busy_bar.setTextVisible(False)
        self.run_busy_bar.setFixedSize(100, 8)
        self.run_busy_bar.hide()
        run_status.addWidget(self.run_busy_bar)
        content.addLayout(run_status)
        self._run_started_at = None
        self._run_tick = 0
        self._run_timer = QTimer(self)
        self._run_timer.setInterval(250)
        self._run_timer.timeout.connect(self._refresh_running_indicator)
        self.scan_status.setWordWrap(True)
        self.scan_status.setMaximumHeight(54)
        self.scan_status.setMinimumWidth(0)
        content.addWidget(self.scan_status)
        self.progress.setFixedHeight(22)
        content.addWidget(self.progress)
        self.workspace_scroll = QScrollArea()
        self.workspace_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.workspace_scroll.setWidgetResizable(True)
        self.workspace_scroll.setAccessibleName("Scrollable audit workspace")
        self.workspace_scroll.setWidget(workspace)
        root.addWidget(self.workspace_scroll, 1)
        self.setCentralWidget(central)
        self.menubar.hide()
        self.statusbar.hide()
        old_central.deleteLater()
        self._configure_columns(self._parameter_columns)
        self.update_progress()

    def showEvent(self, event):
        super().showEvent(event)
        if not self._workspace_revealed:
            self._workspace_revealed = True
            self._animate_workspace()

    def _animate_workspace(self, _index=None):
        self._tab_animation.stop()
        self._tab_animation.setStartValue(0.65)
        self._tab_animation.setEndValue(1.0)
        self._tab_animation.start()

    def _toggle_parameter_columns(self):
        for column in range(4, 4 + self._parameter_columns):
            self.bom_list.setColumnHidden(column, not self.parameter_toggle.isChecked())

    def _apply_audit_filters(self):
        query = self.audit_search.text().strip().casefold()
        status = self.status_filter.currentText()
        visible = 0
        row_count = 0
        filtered_styles = set()
        coverage = self.audit_coverage_filter.currentText()
        suffix_scope = self.audit_suffix_filter.currentIndex()
        outcomes = self._eligible_assembly_results()
        for section_index in range(self.bom_list.topLevelItemCount()):
            section = self.bom_list.topLevelItem(section_index)
            section_visible = False
            for index in range(section.childCount()):
                item = section.child(index)
                row_count += 1
                matches = (not query or query in " ".join(item.text(column) for column in range(item.columnCount())).casefold())
                style = item.text(0).strip().upper()
                overall = self._overall_assembly_status(style, outcomes)
                matches = matches and (status == "All statuses" or overall == status)
                root = self._coverage_root(item.text(0))
                matches = matches and (
                    suffix_scope == 0 or suffix_scope == 1 and root is not None
                    or suffix_scope == 2 and root is None
                )
                root_status = self._index2_root_status(root) if root else "Excluded (no G/H suffix)"
                matches = matches and self._matches_index2_filter(root_status, coverage)
                item.setHidden(not matches)
                section_visible = section_visible or matches
                visible += int(matches)
                if matches:
                    filtered_styles.add(item.text(0))
            section.setHidden(not section_visible)
            if section_visible and (query or status != "All statuses" or coverage != "All Index2 coverage" or suffix_scope):
                section.setExpanded(True)
        if row_count == len(self.checkbox_items):
            self.visible_rows.setText(f"{visible} of {row_count} assembly rows")
        else:
            self.visible_rows.setText(
                f"{visible} of {row_count} PDF rows ({len(self.checkbox_items)} unique assembly/structure pairs)"
            )
        self.audit_coverage_summary.setText(self._assembly_coverage_summary(filtered_styles))

    def _show_row_details(self, item=None, previous=None):
        if item is None or item.parent() is None:
            self.row_details.clear()
            return
        status_column = item.columnCount() - 1
        lines = [f"{item.text(0)}  |  {item.text(1)}", item.toolTip(status_column)]
        overall = self._overall_assembly_status(item.text(0).strip().upper(), self._eligible_assembly_results())
        lines.append(
            f"Overall assembly result: {overall}. Any structure FAIL gives FAIL; "
            "all checks must PASS; otherwise UNRESOLVED. Table results are per structure."
        )
        root = self._coverage_root(item.text(0))
        lines.append(
            f"Index2 root {root}: {self._index2_root_status(root)}" if root
            else "Excluded from root coverage: no G/H suffix."
        )
        links = sorted({entry.link for entry in self._coverage_catalog or () if entry.root == root and entry.link})
        lines.extend(links)
        lines.extend(item.text(column) for column in range(4, 4 + self._parameter_columns) if item.text(column))
        self.row_details.setPlainText("\n".join(lines))

    # ------------------------------------------------------------------
    # Testing tab: ask for an extracted parameter and print its value
    # ------------------------------------------------------------------

    def _build_testing_tab(self):

        page = QWidget()
        layout = QVBoxLayout(page)

        row = QHBoxLayout()
        self.test_query = QLineEdit()
        self.test_query.setPlaceholderText("Parameter name")
        completer = QCompleter(self.test_query)
        completer.setModel(QStringListModel(list(EXCEL_PARAMETERS), completer))
        completer.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        completer.setFilterMode(Qt.MatchFlag.MatchContains)
        self.test_query.setCompleter(completer)
        self.test_get_btn = QPushButton("Get Value")
        self.test_get_btn.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_ArrowForward))
        row.addWidget(self.test_query, 1)
        row.addWidget(self.test_get_btn)

        self.test_output = QPlainTextEdit()
        self.test_output.setReadOnly(True)
        self.test_output.setFont(QFont("Consolas", 11))

        layout.addLayout(row)
        layout.addWidget(self.test_output, 1)

        self.test_query.returnPressed.connect(self.print_parameter_value)
        self.test_get_btn.clicked.connect(self.print_parameter_value)

        # The checklist group moves into the first tab; Testing is the second.
        parent_layout = self.centerGroup.parentWidget().layout()
        index = parent_layout.indexOf(self.centerGroup)
        parent_layout.removeWidget(self.centerGroup)
        self.main_tabs = QTabWidget()
        self.main_tabs.addTab(self.centerGroup, "Checklist")
        self.main_tabs.addTab(page, "Testing")
        parent_layout.insertWidget(index, self.main_tabs)

    def _set_index2_job_sections(self, sections):
        self._coverage_job_roots = {}
        self._coverage_excluded_styles = set()
        self._coverage_expected_checks.clear()
        self._validation_checks.clear()
        for number, _title, items in sections:
            for style, _description, _width, _depth in items:
                style = style.strip().upper()
                base = self._coverage_root(style)
                if base is None:
                    self._coverage_excluded_styles.add(style)
                    continue
                entry = self._coverage_job_roots.setdefault(base, {"styles": set(), "sections": set()})
                entry["styles"].add(style)
                entry["sections"].add(str(number))
                self._coverage_expected_checks[style] = self._coverage_expected_checks.get(style, 0) + 1
        self._update_index2_coverage()

    @staticmethod
    def _coverage_root(style):
        base, suffix = split_style_number(style.strip().upper())
        return base if re.fullmatch(r"[GH]\d+", suffix) else None

    def _index2_root_status(self, root):
        if self._coverage_catalog is None:
            return "Not checked"
        if root in self._coverage_linked_roots:
            return "Present with link"
        if root in self._coverage_indexed_roots:
            return "Listed without link"
        return "Not present in Index2"

    @staticmethod
    def _matches_index2_filter(actual, selected):
        return (
            selected in ("All coverage", "All Index2 coverage") or actual == selected
            or selected == "Present in Index2" and actual in ("Present with link", "Listed without link")
            or selected == "Not linked" and actual in ("Listed without link", "Not present in Index2")
        )

    def _assembly_coverage_counts(self, styles):
        styles = {style.strip().upper() for style in styles}
        grouped = {style: root for style in styles if (root := self._coverage_root(style)) is not None}
        present = sum(root in self._coverage_indexed_roots for root in grouped.values())
        return len(styles), len(grouped), present

    def _assembly_coverage_summary(self, styles):
        total, grouped, present = self._assembly_coverage_counts(styles)
        loaded = self._coverage_catalog is not None
        rate = f"{present / grouped:.1%}" if loaded and grouped else "--"
        return (
            f"Filtered unique assemblies: {total} | G/H: {grouped} | Without G/H: {total - grouped} | "
            f"In Index2: {present if loaded else '--'} | "
            f"Not present: {grouped - present if loaded else '--'} | Presence rate: {rate}"
        )

    def _eligible_assembly_results(self):
        if self._coverage_catalog is None:
            return None
        outcomes = {}
        for root, entry in self._coverage_job_roots.items():
            if root not in self._coverage_indexed_roots:
                continue
            for style in entry["styles"]:
                checks = self._validation_checks.get(style, ())
                if "FAIL" in checks:
                    outcomes[style] = "FAIL"
                elif len(checks) == self._coverage_expected_checks[style] and all(check == "PASS" for check in checks):
                    outcomes[style] = "PASS"
                else:
                    outcomes[style] = "UNRESOLVED"
        return outcomes

    def _overall_assembly_status(self, style, outcomes):
        if self._coverage_root(style) is None:
            return "EXCLUDED"
        if outcomes is None:
            return "UNRESOLVED"
        return outcomes.get(style, "EXCLUDED")

    def _validation_scope_result(self, style):
        root = self._coverage_root(style)
        expected = split_style_number(style)[1]
        if root is None:
            return ValidationResult(expected, "", "EXCLUDED", "Outside validation scope: no G/H group suffix.")
        if self._coverage_catalog is None:
            return ValidationResult(expected, "", "UNRESOLVED", "Index2 presence has not been checked; validation not run.")
        if root not in self._coverage_indexed_roots:
            return ValidationResult(expected, "", "EXCLUDED", f"Outside validation scope: root {root} is not present in Index2.")
        if root not in self._coverage_linked_roots:
            return ValidationResult(expected, "", "UNRESOLVED", f"Root {root} is present in Index2 but has no workbook link.")
        return None

    def _update_validation_metrics(self):
        outcomes = self._eligible_assembly_results()
        loaded = outcomes is not None
        statuses = list(outcomes.values()) if loaded else []
        total = len(statuses)
        values = {
            "validation_total": str(total) if loaded else "--",
            "validation_rate": f"{statuses.count('PASS') / total:.1%}" if loaded and total else "--",
            **{status: str(statuses.count(status)) if loaded else "--" for status in ("PASS", "FAIL", "UNRESOLVED")},
        }
        explanation = (
            "Unique full G/H assembly numbers whose roots are present in Index2. "
            "Any failed structure check gives FAIL. All expected checks must PASS to give PASS. "
            "Missing, pending, or ambiguous checks give UNRESOLVED. "
            "PASS + FAIL + UNRESOLVED equals G/H in Index2. Pass rate = PASS / this total. "
            "Assemblies absent from Index2 or without G/H are excluded."
        )
        for key, value in values.items():
            self.summary_labels[key].setText(value)
            self.summary_labels[key].setToolTip(explanation)

    def _update_coverage_metrics(self):
        if not hasattr(self, "summary_labels"):
            return
        styles = self._coverage_excluded_styles | {
            style for entry in self._coverage_job_roots.values() for style in entry["styles"]
        }
        total, grouped, present = self._assembly_coverage_counts(styles)
        loaded = self._coverage_catalog is not None
        values = {
            "all_assemblies": str(total),
            "gh_assemblies": str(grouped),
            "present": str(present) if loaded else "--",
            "missing": str(grouped - present) if loaded else "--",
            "coverage_rate": f"{present / grouped:.1%}" if loaded and grouped else "--",
        }
        explanation = (
            "Unique full assembly numbers from bold top-level PDF rows in the selected job scope; "
            "repeated occurrences across structures count once. G04 and G05 count separately, "
            "but their common root is searched in Index2 column K. "
            "Presence rate = G/H assemblies whose root is listed / all G/H assemblies. "
            "A listed root counts as present even without a workbook link; this is not PASS/FAIL. "
            f"{total - grouped} unique assemblies without a G/H suffix are included in All, not in G/H. "
            "Top cards are independent of table filters."
        )
        for key, value in values.items():
            self.summary_labels[key].setText(value)
            self.summary_labels[key].setToolTip(explanation)
        self.summary_labels["missing"].setToolTip(
            "Unique full G/H assembly numbers whose roots are absent from Index2 column K. "
            "A listed root without a workbook link is not counted here. "
            "Choose 'Not present in Index2' to show the corresponding full assemblies."
        )
        self._update_validation_metrics()

    def refresh_index2_coverage(self):
        """Refresh the Checklist's Index2 inventory without changing validation caches."""
        controls = (self.start_btn, self.browse_pdf_btn)
        enabled = [control.isEnabled() for control in controls]
        for control in controls:
            control.setEnabled(False)
        try:
            self._set_status("Reading Index2 root coverage...")
            self._coverage_catalog = self._excel_com.get_index2_catalog(refresh=True)
            root_count = len({entry.root for entry in self._coverage_catalog})
            self._set_status(f"Index2 coverage updated: {root_count} unique roots.")
        except Exception as exc:
            self._coverage_catalog = None
            message = f"Index2 coverage unavailable: {type(exc).__name__}: {exc}"
            self._set_status(message)
        finally:
            self._update_index2_coverage()
            for control, was_enabled in zip(controls, enabled):
                control.setEnabled(was_enabled)

    def _update_index2_coverage(self):
        catalog = {}
        for entry in self._coverage_catalog or ():
            catalog.setdefault(entry.root, []).append(entry)
        indexed = set(catalog)
        no_links = {root for root, entries in catalog.items() if not any(entry.link for entry in entries)}
        self._coverage_indexed_roots = indexed
        self._coverage_linked_roots = indexed - no_links
        self._update_coverage_metrics()
        if hasattr(self, "audit_coverage_filter"):
            self._apply_audit_filters()
            self._show_row_details(self.bom_list.currentItem())

    def print_parameter_value(self):

        query = self.test_query.text().strip()
        variables = self.parameter_variables

        if not any(variables.values.values()) and not self.checkbox_items:
            self.test_output.appendPlainText("No data yet - click Start Automation first.")
            return

        if not query:
            names = tuple(variables.values)
        else:
            name = variables.resolve(query)
            if name is None:
                self.test_output.appendPlainText(f"'{query}' is not one of the Excel parameters.")
                return
            names = (name,)

        self.test_output.appendPlainText("\n".join(variables.describe(name) for name in names))

    def _configure_columns(self, parameter_count):
        """Assembly, Description, Name, Code, Parameter 1..N, Expected/Actual Group, Status."""

        self._parameter_columns = parameter_count
        self.bom_list.setHeaderLabels([
            "Assembly", "Description", "Name", "Code",
            *(f"Parameter {number}" for number in range(1, parameter_count + 1)),
            "Expected Group", "Actual Group", "Structure result",
        ])

        for column, width in enumerate((155, 210, 220, 130)):
            self.bom_list.setColumnWidth(column, width)
        for column in range(4, 4 + parameter_count):
            self.bom_list.setColumnWidth(column, 210)
        for column in range(4 + parameter_count, 7 + parameter_count):
            self.bom_list.setColumnWidth(column, 110)
        header = self.bom_list.header()
        for position, column in enumerate(range(4 + parameter_count, 7 + parameter_count), start=2):
            header.moveSection(header.visualIndex(column), position)
        if hasattr(self, "parameter_toggle"):
            self._toggle_parameter_columns()

    def _ensure_parameter_columns(self, parameter_count):
        if parameter_count <= self._parameter_columns:
            return
        old_offset = 4 + self._parameter_columns
        previous = [
            (item, tuple((item.text(i), item.toolTip(i), item.foreground(i)) for i in range(old_offset, old_offset + 3)))
            for item in self.checkbox_items.values()
        ]
        self._configure_columns(parameter_count)
        for item, status_cells in previous:
            for column in range(old_offset, 4 + parameter_count):
                item.setText(column, "")
                item.setToolTip(column, "")
                item.setForeground(column, QBrush())
            for column, (text, tip, brush) in enumerate(status_cells, start=4 + parameter_count):
                item.setText(column, text)
                item.setToolTip(column, tip)
                item.setForeground(column, brush)

    # ------------------------------------------------------------------
    # PDF reading
    # ------------------------------------------------------------------

    def _read_pdf(self, path):
        extraction = extract_pdf_pages(path)
        self._pdf_ocr_page_count += len(extraction.ocr_pages)
        self._pdf_read_issues.extend(
            f"{os.path.basename(path)}: {issue}" for issue in extraction.issues
        )
        return extraction

    def _find_named_pdf(self, folder, keyword):

        keyword = keyword.upper()

        for name in sorted(os.listdir(folder)):

            if name.upper().endswith(".PDF") and keyword in name.upper():
                return os.path.join(folder, name)

        return None

    def _set_status(self, text):
        """Shows what the app is currently doing so the user knows to wait."""

        self.scan_status.setText(text)
        self.scan_status.setToolTip(text)
        if hasattr(self, "automation_log"):
            self.automation_activity.setText(text)
            self.automation_activity.setToolTip(text)
            scrollbar = self.automation_log.verticalScrollBar()
            self.automation_log.appendPlainText(f"[{time.strftime('%H:%M:%S')}] {text}")
            scrollbar.setValue(scrollbar.maximum())
        QApplication.processEvents()

    # ------------------------------------------------------------------
    # Structured BOM checklist (section-wise, from a browsed job folder)
    # ------------------------------------------------------------------

    def browse_job_folder(self):

        folder = QFileDialog.getExistingDirectory(self, "Select Job Folder", "")

        if not folder:
            return

        if pdfplumber is None:
            QMessageBox.warning(
                self,
                "Missing Dependency",
                "pdfplumber is required to read PDFs. Install it with 'pip install pdfplumber'.",
            )
            return

        bom_path = self._find_named_pdf(folder, "STRUCTURED BOM")

        if not bom_path:
            QMessageBox.warning(
                self,
                "Structured BOM Not Found",
                f"No '...STRUCTURED BOM....pdf' file was found in:\n{folder}",
            )
            self.start_btn.setEnabled(False)
            self._set_status("Structured BOM not found in that folder.")
            return

        quickprint_path = self._find_named_pdf(folder, "QUICKPRINT")

        self._pending_folder = folder
        found = os.path.basename(bom_path)
        if quickprint_path:
            found += f", {os.path.basename(quickprint_path)}"

        self.bom_pdf_label.setText(f"Folder ready — found: {found}")
        self._set_status("Folder selected. Click 'Start Automation' to begin.")
        self.start_btn.setEnabled(True)

    def start_automation(self):

        if not self._pending_folder or self._run_timer.isActive():
            return

        self.browse_pdf_btn.setEnabled(False)
        self.start_btn.setEnabled(False)
        self.start_btn.setText("Running...")
        self.test_scope_checkbox.setEnabled(False)
        self._run_started_at = time.monotonic()
        self._run_tick = 0
        self.run_busy_bar.show()
        self._run_timer.start()
        self._refresh_running_indicator()
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)

        try:
            self._set_status("Starting automation. Please wait...")
            completed = self._run_automation(self._pending_folder)
            self._set_run_outcome(
                "Completed - review the audit results" if completed
                else "Needs attention - automation did not complete; see activity log",
                "completed" if completed else "attention",
            )
        except Exception as exc:
            self._set_run_outcome("Failed - see activity log for the error", "failed")
            self._set_status(f"Automation failed: {type(exc).__name__}: {exc}")
            raise
        finally:
            self._run_timer.stop()
            self.run_busy_bar.hide()
            QApplication.restoreOverrideCursor()
            self.browse_pdf_btn.setEnabled(True)
            self.start_btn.setEnabled(True)
            self.start_btn.setText("Start Automation")
            self.test_scope_checkbox.setEnabled(True)

    def _run_elapsed(self):
        seconds = int(time.monotonic() - self._run_started_at) if self._run_started_at is not None else 0
        minutes, seconds = divmod(seconds, 60)
        return f"{minutes:02d}:{seconds:02d}"

    def _refresh_running_indicator(self):
        dots = "." * (self._run_tick % 4)
        self._run_tick += 1
        self._set_run_indicator(f"Running{dots}  |  Elapsed {self._run_elapsed()}  |  Please wait", "running")

    def _set_run_indicator(self, text, state):
        self.run_indicator.setText(text)
        if self.run_indicator.property("runState") != state:
            self.run_indicator.setProperty("runState", state)
            self.run_indicator.style().unpolish(self.run_indicator)
            self.run_indicator.style().polish(self.run_indicator)

    def _set_run_outcome(self, text, state):
        self._run_timer.stop()
        self._set_run_indicator(f"{text}  |  Elapsed {self._run_elapsed()}", state)

    def _run_automation(self, folder):

        self._pdf_read_issues.clear()
        self._pdf_ocr_page_count = 0
        self._set_status("Searching folder for Structured BOM PDF...")

        bom_path = self._find_named_pdf(folder, "STRUCTURED BOM")

        if not bom_path:
            QMessageBox.warning(
                self,
                "Structured BOM Not Found",
                f"No '...STRUCTURED BOM....pdf' file was found in:\n{folder}",
            )
            self._set_status("Structured BOM not found.")
            return

        self._set_status(f"Reading {os.path.basename(bom_path)}...")
        bom_extraction = self._read_pdf(bom_path)
        pdf_text = bom_extraction.text

        if not pdf_text.strip():
            details = "\n".join(self._pdf_read_issues)
            message = f"No readable text was found in:\n{bom_path}"
            if details:
                message += f"\n\nRead details:\n{details}"
            QMessageBox.warning(self, "Unreadable PDF", message)
            self._set_status("Structured BOM PDF was unreadable.")
            return

        sections = parse_structured_bom(pdf_text)

        if not sections:
            QMessageBox.warning(
                self,
                "No Structures Found",
                "No 'Structure N:' sections with bold assembly rows were found in this PDF.",
            )
            self._set_status("No structures found in Structured BOM PDF.")
            return

        self.bom_pdf_path = bom_path
        self.bom_pdf_label.setText(f"Loaded: {os.path.basename(bom_path)}")

        if self.test_scope_checkbox.isChecked():
            sections = [section for section in sections if section[0] == SCOPED_STRUCTURE]
            if not sections or not sections[0][2]:
                self._set_status(f"Structure {SCOPED_STRUCTURE} has no assembly rows.")
                return

        self._set_index2_job_sections(sections)
        quickprint_path = self._find_named_pdf(folder, "QUICKPRINT")
        parameter_sources = self._pdf_parameter_sources(bom_extraction, bom_path)

        if quickprint_path:
            self._set_status(f"Reading {os.path.basename(quickprint_path)}...")
            quickprint_extraction = self._read_pdf(quickprint_path)
            self.enclosure_type, self.enclosure_sheet = extract_enclosure_type(
                quickprint_extraction.text,
            )
            parameter_sources.extend(
                self._pdf_parameter_sources(quickprint_extraction, quickprint_path)
            )
        else:
            self.enclosure_type, self.enclosure_sheet = None, None

        known_paths = {
            os.path.normcase(os.path.abspath(path))
            for path in (bom_path, quickprint_path)
            if path
        }
        for name in sorted(os.listdir(folder)):
            if not name.upper().endswith(".PDF"):
                continue
            path = os.path.join(folder, name)
            if os.path.normcase(os.path.abspath(path)) in known_paths:
                continue
            self._set_status(f"Reading drawing PDF {name}...")
            drawing_extraction = self._read_pdf(path)
            parameter_sources.extend(self._pdf_parameter_sources(drawing_extraction, path))

        self._dimension_lookup_error = None
        self._dimension_results.clear()
        self._assembly_families.clear()
        self._assembly_links.clear()
        self._validation_results.clear()
        self._dimension_cache.clear()
        self._index2_map = {}
        self._excel_com.close()
        self._coverage_catalog = None
        self._update_index2_coverage()

        self._code_mappings, mapping_note = self._load_code_mappings()
        if self._code_mappings:
            self._configuration_repository.use_code_mappings(self._code_mappings)
        else:
            self._set_status(mapping_note)
            QMessageBox.warning(self, "Checklist Unavailable", mapping_note)
            return

        try:
            self._index2_map = self._excel_com.get_index2_map()
            if not self._index2_map:
                raise RuntimeError("Index2 column K returned no assembly mappings.")
        except Exception as exc:
            message = (
                f"Index2 could not be read: {type(exc).__name__}: {exc}. "
                "Close any Excel error dialogs, then click Start Automation again. "
                "The checklist has not been rebuilt."
            )
            self._set_status(message)
            QMessageBox.warning(self, "Index2 Unavailable", message)
            return

        self._parameter_repository.clear()
        self.refresh_index2_coverage()
        if self._coverage_catalog is None:
            QMessageBox.warning(self, "Index2 Coverage Unavailable", "Validation was not run because Index2 presence could not be checked.")
            self._set_status("Validation not run: Index2 presence could not be checked.")
            return
        self._section_bus_requirements.clear()
        self._job_dimensions = ParameterRepository()
        self._section_dimensions = {}
        self._section_parameters = {}
        variable_sources = [source_text for source_text, _source_name in parameter_sources]
        for source_text, source_name in parameter_sources:
            extracted = self._configuration_repository.extract_parameters(
                source_text, source_name,
            )
            self._parameter_repository.merge_sources(extracted)

        dwg_failures = []
        dwg_count = 0
        for name in sorted(os.listdir(folder)):
            if not name.upper().endswith(".DWG"):
                continue
            dwg_count += 1
            path = os.path.join(folder, name)
            self._set_status(f"Reading native DWG {name} through AutoCAD...")
            try:
                dwg_file = DWGFile(path)
                drawing_text = dwg_file.read_text()
                for number, parameters in dwg_file.section_parameters.items():
                    self._section_parameters.setdefault(number, ParameterRepository()).merge_sources(parameters)
                    self._section_dimensions.setdefault(number, ParameterRepository()).merge_sources(parameters)
                extracted = dwg_file.extract_parameters(self._configuration_repository, drawing_text)
                self._parameter_repository.merge_sources(extracted)
                variable_sources.append(drawing_text)
                parameter_sources.append((drawing_text, name))
            except Exception as exc:
                dwg_failures.append(f"{name}: {type(exc).__name__}: {exc}")

        for source_text, source_name in parameter_sources:
            global_dimensions, section_dimensions = extract_job_dimensions(source_text, source_name)
            self._job_dimensions.merge_sources(global_dimensions)
            self._parameter_repository.merge_sources(global_dimensions)
            for number, dimensions in section_dimensions.items():
                self._section_dimensions.setdefault(number, ParameterRepository()).merge_sources(dimensions)
                self._parameter_repository.merge_sources(dimensions)
            self._parameter_repository.merge_sources(extract_table_selectors(source_text, source_name))
            for number, requirements in section_bus_requirements(source_text, source_name).items():
                self._section_bus_requirements.setdefault(number, ParameterRepository()).merge_sources(requirements)

        self.parameter_variables = ExcelParameterVariables.from_sources(
            variable_sources,
            [style for _num, _title, items in sections for style, _desc, _w, _d in items],
            self._common_parameters,
        )
        self.test_query.completer().setModel(QStringListModel(list(self._common_parameters), self.test_query.completer()))

        if self.enclosure_type:
            self.enclosure_label.setText(
                f"Enclosure Type: {self.enclosure_type}  |  BASICSTR sheet: {self.enclosure_sheet}"
            )
        elif quickprint_path:
            self.enclosure_label.setText("Enclosure Type: not found in QuickPrint PDF")
        else:
            self.enclosure_label.setText("Enclosure Type: — (no QuickPrint PDF found in folder)")

        self._set_status("Resolving assembly families via live Excel (Index2)...")

        self.build_bom_checklist(sections)

        total = len(self.checkbox_items)
        mapped = sum(bool(item.text(3)) for item in self.checkbox_items.values())
        assembly_outcomes = self._eligible_assembly_results()
        if assembly_outcomes is None:
            raise RuntimeError("Index2 presence became unavailable before validation completed.")
        outcomes = list(assembly_outcomes.values())
        scope = f"Structure {SCOPED_STRUCTURE}" if self.test_scope_checkbox.isChecked() else "All structures"
        self._set_status(
            f"{scope}: {total} assembly/structure pairs; {mapped} matched to a checklist Code ({mapping_note}); "
            f"validation of {len(outcomes)} unique G/H assemblies present in Index2: "
            f"{outcomes.count('PASS')} pass / {outcomes.count('FAIL')} fail / "
            f"{outcomes.count('UNRESOLVED')} unresolved; "
            f"{self._parameter_repository.value_count} labeled parameter values from source files; "
            f"{dwg_count - len(dwg_failures)}/{dwg_count} DWGs read; "
            f"{self._parameter_repository.conflict_count} conflicting parameters; "
            f"{self._pdf_ocr_page_count} PDF pages OCR-read; "
            f"{len(self._pdf_read_issues)} PDF read/OCR issues."
        )
        if dwg_failures:
            QMessageBox.warning(self, "DWG Extraction Issues", "\n".join(dwg_failures))
        if self._pdf_read_issues:
            QMessageBox.warning(self, "PDF Read/OCR Issues", "\n".join(self._pdf_read_issues))
        return True

    def _load_code_mappings(self):
        """Use the supplied checklist, retaining Python-only family definitions."""

        self._set_status("Reading checklist_parameters.py...")
        self._common_parameters = tuple(dict.fromkeys(checklist_parameters.COMMON_PARAMETERS))
        mappings = parse_code_mappings(checklist_parameters.CODE_WISE_PARAMETERS)
        repository = getattr(self, "_configuration_repository", None)
        if repository is not None:
            repository.reload()
            checklist_mappings = repository.code_mappings()
            for code, mapping in checklist_mappings.items():
                existing = mappings.get(code)
                if not mapping.parameters and existing is not None:
                    mapping = replace(
                        existing, manually_mapped=mapping.manually_mapped, checklist_row=mapping.checklist_row,
                    )
                mappings[code] = mapping
            self._common_parameters = tuple(dict.fromkeys((
                *self._common_parameters,
                *(name for mapping in mappings.values() for name in mapping.parameters if name),
            )))
            self._set_status(
                f"Loaded {len(checklist_mappings)} checklist families; "
                f"{sum(mapping.manually_mapped is True for mapping in checklist_mappings.values())} marked TRUE in column N. "
                "Automatic extraction remains enabled for other families."
            )
        return mappings, "parameters: checklist_parameters.py + MCOE Checklist column N"

    def _code_mapping_for(self, style_number, tables=()):
        """Sheet1 row whose Code equals the Index2 family workbook name, or a sheet of it
        that holds this assembly."""

        base = split_style_number(style_number)[0]
        result = self._dimension_results.get(style_number, DimensionResult())
        candidates = [self._assembly_families.get(style_number, "")]

        for part in result.source.split("; "):
            if " / " in part:
                candidates.append(part.rsplit(" / ", 1)[1])

        candidates.extend(table.sheet_name for table in tables if table.rows_for_base(base))

        for candidate in candidates:
            mapping = self._code_mappings.get(normalize_code(candidate))
            if mapping is not None:
                return mapping

        names = {}
        for table in tables:
            if table.rows_for_base(base):
                for name in table.parameter_names:
                    names.setdefault(ClassifierWorkbookParser.parameter_key(name), name)
        if names:
            family = self._assembly_families.get(style_number, "")
            parameters = tuple(names.values())
            return CodeMapping(family, family, parameters, parameters, False)
        return None

    @staticmethod
    def _pdf_parameter_sources(extraction, path):
        basename = os.path.basename(path)
        sources = []
        for page_number, text in enumerate(extraction.pages, start=1):
            if not text.strip():
                continue
            ocr_marker = ", OCR" if page_number in extraction.ocr_pages else ""
            sources.append((text, f"{basename} (page {page_number}{ocr_marker})"))
        return sources

    def lookup_dimensions(self, style_number):
        """Resolve dimensions and retain the workbook evidence or failure reason."""

        def record(result):
            self._dimension_results[style_number] = result
            return result.width, result.depth

        if win32 is None:
            return record(DimensionResult(message="pywin32 (win32com) is not installed."))

        base, suffix = split_style_number(style_number)
        if not self._index2_map:
            if self._dimension_lookup_error:
                return record(DimensionResult(message=self._dimension_lookup_error))
            try:
                self._index2_map = self._excel_com.get_index2_map()
            except Exception as exc:
                self._dimension_lookup_error = f"Index2 unavailable: {type(exc).__name__}: {exc}"
                return record(DimensionResult(message=self._dimension_lookup_error))

        if not self._index2_map:
            return record(DimensionResult(message="Live Index2 column K returned no mappings."))

        link_target = self._index2_map.get(base.upper())

        if not link_target:
            return record(DimensionResult(message=f"{base} not found in Index2 column K."))

        self._assembly_families[style_number] = os.path.splitext(link_file_name(link_target))[0].upper()
        self._assembly_links[style_number] = link_target
        if not suffix:
            try:
                self._excel_com.get_group_tables(link_target, assembly_base=base)
            except Exception as exc:
                return record(DimensionResult(message=f"Lookup failed: {type(exc).__name__}: {exc}", source=link_target))
            return record(DimensionResult(
                message="Root table extracted; no BOM group was supplied for comparison.", source=link_target,
            ))

        cache_key = (link_target.upper(), self.enclosure_sheet, base.upper(), suffix.upper())

        if cache_key in self._dimension_cache:
            return record(self._dimension_cache[cache_key])

        try:
            self._set_status(f"Looking up {style_number} in {link_target}...")
            result = self._excel_com.lookup_dimensions(link_target, self.enclosure_sheet, base, suffix)
        except Exception as exc:
            result = DimensionResult(message=f"Lookup failed: {type(exc).__name__}: {exc}", source=link_target)

        if result.width or result.depth or not result.message.startswith("Lookup failed"):
            self._dimension_cache[cache_key] = result

        return record(result)

    def build_bom_checklist(self, sections):

        self.bom_list.blockSignals(True)
        try:
            self.bom_list.clear()
            self._configure_columns(max(
                [PARAMETER_COLUMN_COUNT] + [len(mapping.parameters) for mapping in self._code_mappings.values()]
            ))
            self.checkbox_items.clear()
            self._validation_results.clear()
            self._validation_checks.clear()
            if hasattr(self, "_excel_com"):
                targets = [
                    (num, title, [item for item in items if self._validation_scope_result(item[0]) is None])
                    for num, title, items in sections
                ]
                self._excel_com.prepare_assembly_targets(targets, self._index2_map)
            total = sum(len(items) for _num, _title, items in sections)
            completed = 0
            next_update = time.monotonic()
            for num, title, items in sections:
                if not items:
                    continue
                section_item = self.add_section(f"Structure {num}: {title}")
                section_item.setExpanded(True)
                structure_width, structure_depth = structure_dimensions_of(items)
                for style_number, description, _width, _depth in items:
                    self._set_status(f"Processing assembly {completed + 1}/{total}: {style_number}")
                    width, depth = self.lookup_dimensions(style_number) if self._validation_scope_result(style_number) is None else ("", "")
                    key = f"STRUCT{num}::{style_number}"
                    self.add_checkbox(section_item, key, style_number, description, width, depth, structure_width, structure_depth, num)
                    completed += 1
                    if time.monotonic() >= next_update:
                        self.update_progress()
                        next_update = time.monotonic() + 0.25
        finally:
            self.bom_list.expandAll()
            self.bom_list.blockSignals(False)
            self.update_progress()

    def add_section(self, text):

        item = QTreeWidgetItem(self.bom_list, [text, "", "", ""])
        item.setFirstColumnSpanned(True)

        font = QFont()
        font.setBold(True)
        font.setPointSize(11)
        item.setFont(0, font)
        item.setForeground(0, QBrush(QColor("#1a73e8")))

        return item

    def checklist_summary(self, style_number, width, depth):
        """Compares the family's checklist parameters against what was extracted."""

        family = self._assembly_families.get(style_number)

        if not family:
            return "Family not resolved from Index2.", ""

        record = self._configuration_repository.get_configuration(family)

        if record is None:
            return f"{family}: no row in MCOE checklist.", ""

        found, missing = [], []

        for rule in record.parameters:
            label = ConfigurationRepository._normalized(rule.parameter_name)

            if label in ASSEMBLY_LABELS:
                value = style_number
            elif label in WIDTH_LABELS:
                value = width
            elif label in DEPTH_LABELS:
                value = depth
            elif label in ENCLOSURE_LABELS:
                value = self.enclosure_type
            elif self._parameter_repository.has_conflict(rule.parameter_name):
                evidence = "; ".join(
                    f"{value} [{source}]"
                    for value, source in self._parameter_repository.get_entries(
                        rule.parameter_name,
                    )
                )
                missing.append(f"{rule.parameter_name} (conflicting values: {evidence})")
                continue
            else:
                value = self._parameter_repository.get_parameter(rule.parameter_name)

            if value:
                sources = ", ".join(dict.fromkeys(
                    source for entry_value, source
                    in self._parameter_repository.get_entries(rule.parameter_name)
                    if entry_value == value
                ))
                source_detail = f" [{sources}]" if sources else ""
                found.append(f"{rule.parameter_name} = {value}{source_detail}")
            else:
                missing.append(rule.parameter_name)

        total = len(record.parameters)
        summary = f"{family}: {len(found)}/{total} parameters found"
        detail = f"Found: {'; '.join(found) or 'none'}\nMissing: {'; '.join(missing) or 'none'}"

        return summary, detail

    def _resolved_value(self, style_number, name, width, depth, section_number=None):
        """Resolve job values without deriving validation inputs from the existing group."""

        label = ConfigurationRepository._normalized(name)
        canonical = ClassifierWorkbookParser.parameter_key(name)
        section = getattr(self, "_section_parameters", {}).get(str(section_number))
        if section:
            entries = [value for parameter, found in section.parameters.items()
                       if ClassifierWorkbookParser.parameter_key(parameter) == canonical
                       for value, _source in found]
            if entries:
                return ClassifierWorkbookParser.consistent_value(entries)

        if label in ASSEMBLY_LABELS:
            return style_number
        if label in WIDTH_LABELS:
            return width or dimension_value(self._dimension_entries(name, section_number))
        if label in DEPTH_LABELS:
            return depth or dimension_value(self._dimension_entries(name, section_number))
        if label in ENCLOSURE_LABELS:
            return self.enclosure_type or ""

        entries = [value for parameter, found in self._parameter_repository.parameters.items()
                   if ClassifierWorkbookParser.parameter_key(parameter) == canonical
                   for value, _source in found]
        if entries:
            return ClassifierWorkbookParser.consistent_value(entries)
        variables = self.parameter_variables
        found = [value for parameter, values in variables.values.items()
                 if ClassifierWorkbookParser.parameter_key(parameter) == canonical for value in values]
        return ClassifierWorkbookParser.consistent_value(found)

    def _dimension_entries(self, name, section_number=None):
        section = getattr(self, "_section_dimensions", {}).get(str(section_number))
        entries = dimension_entries(section, name) if section else ()
        if entries:
            return entries
        repository = getattr(self, "_job_dimensions", self._parameter_repository)
        return dimension_entries(repository, name)

    def _family_tables(self, style_number):
        """(tables, error) of the Index2-linked family workbook for this assembly."""

        link_target = self._assembly_links.get(style_number)
        if not link_target:
            return (), "Family not resolved from Index2."

        try:
            self._set_status(f"Reading family tables for {style_number}...")
            base, _group = split_style_number(style_number)
            tables = self._excel_com.get_group_tables(link_target, assembly_base=base)
            if any(table.classifier_type for table in tables) and self._assembly_families.get(style_number) == "BASICSTR":
                if not self.enclosure_sheet:
                    return (), "BASICSTR requires an enclosure type from QuickPrint."
                selected = tuple(table for table in tables if table.sheet_name.casefold() == self.enclosure_sheet.casefold())
                if selected:
                    tables = selected
                else:
                    return (), f"Classifier: enclosure sheet {self.enclosure_sheet} was not found for {base}."
            return tables, ""
        except Exception as exc:
            return (), f"Family tables unavailable: {type(exc).__name__}: {exc}"

    def validate_assembly(self, style_number, mapping, width, depth, tables, tables_error, description="", section_number=None):
        """Compares the group in the style number with the group the family workbook gives
        for section/assembly inputs, retaining the job's specified bus test requirement."""

        base, expected = split_style_number(style_number)

        if tables_error:
            return ValidationEngine.compare(expected, "", tables_error)
        dimension_sources = []
        dimension_errors = {}
        for name, supplied in (("Width (Wide)", width), ("Depth", depth)):
            if supplied:
                continue
            entries = self._dimension_entries(name, section_number) if hasattr(self, "_dimension_entries") else ()
            if entries and not dimension_value(entries):
                evidence = "; ".join(f"{value} [{source}]" for value, source in entries)
                dimension_errors[normalize_code(name)] = f"Conflicting or invalid job values for {name}: {evidence}."
            value = dimension_value(entries)
            if value:
                if name == "Depth":
                    depth = value
                else:
                    width = value
                dimension_sources.append(f"{name}={value} [{'; '.join(dict.fromkeys(source for _value, source in entries))}]")
        scoped = assembly_bus_requirements(getattr(self, "_section_bus_requirements", {}), section_number, description)
        if scoped and scoped.has_conflict("Heat Test/Density"):
            sources = "; ".join(f"{value}: {source}" for value, source in scoped.get_entries("Heat Test/Density"))
            return ValidationEngine.compare(expected, "", f"Conflicting section {section_number} bus requirements: {sources}")

        from_description = description_parameter_values(description)
        test_requirement = ""
        if is_bus_description(description):
            test_requirement = self._resolved_value(style_number, "Heat Test/Density", width, depth, section_number)
            if "Amperage" not in from_description:
                bus_amperage = self._resolved_value(style_number, "Bus Amperage", width, depth, section_number)
                if bus_amperage:
                    from_description["Amperage"] = bus_amperage
        if scoped:
            test_requirement = scoped.get_parameter("Heat Test/Density")
        if test_requirement:
            for name in tuple(from_description):
                if ClassifierWorkbookParser.parameter_key(name) == "TEST":
                    from_description[name] = test_requirement
        job_material = self._resolved_value(style_number, "Material", width, depth, section_number)
        if job_material and from_description.get("Material") in ("CU", "COPPER"):
            from_description["Material"] = job_material  # the job's plating is more specific than 'CU'
        by_key = {ClassifierWorkbookParser.parameter_key(name): value for name, value in from_description.items()}
        values = {}
        for name in mapping.parameters if mapping else ():
            if not name or ConfigurationRepository._normalized(name) in ASSEMBLY_LABELS:
                continue
            value = by_key.get(ClassifierWorkbookParser.parameter_key(name), "") or self._resolved_value(style_number, name, width, depth, section_number)
            if value:
                values[name] = value

        if scoped:
            values["Heat Test/Density"] = scoped.get_parameter("Heat Test/Density")

        classified = tuple(table for table in tables if table.classifier_type)
        if classified:
            needed = {name for table in classified for name in table.parameter_names}
            inputs = {**from_description, **values}
            if width:
                inputs["Width (Wide)"] = width
            if depth:
                inputs["Depth"] = depth
            for name in needed:
                canonical = ClassifierWorkbookParser.parameter_key(name)
                for dimension_name, error in dimension_errors.items():
                    if ClassifierWorkbookParser.parameter_key(dimension_name) == canonical:
                        return ValidationEngine.compare(expected, "", error)
                repository = getattr(self, "_parameter_repository", None)
                supplied_dimension = (
                    canonical == "STRUCTUREWIDTH" and width
                    or canonical == "STRUCTUREDEPTH" and depth
                )
                described = any(ClassifierWorkbookParser.parameter_key(key) == canonical for key in from_description)
                if repository and not supplied_dimension and not described and not (scoped and canonical == "TEST"):
                    entries = [
                        entry for parameter, found in repository.parameters.items()
                        if ClassifierWorkbookParser.parameter_key(parameter) == canonical
                        for entry in found
                    ]
                    if entries and any(
                        not ClassifierWorkbookParser.values_match(entries[0][0], value)
                        for value, _source in entries[1:]
                    ):
                        return ValidationEngine.compare(expected, "", f"Conflicting job values for required table input: {name}.")
                if any(ClassifierWorkbookParser.parameter_key(key) == canonical for key in inputs):
                    continue
                value = self._resolved_value(style_number, name, width, depth, section_number)
                if value:
                    inputs[name] = value
                    continue
                for parameter, found in self.parameter_variables.values.items():
                    if ClassifierWorkbookParser.parameter_key(parameter) == canonical and len(found) == 1:
                        inputs[name] = found[0]
            actual, message = ClassifierWorkbookParser.identify_group(classified, base, inputs)
            if scoped:
                message += f" Section {section_number} inputs: " + "; ".join(
                    source for _value, source in scoped.get_entries("Heat Test/Density")
                )
            if dimension_sources:
                message += " Job inputs: " + "; ".join(dimension_sources) + "."
            return ValidationEngine.compare(expected, actual, message)

        link_target = self._assembly_links.get(style_number)
        try:
            sheets = self._excel_com.get_sheet_values(link_target) if link_target else ()
        except Exception as exc:
            return ValidationEngine.compare(expected, "", f"Workbook data unavailable: {exc}")
        matches = [match for sheet_name, rows in sheets for match in find_grid_group_matches(rows, sheet_name, base, "")]
        needed = {name for match in matches for name in match.required_parameters}
        compared_dimensions = {
            normalize_code(name) for match in matches
            for name in grid_parameter_values(match)
        }
        if mapping:
            compared_dimensions.update(normalize_code(name) for name in mapping.parameters if name)
        for name, error in dimension_errors.items():
            if name in compared_dimensions or name in {normalize_code(item) for item in needed}:
                return ValidationEngine.compare(expected, "", error)
        needed.update(("Metering Configuration", "Door Vents"))
        for name in needed:
            repository = getattr(self, "_parameter_repository", None)
            is_supplied_dimension = (
                ConfigurationRepository._normalized(name) in WIDTH_LABELS and width
                or ConfigurationRepository._normalized(name) in DEPTH_LABELS and depth
            )
            if repository and repository.has_conflict(name) and not is_supplied_dimension:
                return ValidationEngine.compare(expected, "", f"Conflicting job values for required table input: {name}.")
            value = from_description.get(name, "") or self._resolved_value(style_number, name, width, depth, section_number)
            if value and name not in values:
                values[name] = value
        if width:
            values.setdefault("Width (Wide)", width)
        if depth:
            values.setdefault("Depth", depth)
        actual, message = find_grid_group(
            sheets, base, {**from_description, **values}, self.parameter_variables.values, description,
        )
        if not actual and message == "Base not found in grid-layout sheets.":
            table_group, table_message = GroupDeterminationEngine.identify_group(tables, base, values) if values else ("", "")
            if table_group or message == "Base not found in grid-layout sheets.":
                actual, message = table_group, table_message or message

        if scoped:
            sources = "; ".join(source for _value, source in scoped.get_entries("Heat Test/Density"))
            message += f" Section {section_number} Heat Test/Density = {values['Heat Test/Density']}. Source: {sources}"

        if dimension_sources:
            message += " Job inputs: " + "; ".join(dimension_sources) + "."
        return ValidationEngine.compare(expected, actual, message)

    def _grid_values(self, style_number, mapping, structure_width, description=""):
        """Index2 values for grid-layout family sheets, limited to this row's parameters."""

        link_target = self._assembly_links.get(style_number)
        if not link_target or mapping is None:
            return {}

        try:
            sheets = self._excel_com.get_sheet_values(link_target)
        except Exception:
            return {}

        base, group = split_style_number(style_number)
        matches = [
            match for sheet_name, values in sheets
            for match in find_grid_group_matches(values, sheet_name, base, "")
        ]
        matches, selection_error = select_grid_tables(matches, self.parameter_variables.values, description)
        if selection_error:
            return {}
        matches = [match for match in matches if match.group == group]
        derived = {normalize_code(name): value for name, value in merge_grid_matches(matches, structure_width).items()}

        return {name: derived[normalize_code(name)] for name in mapping.parameters if name and normalize_code(name) in derived}

    def add_checkbox(self, section_item, key, style_number, description, width, depth, structure_width="", structure_depth="", section_number=None):

        result = self._dimension_results.get(style_number, DimensionResult())
        validation = self._validation_scope_result(style_number)
        tables, tables_error = self._family_tables(style_number) if validation is None else ((), validation.message)
        mapping = self._code_mapping_for(style_number, tables) if validation is None else None
        if mapping is not None:
            self._ensure_parameter_columns(len(mapping.parameters))
        count = self._parameter_columns
        base, expected_group = split_style_number(style_number)
        description_values = description_parameter_values(description)
        job_width = structure_width or description_values.get("Width (Wide)", "")
        job_depth = structure_depth or description_values.get("Depth", "")
        width = job_width or dimension_value(self._dimension_entries("Width (Wide)", section_number))
        depth = job_depth or dimension_value(self._dimension_entries("Depth", section_number))
        colors = {}
        tips = {}
        scoped = assembly_bus_requirements(self._section_bus_requirements, section_number, description)

        if mapping is not None:
            classified = any(table.classifier_type for table in tables)
            workbook_values = {
                **({} if classified else self._grid_values(style_number, mapping, structure_width, description)),
                **TableValueExtractor.extract(
                    tables, base, expected_group,
                    {
                        name: (name, alias)
                        for name, alias in zip(mapping.parameters, mapping.aliases) if name
                    },
                ),
            }
            parameter_cells = []
            for column, name in enumerate(mapping.parameters, start=4):
                if not name:
                    parameter_cells.append("")
                    continue
                job_value = self._resolved_value(style_number, name, width, depth, section_number)
                scoped_entries = scoped.get_entries(name) if scoped else ()
                if scoped_entries:
                    job_value = scoped.get_parameter(name) or "CONFLICT"
                sheet_value = workbook_values.get(name, "")
                text = f"{name} = {job_value}" if job_value else name
                if sheet_value:
                    text += f" [Index2: {sheet_value}]"
                    tips[column] = f"Job files: {job_value or 'not found'}\nIndex2 workbook ({base} {expected_group}): {sheet_value}"
                    if job_value:
                        matches = ClassifierWorkbookParser.values_match(sheet_value, job_value) if classified else values_match(sheet_value, job_value)
                        colors[column] = "#188038" if matches else "#d93025"
                if scoped_entries:
                    tips[column] = tips.get(column, "") + "\nSection requirement: " + "; ".join(source for _value, source in scoped_entries)
                parameter_cells.append(text)
            parameter_cells += [""] * (count - len(parameter_cells))
            cells = [mapping.name, mapping.code, *parameter_cells]
            mode = (
                "Manually mapped (column N = TRUE)" if mapping.manually_mapped is True
                else "Automatic extraction (column N is not TRUE)" if mapping.manually_mapped is False and mapping.checklist_row
                else "Automatic extraction (no checklist row)" if mapping.manually_mapped is False
                else "Python parameter mapping"
            )
            tip = f"{mode}\nChecklist row {mapping.checklist_row or '-'}: {mapping.code}"
            for column in range(4, 4 + count):
                tips[column] = tip + ("\n" + tips[column] if column in tips else "")
        else:
            cells = [""] * (2 + count)
            family = self._assembly_families.get(style_number)
            tip = validation.message if validation is not None else (
                f"No checklist Code row matches '{family}'." if family
                else f"Family not resolved from Index2: {result.message}"
            )

        if validation is None:
            validation = self.validate_assembly(
                style_number, mapping,
                job_width, job_depth,
                tables, tables_error, description, section_number,
            )
        self._validation_results[key] = validation
        self._validation_checks.setdefault(style_number.strip().upper(), []).append(validation.status)

        item = QTreeWidgetItem(
            section_item,
            [style_number, description, *cells, validation.expected_group, validation.actual_group, validation.status],
        )
        for column in range(item.columnCount()):
            item.setToolTip(column, validation.message if column >= 4 + count else tips.get(column, tip))
        for column, color in colors.items():
            item.setForeground(column, QBrush(QColor(color)))

        status_column = item.columnCount() - 1
        item.setForeground(status_column, QBrush(QColor(VALIDATION_COLORS[validation.status])))

        item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
        item.setCheckState(0, Qt.CheckState.Unchecked)

        self.checkbox_items[key] = item
        location = f"Structure {section_number}: " if section_number is not None else ""
        activity = f"{location}{style_number} - {validation.status} (Expected {validation.expected_group or '-'}, Actual {validation.actual_group or '-'})"
        if validation.status != "PASS" and validation.message:
            activity += f". {validation.message}"
        self._set_status(activity)

    def update_progress(self):

        total = len(self.checkbox_items)
        checked = sum(
            1 for item in self.checkbox_items.values()
            if item.checkState(0) == Qt.CheckState.Checked
        )

        if total > 0:
            pct = int(checked / total * 100)
            self.progress.setValue(pct)

        self.progress.setFormat(f"Manual review: {checked}/{total} assembly/structure rows (%p%)")
        if not total:
            self.progress.setValue(0)
        if hasattr(self, "summary_labels"):
            self._update_coverage_metrics()
            self._apply_audit_filters()

    def closeEvent(self, event):

        self._run_timer.stop()
        self._tab_animation.stop()
        self._excel_com.close()
        super().closeEvent(event)

    # ------------------------------------------------------------------
    # Styling
    # ------------------------------------------------------------------

    def _build_stylesheet(self):

        return """
        QMainWindow, QWidget { background: #f3f6fc; color: #26354b; font-size: 13px; }
        QFrame#sidebar { background: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 #172b50, stop:1 #0c172e); border: none; }
        QFrame#sidebar QLabel, QFrame#sidebar QCheckBox { background: transparent; color: #ffffff; }
        QFrame#sidebar QCheckBox::indicator { width: 12px; height: 12px; border: 1px solid #ffffff; border-radius: 2px; background: transparent; }
        QFrame#sidebar QCheckBox::indicator:checked { width: 8px; height: 8px; border: 3px solid #ffffff; background: #174dc9; }
        QLabel#brand { font-family: 'Bahnschrift'; font-size: 27px; font-weight: 700; padding-bottom: 6px; }
        QFrame#sidebar QLabel#sidebarCaption { font-size: 10px; font-weight: 600; color: #a8bddf; }
        QFrame#sidebar QLabel#activityStatus { color: #c4d4ee; font-size: 11px; }
        QPlainTextEdit#automationLog { background: #101e37; color: #dbe8ff; border: 1px solid #314563; border-radius: 8px; padding: 7px; font-size: 11px; }
        QFrame#sidebar QPushButton#clearLog { padding: 0; min-height: 0; border: 1px solid #506586; border-radius: 6px; }
        QFrame#sidebar QPushButton { background: transparent; color: #dbe8ff; border: 1px solid transparent; border-radius: 8px; text-align: left; padding: 10px; }
        QFrame#sidebar QPushButton:hover { background: #263d60; border-color: #435e88; }
        QFrame#sidebar QPushButton:checked { background: #356bf0; color: #ffffff; border-color: #5786f6; }
        QLabel#eyebrow { font-size: 10px; font-weight: 600; color: #647a9c; }
        QLabel#pageTitle { font-family: 'Bahnschrift'; font-size: 28px; font-weight: 600; color: #172c42; }
        QLabel#jobLabel { color: #637692; padding: 8px 12px; background: #e9eff9; border-radius: 7px; }
        QFrame#summary { background: transparent; border: none; }
        QFrame[metricCard="true"] { background: #ffffff; border: 1px solid #dfe7f3; border-radius: 10px; }
        QFrame#cardPASS, QFrame#cardpresent { background: #effaf5; border-color: #ccebdc; }
        QFrame#cardFAIL { background: #fff3f3; border-color: #f3d6d9; }
        QFrame#cardUNRESOLVED, QFrame#cardmissing { background: #fff9ed; border-color: #efdfbc; }
        QFrame[metricCard="true"] QLabel { background: transparent; }
        QLabel[metric="true"] { font-family: 'Bahnschrift'; font-size: 28px; font-weight: 600; }
        QLabel#metricvalidation_total, QLabel#metricvalidation_rate, QLabel#metricall_assemblies, QLabel#metricgh_assemblies, QLabel#metriccoverage_rate { color: #356bf0; }
        QLabel#metricPASS, QLabel#metricpresent { color: #14804a; }
        QLabel#metricFAIL { color: #c13342; }
        QLabel#metricUNRESOLVED, QLabel#metricmissing { color: #a66a08; }
        QLabel#metricCaption { font-size: 11px; color: #68788b; }
        QGroupBox, QTabWidget::pane { border: none; margin: 0; padding: 0; }
        QPushButton { background: #ffffff; color: #254d87; border: 1px solid #cdd9ea; border-radius: 8px; padding: 8px 14px; font-weight: 600; min-height: 18px; }
        QPushButton:hover { background: #eaf1ff; border-color: #83a7ec; }
        QPushButton:pressed { background: #d9e6ff; }
        QPushButton:disabled { background: #edf0f5; color: #8b99aa; border-color: #dde3ec; }
        QPushButton#start_btn { background: #356bf0; color: #ffffff; border-color: #356bf0; }
        QPushButton#start_btn:hover { background: #2454ca; }
        QPushButton#start_btn:disabled { background: #e0e8f9; color: #7184a7; border-color: #d3ddf0; }
        QPushButton:focus { border-color: #356bf0; }
        QLineEdit, QComboBox { background: #ffffff; border: 1px solid #cdd9ea; border-radius: 8px; padding: 8px; min-height: 18px; }
        QLineEdit:focus, QComboBox:focus { border: 1px solid #205bea; }
        QComboBox::drop-down { border: none; width: 24px; }
        QComboBox QAbstractItemView { background: #ffffff; selection-background-color: #e7efff; selection-color: #173968; }
        QCheckBox { spacing: 7px; }
        QTreeWidget { background: #ffffff; alternate-background-color: #f8fafc; border: 1px solid #dce3ed; border-radius: 9px; outline: none; selection-background-color: #e3edff; selection-color: #143563; }
        QTreeWidget::item { min-height: 31px; padding: 3px 5px; border-bottom: 1px solid #f0f3f7; }
        QTreeWidget::item:hover { background: #f0f5ff; }
        QTreeWidget::item:selected { background: #e3edff; color: #143563; }
        QHeaderView::section { background: #e9f0fc; color: #315274; padding: 10px 6px; border: none; border-right: 1px solid #dce5f3; font-weight: 600; }
        QPlainTextEdit { background: #ffffff; border: 1px solid #dce3ed; border-radius: 9px; padding: 10px; selection-background-color: #dce9ff; }
        QScrollBar:vertical { background: transparent; width: 10px; margin: 2px; }
        QScrollBar::handle:vertical { background: #aabbd4; min-height: 28px; border-radius: 3px; }
        QScrollBar::handle:vertical:hover { background: #7796c5; }
        QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
        QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical { background: transparent; }
        QLabel#scan_status { color: #67758a; font-size: 11px; }
        QLabel#runIndicator { color: #637692; font-weight: 600; font-size: 12px; }
        QLabel#runIndicator[runState="running"] { color: #2454ca; }
        QLabel#runIndicator[runState="completed"] { color: #14804a; }
        QLabel#runIndicator[runState="attention"] { color: #a66a08; }
        QLabel#runIndicator[runState="failed"] { color: #c13342; }
        QProgressBar { background: #e8edf4; color: #29455f; border: none; border-radius: 4px; text-align: center; font-size: 11px; }
        QProgressBar::chunk { background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #8eb8ff, stop:1 #78d8c0); border-radius: 4px; }
        QToolTip { background: #ffffff; color: #243347; border: 1px solid #cdd9ea; padding: 6px; }
        """


if __name__ == "__main__":

    app = QApplication(sys.argv)

    window = AuditReviewDemo()

    window.showMaximized()

    sys.exit(app.exec())

