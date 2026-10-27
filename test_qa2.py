import os
import pickle
import subprocess
import sys
import tempfile
import time
import unittest
from types import SimpleNamespace
from unittest.mock import MagicMock, PropertyMock, patch

import QA2


def excel_sheet(name, values, row=1, column=1, merges=()):
    from openpyxl.utils.cell import range_boundaries

    cells = {}
    for top, left, bottom, right in merges:
        area = SimpleNamespace(
            Row=top, Column=left,
            Rows=SimpleNamespace(Count=bottom - top + 1), Columns=SimpleNamespace(Count=right - left + 1),
        )
        cells.update(((r, c), area) for r in range(top, bottom + 1) for c in range(left, right + 1))

    def range_state(address):
        left, top, right, bottom = range_boundaries(address)
        states = {(r, c) in cells for r in range(top, bottom + 1) for c in range(left, right + 1)}
        return SimpleNamespace(MergeCells=next(iter(states)) if len(states) == 1 else None)

    return SimpleNamespace(
        Name=name,
        UsedRange=SimpleNamespace(
            Row=row, Column=column, Value=values, MergeCells=None if merges else False,
            Rows=SimpleNamespace(Count=len(values or ())),
            Columns=SimpleNamespace(Count=max(map(len, values or ()), default=0)),
        ),
        Range=MagicMock(side_effect=range_state),
        Cells=MagicMock(side_effect=lambda r, c: SimpleNamespace(MergeArea=cells[(r, c)])),
    )


class AnchoredExtractionTests(unittest.TestCase):
    def test_plated_materials_special_density_and_phase_alternatives(self):
        from Classifier import AnchoredTableParser
        attributes = AnchoredTableParser.attributes
        self.assertEqual(attributes("COPPER AG"), {"Material": "Silver Plated Copper"})
        self.assertEqual(attributes("COPPER SN"), {"Material": "Tin Plated Copper"})
        self.assertEqual(attributes("DENSITY 800A/SQ INCH"), {"Test": "Density 800A/SQ INCH"})
        self.assertEqual(attributes("3 PHASE 3 WIRE, 1 PHASE 3 WIRE")["Phase"], "3 Phase/1 Phase")
        self.assertEqual(attributes('3 PH. APPLICATION "A"')["Phase"], "3 Phase")
        self.assertEqual(attributes('1PH. APPLICATION "A"')["Phase"], "1 Phase")
        self.assertFalse(QA2.ClassifierWorkbookParser.values_match("Density 800A/SQ INCH", "DENSITY"))

    def test_explicit_door_inputs_exclude_cover_only_tables(self):
        from Classifier import AnchoredTableParser
        door = AnchoredTableParser.attributes("REVERSE FEED DOUBLE DOOR 1 METER")
        cover = AnchoredTableParser.attributes("UGPS COVER ASSEMBLY")
        configs = [{**door, "Structure Width": "38", "Group No.": "G79"},
                   {**cover, "Structure Width": "38", "Group No.": "G11"}]
        actual, _ = QA2.ClassifierWorkbookParser.identify_group(self.tables(configs), "42C5495", {
            "Width": "38", "Door Type": "DOUBLE", "Feed": "REVERSE", "Meter Count": "1",
        })
        self.assertEqual(actual, "G79")

    def test_shared_merged_headers_keep_width_test_and_local_phase(self):
        from Classifier import AnchoredTableParser
        workbook = QA2.Workbook()
        sheet = workbook.active
        for row in (
            ("ASSY DWG", "AMPS", "20W", None, "36W", None),
            (None, None, "Heat Tested", "Density", "Heat Tested", "Density"),
            (None, None, "COPPER AG", "COPPER AG", "COPPER AG", "COPPER AG"),
            ("6552C52", "3000A", "G01", "G02", "G03", "G04"),
            ("THREE PHASE CROSS BUS",),
            ("42C5495", "3000A", "G11", "G12", "G13", "G14"),
            ("NEUTRAL CROSS BUS",),
            ("42C5496", "3000A", "G21", "G22", "G23", "G24"),
        ):
            sheet.append(row)
        sheet.merge_cells("A1:A3")
        sheet.merge_cells("B1:B3")
        configs = AnchoredTableParser(sheet).parse("42C5495")["Configurations"]
        actual, _ = QA2.ClassifierWorkbookParser.identify_group(self.tables(configs), "42C5495", {
            "Amperage": "3000", "Width": "36", "Test": "DENSITY", "Material": "CUAG", "Phase": "3",
        })
        self.assertEqual(actual, "G14")
        neutral = AnchoredTableParser(sheet).parse("42C5496")["Configurations"]
        self.assertTrue(neutral)
        self.assertTrue(all("Phase" not in config for config in neutral))

    def test_discrete_amperage_range_requires_consistent_groups_and_endpoints(self):
        configs = [{"Amperage": str(amps), "Material": "Copper", "Group No.": group}
                   for amps, group in ((400, "G01"), (600, "G01"), (800, "G01"), (1000, "G02"))]
        identify = QA2.ClassifierWorkbookParser.identify_group
        tables = self.tables(configs)
        self.assertEqual(identify(tables, "42C5495", {"Amperage": "400-800"})[0], "G01")
        self.assertEqual(identify(tables, "42C5495", {"Amperage": "400-1000"})[0], "")
        self.assertEqual(identify(tables, "42C5495", {"Amperage": "400-900"})[0], "")
        configs[1]["Group No."] = "G03"
        self.assertEqual(identify(self.tables(configs), "42C5495", {"Amperage": "400-800"})[0], "")

    @staticmethod
    def configurations(rows, root="42C5495", title="Family"):
        from Classifier import AnchoredTableParser
        workbook = QA2.Workbook()
        workbook.active.title = title
        for row in rows:
            workbook.active.append(row)
        return AnchoredTableParser(workbook.active).parse(root)["Configurations"]

    def test_ref_and_full_style_labels_are_not_note_or_partial_matches(self):
        from Classifier import root_reference
        self.assertEqual(root_reference("REF - 8712C42"), "8712C42")
        self.assertEqual(root_reference("8712C42G02"), "8712C42")
        self.assertNotEqual(root_reference("8712C420"), "8712C42")
        self.assertEqual(root_reference("SEE 8712C42 FOR DETAILS"), "")
        self.assertEqual(root_reference("COPPER"), "")

    def test_inline_current_phase_and_sheet_width(self):
        configs = self.configurations([
            ("THREE PHASE",),
            ("CURRENT RATING", "ASSEMBLY DRAWING", "ALUM", "CUAG"),
            (None, None, "DENSITY", "DENSITY"),
            ("800A, MD/PD4 FRAME", "42C5495", "G01", "G02"),
            (None,),
            ("SINGLE PHASE",),
            ("CURRENT RATING", "ASSEMBLY DRAWING", "ALUM", "CUAG"),
            (None, None, "DENSITY", "DENSITY"),
            ("800A, MD/PD4 FRAME", "42C5495", "G11", "G12"),
        ], title="400A - 1000A (36W)")
        self.assertEqual(len(configs), 4)
        for config in configs:
            self.assertEqual(config["Amperage"], "800")
            self.assertEqual(config["Structure Width"], "36")
            self.assertEqual(config["Test"], "Density")
            self.assertEqual(config["Phase"], "3 Phase" if config["Group No."] in ("G01", "G02") else "1 Phase")

    def test_adjacent_root_does_not_supply_groups(self):
        configs = self.configurations([
            ("REF - 42C5495", None, None, None, "REF - 6552C52"),
            ("AMPS", "AL", "CU", None, "AMPS", "AL", "CU"),
            (800, "G01", "G02", None, 800, "G91", "G92"),
        ])
        self.assertEqual({c["Group No."] for c in configs}, {"G01", "G02"})

    def test_vertically_merged_root_does_not_end_its_own_table(self):
        from Classifier import AnchoredTableParser
        workbook = QA2.Workbook()
        sheet = workbook.active
        sheet.append(["WIDTH", "ASSEMBLY", "CU"])
        sheet.append([12, "42C5495", "G01"])
        sheet.append([30, None, "G02"])
        sheet.append([36, None, "G03"])
        sheet.merge_cells("B2:B4")
        # The root's first width has no group in this layout.
        sheet["C2"] = "---"
        configs = AnchoredTableParser(sheet).parse("42C5495")["Configurations"]
        self.assertEqual([(c["Structure Width"], c["Group No."]) for c in configs], [("30", "G02"), ("36", "G03")])

    def test_inline_root_keeps_blank_and_populated_meter_socket_rows(self):
        configs = self.configurations([
            ("METER SOCKETS / PROVISIONS", "ASSEMBLY DRAWING", "WIDTH", None),
            (None, None, "26W", "36W"),
            ("(1) Blank", "42C5495", "G02", "G04"),
            ("(1) Socket", None, "G06", "N/A"),
            ("(2) Sockets", None, "N/A", "G10"),
        ])
        tables = self.tables(configs)
        identify = QA2.ClassifierWorkbookParser.identify_group
        self.assertEqual(identify(tables, "42C5495", {"Width": 36})[0], "")
        self.assertEqual(identify(tables, "42C5495", {"Width": 36, "Meter Socket Cover": "2 SOCKETS"})[0], "G10")
        self.assertEqual(identify(tables, "42C5495", {"Width": 26, "Meter Socket Cover": "1 SOCKETS"})[0], "G06")

    def test_circuit_headers_are_not_replaced_by_unrelated_material_headers(self):
        workbook = QA2.Workbook()
        sheet = workbook.active
        sheet.append(["ASSEMBLY DRAWING", "ALUM", "CUAG"])
        sheet.append(["6552C52", "G01", "G02"])
        sheet.append([])
        sheet.append(["ASSEMBLY DRAWING", "Phasing", "Circuit Numbers", None])
        sheet.append([None, None, "1, 2 or 3", "1/4, 2/5 or 3/6"])
        sheet.append(["42C5495", "Single or Three", "G01", "G02"])
        sheet.merge_cells("C4:D4")
        tables = QA2.ClassifierWorkbookParser.extract_tables(workbook, "SCMM.xlsx", "42C5495")
        self.assertTrue(tables)
        self.assertTrue(all("Circuit Numbers" in t.parameter_names for t in tables))
        self.assertTrue(all("Material" not in t.parameter_names for t in tables))
        group, message = QA2.ClassifierWorkbookParser.identify_group(
            tables, "42C5495", {"Material": "AL", "Phase": "3"},
        )
        self.assertEqual(group, "")
        self.assertIn("Circuit Numbers", message)

    def test_unparsed_local_table_does_not_borrow_previous_material_axes(self):
        workbook = QA2.Workbook()
        sheet = workbook.active
        for row in (
            ("ASSEMBLY DRAWING", "ALUM", "CUAG"),
            ("6552C52", "G21", "G22"),
            (),
            ("ASSEMBLY DRAWING", "Unknown selector", None),
            ("42C5495", "G01", "G02"),
        ):
            sheet.append(row)
        tables = QA2.ClassifierWorkbookParser.extract_tables(workbook, "SCMM.xlsx", "42C5495")
        self.assertTrue(tables)
        self.assertTrue(all("Material" not in table.parameter_names for table in tables))
        self.assertEqual(QA2.ClassifierWorkbookParser.identify_group(tables, "42C5495", {"Material": "AL"})[0], "")

    def test_centered_unmerged_headers_keep_own_columns(self):
        configs = self.configurations([
            ("REF - 42C5495",),
            ("AMPS", "DEPTH", "26W", None, "36W", None),
            (None, None, "AL", "CU", "AL", "CU"),
            (400, 24, "G01", "G02", "G03", "G04"),
        ])
        self.assertEqual([(c["Structure Width"], c["Material"]) for c in configs], [
            ("26", "Aluminum"), ("26", "Copper"), ("36", "Aluminum"), ("36", "Copper"),
        ])

    def test_depth_drillings_and_current_range(self):
        configs = self.configurations([
            ("DRAWING", "AMPACITY", "WIDTH", "BUS", "30D", "36D"),
            (None, None, None, None, "(8 DRILLINGS)", "(10 DRILLINGS)"),
            ("42C5495", "2500-3000A", "51W", "CU DEN", "G02", "G03"),
        ])
        self.assertEqual(len(configs), 2)
        self.assertEqual(configs[1]["Number of Drillings"], "10")
        self.assertEqual(configs[1]["Structure Depth"], "36")
        self.assertEqual(configs[1]["Amperage"], "2500-3000")
        self.assertEqual(configs[1]["Test"], "Density")

    def test_internal_header_change_is_not_silently_reused(self):
        configs = self.configurations([
            ("REF - 42C5495",),
            ("3 PHASE 4 WIRE",),
            ("AMPS", "DEPTH", "CU"),
            (800, 36, "G26"),
            ("3 PHASE 3 WIRE",),
            ("AMPS", "DEPTH", "AL/CU"),
            (800, 36, "G30"),
        ])
        by_group = {c["Group No."]: c for c in configs}
        self.assertNotIn("Unparsed Table Context", by_group["G26"])
        self.assertIn("3 PHASE 3 WIRE", by_group["G30"]["Unparsed Table Context"])
        self.assertEqual(by_group["G30"]["Wires"], "3")
        self.assertNotIn("Material", by_group["G30"])

    @staticmethod
    def tables(configs):
        parsed = [{"sheet_name": "Family", "parser_type": "HIERARCHICAL",
                   "result": {"Assembly": "42C5495", "Configurations": configs}}]
        with patch("Classifier.WorkbookAssemblyParser.parse", return_value=parsed):
            return QA2.ClassifierWorkbookParser.extract_tables(QA2.Workbook(), "FAMILY.xlsx", "42C5495")

    def test_only_viable_alternatives_require_missing_inputs(self):
        tables = self.tables([
            {"Group No.": "G01", "Structure Width": "36", "Material": "Copper"},
            {"Group No.": "G02", "Structure Width": "45", "Feed": "REVERSE"},
        ])
        group, _ = QA2.ClassifierWorkbookParser.identify_group(tables, "42C5495", {"Width": 36, "Material": "CU"})
        self.assertEqual(group, "G01")

    def test_invariant_inputs_and_supplied_conflicts(self):
        tables = self.tables([
            {"Group No.": "G01", "Structure Width": "36", "Material": "Copper"},
            {"Group No.": "G02", "Structure Width": "45", "Material": "Copper"},
        ])
        select = QA2.ClassifierWorkbookParser.identify_group
        self.assertEqual(select(tables, "42C5495", {"Width": 36})[0], "G01")
        self.assertEqual(select(tables, "42C5495", {"Width": 36, "Material": "AL"})[0], "")
        for config in ({"Group No.": "G01"}, {"Group No.": "G01", "Unparsed Table Context": "3 WIRE"}):
            self.assertEqual(select(self.tables([config]), "42C5495", {})[0], "")

    def test_cover_table_cannot_win_with_only_width(self):
        tables = self.tables([
            {"Group No.": "G06", "Structure Width": "36", "Assembly Type": "UGPS COVER"},
            {"Group No.": "G46", "Structure Width": "36", "Door Type": "DOUBLE", "Hinge Side": "RIGHT"},
        ])
        group, message = QA2.ClassifierWorkbookParser.identify_group(tables, "42C5495", {"Width": 36})
        self.assertEqual(group, "")
        self.assertIn("missing job inputs", message)

    def test_minimum_range_and_all_match_semantics(self):
        for actual, wanted, expected in (
            ("30D MIN.", "36", True), ("30 MIN", "24", False),
            ("2500-3000A", "2800A", True), ("2500-3000", "3200", False),
            ("2500-3000", "2500-2800", True), ("2500", "2500-3000", False),
            ("ALL", "36", True), ("ALL", "", False),
            ("Heat Test", "HEAT", True), ("HEAT", "HT", True),
            ("Heat Test", "DEN", False),
        ):
            with self.subTest(actual=actual, wanted=wanted):
                self.assertEqual(QA2.ClassifierWorkbookParser.values_match(actual, wanted), expected)

    def test_browser_and_local_excel_links(self):
        folder = "https://example.com/sites/design/Spreadsheets/PRLC"
        target = "../../../../../:x:/r/sites/design/Test.xlsx?d=123&web=1"
        self.assertEqual(QA2.excel_document_url(folder, target), "https://example.com/sites/design/Test.xlsx")
        self.assertEqual(QA2.excel_document_url(folder, "EUSERC Sheets/Test.xlsx"), folder + "/EUSERC Sheets/Test.xlsx")
        self.assertEqual(QA2.excel_document_url(folder, r"C:\data\Test.xlsx"), r"C:\data\Test.xlsx")
        self.assertEqual(QA2.link_file_name(target), "Test.xlsx")

    def test_drawing_doors_require_exact_unique_section_geometry(self):
        sections = [("1", 32, -144, 36, 90, None), ("2", 68, -144, 36, 90, None)]
        doors = [("Config5_36_Rt", 32, -54, 36, "RIGHT"), ("Config5_36_Lft", 68, -54, 36, "LEFT")]
        result = QA2.DWGFile.section_block_parameters(sections, doors, "outline.dwg")
        self.assertEqual(result["1"].get_parameter("Hinge Side"), "RIGHT")
        self.assertEqual(result["2"].get_parameter("Hinge Side"), "LEFT")
        self.assertEqual(result["1"].get_parameter("Width (Wide)"), "36")
        ambiguous = QA2.DWGFile.section_block_parameters(sections + [("3", 32, -144, 36, 90, None)], doors, "outline.dwg")
        self.assertFalse(ambiguous["1"].get_parameter("Hinge Side"))
        unmatched = QA2.DWGFile.section_block_parameters(sections, [("Config5_38_Rt", 32, -54, 38, "RIGHT")], "outline.dwg")
        self.assertFalse(unmatched["1"].get_parameter("Hinge Side"))

    def test_description_uses_current_not_maximum_capacity(self):
        self.assertNotIn("Amperage", QA2.description_parameter_values("EUSERC 1000A MAX REAR ISO BARR"))
        self.assertEqual(QA2.description_parameter_values("800A CUAG DEN 3PH")["Amperage"], "800")
        selectors = QA2.description_parameter_values("REV FEED DBL DOOR W/1 MTR 38W")
        self.assertEqual((selectors["Feed"], selectors["Door Type"], selectors["Meter Count"]), ("REVERSE", "DOUBLE", "1"))

    def test_assembly_current_overrides_global_current_without_using_expected_group(self):
        repository = QA2.ParameterRepository()
        repository.add_parameter("Amperage", "800", "job.pdf")
        window = SimpleNamespace(
            _parameter_repository=repository,
            parameter_variables=QA2.ExcelParameterVariables.from_sources(["Amperage: 800"]),
        )
        window._resolved_value = lambda *args: QA2.AuditReviewDemo._resolved_value(window, *args)
        tables = self.tables([
            {"Group No.": "G01", "Amperage": "800"},
            {"Group No.": "G02", "Amperage": "1200"},
        ])
        for suffix, status in (("G01", "FAIL"), ("G02", "PASS")):
            result = QA2.AuditReviewDemo.validate_assembly(
                window, "42C5495" + suffix, None, "", "", tables, "", "1200A CU", "7",
            )
            self.assertEqual((result.actual_group, result.status), ("G02", status))

    def test_component_test_description_is_not_a_global_bus_test(self):
        window = SimpleNamespace(
            _parameter_repository=QA2.ParameterRepository(),
            parameter_variables=QA2.ExcelParameterVariables.from_sources(["Heat Test"]),
        )
        window._resolved_value = lambda *args: QA2.AuditReviewDemo._resolved_value(window, *args)
        tables = self.tables([
            {"Group No.": "G01", "Test": "Density"},
            {"Group No.": "G02", "Test": "Heat Test"},
        ])
        result = QA2.AuditReviewDemo.validate_assembly(
            window, "42C5495G01", None, "", "", tables, "", "800A CUAG DEN 3PH", "1",
        )
        self.assertEqual((result.actual_group, result.status), ("G01", "PASS"))

    def test_com_dates_can_be_serialized_without_changing_wall_time(self):
        from datetime import datetime, timezone
        from io import BytesIO
        import pickle
        class ComDate(datetime):
            pass
        value = ComDate(2024, 1, 2, 3, 4, tzinfo=timezone.utc)
        rows = QA2.normalize_excel_rows(((value,),))
        self.assertIs(type(rows[0][0]), datetime)
        self.assertEqual(pickle.loads(pickle.dumps(rows))[0][0], datetime(2024, 1, 2, 3, 4))
        sheet = excel_sheet("Dated", (("42C5495", value),))
        workbook = QA2.ExcelComSession._snapshot_workbook(SimpleNamespace(Worksheets=[sheet]))
        workbook.save(BytesIO())
        self.assertEqual(workbook.active["B1"].value, value.replace(tzinfo=None))
        small = QA2.Workbook()
        self.assertTrue(QA2.ExcelComSession()._copy_small_sheet(sheet, small))
        small.save(BytesIO())
        self.assertEqual(small["Dated"]["B1"].value, value.replace(tzinfo=None))

    def test_section_values_and_aliases_resolve_before_global_values(self):
        global_values = QA2.ParameterRepository()
        global_values.add_parameter("Phase", "3", "job.pdf")
        section = QA2.ParameterRepository()
        section.add_parameter("Hinge Side", "RIGHT", "outline.dwg")
        window = SimpleNamespace(
            _parameter_repository=global_values, _section_parameters={"1": section},
            parameter_variables=QA2.ExcelParameterVariables.from_sources([]),
        )
        resolve = QA2.AuditReviewDemo._resolved_value
        self.assertEqual(resolve(window, "42C5495G01", "3/1 Phase", "", "", "1"), "3")
        self.assertEqual(resolve(window, "42C5495G01", "Hinge Side", "", "", "1"), "RIGHT")
        section.add_parameter("Hinge Side", "LEFT", "other.dwg")
        self.assertEqual(resolve(window, "42C5495G01", "Hinge Side", "", "", "1"), "")


class ClassifierIntegrationTests(unittest.TestCase):
    def setUp(self):
        capture = patch.object(QA2.ExcelComSession, "_capture_family", QA2.ExcelComSession._capture_live_family)
        capture.start()
        self.addCleanup(capture.stop)

    @staticmethod
    def matrix_workbook():
        workbook = QA2.Workbook()
        sheet = workbook.active
        sheet.title = "Matrix"
        sheet.append(["Assembly Drawing", "Depth", "Structure Width", None])
        sheet.append([None, None, 30, 36])
        sheet.append(["42C1084", 24, "G01", "G02"])
        sheet.append([None, 36, "G03", "G04"])
        return workbook

    @staticmethod
    def hierarchy_workbook():
        workbook = QA2.Workbook()
        sheet = workbook.active
        sheet.title = "Hierarchy"
        sheet.append(["Assembly Drawing", "Material", "Copper", None])
        sheet.append([None, "Amperage", 2000, 3000])
        sheet.append(["42C1084", None, "G01", "G02"])
        sheet.merge_cells("C1:D1")
        return workbook

    def test_import_has_no_workbook_or_output_side_effects(self):
        import importlib
        import Classifier

        with patch("openpyxl.load_workbook", side_effect=AssertionError("opened on import")), \
                patch("pathlib.Path.mkdir", side_effect=AssertionError("created output on import")):
            importlib.reload(Classifier)
        importlib.reload(Classifier)

    def test_matrix_configurations_map_dimensions_and_groups(self):
        tables = QA2.ClassifierWorkbookParser.extract_tables(self.matrix_workbook(), "FAMILY.xlsx", "42C1084")
        self.assertEqual(len(tables), 4)
        self.assertTrue(all(table.classifier_type == "MATRIX" for table in tables))
        group, message = QA2.ClassifierWorkbookParser.identify_group(
            tables, "42C1084", {"Width (Wide)": "36", "Depth": "24"},
        )
        self.assertEqual(group, "G02")
        self.assertIn("Matrix (MATRIX)", message)
        self.assertEqual(QA2.TableValueExtractor.extract(
            tables, "42C1084", "G02", {"Width (Wide)": ("Width (Wide)",), "Depth": ("Depth",)},
        ), {"Width (Wide)": "36", "Depth": "24"})

    def test_hierarchical_merged_headers_map_numeric_values_and_aliases(self):
        tables = QA2.ClassifierWorkbookParser.extract_tables(self.hierarchy_workbook(), "FAMILY.xlsx", "42C1084")
        self.assertEqual(len(tables), 2)
        group, _message = QA2.ClassifierWorkbookParser.identify_group(
            tables, "42C1084", {"Material": "COPPER", "Ampacity": "3000A"},
        )
        self.assertEqual(group, "G02")
        self.assertEqual(QA2.TableValueExtractor.extract(
            tables, "42C1084", "G01", {"Amperage": ("Amperage",), "Material": ("Material",)},
        ), {"Amperage": "2000", "Material": "Copper"})

    def test_missing_conflicting_and_nonmatching_inputs_never_resolve(self):
        tables = QA2.ClassifierWorkbookParser.extract_tables(self.matrix_workbook(), "FAMILY.xlsx", "42C1084")
        for inputs in (
            {"Width": "36"},
            {"Width": "36", "Depth": "90"},
            {"Width": "36", "Section Width": "30", "Depth": "24"},
        ):
            with self.subTest(inputs=inputs):
                group, message = QA2.ClassifierWorkbookParser.identify_group(tables, "42C1084", inputs)
                self.assertEqual(group, "")
                self.assertIn("Classifier:", message)

    def test_every_active_sheet_is_parsed_without_overwriting_assemblies(self):
        workbook = self.matrix_workbook()
        second = workbook.copy_worksheet(workbook.active)
        second.title = "Second"
        second.cell(3, 4, "G09")
        obsolete = workbook.copy_worksheet(workbook.active)
        obsolete.title = "Obsolete_old"
        tables = QA2.ClassifierWorkbookParser.extract_tables(workbook, "FAMILY.xlsx", "42C1084")
        self.assertEqual({table.sheet_name for table in tables}, {"Matrix", "Second"})
        group, message = QA2.ClassifierWorkbookParser.identify_group(tables, "42C1084", {"Width": 36, "Depth": 24})
        self.assertEqual(group, "")
        self.assertIn("multiple groups", message)

    def test_matrix_metadata_uses_family_column_in_database_index(self):
        from Classifier import GeneralizedMatrixParser

        parser = GeneralizedMatrixParser(
            "CABLEPULL.xlsx", QA2.METADATA_PATH, QA2.SEMANTIC_PATH, workbook=self.matrix_workbook(),
        )
        self.assertIn("Material", parser.metadata)
        self.assertIn("Assembly Drawing", parser.metadata)

    def test_live_snapshot_preserves_values_offsets_and_merges(self):
        sheet = excel_sheet(
            "Live", (("Material", "Copper", None), ("Amperage", 2000, 3000)),
            row=4, column=4, merges=((4, 5, 4, 6),),
        )
        snapshot = QA2.ExcelComSession._snapshot_workbook(SimpleNamespace(Worksheets=[sheet]))
        self.assertEqual(snapshot["Live"]["D4"].value, "Material")
        self.assertEqual(snapshot["Live"]["F5"].value, 3000)
        self.assertIn("E4:F4", {str(merged) for merged in snapshot["Live"].merged_cells.ranges})

    def test_classifier_cache_and_legacy_fallback_are_reset_on_close(self):
        session = QA2.ExcelComSession()
        session._group_tables["FAMILY.XLSX"] = ("legacy",)
        session._source_names["FAMILY.XLSX"] = "FAMILY.xlsx"
        session._classifier_snapshots["FAMILY.XLSX"] = self.matrix_workbook()
        with patch.object(session, "_snapshot_workbook", side_effect=AssertionError("recaptured")):
            tables = session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
            self.assertIs(session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084"), tables)
            self.assertEqual(len(tables), 4)
            self.assertEqual(session.get_group_tables("FAMILY.xlsx", assembly_base="MISSING"), ("legacy",))
        session.close()
        self.assertFalse(session._classifier_tables)
        self.assertFalse(session._classifier_snapshots)
        self.assertFalse(session._source_names)
        self.assertFalse(session._classifier_parsers)

    def test_parser_failure_is_reported_instead_of_using_legacy_tables(self):
        session = QA2.ExcelComSession()
        session._group_tables["FAMILY.XLSX"] = ("legacy",)
        session._source_names["FAMILY.XLSX"] = "FAMILY.xlsx"
        session._classifier_snapshots["FAMILY.XLSX"] = self.matrix_workbook()
        with patch.object(QA2, "METADATA_PATH", "missing_metadata.xlsx"):
            with self.assertRaises(FileNotFoundError):
                session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
        self.assertFalse(session._classifier_tables)

    def test_live_workbook_flows_through_snapshot_classifier_and_mapping(self):
        source = self.matrix_workbook().active
        workbook = SimpleNamespace(Name="FAMILY.xlsx", Worksheets=[SimpleNamespace(
            Name=source.title, UsedRange=SimpleNamespace(
                Row=1, Column=1, MergeCells=False, Value=tuple(source.values),
            ),
        )])
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", return_value=workbook):
            tables = session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
        group, _message = QA2.ClassifierWorkbookParser.identify_group(
            tables, "42C1084", {"Width": 36, "Depth": 24},
        )
        self.assertEqual(group, "G02")
        self.assertEqual(len(tables), 4)
        session.close()

    def test_family_parser_error_produces_explicit_unresolved_validation(self):
        reader = MagicMock(side_effect=ValueError("invalid metadata"))
        window = SimpleNamespace(
            _assembly_links={"42C1084G01": "FAMILY.xlsx"},
            _set_status=lambda text: None,
            _excel_com=SimpleNamespace(get_group_tables=reader),
        )
        tables, error = QA2.AuditReviewDemo._family_tables(window, "42C1084G01")
        result = QA2.AuditReviewDemo.validate_assembly(
            window, "42C1084G01", None, "", "", tables, error,
        )
        self.assertEqual(result.status, "UNRESOLVED")
        self.assertIn("ValueError: invalid metadata", result.message)

    def test_mixed_matrix_and_hierarchy_sheets_keep_separate_parser_selection(self):
        workbook = self.matrix_workbook()
        source = self.hierarchy_workbook().active
        target = workbook.create_sheet(source.title)
        for row in source.iter_rows(values_only=True):
            target.append(row)
        target.merge_cells("C1:D1")
        tables = QA2.ClassifierWorkbookParser.extract_tables(workbook, "FAMILY.xlsx", "42C1084")
        self.assertEqual({(table.sheet_name, table.classifier_type) for table in tables}, {
            ("Matrix", "MATRIX"), ("Hierarchy", "HIERARCHICAL"),
        })


class ExcelPerformanceTests(unittest.TestCase):
    def test_only_excel_busy_hresult_is_retried_with_a_bound(self):
        class ComBusy(Exception):
            hresult = -2147352567
            excepinfo = (0, "", "", "", 0, -2146777998)

        updates = []
        session = QA2.ExcelComSession(status_callback=updates.append)
        operation = MagicMock(side_effect=[ComBusy(), "captured"])
        with patch.object(QA2.time, "sleep"):
            self.assertEqual(session._excel_call("finding root", operation), "captured")
            self.assertEqual(operation.call_count, 2)
            self.assertIn("retry 1/3", updates[0])
            operation = MagicMock(side_effect=ComBusy())
            with self.assertRaises(ComBusy):
                session._excel_call("closing", operation)
            self.assertEqual(operation.call_count, 4)
        operation = MagicMock(side_effect=ValueError("not a busy error"))
        with self.assertRaises(ValueError):
            session._excel_call("finding root", operation)
        self.assertEqual(operation.call_count, 1)

    def setUp(self):
        capture = patch.object(QA2.ExcelComSession, "_capture_family", QA2.ExcelComSession._capture_live_family)
        capture.start()
        self.addCleanup(capture.stop)

    def family(self):
        source = ClassifierIntegrationTests.matrix_workbook().active
        return SimpleNamespace(
            Name="FAMILY.xlsx", Worksheets=[excel_sheet(source.title, tuple(source.values))],
            Close=MagicMock(),
        )

    def test_owned_family_closes_once_and_all_later_operations_use_cached_data(self):
        session = QA2.ExcelComSession()
        index = SimpleNamespace(Name="Index2.xlsx", Close=MagicMock())
        session._index2_wb = index
        workbook = self.family()

        def opened(link):
            session._family_workbooks[link.upper()] = workbook
            return workbook

        with patch.object(session, "get_family_workbook", side_effect=opened) as open_workbook:
            tables = session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
            workbook.Close.assert_called_once_with(False)
            session.lookup_dimensions("FAMILY.xlsx", None, "42C1084", "G02")
            self.assertIs(tables, session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084"))
            self.assertTrue(session.get_sheet_values("FAMILY.xlsx"))
            open_workbook.assert_called_once()
        session.close()
        workbook.Close.assert_called_once_with(False)
        index.Close.assert_not_called()

    def test_preexisting_user_workbook_is_not_closed(self):
        session = QA2.ExcelComSession()
        workbook = self.family()
        with patch.object(session, "get_family_workbook", return_value=workbook):
            session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
        session.close()
        workbook.Close.assert_not_called()

    def test_capture_failure_still_closes_owned_workbook_and_preserves_index2(self):
        session = QA2.ExcelComSession()
        workbook = self.family()
        index = SimpleNamespace(Name="Index2.xlsx", Close=MagicMock())
        session._index2_wb = index
        session._family_workbooks["FAMILY.XLSX"] = workbook
        with patch.object(session, "get_family_workbook", return_value=workbook), \
                patch.object(session, "_snapshot_workbook", side_effect=RuntimeError("read failed")):
            with self.assertRaisesRegex(RuntimeError, "read failed"):
                session.get_group_tables("FAMILY.xlsx")
        workbook.Close.assert_called_once_with(False)
        index.Close.assert_not_called()
        self.assertNotIn("FAMILY.XLSX", session._group_tables)

    def test_close_failure_is_reported_and_can_be_retried(self):
        messages = []
        session = QA2.ExcelComSession(status_callback=messages.append)
        workbook = self.family()
        workbook.Close.side_effect = RuntimeError("Excel is busy")
        session._family_workbooks["FAMILY.XLSX"] = workbook
        with self.assertRaisesRegex(RuntimeError, "Excel is busy"):
            session._close_owned_family("FAMILY.xlsx")
        self.assertIn("FAMILY.XLSX", session._family_workbooks)
        self.assertTrue(any("Could not close FAMILY.xlsx" in message for message in messages))
        workbook.Close.side_effect = None
        session.close()
        self.assertEqual(workbook.Close.call_count, 2)

    def test_failed_family_is_not_reopened_for_every_assembly_and_next_run_can_retry(self):
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", side_effect=RuntimeError("unavailable")) as opened:
            for base in ("42C1084", "42C1085"):
                with self.assertRaisesRegex(RuntimeError, "unavailable"):
                    session.get_group_tables("FAMILY.xlsx", assembly_base=base)
            opened.assert_called_once()
            session.close()
            with self.assertRaisesRegex(RuntimeError, "unavailable"):
                session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
            self.assertEqual(opened.call_count, 2)

    def test_retired_workbook_closes_before_replacement_opens(self):
        session = QA2.ExcelComSession()
        old = SimpleNamespace(
            Name="OLD.xlsx", Close=MagicMock(),
            Worksheets=[excel_sheet("Notice", (("Will be maintained on spreadsheet: FAMILY.xls (Index 1)",),))],
        )
        current = self.family()
        index = SimpleNamespace(
            Name="Index2.xlsx", Close=MagicMock(),
            Worksheets=lambda number: SimpleNamespace(Hyperlinks=[SimpleNamespace(
                Range=SimpleNamespace(Column=5, Value="FAMILY"), Address="FAMILY.xlsx",
            )]),
        )

        def opened(link):
            if link == "FAMILY.xlsx":
                old.Close.assert_called_once_with(False)
            workbook = old if link == "OLD.xlsx" else current
            session._family_workbooks[link.upper()] = workbook
            return workbook

        with patch.object(session, "get_family_workbook", side_effect=opened), \
                patch.object(session, "_get_index2_workbook", return_value=index):
            tables = session.get_group_tables("OLD.xlsx", assembly_base="42C1084")
        self.assertTrue(tables)
        current.Close.assert_called_once_with(False)
        index.Close.assert_not_called()
        self.assertEqual(session._source_names["OLD.XLSX"], "FAMILY.xlsx")

    def test_each_live_sheet_value_is_read_only_once(self):
        workbook = self.family()
        used = MagicMock(Row=1, Column=1, MergeCells=False)
        values = PropertyMock(return_value=workbook.Worksheets[0].UsedRange.Value)
        type(used).Value = values
        workbook.Worksheets[0].UsedRange = used
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", return_value=workbook):
            session.get_sheet_values("FAMILY.xlsx")
            session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
            session.get_sheet_values("FAMILY.xlsx")
        values.assert_called_once()

    def test_merge_capture_uses_block_probes_not_one_com_call_per_populated_cell(self):
        rows = tuple(tuple(None if (r, c) == (0, 1) else "value" for c in range(20)) for r in range(100))
        sheet = excel_sheet("Sparse merges", rows, merges=((1, 1, 1, 2),))
        snapshot = QA2.ExcelComSession._snapshot_workbook(SimpleNamespace(Worksheets=[sheet]))
        self.assertIn("A1:B1", {str(area) for area in snapshot.active.merged_cells.ranges})
        self.assertLessEqual(sheet.Range.call_count, 40)
        sheet.Cells.assert_called_once_with(1, 1)

    def test_adjacent_merged_areas_are_preserved(self):
        sheet = excel_sheet(
            "Adjacent", (("Copper", None, "Aluminum", None),),
            merges=((1, 1, 1, 2), (1, 3, 1, 4)),
        )
        snapshot = QA2.ExcelComSession._snapshot_workbook(SimpleNamespace(Worksheets=[sheet]))
        self.assertEqual({str(area) for area in snapshot.active.merged_cells.ranges}, {"A1:B1", "C1:D1"})

    def test_matrix_is_prepared_once_for_multiple_assemblies(self):
        from Classifier import GeneralizedMatrixParser, WorkbookAssemblyParser

        workbook = ClassifierIntegrationTests.matrix_workbook()
        workbook.active.append(["42C1085", 48, "G05", "G06"])
        parser = WorkbookAssemblyParser("FAMILY.xlsx", QA2.METADATA_PATH, QA2.SEMANTIC_PATH, workbook)
        original = GeneralizedMatrixParser.parse_matrix
        with patch.object(GeneralizedMatrixParser, "parse_matrix", autospec=True, side_effect=original) as parse:
            first = parser.parse("42C1084")
            second = parser.parse("42C1085")
            self.assertEqual(first[0]["result"]["Assembly"], "42C1084")
            self.assertEqual(second[0]["result"]["Assembly"], "42C1085")
            parse.assert_called_once()

    def test_basicstr_dimensions_are_read_from_snapshot_without_reopening_excel(self):
        source = QA2.Workbook()
        sheet = source.active
        sheet.title = "BASICSTR-NEMA 1"
        sheet.append([None, None, None, 24, 36])
        sheet.append([None, 30, "42C1084", "G01", "G02"])
        sheet.append([None, 36, None, "G03", "G04"])
        sheet.merge_cells("C2:C3")
        workbook = SimpleNamespace(
            Name="BASICSTR.xlsx", Close=MagicMock(),
            Worksheets=[excel_sheet(sheet.title, tuple(sheet.values), merges=((2, 3, 3, 3),))],
        )
        session = QA2.ExcelComSession()
        session._family_workbooks["BASICSTR.XLSX"] = workbook
        with patch.object(session, "get_family_workbook", return_value=workbook) as opened:
            first = session.lookup_dimensions("BASICSTR.xlsx", sheet.title, "42C1084", "G02")
            second = session.lookup_dimensions("BASICSTR.xlsx", sheet.title, "42C1084", "G03")
        self.assertEqual((first.width, first.depth), ("30", "36"))
        self.assertEqual((second.width, second.depth), ("36", "24"))
        opened.assert_called_once()
        workbook.Close.assert_called_once_with(False)


class ExcelCaptureTimeoutTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.output = os.path.join(self.directory.name, "capture.pickle")
        self.progress = os.path.join(self.directory.name, "progress.log")
        self.errors = os.path.join(self.directory.name, "errors.log")
        with open(self.progress, "w", encoding="utf-8"):
            pass

    def run_reader(self, script, timeout=5, callback=None):
        from excel_capture import run_capture_command

        return run_capture_command(
            [sys.executable, "-B", "-c", script, self.output, self.progress],
            self.output, self.progress, self.errors, timeout, callback,
        )

    def test_real_blocked_reader_is_terminated_within_deadline_without_killing_excel(self):
        processes = []
        original = subprocess.Popen

        def started(*args, **kwargs):
            process = original(*args, **kwargs)
            processes.append(process)
            return process

        messages = []
        begin = time.monotonic()
        with patch("excel_capture.subprocess.Popen", side_effect=started):
            with self.assertRaisesRegex(TimeoutError, "Excel and Index2 were not terminated"):
                self.run_reader("import time; time.sleep(60)", timeout=1.1, callback=messages.append)
        self.assertLess(time.monotonic() - begin, 4)
        self.assertEqual(len(processes), 1)
        self.assertIsNotNone(processes[0].poll())
        self.assertTrue(any("1.1s limit" in message for message in messages))

    def test_worker_progress_and_result_are_delivered(self):
        messages = []
        result = self.run_reader(
            "import pickle,sys,time; "
            "open(sys.argv[2],'a',encoding='utf-8').write('Reading fixture\\n'); "
            "time.sleep(0.1); "
            "pickle.dump(('FAMILY.xlsx', (), b'snapshot'), open(sys.argv[1],'wb'))",
            callback=messages.append,
        )
        self.assertEqual(result, ("FAMILY.xlsx", (), b"snapshot"))
        self.assertIn("Reading fixture", messages)

    def test_worker_errors_are_reported_instead_of_returning_empty_data(self):
        with self.assertRaisesRegex(RuntimeError, "Unable to read fixture"):
            self.run_reader("raise RuntimeError('Unable to read fixture')")

    def test_worker_exiting_without_output_is_an_error(self):
        with self.assertRaisesRegex(RuntimeError, "without captured workbook data"):
            self.run_reader("pass")

    def test_capture_payload_is_loaded_and_cached_without_parent_com_calls(self):
        from io import BytesIO

        buffer = BytesIO()
        workbook = ClassifierIntegrationTests.matrix_workbook()
        workbook.save(buffer)
        payload = ("FAMILY.xlsx", (("Matrix", tuple(workbook.active.values)),), buffer.getvalue())
        session = QA2.ExcelComSession()
        session._index2_folder = "https://example.invalid/handbook"
        with patch("excel_capture.capture_family", return_value=payload) as capture, \
                patch.object(session, "_get_app", side_effect=AssertionError("COM on UI thread")):
            tables = session.get_group_tables("FAMILY.xlsx", assembly_base="42C1084")
            group, _message = QA2.ClassifierWorkbookParser.identify_group(
                tables, "42C1084", {"Width": 36, "Depth": 24},
            )
            self.assertEqual(group, "G02")
            session.get_sheet_values("FAMILY.xlsx")
            capture.assert_called_once_with(
                "https://example.invalid/handbook/FAMILY.xlsx", 45, session._report,
                assemblies=("42C1084",),
            )
        session.close()

    def test_timeout_skips_further_uncached_reads_but_keeps_cached_data_and_allows_next_run(self):
        session = QA2.ExcelComSession()
        session._index2_folder = "https://example.invalid/handbook"
        session._group_tables["CACHED.XLSX"] = ("cached table",)
        with patch("excel_capture.capture_family", side_effect=TimeoutError("Excel is blocked")) as capture, \
                patch.object(session, "_get_app", side_effect=AssertionError("COM on UI thread")):
            with self.assertRaises(TimeoutError):
                session.get_group_tables("PRINGLE SWTCH.xls")
            with self.assertRaisesRegex(RuntimeError, "uncached family reads were skipped"):
                session.get_group_tables("OTHER.xlsx")
            self.assertEqual(session.get_group_tables("CACHED.xlsx"), ("cached table",))
            capture.assert_called_once()
            session.close()
            with self.assertRaises(TimeoutError):
                session.get_group_tables("https://example.invalid/handbook/PRINGLE SWTCH.xls")
            self.assertEqual(capture.call_count, 2)

    def test_index2_folder_and_replacement_links_are_cached_before_family_operations(self):
        sheet = SimpleNamespace(Hyperlinks=[SimpleNamespace(
            Range=SimpleNamespace(Column=5, Value="FAMILY"), Address="FAMILY.xls",
        )])
        workbook = SimpleNamespace(
            Path="https://example.invalid/handbook", Sheets=lambda number: sheet, Worksheets=lambda number: sheet,
        )
        session = QA2.ExcelComSession()
        with patch.object(session, "_get_index2_workbook", return_value=workbook), \
                patch.object(QA2, "build_index2_map_from_com", return_value={"42C1084": "FAMILY.xls"}):
            self.assertEqual(session.get_index2_map(), {"42C1084": "FAMILY.xls"})
        with patch.object(session, "_get_index2_workbook", side_effect=AssertionError("Index2 reread")):
            self.assertEqual(session._replacement_link("FAMILY.xlsx", "retired"), "FAMILY.xls")
        self.assertEqual(session._index2_folder, workbook.Path)

    @staticmethod
    def checklist_window():
        counts = []
        window = SimpleNamespace(
            bom_list=MagicMock(), checkbox_items={}, _code_mappings={},
            _configure_columns=MagicMock(), add_section=MagicMock(), _set_status=MagicMock(),
            _validation_results={}, _validation_checks={}, _validation_scope_result=lambda style: None,
        )
        window.update_progress = lambda: counts.append(len(window.checkbox_items))
        window.add_checkbox = lambda section, key, *args: window.checkbox_items.update({key: object()})
        return window, counts

    def test_checklist_counts_update_before_the_next_workbook_read(self):
        window, counts = self.checklist_window()

        def dimensions(style):
            if style == "42C1085G01":
                self.assertIn(1, counts)
            return "", ""

        window.lookup_dimensions = dimensions
        sections = [("1", "Test", [("42C1084G01", "", "", ""), ("42C1085G01", "", "", "")])]
        with patch.object(QA2.time, "monotonic", return_value=0):
            QA2.AuditReviewDemo.build_bom_checklist(window, sections)
        self.assertEqual(counts[-1], 2)
        self.assertEqual(window.bom_list.blockSignals.call_args_list[-1].args, (False,))

    def test_checklist_signals_are_restored_when_a_read_fails(self):
        window, _counts = self.checklist_window()
        window.lookup_dimensions = MagicMock(side_effect=RuntimeError("fixture failed"))
        with self.assertRaisesRegex(RuntimeError, "fixture failed"):
            QA2.AuditReviewDemo.build_bom_checklist(window, [("1", "Test", [("42C1084G01", "", "", "")])])
        self.assertEqual(window.bom_list.blockSignals.call_args_list[-1].args, (False,))


class RootTableCaptureTests(unittest.TestCase):
    @staticmethod
    def sheet(name, values, regions, merges=()):
        from openpyxl.utils.cell import range_boundaries

        sheet = MagicMock(Name=name, Rows=SimpleNamespace(Count=65536), Columns=SimpleNamespace(Count=256))
        used = SimpleNamespace(Rows=SimpleNamespace(Count=100000), Columns=SimpleNamespace(Count=256))
        sheet.UsedRange = used
        sheet.value_reads = []
        merged = {(r, c): bounds for bounds in merges
                  for r in range(bounds[0], bounds[2] + 1) for c in range(bounds[1], bounds[3] + 1)}

        def inside(coordinate, bounds):
            return bounds[0] <= coordinate[0] <= bounds[2] and bounds[1] <= coordinate[1] <= bounds[3]

        def find(bounds, What, After=None, **_options):
            positions = sorted(
                coordinate for coordinate, value in values.items()
                if inside(coordinate, bounds) and value is not None
                and (What == "*" or str(What).upper() in str(value).upper())
            )
            if not positions:
                return None
            after = (After.Row, After.Column) if After is not None else (0, 0)
            coordinate = next((position for position in positions if position > after), positions[0])
            return cell(*coordinate)

        def region(bounds):
            top, left, bottom, right = bounds
            covered = sum(inside(coordinate, bounds) for coordinate in merged)
            count = (bottom - top + 1) * (right - left + 1)
            result = MagicMock(
                Row=top, Column=left, Rows=SimpleNamespace(Count=bottom - top + 1),
                Columns=SimpleNamespace(Count=right - left + 1),
                MergeCells=(True if covered == count else None) if covered else False,
            )

            def read():
                sheet.value_reads.append(bounds)
                return tuple(tuple(values.get((r, c)) for c in range(left, right + 1)) for r in range(top, bottom + 1))

            type(result).Value = PropertyMock(side_effect=read)
            result.Find.side_effect = lambda **options: find(bounds, **options)
            return result

        def cell(row, column):
            coordinate = row, column
            current = next((bounds for bounds in regions if inside(coordinate, bounds)), (row, column, row, column))
            return SimpleNamespace(
                Row=row, Column=column, Value=values.get(coordinate),
                CurrentRegion=region(current), MergeArea=region(merged.get(coordinate, (row, column, row, column))),
            )

        def address(text):
            left, top, right, bottom = range_boundaries(text)
            return region((top, left, bottom, right))

        sheet.Cells.side_effect = cell
        sheet.Cells.Find.side_effect = lambda **options: find((1, 1, 65536, 256), **options)
        sheet.Range.side_effect = address
        return sheet

    @staticmethod
    def matrix_values(start=100):
        source = ClassifierIntegrationTests.matrix_workbook().active
        return {(r + start, c): cell.value for r, row in enumerate(source) for c, cell in enumerate(row, start=1)
                if cell.value is not None}

    def test_only_root_table_is_read_without_group_filter_or_full_sheet_scanning(self):
        wanted = self.sheet("Wanted", self.matrix_values(), ((100, 1, 103, 4),))
        unrelated = self.sheet("Unrelated", {(1, 1): "42C10840", (2, 1): "G99"}, ((1, 1, 2, 1),))
        obsolete = self.sheet("Obsolete_old", {(1, 1): "42C1084"}, ((1, 1, 1, 1),))
        workbook = SimpleNamespace(Name="FAMILY.xls", Worksheets=[wanted, unrelated, obsolete])
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", return_value=workbook):
            _name, sheets, snapshot = session._capture_live_family("FAMILY.xls", ("42C1084G99",))
        self.assertEqual([name for name, _rows in sheets], ["Wanted"])
        self.assertLess(snapshot.active.max_row, 20)
        self.assertEqual(len(wanted.value_reads), 1)
        self.assertFalse(unrelated.value_reads)
        obsolete.Cells.Find.assert_not_called()
        self.assertEqual({call.kwargs["What"] for call in wanted.Cells.Find.call_args_list}, {"42C1084"})
        tables = QA2.ClassifierWorkbookParser.extract_tables(snapshot, "FAMILY.xls", "42C1084")
        group, _message = QA2.ClassifierWorkbookParser.identify_group(tables, "42C1084", {"Width": 36, "Depth": 24})
        self.assertEqual(group, "G02")

    def test_all_job_roots_in_the_same_family_are_batched_without_duplicate_group_requests(self):
        session = QA2.ExcelComSession()
        sections = [("1", "", [
            ("42C1084G01", "", "", ""), ("42C1084G02", "", "", ""), ("42C1085G03", "", "", ""),
        ])]
        session.prepare_assembly_targets(sections, {"42C1084": "FAMILY.xls", "42C1085": "FAMILY.xls"})
        self.assertEqual(session._assembly_targets["FAMILY.XLS"], {"42C1084", "42C1085"})
        values = self.matrix_values()
        values[(103, 1)] = "42C1085"
        sheet = self.sheet("Wanted", values, ((100, 1, 103, 4),))
        workbook = SimpleNamespace(Name="FAMILY.xls", Worksheets=[sheet])
        with patch.object(session, "get_family_workbook", return_value=workbook):
            _name, _rows, snapshot = session._capture_live_family("FAMILY.xls", ("42C1084", "42C1085"))
        self.assertEqual(len(sheet.value_reads), 1)
        self.assertTrue(QA2.ClassifierWorkbookParser.extract_tables(snapshot, "FAMILY.xls", "42C1085"))

    def test_isolated_root_finds_the_following_table_and_supports_legacy_xls_column_limits(self):
        values = {(1, 1): "42C1084", (4, 1): "Depth", (4, 2): "Group", (5, 1): 24, (5, 2): "G01"}
        sheet = self.sheet("External", values, ((1, 1, 1, 1), (4, 1, 5, 2)))
        session = QA2.ExcelComSession()
        snapshot = session._capture_root_tables(SimpleNamespace(Name="FAMILY.xls", Worksheets=[sheet]), ("42C1084",))
        self.assertEqual(snapshot.active["B5"].value, "G01")
        self.assertTrue(any("IV" in call.args[0] for call in sheet.Range.call_args_list))
        self.assertFalse(any("XFD" in call.args[0] for call in sheet.Range.call_args_list))

    def test_selected_merged_ranges_survive_row_compaction(self):
        values = self.matrix_values()
        sheet = self.sheet("Merged", values, ((100, 1, 103, 4),), merges=((102, 1, 103, 1),))
        session = QA2.ExcelComSession()
        snapshot = session._capture_root_tables(SimpleNamespace(Name="FAMILY.xls", Worksheets=[sheet]), ("42C1084",))
        first_row = QA2.CONTEXT_LOOKBACK + 3
        self.assertIn(f"A{first_row}:A{first_row + 1}", {str(area) for area in snapshot.active.merged_cells.ranges})

    def test_missing_root_is_reported_explicitly(self):
        sheet = self.sheet("Different", {(1, 1): "42C10840"}, ((1, 1, 1, 1),))
        session = QA2.ExcelComSession()
        with self.assertRaisesRegex(LookupError, "42C1084.*not found"):
            session._capture_root_tables(SimpleNamespace(Name="FAMILY.xls", Worksheets=[sheet]), ("42C1084",))

    def test_full_style_root_can_be_captured(self):
        values = self.matrix_values()
        values[(102, 1)] = "42C1084G01"
        sheet = self.sheet("Family", values, ((100, 1, 103, 4),))
        snapshot = QA2.ExcelComSession()._capture_root_tables(
            SimpleNamespace(Name="Family.xlsx", Worksheets=[sheet]), ("42C1084",),
        )
        self.assertTrue(any(cell.value == "42C1084G01" for row in snapshot.active for cell in row))

    def test_small_sheet_retains_original_coordinates_and_adjacent_tables(self):
        values = self.matrix_values()
        values[(110, 8)] = "6552C52"
        sheet = self.sheet("Family", values, ((100, 1, 103, 4), (110, 8, 110, 8)))
        sheet.UsedRange = SimpleNamespace(
            Row=100, Column=1, Rows=SimpleNamespace(Count=11), Columns=SimpleNamespace(Count=8),
            MergeCells=False,
            Value=tuple(tuple(values.get((r, c)) for c in range(1, 9)) for r in range(100, 111)),
        )
        snapshot = QA2.ExcelComSession()._capture_root_tables(
            SimpleNamespace(Name="Family.xlsx", Worksheets=[sheet]), ("42C1084",),
        )
        self.assertEqual(snapshot.active["A102"].value, "42C1084")
        self.assertEqual(snapshot.active["H110"].value, "6552C52")
        self.assertFalse(sheet.value_reads)

    def test_root_numbers_are_not_mistaken_for_group_suffixes(self):
        for root in ("42C1084", "6575C88", "43C1696", "1A85142"):
            self.assertEqual(QA2.split_style_number(root), (root, ""))
            self.assertEqual(QA2.split_style_number(root + "G01"), (root, "G01"))


class ManualMappingMetadataTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = os.path.join(self.directory.name, "Metadata.xlsx")
        workbook = QA2.Workbook()
        checklist = workbook.active
        checklist.title = "MCOE Checklist"
        for number, (code, flag) in enumerate((("MANUAL", True), ("AUTO", False), ("BLANK", None), ("TEXT", "TRUE")), start=2):
            checklist.cell(number, 1, "Category")
            checklist.cell(number, 2, code)
            checklist.cell(number, 3, "Amperage")
            checklist.cell(number, 14, flag)
        database = workbook.create_sheet("Database_Index2")
        database.append(["Parameter", "Data Point1"])
        database.append(["AUTO", "Ampacity"])
        database.append(["MANUAL", "AMP"])
        workbook.save(self.path)
        workbook.close()

    def test_column_n_is_a_manual_mapping_flag_not_a_filter_or_pass_result(self):
        repository = QA2.ConfigurationRepository(self.path)
        mappings = repository.code_mappings()
        self.assertTrue(mappings["MANUAL"].manually_mapped)
        self.assertTrue(mappings["TEXT"].manually_mapped)
        self.assertFalse(mappings["AUTO"].manually_mapped)
        self.assertFalse(mappings["BLANK"].manually_mapped)
        self.assertEqual(mappings["AUTO"].aliases, ("Ampacity",))
        self.assertEqual(mappings["MANUAL"].aliases, ("AMP",))
        self.assertEqual(mappings["MANUAL"].checklist_row, 2)

    def test_configuration_loading_never_performs_individual_read_only_cell_lookups(self):
        with patch("openpyxl.worksheet._read_only.ReadOnlyWorksheet.cell", side_effect=AssertionError("Slow cell lookup")):
            repository = QA2.ConfigurationRepository(self.path)
        self.assertEqual(len(repository.code_mappings()), 4)

    def test_manual_and_automatic_checklist_families_both_reach_the_app_mapping(self):
        repository = QA2.ConfigurationRepository(self.path)
        window = SimpleNamespace(
            _set_status=MagicMock(), _configuration_repository=repository, _common_parameters=(),
        )
        mappings, _note = QA2.AuditReviewDemo._load_code_mappings(window)
        self.assertTrue(mappings["MANUAL"].manually_mapped)
        self.assertFalse(mappings["AUTO"].manually_mapped)
        self.assertIn("Ampacity", mappings["AUTO"].aliases)
        self.assertIn("Amperage", window._common_parameters)
        repository.use_code_mappings(mappings)
        self.assertFalse(repository.get_configuration("AUTO").manually_mapped)
        self.assertTrue(repository.get_configuration("MANUAL").manually_mapped)


class Index2CatalogTests(unittest.TestCase):
    def test_inventory_includes_unlinked_roots_and_preserves_source_rows(self):
        sheet = SimpleNamespace(
            UsedRange=SimpleNamespace(Row=10, Column=3, Columns=SimpleNamespace(Count=9), Value=(
                (None, None, "BASICSTR", None, None, None, None, None, "43C1696, 43C1696G50, 43C1673"),
                (None, None, "OTHER", None, None, None, None, None, "43C1696"),
                (None, None, "NO LINK", None, None, None, None, None, "6575C01"),
                (None, None, "HEADER", None, None, None, None, None, "Assembly drawing numbers"),
            )),
            Hyperlinks=[
                SimpleNamespace(Range=SimpleNamespace(Row=10, Column=5), Address="BASICSTR.xlsx"),
                SimpleNamespace(Range=SimpleNamespace(Row=11, Column=5), Address="OTHER.xlsx"),
            ],
        )
        entries = QA2.read_index2_root_catalog(sheet)
        self.assertEqual({e.root for e in entries}, {"43C1696", "43C1673", "6575C01"})
        self.assertEqual(len(entries), 4)
        self.assertEqual({e.row for e in entries if e.root == "43C1696"}, {10, 11})
        self.assertEqual(next(e for e in entries if e.root == "6575C01").link, "")
        self.assertEqual(entries[0].family, "BASICSTR")

    def test_catalog_cache_refresh_does_not_open_family_workbooks(self):
        session = QA2.ExcelComSession()
        workbook = SimpleNamespace(Sheets=lambda number: "index sheet")
        entries = (QA2.Index2RootEntry("43C1696", 10, "BASICSTR", "BASICSTR.xlsx"),)
        with patch.object(session, "_get_index2_workbook", return_value=workbook), \
                patch.object(QA2, "read_index2_root_catalog", return_value=entries) as reader, \
                patch.object(session, "get_family_workbook", side_effect=AssertionError("opened family")):
            self.assertEqual(session.get_index2_catalog(), entries)
            self.assertEqual(session.get_index2_catalog(), entries)
            self.assertEqual(reader.call_count, 1)
            session.get_index2_catalog(refresh=True)
            self.assertEqual(reader.call_count, 2)
        session.close()
        self.assertIsNone(session._index2_catalog)


class AuditWorkspaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        cls.application = QA2.QApplication.instance() or QA2.QApplication([])

    def setUp(self):
        self.window = QA2.AuditReviewDemo()
        self.window._coverage_catalog = (
            QA2.Index2RootEntry("6575C88", 1, "TOPCOVER", "TOPCOVER.xlsx"),
            QA2.Index2RootEntry("42C2066", 2, "REARCOVER", "REARCOVER.xlsx"),
        )
        self.window._set_index2_job_sections([
            ("1", "", [("6575C88G16", "", "", ""), ("42C2066G31", "", "", "")]),
        ])
        section = self.window.add_section("Structure 1: Main device")
        for style, status in (("6575C88G16", "PASS"), ("42C2066G31", "UNRESOLVED")):
            item = QA2.QTreeWidgetItem(section, [
                style, "Test assembly", "Cover", "TOPCOVER", "Width = 30", "", "", "", "",
                "G16", "G16" if status == "PASS" else "", status,
            ])
            item.setToolTip(11, "Validation evidence")
            item.setCheckState(0, QA2.Qt.CheckState.Unchecked)
            self.window.checkbox_items[style] = item
            self.window._validation_checks[style] = [status]
        self.window.update_progress()

    def tearDown(self):
        self.window.close()

    def test_navigation_has_only_checklist_and_testing(self):
        window = self.window
        self.assertEqual(window.main_tabs.count(), 2)
        self.assertEqual([window.main_tabs.tabText(i) for i in range(2)], ["Checklist", "Testing"])
        self.assertEqual(len(window.navigation_group.buttons()), 2)
        self.assertFalse(hasattr(window, "coverage_tree"))
        for index in (1, 0):
            window.navigation_group.button(index).click()
            self.assertEqual(window.main_tabs.currentIndex(), index)
            self.assertTrue(window.navigation_group.button(index).isChecked())

    def test_index2_inventory_keeps_checklist_counts_without_coverage_tab(self):
        window = self.window
        window._coverage_catalog = (
            QA2.Index2RootEntry("43C1696", 10, "BASICSTR", "BASICSTR.xlsx"),
            QA2.Index2RootEntry("43C1696", 11, "BASICSTR", "BASICSTR.xlsx"),
            QA2.Index2RootEntry("43C1696", 14, "BASICSTR", ""),
            QA2.Index2RootEntry("6575C01", 12, "NO LINK", ""),
            QA2.Index2RootEntry("43C1673", 13, "REARCOVER", "REARCOVER.xlsx"),
        )
        window._set_index2_job_sections([
            ("1", "", [("43C1696G50", "", "", ""), ("43C1696G52", "", "", ""),
                       ("6575C01G16", "", "", ""), ("8805C17G26", "", "", ""), ("30-47592", "", "", "")]),
            ("2", "", [("43C1696G50", "", "", "")]),
        ])
        self.assertEqual(len(window._coverage_indexed_roots), 3)
        self.assertEqual(len(window._coverage_job_roots), 3)
        self.assertEqual(window._coverage_indexed_roots - window._coverage_linked_roots, {"6575C01"})
        self.assertEqual(window._index2_root_status("43C1696"), "Present with link")
        self.assertEqual(window._index2_root_status("6575C01"), "Listed without link")
        self.assertEqual(window._index2_root_status("8805C17"), "Not present in Index2")
        self.assertEqual(window.summary_labels["all_assemblies"].text(), "5")
        self.assertEqual(window.summary_labels["gh_assemblies"].text(), "4")
        self.assertEqual(window.summary_labels["present"].text(), "3")
        self.assertEqual(window.summary_labels["missing"].text(), "1")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "75.0%")

    def test_checklist_suffix_and_presence_filters_use_unique_full_assemblies(self):
        window = self.window
        styles = ["6575C88G16", "6575C88G17", "42C2066G31", "8805C17G26", "30-47592"]
        section = window.add_section("Structure 2")
        template = window.checkbox_items["6575C88G16"]
        for style in (styles[0], styles[1], styles[3], styles[4]):
            cells = [template.text(c) for c in range(template.columnCount())]
            cells[0] = style
            item = QA2.QTreeWidgetItem(section, cells)
            window.checkbox_items[f"STRUCT2::{style}"] = item
        window._coverage_catalog = (
            QA2.Index2RootEntry("6575C88", 10, "TOPCOVER", ""),
            QA2.Index2RootEntry("6575C88", 11, "TOPCOVER", "TOPCOVER.xlsx"),
            QA2.Index2RootEntry("42C2066", 12, "NO LINK", ""),
            QA2.Index2RootEntry("43C1696", 13, "NOT IN JOB", "BASICSTR.xlsx"),
        )
        window._set_index2_job_sections([
            ("1", "", [(style, "", "", "") for style in styles]),
            ("2", "", [(styles[0], "", "", "")]),
        ])
        window.update_progress()
        self.assertEqual(window.summary_labels["all_assemblies"].text(), "5")
        self.assertEqual(window.summary_labels["gh_assemblies"].text(), "4")
        self.assertEqual(window.summary_labels["present"].text(), "3")
        self.assertEqual(window.summary_labels["missing"].text(), "1")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "75.0%")
        self.assertIn("1 unique assemblies without a G/H suffix", window.summary_labels["coverage_rate"].toolTip())
        self.assertEqual(window.visible_rows.text(), "6 of 6 assembly rows")
        self.assertIn("Filtered unique assemblies: 5 | G/H: 4 | Without G/H: 1", window.audit_coverage_summary.text())
        window.audit_suffix_filter.setCurrentText("G/H assemblies")
        self.assertEqual(window.visible_rows.text(), "5 of 6 assembly rows")
        window.audit_coverage_filter.setCurrentText("Present in Index2")
        self.assertEqual(window.visible_rows.text(), "4 of 6 assembly rows")
        self.assertIn("Filtered unique assemblies: 3", window.audit_coverage_summary.text())
        window.audit_suffix_filter.setCurrentText("Without G/H suffix")
        self.assertEqual(window.visible_rows.text(), "0 of 6 assembly rows")
        window.audit_coverage_filter.setCurrentText("All Index2 coverage")
        self.assertEqual(window.visible_rows.text(), "1 of 6 assembly rows")
        self.assertIn("G/H: 0 | Without G/H: 1", window.audit_coverage_summary.text())
        self.assertIn("Presence rate: --", window.audit_coverage_summary.text())
        window.audit_suffix_filter.setCurrentText("All bold assemblies")
        window.audit_coverage_filter.setCurrentText("Present with link")
        self.assertEqual(window.visible_rows.text(), "3 of 6 assembly rows")
        self.assertIn("Filtered unique assemblies: 2", window.audit_coverage_summary.text())
        self.assertIn("Presence rate: 100.0%", window.audit_coverage_summary.text())
        window.audit_search.setText("6575C88G17")
        self.assertEqual(window.visible_rows.text(), "1 of 6 assembly rows")
        window.status_filter.setCurrentText("FAIL")
        self.assertEqual(window.visible_rows.text(), "0 of 6 assembly rows")
        self.assertIn("Presence rate: --", window.audit_coverage_summary.text())
        window.audit_search.clear()
        window.status_filter.setCurrentText("All statuses")
        window.audit_coverage_filter.setCurrentText("Not linked")
        self.assertEqual(window.visible_rows.text(), "2 of 6 assembly rows")
        self.assertIn("Presence rate: 50.0%", window.audit_coverage_summary.text())
        window.audit_coverage_filter.setCurrentText("Listed without link")
        self.assertEqual(window.visible_rows.text(), "1 of 6 assembly rows")
        window.audit_coverage_filter.setCurrentText("Not present in Index2")
        self.assertEqual(window.visible_rows.text(), "1 of 6 assembly rows")
        window.audit_coverage_filter.setCurrentText("Excluded (no G/H suffix)")
        self.assertEqual(window.visible_rows.text(), "1 of 6 assembly rows")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "75.0%")
        window._show_row_details(template)
        self.assertIn("Index2 root 6575C88: Present with link", window.row_details.toPlainText())
        self.assertIn("TOPCOVER.xlsx", window.row_details.toPlainText())

    def test_repeated_pdf_rows_have_consistent_visible_counts(self):
        window = self.window
        item = window.checkbox_items["6575C88G16"]
        QA2.QTreeWidgetItem(item.parent(), [item.text(c) for c in range(item.columnCount())])
        window.update_progress()
        self.assertEqual(window.visible_rows.text(), "3 of 3 PDF rows (2 unique assembly/structure pairs)")
        self.assertEqual(window.summary_labels["validation_total"].text(), "2")
        window.status_filter.setCurrentText("PASS")
        self.assertEqual(window.visible_rows.text(), "2 of 3 PDF rows (2 unique assembly/structure pairs)")
        self.assertIn("Filtered unique assemblies: 1", window.audit_coverage_summary.text())

    def test_bold_assembly_groups_use_exact_root_for_missing_count_and_filters(self):
        window = self.window
        sections = QA2.parse_structured_bom(
            "Structure 1: Main\n"
            "b 6573C44G04: SWBC, WRP CROSS BUS MTG 30W\n"
            "\u25cb (0001) 6550C14H01: VERTICAL CHANNEL, WRP, 90 HIGH\n"
            "\u25cb (0002) 6551C13H01: SWBC,SUPPT,MAIN DEV 30\n"
            "b 6575C88G16: SWBC, TOP COVER ASSY\n"
            "b 30-47592: TIP HAZARD WARNING LABEL\n"
            "Structure 2: Main\n"
            "b 6573C44G05: SWBC, WRP CROSS BUS MTG 36W\n"
        )
        self.assertEqual(
            [style for _, _, items in sections for style, *_ in items],
            ["6573C44G04", "6575C88G16", "30-47592", "6573C44G05"],
        )
        window._coverage_catalog = (
            QA2.Index2RootEntry("6573C440", 10, "DIFFERENT ROOT", "OTHER.xlsx"),
            QA2.Index2RootEntry("6575C88", 11, "TOPCOVER", ""),
        )
        window._set_index2_job_sections(sections)
        with patch.object(window, "lookup_dimensions", return_value=("", "")), \
                patch.object(window, "_family_tables", return_value=((), "No table")):
            window.build_bom_checklist(sections)
        self.assertEqual(set(window._coverage_job_roots), {"6573C44", "6575C88"})
        self.assertEqual(window.summary_labels["all_assemblies"].text(), "4")
        self.assertEqual(window.summary_labels["gh_assemblies"].text(), "3")
        self.assertEqual(window.summary_labels["missing"].text(), "2")
        self.assertEqual(window.summary_labels["present"].text(), "1")
        window.audit_coverage_filter.setCurrentText("Not present in Index2")
        self.assertEqual(window.visible_rows.text(), "2 of 4 assembly rows")
        self.assertIn("Filtered unique assemblies: 2", window.audit_coverage_summary.text())
        self.assertIn("Not present: 2", window.audit_coverage_summary.text())
        shown = {item.text(0) for item in window.checkbox_items.values() if not item.isHidden()}
        self.assertEqual(shown, {"6573C44G04", "6573C44G05"})
        window._coverage_catalog += (QA2.Index2RootEntry("6573C44", 178, "WRPCBMTG", ""),)
        window._update_index2_coverage()
        self.assertEqual(window.summary_labels["missing"].text(), "0")
        self.assertEqual(window.summary_labels["present"].text(), "3")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "100.0%")
        self.assertEqual(window.visible_rows.text(), "0 of 4 assembly rows")

    def test_root_summary_and_tables_remain_accessible_at_minimum_window_size(self):
        window = self.window
        window.resize(960, 640)
        window.show()
        for index, tree in ((0, window.bom_list), (1, window.test_output)):
            window.navigation_group.button(index).click()
            self.application.processEvents()
            self.assertGreaterEqual(tree.height(), tree.minimumSizeHint().height())
            area = window.workspace_scroll
            self.assertGreaterEqual(area.widget().height(), area.widget().minimumSizeHint().height())
            if area.widget().height() > area.viewport().height():
                self.assertGreater(area.verticalScrollBar().maximum(), 0)

    def test_presence_rate_is_unknown_until_read_and_empty_job_has_no_rate(self):
        window = self.window
        window._coverage_catalog = None
        window._set_index2_job_sections([("1", "", [("6575C88G16", "", "", "")])])
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "--")
        self.assertEqual(window.summary_labels["present"].text(), "--")
        self.assertEqual(window.summary_labels["missing"].text(), "--")
        window.audit_coverage_filter.setCurrentText("Not linked")
        self.assertEqual(window.visible_rows.text(), "0 of 2 assembly rows")
        window.audit_coverage_filter.setCurrentText("Not checked")
        self.assertEqual(window.visible_rows.text(), "2 of 2 assembly rows")
        window._coverage_catalog = ()
        window._update_index2_coverage()
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "0.0%")
        window._set_index2_job_sections([])
        self.assertEqual(window.summary_labels["all_assemblies"].text(), "0")
        self.assertEqual(window.summary_labels["gh_assemblies"].text(), "0")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "--")

    def test_index2_read_failure_is_not_reported_as_missing_roots(self):
        window = self.window
        window._coverage_catalog = None
        window._set_index2_job_sections([("1", "", [("43C1696G50", "", "", "")])])
        self.assertEqual(window._index2_root_status("43C1696"), "Not checked")
        window._coverage_catalog = (QA2.Index2RootEntry("43C1696", 10, "BASICSTR", "BASICSTR.xlsx"),)
        window._update_index2_coverage()
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "100.0%")
        with patch.object(window._excel_com, "get_index2_catalog", side_effect=RuntimeError("Excel unavailable")), \
                patch.object(window, "_set_status") as status:
            window.refresh_index2_coverage()
        status.assert_called_with("Index2 coverage unavailable: RuntimeError: Excel unavailable")
        self.assertEqual(window._index2_root_status("43C1696"), "Not checked")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "--")
        self.assertEqual(window.summary_labels["present"].text(), "--")
        self.assertEqual(window.summary_labels["missing"].text(), "--")

    def test_index2_refresh_replaces_inventory_and_keeps_audit_results(self):
        window = self.window
        entries = (QA2.Index2RootEntry("43C1696", 10, "BASICSTR", "BASICSTR.xlsx"),)
        window._set_index2_job_sections([("1", "", [("43C1696G50", "", "", "")])])
        original_items = dict(window.checkbox_items)
        with patch.object(window._excel_com, "get_index2_catalog", return_value=entries) as reader, \
                patch.object(window, "_set_status") as status:
            window.refresh_index2_coverage()
        reader.assert_called_once_with(refresh=True)
        status.assert_called_with("Index2 coverage updated: 1 unique roots.")
        self.assertEqual(window.checkbox_items, original_items)
        self.assertEqual(window._index2_root_status("43C1696"), "Present with link")
        self.assertEqual(window.summary_labels["missing"].text(), "0")
        window._set_index2_job_sections([("2", "", [("6575C01G16", "", "", "")])])
        self.assertEqual(window.summary_labels["missing"].text(), "1")
        self.assertNotIn("43C1696", window._coverage_job_roots)
        self.assertEqual(window.summary_labels["all_assemblies"].text(), "1")
        self.assertEqual(window.summary_labels["present"].text(), "0")

    def test_full_assembly_counts_include_g_and_h_but_not_embedded_letters(self):
        window = self.window
        window._coverage_catalog = (QA2.Index2RootEntry("6573C44", 178, "WRPCBMTG", ""),)
        window._set_index2_job_sections([
            ("1", "", [(style, "", "", "") for style in (
                "6573C44G04", "6573C44G05", "7497A69H01", "30-47592", "PDG2XPLKT",
            )]),
            ("2", "", [("6573C44G04", "", "", "")]),
        ])
        self.assertEqual(window.summary_labels["all_assemblies"].text(), "5")
        self.assertEqual(window.summary_labels["gh_assemblies"].text(), "3")
        self.assertEqual(window.summary_labels["present"].text(), "2")
        self.assertEqual(window.summary_labels["missing"].text(), "1")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "66.7%")
        window._set_index2_job_sections([("1", "", [("PDG2XPLKT", "", "", "")])])
        self.assertEqual(window.summary_labels["all_assemblies"].text(), "1")
        self.assertEqual(window.summary_labels["gh_assemblies"].text(), "0")
        self.assertEqual(window.summary_labels["coverage_rate"].text(), "--")

    def test_validation_totals_and_filters_use_only_unique_eligible_assemblies(self):
        window = self.window
        styles = ["6575C88G16", "6575C88G17", "42C2066G31", "42C2066G32",
                  "43C1696G50", "43C1696G51", "8805C17G26", "30-47592"]
        window._coverage_catalog += (QA2.Index2RootEntry("43C1696", 3, "BASICSTR", ""),)
        sections = [
            ("1", "", [(style, "", "", "") for style in styles]),
            ("2", "", [(style, "", "", "") for style in styles[:4]]),
        ]
        window._set_index2_job_sections(sections)
        self.assertEqual(window.summary_labels["validation_total"].text(), "6")
        self.assertEqual(window.summary_labels["UNRESOLVED"].text(), "6")
        statuses = {
            "6575C88G16": ("PASS", "PASS"),
            "6575C88G17": ("PASS", "FAIL"),
            "42C2066G31": ("PASS", "UNRESOLVED"),
            "42C2066G32": ("FAIL", "UNRESOLVED"),
        }

        def validate(style, mapping, width, depth, tables, error, description, section):
            status = statuses[style][int(section) - 1]
            group = QA2.split_style_number(style)[1]
            return QA2.ValidationResult(group, group if status == "PASS" else "", status, "Fixture evidence")

        with patch.object(window, "lookup_dimensions", return_value=("", "")) as lookup, \
                patch.object(window, "_family_tables", return_value=((), "")) as tables, \
                patch.object(window, "validate_assembly", side_effect=validate) as compare, \
                patch.object(window._excel_com, "prepare_assembly_targets") as prepare:
            window.build_bom_checklist(sections)
        self.assertEqual(lookup.call_count, 8)
        self.assertEqual(tables.call_count, 8)
        self.assertEqual(compare.call_count, 8)
        self.assertEqual(
            {style for _, _, items in prepare.call_args.args[0] for style, *_ in items},
            set(styles[:4]),
        )
        self.assertEqual(window._eligible_assembly_results(), {
            "6575C88G16": "PASS", "6575C88G17": "FAIL",
            "42C2066G31": "UNRESOLVED", "42C2066G32": "FAIL",
            "43C1696G50": "UNRESOLVED", "43C1696G51": "UNRESOLVED",
        })
        self.assertEqual(window.summary_labels["present"].text(), "6")
        self.assertEqual(window.summary_labels["validation_total"].text(), "6")
        self.assertEqual(window.summary_labels["PASS"].text(), "1")
        self.assertEqual(window.summary_labels["FAIL"].text(), "2")
        self.assertEqual(window.summary_labels["UNRESOLVED"].text(), "3")
        self.assertEqual(window.summary_labels["validation_rate"].text(), "16.7%")
        self.assertEqual(window._validation_results["STRUCT1::8805C17G26"].status, "EXCLUDED")
        self.assertEqual(window._validation_results["STRUCT1::30-47592"].status, "EXCLUDED")
        self.assertIn("no workbook link", window._validation_results["STRUCT1::43C1696G50"].message)
        for status, visible in (("PASS", 2), ("FAIL", 4), ("UNRESOLVED", 4), ("EXCLUDED", 2)):
            window.status_filter.setCurrentText(status)
            self.assertEqual(window.visible_rows.text(), f"{visible} of 12 assembly rows")
        window.status_filter.setCurrentText("FAIL")
        window.audit_search.setText("6575C88G17")
        self.assertEqual(window.visible_rows.text(), "2 of 12 assembly rows")
        window._show_row_details(window.checkbox_items["STRUCT1::6575C88G17"])
        self.assertIn("Overall assembly result: FAIL", window.row_details.toPlainText())

    def test_pending_checks_and_job_changes_cannot_produce_a_premature_pass(self):
        window = self.window
        sections = [
            ("1", "", [("6575C88G16", "", "", "")]),
            ("2", "", [("6575C88G16", "", "", "")]),
        ]
        window._set_index2_job_sections(sections)
        window._validation_checks["6575C88G16"] = ["PASS"]
        window.update_progress()
        self.assertEqual(window.summary_labels["PASS"].text(), "0")
        self.assertEqual(window.summary_labels["UNRESOLVED"].text(), "1")
        window._validation_checks["6575C88G16"].append("PASS")
        window.update_progress()
        self.assertEqual(window.summary_labels["PASS"].text(), "1")
        self.assertEqual(window.summary_labels["validation_rate"].text(), "100.0%")
        window._set_index2_job_sections(sections)
        self.assertEqual(window.summary_labels["PASS"].text(), "0")
        self.assertEqual(window.summary_labels["UNRESOLVED"].text(), "1")
        window._validation_checks["6575C88G16"] = ["FAIL"]
        window.update_progress()
        self.assertEqual(window.summary_labels["FAIL"].text(), "1")
        window._coverage_catalog = None
        window._update_index2_coverage()
        for key in ("validation_total", "PASS", "FAIL", "UNRESOLVED", "validation_rate"):
            self.assertEqual(window.summary_labels[key].text(), "--")

    def test_repeated_pdf_check_failure_is_not_overwritten_by_a_later_pass(self):
        window = self.window
        sections = [("1", "", [("6575C88G16", "", "", ""), ("6575C88G16", "", "", "")])]
        window._set_index2_job_sections(sections)
        with patch.object(window, "lookup_dimensions", return_value=("", "")), \
                patch.object(window, "_family_tables", return_value=((), "")), \
                patch.object(window, "validate_assembly", side_effect=[
                    QA2.ValidationResult("G16", "G17", "FAIL", "First PDF occurrence"),
                    QA2.ValidationResult("G16", "G16", "PASS", "Second PDF occurrence"),
                ]):
            window.build_bom_checklist(sections)
        self.assertEqual(window.summary_labels["validation_total"].text(), "1")
        self.assertEqual(window.summary_labels["FAIL"].text(), "1")
        self.assertEqual(window.summary_labels["PASS"].text(), "0")
        window.status_filter.setCurrentText("FAIL")
        self.assertEqual(window.visible_rows.text(), "2 of 2 PDF rows (1 unique assembly/structure pairs)")

    def test_automatically_discovered_columns_preserve_existing_validation_cells(self):
        window = self.window
        item = window.checkbox_items["6575C88G16"]
        window._ensure_parameter_columns(8)
        self.assertEqual(window._parameter_columns, 8)
        self.assertEqual([item.text(column) for column in range(12, 15)], ["G16", "G16", "PASS"])
        self.assertEqual([item.text(column) for column in range(9, 12)], ["", "", ""])
        self.assertEqual(item.toolTip(14), "Validation evidence")
        window.update_progress()
        self.assertEqual(window.summary_labels["PASS"].text(), "1")

    def test_manual_and_automatic_mapping_modes_are_visible_without_forcing_a_pass(self):
        window = self.window
        window._coverage_catalog += (QA2.Index2RootEntry("42C1084", 3, "FAMILY", "FAMILY.xlsx"),)
        window._update_index2_coverage()
        tables = QA2.ClassifierWorkbookParser.extract_tables(
            ClassifierIntegrationTests.matrix_workbook(), "FAMILY.xlsx", "42C1084",
        )
        style = "42C1084G01"
        window._assembly_families[style] = "FAMILY"
        window._assembly_links[style] = "FAMILY.xlsx"
        section = window.add_section("Structure 2")
        for manual in (True, False):
            key = str(manual)
            window._code_mappings = {"FAMILY": QA2.CodeMapping(
                "Family", "FAMILY", ("Width (Wide)", "Depth"), ("Width", "Depth"),
                manually_mapped=manual, checklist_row=2,
            )}
            with patch.object(window._excel_com, "get_group_tables", return_value=tables):
                window.add_checkbox(section, key, style, "", "", "", "36", "24", "2")
            expected = "column N = TRUE" if manual else "column N is not TRUE"
            self.assertIn(expected, window.checkbox_items[key].toolTip(2))
            self.assertIn(expected, window.checkbox_items[key].toolTip(4))
            self.assertEqual(window._validation_results[key].status, "FAIL")

    def test_classifier_drives_validation_and_parameter_display(self):
        window = self.window
        window._coverage_catalog += (QA2.Index2RootEntry("42C1084", 3, "FAMILY", "FAMILY.xlsx"),)
        window._update_index2_coverage()
        workbook = ClassifierIntegrationTests.matrix_workbook()
        tables = QA2.ClassifierWorkbookParser.extract_tables(workbook, "FAMILY.xlsx", "42C1084")
        mapping = QA2.CodeMapping("Family", "FAMILY", ("Width (Wide)", "Depth"), ("Width", "Depth"))
        window._assembly_families["42C1084G01"] = "FAMILY"
        window._assembly_links["42C1084G01"] = "FAMILY.xlsx"
        window._code_mappings = {"FAMILY": mapping}
        section = window.add_section("Structure 2")
        with patch.object(window._excel_com, "get_group_tables", return_value=tables) as read, \
                patch.object(window._excel_com, "get_sheet_values", side_effect=AssertionError("legacy reader used")):
            window.add_checkbox(section, "classifier", "42C1084G01", "", "", "", "36", "24", "2")
        read.assert_called_once_with("FAMILY.xlsx", assembly_base="42C1084")
        result = window._validation_results["classifier"]
        self.assertEqual((result.status, result.expected_group, result.actual_group), ("FAIL", "G01", "G02"))
        item = window.checkbox_items["classifier"]
        self.assertIn("Width (Wide) = 36 [Index2: 30]", item.text(4))
        self.assertIn("Classifier", result.message)

    def test_classifier_respects_basicstr_enclosure_selection(self):
        window = self.window
        workbook = ClassifierIntegrationTests.matrix_workbook()
        workbook.active.title = "BASICSTR-NEMA 1"
        second = workbook.copy_worksheet(workbook.active)
        second.title = "BASICSTR-NEMA 3R"
        second.cell(3, 4, "G09")
        tables = QA2.ClassifierWorkbookParser.extract_tables(workbook, "BASICSTR.xlsx", "42C1084")
        style = "42C1084G02"
        window._assembly_families[style] = "BASICSTR"
        window._assembly_links[style] = "BASICSTR.xlsx"
        with patch.object(window._excel_com, "get_group_tables", return_value=tables):
            selected, error = window._family_tables(style)
            self.assertFalse(selected)
            self.assertIn("requires an enclosure", error)
            window.enclosure_sheet = "BASICSTR-NEMA 1"
            selected, error = window._family_tables(style)
        self.assertFalse(error)
        self.assertEqual({table.sheet_name for table in selected}, {"BASICSTR-NEMA 1"})
        group, _message = QA2.ClassifierWorkbookParser.identify_group(selected, "42C1084", {"Width": 36, "Depth": 24})
        self.assertEqual(group, "G02")

    def test_search_and_status_filters_preserve_audit_rows(self):
        self.window.status_filter.setCurrentText("PASS")
        self.assertEqual(self.window.visible_rows.text(), "1 of 2 assembly rows")
        self.window.audit_search.setText("42C2066")
        self.assertEqual(self.window.visible_rows.text(), "0 of 2 assembly rows")
        self.window.status_filter.setCurrentText("All statuses")
        self.assertEqual(self.window.visible_rows.text(), "1 of 2 assembly rows")
        self.assertEqual(len(self.window.checkbox_items), 2)
        self.assertEqual(self.window.summary_labels["PASS"].text(), "1")

    def test_navigation_column_toggle_and_details(self):
        self.window.navigation_group.button(1).click()
        self.assertEqual(self.window.main_tabs.currentIndex(), 1)
        self.assertEqual(self.window.bom_list.header().visualIndex(10), 3)
        self.assertTrue(self.window.bom_list.isColumnHidden(4))
        self.window.parameter_toggle.setChecked(True)
        self.assertFalse(self.window.bom_list.isColumnHidden(4))
        self.window._show_row_details(self.window.checkbox_items["6575C88G16"])
        self.assertIn("Validation evidence", self.window.row_details.toPlainText())
        self.assertIn("Width = 30", self.window.row_details.toPlainText())

    def test_workspace_animation_finishes_and_restarts_on_navigation(self):
        window = self.window
        window.show()
        QA2.QApplication.processEvents()
        self.assertTrue(window._workspace_revealed)
        window._tab_animation.setCurrentTime(window._tab_animation.duration())
        self.assertEqual(window._tab_opacity.opacity(), 1.0)
        window.navigation_group.button(1).click()
        self.assertEqual(window.main_tabs.currentIndex(), 1)
        self.assertLess(window._tab_opacity.opacity(), 1.0)
        window._tab_animation.setCurrentTime(window._tab_animation.duration())
        self.assertEqual(window._tab_opacity.opacity(), 1.0)

    def test_running_indicator_animates_and_finishes_with_explicit_outcome(self):
        window = self.window
        window._pending_folder = "job"

        def run(_folder):
            self.assertTrue(window._run_timer.isActive())
            self.assertFalse(window.run_busy_bar.isHidden())
            self.assertFalse(window.start_btn.isEnabled())
            self.assertEqual(window.start_btn.text(), "Running...")
            first = window.run_indicator.text()
            window._refresh_running_indicator()
            self.assertNotEqual(window.run_indicator.text(), first)
            self.assertIn("Elapsed", window.run_indicator.text())
            window.start_automation()
            return True

        with patch.object(window, "_run_automation", side_effect=run) as automation:
            window.start_automation()
        automation.assert_called_once_with("job")
        self.assertEqual(window.run_indicator.property("runState"), "completed")
        self.assertIn("Completed", window.run_indicator.text())
        self.assertFalse(window._run_timer.isActive())
        self.assertTrue(window.run_busy_bar.isHidden())
        self.assertTrue(window.start_btn.isEnabled())
        self.assertTrue(window.test_scope_checkbox.isEnabled())
        self.assertEqual(window.start_btn.text(), "Start Automation")

    def test_incomplete_and_failed_runs_do_not_show_completed(self):
        window = self.window
        window._pending_folder = "job"
        with patch.object(window, "_run_automation", return_value=None):
            window.start_automation()
        self.assertEqual(window.run_indicator.property("runState"), "attention")
        self.assertFalse(window._run_timer.isActive())

        with patch.object(window, "_run_automation", side_effect=RuntimeError("Excel unavailable")):
            with self.assertRaisesRegex(RuntimeError, "Excel unavailable"):
                window.start_automation()
        self.assertEqual(window.run_indicator.property("runState"), "failed")
        self.assertIn("Excel unavailable", window.automation_log.toPlainText())
        self.assertFalse(window._run_timer.isActive())
        self.assertTrue(window.run_busy_bar.isHidden())
        self.assertTrue(window.start_btn.isEnabled())
        self.assertTrue(window.browse_pdf_btn.isEnabled())
        self.assertIsNone(QA2.QApplication.overrideCursor())

    def test_status_cards_and_scope_fit_at_minimum_window_size(self):
        window = self.window
        window.resize(960, 640)
        window.show()
        QA2.QApplication.processEvents()
        for key, label in window.summary_labels.items():
            self.assertEqual(label.parentWidget().objectName(), "card" + key)
            self.assertGreaterEqual(label.width(), label.sizeHint().width())
        sidebar = window.test_scope_checkbox.parentWidget()
        for widget in (window.automation_log, window.test_scope_checkbox, window.enclosure_label):
            self.assertTrue(sidebar.rect().contains(widget.geometry()))
        self.assertLess(window.automation_log.geometry().bottom(), window.test_scope_checkbox.y())

    def test_activity_log_tracks_status_and_limits_history(self):
        window = self.window
        self.assertTrue(window.automation_log.isReadOnly())
        self.assertEqual(window.automation_log.maximumBlockCount(), 1000)
        window._set_status("Reading Index2...")
        window._set_status("Opening Enclosure.xlsx...")
        text = window.automation_log.toPlainText()
        self.assertRegex(text, r"\[\d{2}:\d{2}:\d{2}\] Reading Index2")
        self.assertLess(text.index("Reading Index2"), text.index("Opening Enclosure"))
        self.assertEqual(window.automation_activity.text(), "Opening Enclosure.xlsx...")
        self.assertEqual(window.scan_status.text(), window.automation_activity.text())
        window.automation_log.setMaximumBlockCount(2)
        window._set_status("Validation complete")
        self.assertNotIn("Reading Index2", window.automation_log.toPlainText())
        window.clear_log_button.click()
        self.assertEqual(window.automation_log.toPlainText(), "")
        self.assertEqual(window.automation_activity.text(), "Validation complete")

    def test_activity_log_scrolls_to_latest_message(self):
        window = self.window
        window.show()
        QA2.QApplication.processEvents()
        for index in range(100):
            window._set_status(f"Log message {index}")

        scrollbar = window.automation_log.verticalScrollBar()
        self.assertGreater(scrollbar.maximum(), 0)
        scrollbar.setValue(0)
        window._set_status("Latest log message")

        self.assertEqual(scrollbar.value(), scrollbar.maximum())

    def test_activity_log_does_not_push_scope_controls_out_of_view(self):
        window = self.window
        window.resize(960, 640)
        window.show()
        QA2.QApplication.processEvents()

        self.assertLessEqual(window.automation_log.height(), 220)
        self.assertTrue(window.test_scope_checkbox.isVisible())
        self.assertTrue(window.test_scope_checkbox.isEnabled())
        previous_state = window.test_scope_checkbox.isChecked()
        window.test_scope_checkbox.click()
        self.assertNotEqual(window.test_scope_checkbox.isChecked(), previous_state)

    def test_barrier_validation_uses_job_depth_not_existing_group(self):
        window = self.window
        sheets = (("WCVTBARA", (
            ("REF. - 6552C34",), ("DEPTH", "GROUP"),
            ("24D", "G02"), ("30D", "G03"), ("36D", "G04"),
        )),)
        window._assembly_links["6552C34G04"] = "WCVTBAR1.xlsx"
        with patch.object(window._excel_com, "get_sheet_values", return_value=sheets):
            for job_depth, actual, status in (("36D", "G04", "PASS"), ("30D", "G03", "FAIL")):
                with self.subTest(job_depth=job_depth):
                    window._parameter_repository.clear()
                    window._parameter_repository.add_parameter("Structure Depth", job_depth, "outline.pdf")
                    result = window.validate_assembly("6552C34G04", None, "", "", (), "")
                    self.assertEqual((result.actual_group, result.status), (actual, status))
                    self.assertIn("outline.pdf", result.message)
            window._parameter_repository.add_parameter("Depth", "36", "other.pdf")
            result = window.validate_assembly("6552C34G04", None, "", "", (), "")
            self.assertEqual(result.status, "UNRESOLVED")
            self.assertIn("Conflicting", result.message)

    def test_section_depths_do_not_leak_between_barrier_assemblies(self):
        window = self.window
        sheets = (("WCVTBARA", (
            ("REF. - 6552C34",), ("DEPTH", "GROUP"),
            ("24D", "G02"), ("30D", "G03"), ("36D", "G04"),
        )),)
        window._assembly_links["6552C34G04"] = "WCVTBAR1.xlsx"
        window._job_dimensions, window._section_dimensions = QA2.extract_job_dimensions(
            "Structure 1: Utility\nDepth: 36D\nStructure 2: Utility\nDepth: 30D",
            "outline.pdf",
        )
        window._parameter_repository.add_parameter("Depth", "30", "outline.pdf")
        window._parameter_repository.add_parameter("Depth", "36", "outline.pdf")
        with patch.object(window._excel_com, "get_sheet_values", return_value=sheets):
            for number, group, status in ((1, "G04", "PASS"), (2, "G03", "FAIL"), (3, "", "UNRESOLVED")):
                with self.subTest(section=number):
                    result = window.validate_assembly(
                        "6552C34G04", None, "", "", (), "", section_number=number,
                    )
                    self.assertEqual((result.actual_group, result.status), (group, status))

    def test_display_and_validation_do_not_borrow_existing_group_depth(self):
        window = self.window
        window._coverage_catalog += (QA2.Index2RootEntry("6552C34", 3, "WCVTBAR1", "WCVTBAR1.xlsx"),)
        window._update_index2_coverage()
        sheets = (("WCVTBARA", (
            ("REF. - 6552C34",), ("DEPTH", "GROUP"), ("36D", "G04"),
        )),)
        window._assembly_links["6552C34G04"] = "WCVTBAR1.xlsx"
        window._assembly_families["6552C34G04"] = "WCVTBAR1"
        window._code_mappings = QA2.parse_code_mappings((
            ("Name", "Code", "Parameter 1"), ("Barrier", "WCVTBAR1", "Depth"),
        ))
        section = window.bom_list.topLevelItem(0)
        with patch.object(window._excel_com, "get_group_tables", return_value=()), \
                patch.object(window._excel_com, "get_sheet_values", return_value=sheets):
            window.add_checkbox(section, "barrier", "6552C34G04", "VERTICAL SIDE BARRIER", "", "36")
            item = window.checkbox_items["barrier"]
            self.assertEqual(window._validation_results["barrier"].status, "UNRESOLVED")
            self.assertIn("Depth [Index2: 36", item.text(4))
            self.assertNotIn("Depth =", item.text(4))
            window._parameter_repository.add_parameter("Depth", "36", "outline.pdf")
            window.add_checkbox(section, "barrier-job", "6552C34G04", "VERTICAL SIDE BARRIER", "", "36")
            item = window.checkbox_items["barrier-job"]
            self.assertEqual(window._validation_results["barrier-job"].status, "PASS")
            self.assertIn("Depth = 36", item.text(4))

    def test_section_requirement_is_used_in_display_and_comparison(self):
        window = self.window
        window._coverage_catalog += (QA2.Index2RootEntry("2C11531", 3, "WCCTAUX3", "WCCTAUX3.xlsx"),)
        window._update_index2_coverage()
        window._code_mappings = QA2.parse_code_mappings((
            ("Name", "Code", "Parameter 1", "Parameter 2", "Parameter 3"),
            ("Bus", "WCCTAUX3", "Amperage", "Material", "Heat Test/Density"),
        ))
        window._assembly_families["2C11531G03"] = "WCCTAUX3"
        window._assembly_links["2C11531G03"] = "WCCTAUX3.xlsx"
        window.parameter_variables = QA2.ExcelParameterVariables.from_sources([
            "Bus Amps: 3000\nBus Material: Copper\nHeat Test",
        ])
        window._section_bus_requirements = QA2.section_bus_requirements(
            "SECTION 8 NEEDS DENSITY RATED BUS (998P160)", "outline.dwg",
        )
        with patch.object(window._excel_com, "get_group_tables", return_value=()), \
                patch.object(window._excel_com, "get_sheet_values", return_value=WorkbookValueTests.wcctaux3_sheet()):
            for number in (6, 8):
                section = window.add_section(f"Structure {number}")
                window.add_checkbox(section, f"STRUCT{number}::2C11531G03", "2C11531G03", "XBUS,UGPS,25-3000,CU DEN,51W", "51", "36", "51", "36", number)
        item = window.checkbox_items["STRUCT8::2C11531G03"]
        self.assertEqual([item.text(column) for column in (9, 10, 11)], ["G03", "G03", "PASS"])
        self.assertIn("Heat Test/Density = DENSITY", item.text(6))
        self.assertIn("998P160", item.toolTip(6))
        self.assertEqual(window._validation_results["STRUCT6::2C11531G03"].status, "FAIL")
        self.assertEqual(window._validation_results["STRUCT8::2C11531G03"].status, "PASS")


class CodeMappingTests(unittest.TestCase):
    def test_index2_failure_does_not_replace_checklist_with_empty_rows(self):
        window = MagicMock()
        window._find_named_pdf.side_effect = ["bom.pdf", None]
        window._read_pdf.return_value = SimpleNamespace(text="BOM")
        window._pdf_parameter_sources.return_value = []
        window.test_scope_checkbox.isChecked.return_value = False
        window._load_code_mappings.return_value = ({"TEST": object()}, "local checklist")
        window._excel_com.get_index2_map.side_effect = RuntimeError("Excel dialog is open")
        with patch.object(QA2, "parse_structured_bom", return_value=[("1", "Structure", [])]), \
                patch.object(QA2.os, "listdir", return_value=[]), \
                patch.object(QA2.QMessageBox, "warning") as warning:
            QA2.AuditReviewDemo._run_automation(window, "job")
        window.build_bom_checklist.assert_not_called()
        self.assertEqual(warning.call_args.args[1], "Index2 Unavailable")
        self.assertIn("Close any Excel error dialogs", warning.call_args.args[2])

    def test_parameter_module_is_used_without_excel_or_local_workbook(self):
        window = SimpleNamespace(_excel_com=MagicMock(), _set_status=MagicMock())
        with patch.object(QA2, "read_local_sheet_values") as local_read, \
                patch.object(QA2, "load_workbook") as open_workbook:
            mappings, note = QA2.AuditReviewDemo._load_code_mappings(window)
        local_read.assert_not_called()
        open_workbook.assert_not_called()
        window._excel_com.read_sheets_values.assert_not_called()
        self.assertIn("WCUTO22XMAIN", mappings)
        self.assertEqual(window._common_parameters, tuple(QA2.checklist_parameters.COMMON_PARAMETERS))
        self.assertEqual(QA2.EXCEL_PARAMETERS, window._common_parameters)
        self.assertIn("checklist_parameters.py", note)

    def test_updated_parameter_lists_are_loaded_from_the_module(self):
        rows = [("Name", "Code", "Parameter 1"), ("New family", "NEWCODE", "New Parameter")]
        window = SimpleNamespace(_excel_com=MagicMock(), _set_status=MagicMock())
        with patch.object(QA2.checklist_parameters, "CODE_WISE_PARAMETERS", rows), \
                patch.object(QA2.checklist_parameters, "COMMON_PARAMETERS", ["New Parameter"]):
            mappings, _note = QA2.AuditReviewDemo._load_code_mappings(window)
        self.assertEqual(tuple(mappings), ("NEWCODE",))
        self.assertEqual(mappings["NEWCODE"].parameters[0], "New Parameter")
        self.assertEqual(window._common_parameters, ("New Parameter",))
        window._excel_com.read_sheets_values.assert_not_called()

    def test_parameter_definitions_have_no_duplicate_codes_or_names(self):
        rows = QA2.checklist_parameters.CODE_WISE_PARAMETERS
        self.assertEqual(len(QA2.parse_code_mappings(rows)), len(rows) - 1)
        names = QA2.checklist_parameters.COMMON_PARAMETERS
        self.assertEqual(len(names), len(set(names)))

    def test_open_only_checklist_read_never_launches_a_workbook(self):
        session = QA2.ExcelComSession()
        with patch.object(session, "_get_app", return_value=SimpleNamespace(Workbooks=[])), \
                patch.object(session, "_get_index2_workbook") as index2, \
                patch.object(session, "_open_with_sign_in") as launch:
            with self.assertRaisesRegex(RuntimeError, "Open .* in Excel"):
                session.read_sheets_values(QA2.CHECKLIST_LIVE_URLS[0], (QA2.CODE_SHEET_NAME,), open_if_missing=False)
        index2.assert_not_called()
        launch.assert_not_called()

    def test_sheet1_rows_map_by_normalized_code(self):
        rows = (
            ("Name", "Code", "Parameter 1", "Parameter 2"),
            ("Bus Ground", "GND BOND", "Amperage", "Assembly No.", None),
            ("Utilities", "CTWISPUB", "Amperage"),
        )
        mappings = QA2.parse_code_mappings(rows)

        mapping = mappings[QA2.normalize_code("gndbond")]
        self.assertEqual(mapping.name, "Bus Ground")
        self.assertEqual(mapping.code, "GND BOND")
        self.assertEqual(mapping.parameters, ("Amperage", "Assembly No.", "", "", ""))
        self.assertEqual(mappings["CTWISPUB"].parameters, ("Amperage", "", "", "", ""))


class GroupValidationTests(unittest.TestCase):
    def test_explicit_job_dimensions_keep_scope_and_provenance(self):
        global_values, sections = QA2.extract_job_dimensions(
            "Depth: 48 inches\nWidth (Wide): 42\n"
            "SECTION 1: Utility\nStructure Depth = 36D\nSection Width: 30\n"
            "SECTION 2: Utility\nDepth: 30.00\n"
            "BARRIER 24D\n42C2066G05 COVER 36W",
            "outline.pdf (page 2)",
        )
        self.assertEqual(QA2.dimension_value(QA2.dimension_entries(global_values, "Depth")), "48")
        self.assertEqual(QA2.dimension_value(QA2.dimension_entries(global_values, "Width")), "42")
        self.assertEqual(sections["1"].get_parameter("Depth"), "36")
        self.assertEqual(sections["1"].get_parameter("Width (Wide)"), "30")
        self.assertEqual(QA2.dimension_value(QA2.dimension_entries(sections["2"], "Depth")), "30")
        self.assertEqual(sections["2"].get_entries("Depth"), (("30.00", "outline.pdf (page 2)"),))
        self.assertEqual(sections["2"].get_entries("Width (Wide)"), ())

    def test_dimension_aliases_allow_equivalent_units_but_not_conflicts(self):
        repository = QA2.ParameterRepository()
        repository.add_parameter("Structure Depth", "36D", "outline.pdf")
        repository.add_parameter("Depth", "36.00", "job.pdf")
        self.assertEqual(QA2.dimension_value(QA2.dimension_entries(repository, "Depth")), "36")
        repository.add_parameter("Depth", "30D", "other.pdf")
        self.assertEqual(QA2.dimension_value(QA2.dimension_entries(repository, "Depth")), "")

    values = (
        ("8708C64 ASSEMBLY", None, None),
        (None, None, None),
        ("Amperage", "Material", "Group"),
        (800, "Copper", "G01"),
        (1200, "Copper", "G02"),
        (1200, "Aluminum", "G03"),
        (None, None, None),
        ("Amperage", "Group", None),
        (400, "G99", None),
    )

    def tables(self):
        return QA2.WorkbookParser.extract_tables(self.values, "BUS")

    def test_parser_stops_tables_at_blank_rows(self):
        tables = self.tables()
        self.assertEqual([len(table.rows) for table in tables], [3, 1])

    def test_group_found_from_all_matching_parameters(self):
        group, _ = QA2.GroupDeterminationEngine.identify_group(
            self.tables(), "8708C64", {"Amperage": "1200A", "Material": "aluminum"},
        )
        self.assertEqual(group, "G03")

    def test_ambiguous_parameters_do_not_guess_a_group(self):
        group, message = QA2.GroupDeterminationEngine.identify_group(
            self.tables(), "8708C64", {"Amperage": "1200"},
        )
        self.assertEqual(group, "")
        self.assertIn("Multiple groups", message)

    def test_table_for_other_base_is_ignored(self):
        group, _ = QA2.GroupDeterminationEngine.identify_group(
            self.tables(), "9999C99", {"Amperage": "800"},
        )
        self.assertEqual(group, "")

    def test_compare_reports_pass_fail_and_unresolved(self):
        compare = QA2.ValidationEngine.compare
        self.assertEqual(compare("G03", "g03").status, "PASS")
        self.assertEqual(compare("G03", "G02").status, "FAIL")
        self.assertEqual(compare("G03", "", "why").status, "UNRESOLVED")
        self.assertEqual(compare("", "G02").status, "UNRESOLVED")


class WorkbookValueTests(unittest.TestCase):
    def test_phase_order_does_not_override_system_phase_count(self):
        configuration = QA2.ConfigurationRepository.__new__(QA2.ConfigurationRepository)
        configuration._parameter_rules = (QA2.ParameterRule("3/1 Phase", ("Phase",)),)
        repository = configuration.extract_parameters("Phase: ABC FEEDER BREAKER\nPhase: 3-Phase", "quickprint")
        self.assertEqual(repository.get_parameter("3/1 Phase"), "3")
        self.assertEqual(repository.get_parameter("Phase Order"), "ABC FEEDER BREAKER")
        unknown = configuration.extract_parameters("Phase: UNKNOWN", "quickprint")
        self.assertEqual(unknown.get_parameter("3/1 Phase"), "UNKNOWN")

    def test_paired_door_descriptions_supply_separate_structure_widths(self):
        for description, left, right in (
            ("FRWP FRT EXT DBL DR 36-36", "36", "36"),
            ("ASSY,DR,NEMA/3R,DBL,38-51", "38", "51"),
        ):
            values = QA2.description_parameter_values(description)
            self.assertEqual(values["Structure Width - Left (Or Single) Door"], left)
            self.assertEqual(values["Structure Width - Right Door"], right)
        self.assertNotIn("Structure Width - Right Door", QA2.description_parameter_values("BUS 36-51"))

    def test_wire_aliases_and_specific_consistent_material(self):
        parser = QA2.ClassifierWorkbookParser
        for label in ("3/4 Wire", "3/4 Wire (3/4)", "Wire", "Wires"):
            self.assertEqual(parser.parameter_key(label), "WIRES")
        self.assertEqual(parser.consistent_value(["Copper", "Silver Plated Copper"]), "Silver Plated Copper")
        self.assertEqual(parser.consistent_value(["Copper", "Silver Plated Copper", "Tin Plated Copper"]), "")

    def test_description_ranges_preserve_lower_endpoint(self):
        for description, expected in (
            ("W.C. UTIL.COMP 400-800A CU HT", "400-800"),
            ("W/C NT TAP 400-1000 CU HT 36D", "400-1000"),
            ("BUS,25-3000 CU DEN", "2500-3000"),
        ):
            self.assertEqual(QA2.description_parameter_values(description)["Amperage"], expected)
        self.assertNotIn("Amperage", QA2.description_parameter_values('SPLICE PLS 3-4"'))
        self.assertNotIn("Amperage", QA2.description_parameter_values("1000A MAX"))

    def test_bus_amperage_is_separate_from_device_ratings(self):
        repository = QA2.extract_table_selectors("Bus Amps: 3000 Bus Bracing Rating: 65kA\nBreaker Amps: 800", "drawing")
        self.assertEqual(repository.get_parameter("Bus Amperage"), "3000")
        self.assertFalse(repository.get_parameter("Amperage"))
        self.assertTrue(QA2.is_bus_description("EUSERC, UGPS, 2500-3000 CUAG D"))
        self.assertFalse(QA2.is_bus_description("EUSERC, UGPS SIDE PANEL"))

    def test_table_selector_labels_retain_source_and_conflicts(self):
        repository = QA2.extract_table_selectors("NG Drawout: No\nStructure Material: Galvanized (Painted)\nTest Plug: WITH", "drawing.pdf page 2")
        self.assertEqual(repository.get_parameter("NG Drawout"), "No")
        self.assertEqual(repository.get_entries("Structure Material")[0][1], "drawing.pdf page 2")
        repository.merge_sources(QA2.extract_table_selectors("NG Drawout: Yes", "other drawing"))
        self.assertTrue(repository.has_conflict("NG Drawout"))

    def test_recognized_missing_inputs_do_not_use_conventional_fallback(self):
        sheets = (("WCVTBARA", (("REF. - 6552C34",), ("DEPTH", "GROUP"), ("36D", "G04"))),)
        window = SimpleNamespace(
            parameter_variables=QA2.ExcelParameterVariables({}),
            _resolved_value=lambda *args: "", _assembly_links={"6552C34G04": "barrier.xlsx"},
            _excel_com=SimpleNamespace(get_sheet_values=lambda link: sheets),
        )
        with patch.object(QA2.GroupDeterminationEngine, "identify_group", return_value=("G04", "wrong fallback")) as fallback:
            result = QA2.AuditReviewDemo.validate_assembly(window, "6552C34G04", None, "", "", (), "")
        self.assertEqual(result.status, "UNRESOLVED")
        self.assertIn("Depth", result.message)
        fallback.assert_not_called()

    def test_enclosure_smart_table_requires_dimensions_not_fixed_applicability(self):
        sheets = (("Enclosure.xlsx / Basic Structure - Type 1", (
            ("INPUT", None, None, None, None, "OUTPUT", None, "REFERENCE INFORMATION"),
            ("NG DRAWOUT", "STRUCTURE HEIGHT", "STRUCTURE WIDTH", "STRUCTURE DEPTH", "SEISMIC", "BASIC STRUCTURE"),
            ("No", "90 Only", "36 Only", "30 Only", "Non-Seismic", "6575C01G22"),
            ("No", "90 Only", "36 Only", "36 Only", "Non-Seismic", "6575C01G23"),
        )),)
        wanted = {"NG Drawout": "No", "Height": "90", "Width (Wide)": "36", "Depth": "30", "Seismic/Non-Seismic": "Non-Seismic", "NEMA 1/3R": "1"}
        self.assertEqual(QA2.find_grid_group(sheets, "6575C01", wanted)[0], "G22")
        for name in wanted:
            incomplete = {key: value for key, value in wanted.items() if key != name}
            actual, message = QA2.find_grid_group(sheets, "6575C01", incomplete)
            if name in ("Height", "Width (Wide)", "Depth"):
                self.assertEqual(actual, "")
                self.assertIn(name, message)
            else:
                self.assertEqual(actual, "G22")
        self.assertEqual(QA2.find_grid_group(sheets, "6575C01", {**wanted, "NG Drawout": "Yes"})[0], "")

    def test_retired_workbooks_cannot_supply_groups(self):
        for notice, target in (
            ("Will be maintained on spreadsheet: ENCLOSURE.xls (Index 377)", "ENCLOSURE.XLS"),
            ("All groups will be maintained on spreadsheet: Generic Cross Bus and Mountings.xls (Index 369)", "GENERIC CROSS BUS AND MOUNTINGS.XLS"),
            ("THIS FILE IS OBSOLTED", ""),
        ):
            sheets = (("OLD", ((notice,), ("ASSY DWG", '36" WIDE'), ("42C2066", "G05"))),)
            with self.subTest(notice=notice):
                self.assertEqual(QA2.workbook_migration(sheets)[0], target)
                actual, message = QA2.find_grid_group(sheets, "42C2066", {"Width (Wide)": "36"})
                self.assertEqual(actual, "")
                self.assertIn("retired", message)

    def test_vertical_barrier_depth_table_requires_depth(self):
        sheets = (("WCVTBARA", (
            ("REF. - 6552C34",), ("DEPTH", "GROUP"),
            ("24D", "G02"), ("30D", "G03"), ("36D", "G04"),
        )),)
        self.assertEqual(QA2.find_grid_group(sheets, "6552C34", {"Depth": "36"})[0], "G04")
        self.assertIn("Depth", QA2.find_grid_group(sheets, "6552C34", {"Width (Wide)": "36"})[1])
        self.assertEqual(QA2.find_grid_group(sheets, "6552C34", {"Depth": "40"})[0], "")

    def test_metering_matrix_requires_width_row_and_test_plug(self):
        sheets = (("DORMTGA", (
            ("W/O TEST PLUG",), ("6555C79",), ("DESCRIPTION", "30W", "36W"),
            ("AM & VM & SWITCHES", "G03", "G04"), (),
            ("INDICATING LIGHT (1)", "G11", "G12"), (),
            ("WITH TEST PLUG",), ("6555C79",), ("DESCRIPTION", "30W", "36W"),
            ("AM & VM & SWITCHES", "G62", "G63"),
        )),)
        for text, actual in (("AM & VM & SWITCHES WITHOUT TEST PLUG", "G04"), ("AM & VM & SWITCHES WITH TEST PLUG", "G63")):
            self.assertEqual(QA2.find_grid_group(sheets, "6555C79", {"Width (Wide)": "36"}, description=text)[0], actual)
        self.assertEqual(QA2.find_grid_group(sheets, "6555C79", {"Width (Wide)": "36"})[0], "")
        self.assertEqual(QA2.find_grid_group(sheets, "6555C79", {}, description="AM & VM & SWITCHES WITH TEST PLUG")[0], "")

    @staticmethod
    def wcctaux3_sheet():
        return (("WCCTAUX3", (
            ("WEST COAST UGPS AUX CROSS BUS APPLICATION A",),
            ("REF - 2C11531 (51W STRUCTURE)",),
            ("AMPS", "COPPER", "COPPER"),
            (None, "HEAT TEST", "DENSITY"),
            ("2500/3000", "G02", "G03"),
        )),)

    def test_section_bus_override_drives_independent_group_comparison(self):
        window = SimpleNamespace(
            _parameter_repository=QA2.ParameterRepository(),
            parameter_variables=QA2.ExcelParameterVariables.from_sources([
                "Bus Amps: 3000\nBus Material: Silver Plated Copper\nHeat Test",
            ]),
            _section_bus_requirements=QA2.section_bus_requirements(
                "SECTION 1,2,3,4,5,7&8 NEEDS TO BE DENSITY RATED BUS (998P160)", "outline.dwg",
            ),
            _assembly_links={"2C11531G03": "WCCTAUX3.xlsx", "2C11531G02": "WCCTAUX3.xlsx"},
            _excel_com=SimpleNamespace(get_sheet_values=lambda link: self.wcctaux3_sheet()),
        )
        window._resolved_value = lambda *args: QA2.AuditReviewDemo._resolved_value(window, *args)
        mapping = QA2.parse_code_mappings((
            ("Name", "Code", "Parameter 1", "Parameter 2", "Parameter 3"),
            ("Bus", "WCTAUX5A", "Amperage", "Material", "Heat Test/Density"),
        ))["WCTAUX5A"]
        validate = QA2.AuditReviewDemo.validate_assembly
        for section, style, expected_actual, status in ((8, "2C11531G03", "G03", "PASS"), (6, "2C11531G03", "G02", "FAIL"), (8, "2C11531G02", "G03", "FAIL")):
            with self.subTest(section=section, style=style):
                result = validate(window, style, mapping, "51", "36", (), "", "XBUS,UGPS,25-3000,CU DEN,51W", section)
                self.assertEqual((result.actual_group, result.status), (expected_actual, status))
                if section == 8:
                    self.assertIn("998P160", result.message)
        window._section_bus_requirements[8].add_parameter("Heat Test/Density", "HEAT TEST", "other drawing")
        result = validate(window, "2C11531G03", mapping, "51", "36", (), "", "XBUS,UGPS,25-3000,CU DEN,51W", 8)
        self.assertEqual((result.actual_group, result.status), ("", "UNRESOLVED"))
        self.assertIn("Conflicting section 8", result.message)
        self.assertIsNone(QA2.assembly_bus_requirements(window._section_bus_requirements, 8, "FLAT ROOF,51W"))

    def test_description_selects_cross_bus_table_before_group_parameters(self):
        sheets = [("BUS TABLES", (
            ("SCMM CROSS BUS",),
            ("ASSY DWG", '36" WIDE', '38" WIDE'),
            ("42C7103", "G84", "G85"),
            ("SCMM CROSS BUS MOUNTING",),
            ("ASSY DWG", '36" WIDE', '38" WIDE'),
            ("42C7103", "G14", "G15"),
        ))]
        group, message = QA2.find_grid_group(
            sheets, "42C7103", {"Width (Wide)": "36"},
            description="ASSY, SCMM CROSS BUS, 36W",
        )
        self.assertEqual(group, "G84")
        self.assertIn("SCMM CROSS BUS", message)

    def test_description_expands_mounting_abbreviation(self):
        matches = [
            QA2.GridMatch("Sheet1", (), heading, (), "", group)
            for heading, group in (("WRP CROSS BUS", "G01"), ("WRP CROSS BUS MOUNTING", "G05"))
        ]
        selected, error = QA2.select_grid_tables(matches, {}, "SWBC, WRP CROSS BUS MTG 36W")
        self.assertEqual([match.group for match in selected], ["G05"])
        self.assertEqual(error, "")

    def test_single_mounting_table_needs_no_description_selector(self):
        matches = [QA2.GridMatch("Sheet1", (), "CROSS BUS MOUNTING - AUXILIARY (TOP)", (), "", "G05")]
        self.assertEqual(QA2.select_grid_tables(matches, {}), (matches, ""))

    def test_wrp_full_assembly_columns_keep_height_and_channel_choices(self):
        sheets = [("WRPCBMTG", (
            (None, "WRP CROSS BUS MOUNTING"),
            (None,),
            (None, '90" HIGH STRUCTURE', None, None, '78" HIGH STRUCTURE'),
            (None, "STRUCTURE", "SPLIT CHANNEL", "FULL CHANNEL *", "STRUCTURE", "SPLIT CHANNEL", "FULL CHANNEL *"),
            (None, "WIDTH", "ASSEMBLY", "ASSEMBLY", "WIDTH", "ASSEMBLY", "ASSEMBLY"),
            (None, "36", "6573C44G05", "6573C44G13", "36", "6573C44G22", "6573C44G30"),
        ))]
        wanted = {"Width (Wide)": "36"}
        description = "SWBC, WRP CROSS BUS MTG 36W"
        group, message = QA2.find_grid_group(sheets, "6573C44", wanted, description=description)
        self.assertEqual(group, "")
        self.assertIn("Multiple groups", message)
        wanted.update({"Height": "90", "Channel Type": "SPLIT CHANNEL"})
        self.assertEqual(QA2.find_grid_group(sheets, "6573C44", wanted, description=description)[0], "G05")

    def test_scmm_cross_bus_inherits_only_its_local_parent_headers(self):
        sheet = (
            ("SINGLE PHASE CROSS BUS ASSEMBLY - BUS POSITION 1 ONLY",),
            ("ASSY DWG", "AMPS", "26W"),
            (None, None, "Density"),
            (None, None, "CUAG"),
            ("42C7102", "2500A", "G38"),
            ("THREE PHASE CROSS BUS ASSEMBLY - BUS POSITION 1 ONLY",),
            ("42C7103", "2500A", "G38"),
            (None,),
            ("SINGLE PHASE CROSS BUS ASSEMBLY - BUS POSITION 1 ONLY",),
            ("ASSY DWG", "AMPS", "36W", None),
            (None, None, "Density", "Heat Tested"),
            (None, None, "CUAG", "CUAG"),
            ("42C7102", "2500A", "G84", "G81"),
            ("THREE PHASE CROSS BUS ASSEMBLY - BUS POSITION 1 ONLY",),
            ("42C7103", "2500A", "G84", "G81"),
        )
        wanted = {"Width (Wide)": "36", "Material": "CUAG", "Heat Test/Density": "DENSITY", "Amperage": "2500"}
        self.assertEqual(QA2.find_grid_group([("SCMM", sheet)], "42C7103", wanted, description="ASSY, SCMM CROSS BUS, 36W")[0], "G84")

    def test_ambiguous_cross_bus_descriptions_do_not_use_group_to_pick_table(self):
        matches = [
            QA2.GridMatch("Sheet1", (), heading, (), "", "G05")
            for heading in ("SCMM CROSS BUS", "WRP CROSS BUS")
        ]
        selected, error = QA2.select_grid_tables(matches, {}, "ASSY, CROSS BUS, 36W")
        self.assertEqual(selected, [])
        self.assertIn("Table selection unresolved", error)

    def test_repeated_group_does_not_combine_incompatible_parameter_cells(self):
        sheets = [("SCMM", (
            ("THREE PHASE CROSS BUS ASSEMBLY",),
            ("ASSY DWG", "AMPS", "36W", None),
            (None, None, "Density", "Heat Tested"),
            (None, None, "CUAG", "CUAG"),
            ("42C7103", "1600A", "G80", "G77"),
            (None, "2500A", "G84", "G80"),
        ))]
        wanted = {"Width (Wide)": "36", "Material": "CUAG", "Heat Test/Density": "DENSITY", "Amperage": "2500"}
        self.assertEqual(QA2.find_grid_group(sheets, "42C7103", wanted, description="ASSY, SCMM CROSS BUS, 36W")[0], "G84")

    def test_numeric_top_cover_axes_do_not_borrow_adjacent_drawing_groups(self):
        sheet = (
            ("TYPE 1 TOP COVER", None, None, None, "BUS DUCT TOP COVER"),
            ("STRUCTURE WIDTH", "ASSEMBLY DRAWING", "DEPTH", None, "STRUCTURE WIDTH", "ASSEMBLY DRAWING", "DEPTH"),
            (None, None, 24, 30, None, None, 30),
            (30, "6575C88", "G10", "G16", 30, "42C9673", "G90"),
            (36, None, "G11", "G17", 36, None, "G91"),
        )
        sheets = [("TOPCOVER", sheet)]
        self.assertEqual(QA2.find_grid_group(sheets, "6575C88", {"Width (Wide)": "30", "Depth": "30"})[0], "G16")
        self.assertEqual(QA2.find_grid_group(sheets, "6575C88", {"Width (Wide)": "36", "Depth": "30"})[0], "G17")
        self.assertEqual(QA2.find_grid_group(sheets, "42C9673", {"Width (Wide)": "30", "Depth": "30"})[0], "G90")
        self.assertEqual(QA2.find_grid_group(sheets, "6575C88", {"Width (Wide)": "45", "Depth": "30"})[0], "")

    def test_mda_structure_dimensions_and_standalone_heat_test(self):
        first = [("6575C01G16", "SWBC, 30 X 30 STRUCTURE ASSY 1 0 F NOLOC", "", "")]
        second = [("6575C01G22", "SWBC, 36 X 30 STRUCTURE ASSY 1 0 F NOLOC", "", "")]
        self.assertEqual(QA2.structure_dimensions_of(first), ("30", "30"))
        self.assertEqual(QA2.structure_dimensions_of(second), ("36", "30"))
        self.assertEqual(QA2.structure_width_of(second), "36")
        variables = QA2.ExcelParameterVariables.from_sources([
            "Bus Amps: 1200\nBus Material: Copper    Heat Test\nEnclosure Type: Type 1",
        ])
        self.assertEqual(variables.Heat_Test_Density, "HEAT TEST")
        self.assertEqual(variables.Material, "COPPER")

    def test_group_inputs_use_saved_common_variables_without_guessing_conflicts(self):
        window = SimpleNamespace(
            _parameter_repository=QA2.ParameterRepository(),
            parameter_variables=QA2.ExcelParameterVariables.from_sources([
                "Bus Amps: 1200\nBus Material: Copper",
            ]),
        )
        resolve = QA2.AuditReviewDemo._resolved_value
        self.assertEqual(resolve(window, "6575C20G06", "Amperage", "", ""), "1200")
        self.assertEqual(resolve(window, "6575C20G06", "Material", "", ""), "COPPER")
        window.parameter_variables = QA2.ExcelParameterVariables({"Amperage": ["1200", "800"]})
        self.assertEqual(resolve(window, "6575C20G06", "Amperage", "", ""), "")
        window.parameter_variables = QA2.ExcelParameterVariables({"Material": ["COPPER"]})
        window._parameter_repository.add_parameter("Material", "Copper", "first")
        window._parameter_repository.add_parameter("Material", "Aluminum", "second")
        self.assertEqual(resolve(window, "6575C20G06", "Material", "", ""), "")

    def test_section_bus_requirements_keep_scope_and_source(self):
        note = "2. HAVE A QUICK CHECK OF BUSES & SPLICES IN SECTION 1,2,3,4,5,7&8 IN DA BOM. IT NEEDS TO BE DENSITY RATED BUS AS PER NEW GUIDELINES (998P160)"
        requirements = QA2.section_bus_requirements(note, "outline.dwg")
        self.assertEqual(set(requirements), {1, 2, 3, 4, 5, 7, 8})
        self.assertEqual(requirements[8].get_parameter("Heat Test/Density"), "DENSITY")
        self.assertIn("998P160", requirements[8].get_entries("Heat Test/Density")[0][1])
        self.assertEqual(QA2.section_bus_requirements("Bus Material: Copper Heat Test", "QuickPrint"), {})
        ranged = QA2.section_bus_requirements("STRUCTURES 2-4 AND 8 MUST USE HEAT TESTED BUS", "drawing")
        self.assertEqual(set(ranged), {2, 3, 4, 8})
        self.assertEqual(ranged[8].get_parameter("Heat Test/Density"), "HEAT TEST")

    def test_saved_bus_position_selects_table_before_width(self):
        sheet = (
            ("CROSS BUS MOUNTING - AUXILIARY (TOP)",),
            ("ASSY DWG", '36" WIDE', '38" WIDE'),
            ("42C2066", "G05", "G06"),
            ("CROSS BUS MOUNTING - AUXILIARY (TOP) AND STANDARD (BOTTOM)",),
            ("ASSY DWG", '36" WIDE', '38" WIDE'),
            ("42C2066", "G14", "G15"),
        )
        sheets = [("AUXILCB", sheet)]
        saved = QA2.ExcelParameterVariables.from_sources(["Bus Position: Auxiliary (Top)"])
        self.assertEqual(saved.Bus_Position, "AUXILIARY (TOP)")
        group, message = QA2.find_grid_group(sheets, "42C2066", {"Width (Wide)": "36"}, saved.values)
        self.assertEqual(group, "G05")
        self.assertIn("AUXILIARY (TOP)", message)
        self.assertEqual(QA2.find_grid_group(sheets, "42C2066", {"Width (Wide)": "45"}, saved.values)[0], "")
        for selectors in ({}, {"Bus Position": ["Bottom"]}, {"Bus Position": ["Auxiliary (Top)", "Bottom"]}):
            group, message = QA2.find_grid_group(sheets, "42C2066", {"Width (Wide)": "36"}, selectors)
            self.assertEqual(group, "")
            self.assertIn("Table selection unresolved", message)
        combined = {"Part Name": ["Auxiliary (Top) and Standard (Bottom)"]}
        self.assertEqual(QA2.find_grid_group(sheets, "42C2066", {"Width (Wide)": "36"}, combined)[0], "G14")
        combined_saved = QA2.ExcelParameterVariables.from_sources([
            "Bus Position: Auxiliary (Top) and Standard (Bottom)",
        ])
        self.assertEqual(QA2.find_grid_group(sheets, "42C2066", {"Width (Wide)": "36"}, combined_saved.values)[0], "G14")
        window = SimpleNamespace(
            parameter_variables=saved,
            _resolved_value=lambda style, name, width, depth, section=None: width if name == "Width (Wide)" else "",
            _assembly_links={"42C2066G14": "AUXMAMTG.xlsx"},
            _excel_com=SimpleNamespace(get_sheet_values=lambda link: sheets),
        )
        mapping = QA2.parse_code_mappings((("Name", "Code", "Parameter"), ("", "AUXMAMTG", "Width (Wide)")))["AUXMAMTG"]
        result = QA2.AuditReviewDemo.validate_assembly(window, "42C2066G14", mapping, "36", "", (), "")
        self.assertEqual((result.actual_group, result.status), ("G05", "FAIL"))

    @staticmethod
    def wcuto22xmain_sheet():
        def row(*cells):
            return tuple(cells) + (None,) * (7 - len(cells))

        return (
            row(None, '400A THRU 1000A'),
            row(None, 'WEST COAST UTILITY COMPARTMENT'),
            row(None, 'BUS TO 22X CHASSIS MAIN (36"W)'),
            row(None, '9" BUS CENTERS (2C11905, 2C11907)'),
            row(None, 'THREE PHASE'),
            row(None, 'CURRENT', 'ASSEMBLY', 'ALUM', 'CUAG', 'CUSN'),
            row(None, 'RATING', 'DRAWING', 'DENSITY', 'DENSITY', 'DENSITY'),
            row(None, '400A, K/PD3(400A) FRAME *', '42C5490', 'G01', 'G02', 'G03'),
            row(None, '800A, MD/PD4 FRAME', '42C5495', 'G01', 'G02', 'G03'),
            row(None, '1000A, N/PD5 FRAME', '42C5496', 'G01', 'G02', 'G03'),
            row(None, 'SINGLE PHASE'),
            row(None, 'CURRENT', 'ASSEMBLY', 'ALUM', 'CUAG', 'CUSN'),
            row(None, 'RATING', 'DRAWING', 'DENSITY', 'DENSITY', 'DENSITY'),
            row(None, '800A, MD/PD4 FRAME', '42C5495', 'G04', 'G05', 'G06'),
        )

    def test_grid_layout_gives_material_density_amperage_and_width(self):
        matches = QA2.find_grid_group_matches(self.wcuto22xmain_sheet(), "400A - 1000A (36W)", "42C5495", "G02")
        self.assertEqual(len(matches), 1)
        self.assertEqual(QA2.grid_parameter_values(matches[0]), {
            "Material": "CUAG", "Heat Test/Density": "DENSITY", "3/1 Phase": "3",
            "Amperage": "800", "Width (Wide)": "36",
        })
        single = QA2.find_grid_group_matches(self.wcuto22xmain_sheet(), "400A - 1000A (36W)", "42C5495", "G05")
        self.assertEqual(QA2.grid_parameter_values(single[0])["3/1 Phase"], "1")

    def test_actual_group_comes_from_description_values(self):
        sheets = [("400A - 1000A (36W)", self.wcuto22xmain_sheet())]
        wanted = QA2.description_parameter_values("800A CUAG DEN 3PH 1 0 F NOLOC")
        self.assertEqual(QA2.find_grid_group(sheets, "42C5495", wanted)[0], "G02")
        single = QA2.description_parameter_values("800A CUAG DEN 1PH")
        self.assertEqual(QA2.find_grid_group(sheets, "42C5495", single)[0], "G05")
        group, message = QA2.find_grid_group(sheets, "42C5495", {"Heat Test/Density": "DENSITY"})
        self.assertEqual(group, "")
        self.assertIn("Multiple groups", message)
        self.assertEqual(QA2.description_parameter_values("XBUS,UGPS,25-3000,CU DEN,51W")["Amperage"], "2500-3000")
        self.assertTrue(QA2._grid_value_matches("CUAG", "CU"))
        self.assertTrue(QA2._grid_value_matches("2500/3000", "2500-3000"))
        self.assertFalse(QA2._grid_value_matches("4000", "2500-3000"))

    def test_sheet_is_chosen_by_structure_width(self):
        narrow = QA2.find_grid_group_matches(self.wcuto22xmain_sheet(), "400A - 1000A (36W)", "42C5495", "G02")
        wide = QA2.find_grid_group_matches(self.wcuto22xmain_sheet(), "400A - 1000A (45W)", "42C5495", "G02")
        self.assertEqual(QA2.merge_grid_matches(narrow + wide, "36")["Width (Wide)"], "36")
        self.assertEqual(QA2.merge_grid_matches(narrow + wide, "45")["Width (Wide)"], "45")
        self.assertEqual(QA2.find_grid_group_matches(self.wcuto22xmain_sheet(), "x", "42C9999", "G02"), [])

    def test_merged_headings_and_base_row_above_the_table(self):
        sheet = (
            (None, "1200A THRU 4000A"),
            (None, "1200-2000A"),
            (None, "2C11909", "38 W  x  30D MIN", None, None, None, None, None, "45 W  x  30D MIN"),
            (None, None, "HEAT TEST", None, None, "DENSITY", None, None, "HEAT TEST"),
            (None, "AMP RATING", "CUAG", "AL", "CUSN", "CUAG", "AL", "CUSN", "CUAG"),
            (None, "1200A", "G08", "G02", "G18", "G09", "G03", "G19", "N/A"),
        )
        found = QA2.merge_grid_matches(QA2.find_grid_group_matches(sheet, "Sheet1", "2C11909", "G03"), "38")
        self.assertEqual(found, {
            "Material": "AL", "Heat Test/Density": "DENSITY", "Amperage": "1200",
            "Width (Wide)": "38", "Depth": "30D MIN",
        })
        self.assertEqual(QA2.find_grid_group_matches(sheet, "Sheet1", "2C11999", "G03"), [])

    def test_ugps_row_gives_range_drillings_and_depth(self):
        sheet = (
            (None, "WEST COAST UTILITY"),
            (None, "DRAWING", "AMPACITY", "WIDTH", "BUS", "30D", "36D", "48D"),
            (None, None, None, None, None, "(8 DRILLINGS)", "(10 DRILLINGS)", "(13 DRILLINGS)"),
            (None, "2C11518", "2500-3000A", "51W", "CU DEN", "G02", "G03", "G04"),
        )
        found = QA2.merge_grid_matches(QA2.find_grid_group_matches(sheet, "UGPS", "2C11518", "G03"))
        self.assertEqual(found["Number of Drillings"], "10")
        self.assertEqual(found["Depth"], "36D")
        self.assertEqual(found["Amperage"], "2500-3000")
        self.assertEqual((found["Material"], found["Heat Test/Density"], found["Width (Wide)"]), ("CU", "DENSITY", "51"))

    def test_structure_width_comes_from_the_structure_assembly_row(self):
        items = [
            ("42C5495G02", "800A CUAG DEN 3PH", "", ""),
            ("43C1696G52", 'ASSY.SEISMIC STRUCTURE, 36"W X', "", ""),
        ]
        self.assertEqual(QA2.structure_width_of(items), "36")

    def test_excel_parameters_become_variables_named_after_them(self):
        self.assertEqual(QA2.excel_variable_name("Width (Wide)"), "Width_Wide")
        self.assertEqual(QA2.excel_variable_name("3/1 Phase"), "_3_1_Phase")
        self.assertEqual(
            QA2.excel_variable_name("Service/Non-Service Enterance"), "Service_Non_Service_Enterance",
        )

        text = (
            "480Y/277V 3-PHASE 4-WIRE Bus Amps: 1600 Bus Material: Silver Plated Copper "
            "Heat Test: 65 C Main Breaker - Ind Mtd-3000A 3P (8) EUSERC Press Bolts "
            "TYPE 3R (NONWALK-IN) FLAT ROOF"
        )
        variables = QA2.ExcelParameterVariables.from_sources([text], ["43C1696G50"])

        self.assertEqual(variables.Amperage, "1600")
        self.assertEqual(variables.Material, "SILVER PLATED COPPER")
        self.assertEqual(variables.Voltage, "480Y/277V")
        self.assertEqual(variables._3_1_Phase, "3")
        self.assertEqual(variables.Heat_Test_Density, "65 C")
        self.assertEqual(variables.Number_of_Drillings, "8")
        self.assertEqual(variables.Flat_Roof_Sloped_Roof, "FLAT")
        self.assertEqual(variables.Assembly_No, "43C1696G50")
        self.assertEqual(variables.Height, "")
        self.assertEqual(variables.values["Poles"], ["3"])
        for name in QA2.EXCEL_PARAMETERS:
            self.assertTrue(hasattr(variables, QA2.excel_variable_name(name)), name)
    def test_repeated_header_row_and_xlsx_codes_in_sheet1(self):
        rows = (
            ("Name", "Code", "Parameter 1"),
            (None, "Code", "Parameter 1"),
            (None, "IFSXFMRSEL 225-300kVA.xlsx", "Depth"),
        )
        mappings = QA2.parse_code_mappings(rows)
        self.assertEqual(list(mappings), ["IFSXFMRSEL225300KVA"])
        self.assertEqual(QA2.normalize_code("IFSXFMRSEL 225-300kVA"), "IFSXFMRSEL225300KVA")
    def test_common_parameters_and_code_wise_rows(self):
        common = (("Parameters",), ("Amperage",), ("Part Name",), ("Part Name",), (None,), ("Width (Wide)",))
        self.assertEqual(QA2.parse_common_parameters(common), ("Amperage", "Part Name", "Width (Wide)"))

        rows = (
            ("Name", "Code", "Parameter 1"),
            ("Pull", "CABLE PULL", "Amperage", "Part Name", "A", "B", "C", "Sixth"),
        )
        mapping = QA2.parse_code_mappings(rows)["CABLEPULL"]
        self.assertEqual(mapping.parameters, ("Amperage", "Part Name", "A", "B", "C", "Sixth"))
        self.assertEqual(mapping.aliases, mapping.parameters)

        variables = QA2.ExcelParameterVariables.from_sources(["Bus Amps: 1600"], names=("Amperage", "Part Name"))
        self.assertEqual(tuple(variables.values), ("Amperage", "Part Name"))
        self.assertEqual(variables.Amperage, "1600")

    def test_values_come_from_the_row_of_base_and_group(self):
        values = (
            ("8708C64 ASSEMBLY", None, None),
            ("Ampacity", "Part Name", "Group"),
            (800, "Splice A", "G01"),
            (1200, "Splice B", "G02"),
        )
        tables = QA2.WorkbookParser.extract_tables(values, "SPLICE")
        found = QA2.TableValueExtractor.extract(
            tables, "8708C64", "G02",
            {"Amperage": ("Amperage", "Ampacity"), "Part Name": ("Part Name",), "Depth": ("Depth",)},
        )
        self.assertEqual(found, {"Amperage": "1200", "Part Name": "Splice B"})


class IndexMappingTests(unittest.TestCase):
    def test_local_map_uses_only_column_k_and_accepts_single_base(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "INDEX2"
        sheet["D3"] = "DECOY1, DECOY2"
        sheet["E3"].hyperlink = "REARCOVER.xlsx"
        sheet["K3"] = "43C1673"
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "Index2.xlsx")
            workbook.save(path)
            self.assertEqual(QA2.build_index2_map(path), {"43C1673": "REARCOVER.xlsx"})
        workbook.close()

    def test_live_map_handles_used_range_offset_and_single_base(self):
        sheet = SimpleNamespace(
            UsedRange=SimpleNamespace(
                Row=10, Column=5, Columns=SimpleNamespace(Count=7),
                Value=(("DECOY1, DECOY2", None, None, None, None, None, "43C1673"),),
            ),
            Hyperlinks=[SimpleNamespace(
                Range=SimpleNamespace(Row=10, Column=5), Address="REARCOVER.xlsx",
            )],
        )
        self.assertEqual(QA2.build_index2_map_from_com(sheet), {"43C1673": "REARCOVER.xlsx"})


class FamilyLayoutTests(unittest.TestCase):
    def test_renamed_migration_requires_one_common_indexed_root_destination(self):
        session = QA2.ExcelComSession()
        session._replacement_links = {}
        session._assembly_targets["OLD.XLSX"] = {"42C5495", "42C5496"}
        session._index2_catalog = (
            QA2.Index2RootEntry("42C5495", 1, "OLD", "OLD.xlsx"),
            QA2.Index2RootEntry("42C5495", 2, "NEW", "NEW.xlsx"),
            QA2.Index2RootEntry("42C5496", 2, "NEW", "NEW.xlsx"),
            QA2.Index2RootEntry("42C5496", 3, "OTHER", "OTHER.xlsx"),
        )
        self.assertEqual(session._replacement_link("Renamed.xls", "retired", "OLD.xlsx"), "NEW.xlsx")
        session._index2_catalog += (QA2.Index2RootEntry("42C5495", 3, "OTHER", "OTHER.xlsx"),)
        with self.assertRaisesRegex(RuntimeError, "no unique Index2 link"):
            session._replacement_link("Renamed.xls", "retired", "OLD.xlsx")
        session._assembly_targets["OLD.XLSX"].add("42C5497")
        with self.assertRaisesRegex(RuntimeError, "no unique Index2 link"):
            session._replacement_link("Renamed.xls", "retired", "OLD.xlsx")

    def setUp(self):
        capture = patch.object(QA2.ExcelComSession, "_capture_family", QA2.ExcelComSession._capture_live_family)
        capture.start()
        self.addCleanup(capture.stop)

    def test_excel_values_normalize_to_rectangular_rows(self):
        self.assertEqual(QA2.normalize_excel_rows("single cell"), (("single cell",),))
        self.assertEqual(QA2.normalize_excel_rows(("one", "two")), (("one", "two"),))
        self.assertEqual(
            QA2.normalize_excel_rows((("one",), ("two", "three"))),
            (("one", None), ("two", "three")),
        )

    def test_all_active_sheets_use_the_same_normalized_row_format(self):
        workbook = SimpleNamespace(Name="FAMILY.xlsx", Worksheets=[
            excel_sheet("Single Cell", "single cell"),
            excel_sheet("Single Row", ("Assembly", "Group")),
            excel_sheet("Two Dimensional", (("Assembly", "Group"), ("43C1673", "G07"))),
            SimpleNamespace(Name="Obsolete__Old", UsedRange=SimpleNamespace(Value="ignored")),
        ])
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", return_value=workbook):
            sheets = session.get_sheet_values("FAMILY.xlsx")

        self.assertEqual(sheets, (
            ("Single Cell", (("single cell",),)),
            ("Single Row", (("Assembly", "Group"),)),
            ("Two Dimensional", (("Assembly", "Group"), ("43C1673", "G07"))),
        ))

    def test_migration_uses_unique_index2_link_and_replacement_data(self):
        old_rows = (("Will be maintained on spreadsheet: ENCLOSURE.xls (Index 377)",),)
        current_rows = (("CURRENT DATA",),)
        old = SimpleNamespace(Name="BASICSTR.xlsx", Worksheets=[excel_sheet("Old", old_rows)])
        current = SimpleNamespace(Name="Enclosure.xlsx", Worksheets=[excel_sheet("Basic Structure - Type 1", current_rows)])
        sheet = SimpleNamespace(Hyperlinks=[SimpleNamespace(
            Range=SimpleNamespace(Column=5, Value="ENCLOSURE"), Address="Enclosure.xlsx",
        )])
        index = SimpleNamespace(Worksheets=lambda number: sheet)
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", side_effect=lambda link: old if link == "BASICSTR.xlsx" else current), \
                patch.object(session, "_get_index2_workbook", return_value=index):
            rows = session.get_sheet_values("BASICSTR.xlsx")
        self.assertEqual(rows, (("Enclosure.xlsx / Basic Structure - Type 1", current_rows),))
        self.assertNotIn("Old", rows[0][0])

    def test_empty_and_circular_replacement_workbooks_are_rejected(self):
        empty = SimpleNamespace(Name="Empty.xlsx", Worksheets=[excel_sheet("Empty", None)])
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", return_value=empty):
            with self.assertRaisesRegex(RuntimeError, "no readable active table data"):
                session.get_group_tables("Empty.xlsx")
        rows = (("Will be maintained on spreadsheet: CYCLE.xls (Index 1)",),)
        cycle = SimpleNamespace(Name="CYCLE.xlsx", Worksheets=[excel_sheet("Old", rows)])
        sheet = SimpleNamespace(Hyperlinks=[SimpleNamespace(Range=SimpleNamespace(Column=5, Value="CYCLE"), Address="CYCLE.xlsx")])
        with patch.object(session, "get_family_workbook", return_value=cycle), \
                patch.object(session, "_get_index2_workbook", return_value=SimpleNamespace(Worksheets=lambda number: sheet)):
            with self.assertRaisesRegex(RuntimeError, "cycle"):
                session.get_group_tables("CYCLE.xlsx")

    def test_shared_replacement_recaptures_new_family_roots(self):
        session = QA2.ExcelComSession()
        captures = []

        def capture(link):
            workbook = QA2.Workbook()
            sheet = workbook.active
            if link == "CURRENT.xlsx":
                roots = tuple(sorted(session._assembly_targets[link.upper()]))
                captures.append(roots)
                sheet.append(["Assembly Drawing", "Group", "Width"])
                for root in roots:
                    sheet.append([root, "G02", 30])
            else:
                sheet.append(["Will be maintained on spreadsheet: CURRENT.xls (Index 1)"])
            return link, ((sheet.title, tuple(sheet.values)),), workbook

        with patch.object(session, "_capture_family", side_effect=capture), \
                patch.object(session, "_replacement_link", return_value="CURRENT.xlsx"):
            session.get_group_tables("OLD1.xlsx", assembly_base="42C1084")
            tables = session.get_group_tables("OLD2.xlsx", assembly_base="42C1085")
        self.assertEqual(captures, [("42C1084",), ("42C1084", "42C1085")])
        self.assertTrue(any(table.rows_for_base("42C1085") for table in tables))

    def test_smart_root_invariants_do_not_block_group_selection(self):
        sheets = (("Structure - Type 1", (
            ("INPUT", None, None, "OUTPUT"),
            ("Structure Width", "NG Drawout", "Seismic", "Assembly"),
            (30, "No", "Non-Seismic", "6575C01G16"),
            (36, "No", "Non-Seismic", "6575C01G22"),
        )),)
        self.assertEqual(QA2.find_grid_group(sheets, "6575C01", {"Width (Wide)": 30})[0], "G16")
        self.assertEqual(QA2.find_grid_group(sheets, "6575C01", {"Width (Wide)": 36})[0], "G22")
        self.assertEqual(QA2.find_grid_group(sheets, "6575C01", {"Width (Wide)": 30, "NG Drawout": "Yes"})[0], "")
        self.assertEqual(QA2.find_grid_group(sheets, "6575C01", {})[0], "")

    def test_missing_varying_smart_selector_cannot_choose_a_group(self):
        sheets = (("Structure", (
            ("INPUT", None, "OUTPUT"),
            ("Structure Width", "NG Drawout", "Assembly"),
            (30, "No", "6575C01G16"),
            (30, "Yes", "6575C01G17"),
        )),)
        group, message = QA2.find_grid_group(sheets, "6575C01", {"Width (Wide)": 30})
        self.assertEqual(group, "")
        self.assertIn("NG Drawout", message)

    def test_smart_text_alternatives_accept_only_listed_values(self):
        value = "NON-SEISMIC OR SEISMIC (FREESTANDING)"
        self.assertTrue(QA2._grid_value_matches(value, "NON-SEISMIC"))
        self.assertTrue(QA2._grid_value_matches(value, "SEISMIC (FREESTANDING)"))
        self.assertFalse(QA2._grid_value_matches(value, "SEISMIC (WALL MOUNTED)"))

    def test_generic_aluminum_matches_tin_plated_aluminum_not_copper(self):
        for compare in (QA2._grid_value_matches, QA2.ClassifierWorkbookParser.values_match):
            with self.subTest(compare=compare.__qualname__):
                self.assertTrue(compare("AL", "TIN PLATED ALUMINUM"))
                self.assertTrue(compare("ALUM", "TIN PLATED ALUMINUM"))
                self.assertFalse(compare("CU", "TIN PLATED ALUMINUM"))
                self.assertFalse(compare("TIN PLATED ALUMINUM", "TIN PLATED COPPER"))
                self.assertFalse(compare("TIN PLATED ALUMINUM", "SILVER PLATED ALUMINUM"))

    def test_shared_enclosure_sheet_uses_row_enclosure_not_first_title_type(self):
        rows = (
            ("INPUT", None, "OUTPUT"),
            ("Structure Type", "Structure Width", "Rear Cover"),
            ("Type 3R", 36, "43C1673G07"),
        )
        matches = QA2.find_smart_table_matches(rows, "Rear Covers-Type 1 & 3R", "43C1673", "")
        self.assertEqual(dict(matches[0].axis_values)["NEMA 1/3R"], "3R")
        self.assertEqual(QA2.find_grid_group((("Rear Covers-Type 1 & 3R", rows),),
                                           "43C1673", {"Width (Wide)": 36, "NEMA 1/3R": "3R"})[0], "G07")

    def test_smart_table_does_not_reuse_previous_table_headers(self):
        rows = (
            ("INPUT", None, "OUTPUT"),
            ("Depth", "Description #1", "Assembly"),
            (24, "Lifting assembly", "8708C64G02"),
            (None,),
            ("INPUT", None, "OUTPUT"),
            ("Structure Width", "Structure Type", "Assembly"),
            (36, "Type 3R", "43C1673G07"),
        )
        matches = QA2.find_smart_table_matches(rows, "Misc Tables", "43C1673", "")
        self.assertEqual(len(matches), 1)
        self.assertEqual(dict(matches[0].axis_values)["Width (Wide)"], "36")
        self.assertNotIn("Description #1", dict(matches[0].axis_values))

    def test_end_cap_outputs_keep_their_position_and_structure_material(self):
        rows = (
            ("INPUT", None, None, "OUTPUT", None, None, "REFERENCE INFORMATION"),
            ("Roof Type", "Material", "Structure Depth", "Roof End Cap - Left Side",
             "Roof End Cap - Right Side", "Intermediate Roof End Cap", "Notes"),
            ("Flat Roof", "Galvanized", "24 Only", "43C1682G01", "43C1682G02", "43C1682G03"),
        )
        sheets = (("Misc Tables", rows),)
        matches = QA2.find_smart_table_matches(rows, "Misc Tables", "43C1682", "")
        self.assertEqual(dict(matches[0].axis_values)["Structure Material"], "GALVANIZED")
        for description, expected in (("END CAP 24D LH", "G01"), ("END CAP 24D RH", "G02"),
                                      ("END CAP 24D CTR", "G03")):
            wanted = {**QA2.description_parameter_values(description), "Depth": 24, "Material": "TIN PLATED ALUMINUM"}
            self.assertEqual(QA2.find_grid_group(sheets, "43C1682", wanted)[0], expected)
        self.assertEqual(QA2.find_grid_group(sheets, "43C1682", {"Depth": 24})[0], "")

    def test_seismic_all_mounting_modes_accepts_generic_requirement_only_when_covered(self):
        all_modes = "SEISMIC (FREESTANDING OR WALL SECURED)"
        self.assertTrue(QA2._grid_value_matches(all_modes, "SEISMIC"))
        self.assertTrue(QA2._grid_value_matches(all_modes, "SEISMIC (FREESTANDING)"))
        self.assertFalse(QA2._grid_value_matches(all_modes, "NON-SEISMIC"))
        self.assertFalse(QA2._grid_value_matches("SEISMIC (FREESTANDING)", "SEISMIC"))

    def test_seismic_hardware_is_not_a_global_structure_requirement(self):
        text = "ASSY, MOUNTING SEISMIC ANCHOR\nASSY, SEISMIC BOLT PIPE SPACER\nSEISMIC BRACKET"
        self.assertEqual(QA2.extract_excel_parameter_values(text)["Seismic/Non-Seismic"], [])
        for source, expected in (("Seismic: Yes", "SEISMIC"), ("Seismic: No", "NON-SEISMIC"),
                                 ("Structure Rating: Non-Seismic", "NON-SEISMIC"), ("SEISMIC", "SEISMIC")):
            self.assertEqual(QA2.extract_excel_parameter_values(source)["Seismic/Non-Seismic"], [expected])
        self.assertEqual(QA2.description_parameter_values("ASSY.SEISMIC STRUCTURE")["Seismic/Non-Seismic"], "SEISMIC")

    def test_ambiguous_migration_does_not_reuse_old_table(self):
        rows = (("Will be maintained on spreadsheet: NEW.xls (Index 1)",), ("ASSY DWG", '36" WIDE'), ("42C2066", "G05"))
        old = SimpleNamespace(Name="OLD.xlsx", Worksheets=[excel_sheet("Old", rows)])
        session = QA2.ExcelComSession()
        index = SimpleNamespace(Worksheets=lambda number: SimpleNamespace(Hyperlinks=[]))
        with patch.object(session, "get_family_workbook", return_value=old), patch.object(session, "_get_index2_workbook", return_value=index):
            with self.assertRaisesRegex(RuntimeError, "no unique Index2 link"):
                session.get_group_tables("OLD.xlsx")

    def lifting_table(self, depth="24D-48D", base="8708C64"):
        return (
            (f"{base} ASSEMBLY", None, None),
            (None, None, None),
            ("Depth", None, "Group"),
            (depth, None, "G02"),
        )

    def test_rear_cover_width_is_not_structure_height(self):
        values = (
            ("Structure Height", "Structure Width", "Assembly", "Group"),
            (90, 9, "43C1673", "G01"),
            (None, 36, "43C1673", "G07"),
        )
        result = QA2.table_dimension_lookup(values, "REARCOVER", "43C1673", "G07")
        self.assertEqual((result.width, result.depth), ("36", ""))

    def test_lifting_assembly_preserves_depth_range(self):
        result = QA2.table_dimension_lookup(self.lifting_table(), "WPFLAT", "8708C64", "G02")
        self.assertEqual((result.width, result.depth), ("", "24D-48D"))

    def test_partial_base_match_is_rejected(self):
        result = QA2.table_dimension_lookup(
            self.lifting_table(base="8708C640"), "WPFLAT", "8708C64", "G02",
        )
        self.assertIsNone(result)

    def test_missing_group_does_not_borrow_dimensions(self):
        result = QA2.table_dimension_lookup(self.lifting_table(), "WPFLAT", "8708C64", "G99")
        self.assertEqual((result.width, result.depth), ("", ""))
        self.assertIn("not supported", result.message)

    def test_bus_properties_are_not_width_and_depth(self):
        values = (
            ("6575C43 - 3PH 3W", None, None, None),
            (None, None, None, None),
            ("OFFSET", "BARS PER PHASE", '3" BUS SIZE', None),
            (None, None, "GROUPS", "AMP RATING"),
            (None, None, None, "HT"),
            (None, None, None, None),
            ('6"', 1, "G02", 800),
        )
        result = QA2.table_dimension_lookup(values, "BUSSPLIC", "6575C43", "G02")
        self.assertEqual((result.width, result.depth), ("", ""))
        self.assertIn("not Width/Depth", result.message)

    def test_conflicting_active_sheets_are_unresolved(self):
        workbook = SimpleNamespace(Name="WPFLAT.xlsx", Worksheets=[
            excel_sheet("WPFLAT-Steel", self.lifting_table()),
            excel_sheet("WPFLAT-Galvanized Steel", self.lifting_table("30D-48D")),
        ])
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", return_value=workbook):
            result = session.lookup_dimensions("WPFLAT.xlsx", "BASICSTR-NEMA 3R", "8708C64", "G02")
        self.assertEqual((result.width, result.depth), ("", ""))
        self.assertIn("selection unresolved", result.message)

    def test_equal_active_sheets_agree_and_obsolete_sheet_is_ignored(self):
        workbook = SimpleNamespace(Name="WPFLAT.xlsx", Worksheets=[
            excel_sheet("WPFLAT-Steel", self.lifting_table()),
            excel_sheet("WPFLAT-Galvanized Steel", self.lifting_table()),
            SimpleNamespace(Name="Obsolete__Old", UsedRange=SimpleNamespace(Value=self.lifting_table("30D-48D"))),
        ])
        session = QA2.ExcelComSession()
        with patch.object(session, "get_family_workbook", return_value=workbook):
            result = session.lookup_dimensions("WPFLAT.xlsx", None, "8708C64", "G02")
        self.assertEqual((result.width, result.depth), ("", "24D-48D"))
        self.assertIn("WPFLAT-Steel", result.source)
        self.assertIn("WPFLAT-Galvanized Steel", result.source)
        self.assertNotIn("Obsolete", result.source)


class ChecklistConfigurationTests(unittest.TestCase):
    def test_checklist_summary_reports_found_and_missing_parameters(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "checklist.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "MCOE Checklist"
            for column, label in enumerate(["REARCOVER", "Assembly No.", "Width (Wide)", "Depth", "Material"], 2):
                sheet.cell(2, column, label)
            workbook.save(path)
            workbook.close()

            configuration = QA2.ConfigurationRepository(path)
            parameters = QA2.ParameterRepository()
            parameters.add_parameter("Material", "Steel", "bom.pdf")
            window = SimpleNamespace(
                _assembly_families={"43C1673G07": "REARCOVER"},
                _configuration_repository=configuration,
                _parameter_repository=parameters,
                enclosure_type="3R",
            )
            summary, detail = QA2.AuditReviewDemo.checklist_summary(window, "43C1673G07", "36", "")

            self.assertEqual(summary, "REARCOVER: 3/4 parameters found")
            self.assertIn("Missing: Depth", detail)
            self.assertIn("Material = Steel", detail)

            window._assembly_families = {}
            self.assertEqual(
                QA2.AuditReviewDemo.checklist_summary(window, "6575C43G02", "", "")[0],
                "Family not resolved from Index2.",
            )

    def test_link_file_name_handles_urls_and_relative_links(self):
        self.assertEqual(QA2.link_file_name("https://x.sharepoint.com/a%20b/REARCOVER.xlsx"), "REARCOVER.xlsx")
        self.assertEqual(QA2.link_file_name("sub\\WPFLAT.xlsx"), "WPFLAT.xlsx")

    def test_absolute_index2_link_is_opened_without_prefixing_folder(self):
        url = "https://eaton.sharepoint.com/sites/Team/Spreadsheets/REARCOVER.xlsx"
        opened = []
        workbook = SimpleNamespace(Name="REARCOVER.xlsx")
        app = SimpleNamespace(
            Workbooks=SimpleNamespace(Open=lambda path, **_kwargs: opened.append(path) or workbook),
        )
        index2 = SimpleNamespace(Path="https://eaton.sharepoint.com/sites/Team/Spreadsheets")
        session = QA2.ExcelComSession()
        with patch.object(session, "_get_app", return_value=app), \
                patch.object(session, "_find_open_workbook", return_value=None), \
                patch.object(session, "_get_index2_workbook", return_value=index2):
            self.assertIs(session.get_family_workbook(url), workbook)
        self.assertEqual(opened, [url])

    def create_checklist(self, path):
        from openpyxl import Workbook

        workbook = Workbook()
        checklist = workbook.active
        checklist.title = "MCOE Checklist"
        database = workbook.create_sheet("Database_Index2")
        checklist["A2"] = "Test Category"
        checklist["B2"] = "TESTCODE"
        checklist["C2"] = "Amperage"
        checklist["D2"] = "Material"
        database["B2"] = "TESTCODE"
        database["C2"] = "Ampacity"
        database["D2"] = "Metal"
        workbook.save(path)
        workbook.close()

    def test_configuration_loads_parameter_aliases(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "checklist.xlsx")
            self.create_checklist(path)
            repository = QA2.ConfigurationRepository(path)

            record = repository.get_configuration("testcode")
            self.assertEqual(record.category, "Test Category")
            self.assertEqual(record.parameters[0].parameter_name, "Amperage")
            self.assertEqual(record.parameters[0].keywords, ("Amperage", "Ampacity"))

    def test_parameter_sources_preserve_and_flag_conflicting_values(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "checklist.xlsx")
            self.create_checklist(path)
            configuration = QA2.ConfigurationRepository(path)
            repository = QA2.ParameterRepository()
            repository.merge_sources(
                configuration.extract_parameters("Amperage: 800A", "bom.pdf"),
                configuration.extract_parameters("Ampacity: 800A", "drawing.pdf"),
            )

            self.assertEqual(repository.get_parameter("Amperage"), "800A")
            self.assertEqual(repository.conflict_count, 0)

            repository.merge_sources(
                configuration.extract_parameters("Ampacity: 1000A", "quickprint.pdf"),
            )
            self.assertIsNone(repository.get_parameter("Amperage"))
            self.assertEqual(repository.conflict_count, 1)
            self.assertEqual(repository.value_count, 3)

    def test_checklist_summary_includes_document_evidence(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "checklist.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "MCOE Checklist"
            for column, label in enumerate(["REARCOVER", "Material"], 2):
                sheet.cell(2, column, label)
            workbook.save(path)
            workbook.close()

            configuration = QA2.ConfigurationRepository(path)
            parameters = QA2.ParameterRepository()
            parameters.add_parameter("Material", "Copper", "drawing.pdf (page 2, OCR)")
            parameters.add_parameter("Material", "Steel", "assembly.dwg")
            window = SimpleNamespace(
                _assembly_families={"43C1673G07": "REARCOVER"},
                _configuration_repository=configuration,
                _parameter_repository=parameters,
                enclosure_type="3R",
            )

            summary, detail = QA2.AuditReviewDemo.checklist_summary(
                window, "43C1673G07", "", "",
            )

            self.assertEqual(summary, "REARCOVER: 0/1 parameters found")
            self.assertIn("Copper [drawing.pdf (page 2, OCR)]", detail)
            self.assertIn("Steel [assembly.dwg]", detail)

    def test_native_dwg_reader_extracts_text_and_block_attributes(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "checklist.xlsx")
            self.create_checklist(path)
            configuration = QA2.ConfigurationRepository(path)
            document = SimpleNamespace(
                ModelSpace=[SimpleNamespace(ObjectName="AcDbText", TextString="Amperage: 800A")],
                PaperSpace=[SimpleNamespace(ObjectName="AcDbBlockReference", GetAttributes=lambda: [
                    SimpleNamespace(TagString="Material", TextString="Copper"),
                ])],
                Close=lambda save: None,
            )
            documents = MagicMock()
            documents.__iter__.return_value = iter(())
            documents.Open.return_value = document
            app = SimpleNamespace(Documents=documents)
            com_client = SimpleNamespace(
                GetActiveObject=lambda _prog_id: (_ for _ in ()).throw(RuntimeError("not running")),
                Dispatch=lambda _prog_id: app,
            )

            with patch.object(QA2, "win32", com_client):
                result = QA2.DWGFile(os.path.join(folder, "sample.dwg")).extract_parameters(configuration)

            self.assertEqual(result.get_parameter("Amperage"), "800A")
            self.assertEqual(result.get_parameter("Material"), "Copper")

    def test_native_dwg_reader_preserves_existing_document(self):
        path = os.path.abspath("sample.dwg")
        document = SimpleNamespace(
            FullName=path,
            ModelSpace=[SimpleNamespace(ObjectName="AcDbText", TextString="SECTION 8: DENSITY RATED BUS")],
            PaperSpace=[], Close=MagicMock(),
        )
        documents = MagicMock()
        documents.__iter__.return_value = iter((document,))
        client = SimpleNamespace(GetActiveObject=lambda name: SimpleNamespace(Documents=documents))
        with patch.object(QA2, "win32", client):
            self.assertIn("DENSITY RATED BUS", QA2.DWGFile(path).read_text())
        document.Close.assert_not_called()
        documents.Open.assert_not_called()


class PDFExtractionTests(unittest.TestCase):
    def test_scanned_page_uses_ocr_and_retains_page_provenance(self):
        pdf = MagicMock()
        pdf.__enter__.return_value.pages = [SimpleNamespace(extract_text=lambda: "")]
        rendered_pdf = MagicMock()
        rendered_pdf.load_page.return_value.get_pixmap.return_value.tobytes.return_value = b"png"
        image = MagicMock()

        with patch.object(QA2, "pdfplumber", SimpleNamespace(open=lambda _path: pdf)), \
                patch.object(QA2, "fitz", SimpleNamespace(open=lambda _path: rendered_pdf)), \
                patch.object(QA2, "pytesseract", SimpleNamespace(
                    image_to_string=lambda _image: "Material: Copper",
                )), \
                patch.object(QA2, "Image", SimpleNamespace(
                    open=lambda _image_data: image,
                )):
            extraction = QA2.extract_pdf_pages("drawing.pdf")

        self.assertEqual(extraction.text, "Material: Copper")
        self.assertEqual(extraction.ocr_pages, (1,))
        self.assertEqual(extraction.issues, ())
        self.assertEqual(
            QA2.AuditReviewDemo._pdf_parameter_sources(extraction, "drawing.pdf"),
            [("Material: Copper", "drawing.pdf (page 1, OCR)")],
        )
        rendered_pdf.close.assert_called_once_with()

    def test_scanned_page_reports_missing_ocr_dependencies(self):
        pdf = MagicMock()
        pdf.__enter__.return_value.pages = [SimpleNamespace(extract_text=lambda: "")]
        with patch.object(QA2, "pdfplumber", SimpleNamespace(open=lambda _path: pdf)), \
                patch.object(QA2, "fitz", None), \
                patch.object(QA2, "pytesseract", None), \
                patch.object(QA2, "Image", None):
            extraction = QA2.extract_pdf_pages("scan.pdf")

        self.assertEqual(extraction.pages, ("",))
        self.assertEqual(extraction.ocr_pages, ())
        self.assertIn("OCR unavailable", extraction.issues[0])
        self.assertIn("PyMuPDF, pytesseract, Pillow", extraction.issues[0])
        self.assertIn("Tesseract OCR executable", extraction.issues[0])


if __name__ == "__main__":
    unittest.main()