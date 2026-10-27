"""Generalized Excel parsers shared by the live QA2 audit and a standalone CLI.

Run ``Classifier.py WORKBOOK ASSEMBLY [--output OUTPUT.json]`` for local files.
Metadata.xlsx and semantic_aliases.example.json are resolved beside this module.
Live callers supply an openpyxl snapshot with values and merged ranges preserved.
"""

import re
import logging
from functools import lru_cache
from openpyxl import load_workbook
from pathlib import Path
from pprint import pprint
import json


@lru_cache(maxsize=8)
def _metadata_rows(path, modified_ns, size):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        return tuple(workbook["Database_Index2"].iter_rows(values_only=True))
    finally:
        workbook.close()


def metadata_rows(path):
    path = Path(path)
    stat = path.stat()
    return _metadata_rows(str(path.resolve()), stat.st_mtime_ns, stat.st_size)


@lru_cache(maxsize=8)
def _semantic_data(path, modified_ns, size):
    with open(path, encoding="utf-8") as stream:
        return json.load(stream)


def semantic_data(path):
    path = Path(path)
    stat = path.stat()
    return _semantic_data(str(path.resolve()), stat.st_mtime_ns, stat.st_size)


def root_reference(value):
    """Only a drawing label, not a drawing mentioned in a note or a partial number."""
    text = str(value or "").strip().upper()
    match = re.fullmatch(r"(?:REF\s*[-:]?\s*)?([0-9A-Z]{6,})(?:\s*)", text)
    if not match or not re.search(r"[A-Z]", match[1]) or not re.search(r"\d", match[1]):
        return ""
    style = re.fullmatch(r"([0-9A-Z]{6,}?)[GH]\d+", match[1])
    return style[1] if style else match[1]


class AnchoredTableParser:
    """Read root-labelled tables in place, retaining merged row/column selectors.

    Unlike sheet-wide classification, each root occurrence owns only its local
    header band and body. A neighbouring table must never supply its groups.
    Numeric two-axis matrices remain with GeneralizedMatrixParser.
    Unparsed secondary header bands retain an explicit blocker, not inherited axes.
    """

    GROUP = re.compile(r"[GH]\d+", re.I)
    TESTS = {"HT": "Heat Test", "HEAT": "Heat Test", "HEAT TEST": "Heat Test",
             "DEN": "Density", "DENSITY": "Density"}
    LABELS = {
        "AMPS": "Amperage", "AMP": "Amperage", "AMPACITY": "Amperage",
        "AMPERAGE": "Amperage", "AMPRATING": "Amperage", "CURRENTRATING": "Amperage",
        "RATING": "Amperage", "WIDTH": "Structure Width", "STRUCTUREWIDTH": "Structure Width",
        "DEPTH": "Structure Depth", "STRUCTUREDEPTH": "Structure Depth",
        "HEIGHT": "Structure Height", "MATERIAL": "Material",
        "PHASING": "Phase", "CIRCUITNUMBERS": "Circuit Numbers",
        "NUMBEROFDRILLINGS": "Number of Drillings", "NOOFDRILLINGS": "Number of Drillings",
    }

    def __init__(self, sheet):
        self.sheet = sheet
        self.grid = {
            (cell.row, cell.column): str(cell.value).strip()
            for row in sheet for cell in row if cell.value is not None
        }
        self.original = dict(self.grid)
        for area in sheet.merged_cells.ranges:
            value = self.grid.get((area.min_row, area.min_col), "")
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    self.grid[row, col] = value

    @staticmethod
    def attributes(text):
        text = str(text).upper().strip()
        result = {}
        for name, pattern in (
            ("Structure Width", r'(?<![\d.])(\d+(?:\.\d+)?)\s*"?\s*(?:WIDE|W)\b'),
            ("Structure Depth", r'(?<![\d.])(\d+(?:\.\d+)?)\s*"?\s*D\b(\s*MIN\.?)?'),
            ("Structure Height", r'(?<![\d.])(\d+(?:\.\d+)?)\s*"?\s*(?:HIGH|H)\b'),
        ):
            match = re.search(pattern, text)
            if match:
                result[name] = match[1] + (" MIN" if name == "Structure Depth" and match[2] else "")
        material = {
            "AL": "Aluminum", "ALUM": "Aluminum", "ALUMINUM": "Aluminum",
            "CU": "Copper", "COPPER": "Copper", "CUAG": "Silver Plated Copper",
            "CUSN": "Tin Plated Copper", "SILVER PLATED COPPER": "Silver Plated Copper",
            "TIN PLATED COPPER": "Tin Plated Copper",
            "COPPER AG": "Silver Plated Copper", "COPPER SN": "Tin Plated Copper",
            "CU AG": "Silver Plated Copper", "CU SN": "Tin Plated Copper",
        }.get(text)
        if material:
            result["Material"] = material
        bus = re.fullmatch(r"(AL|CU|CUAG|CUSN)\.?\s+(HT|DEN)", text)
        if bus:
            result.update(AnchoredTableParser.attributes(bus[1]))
            result.update(AnchoredTableParser.attributes(bus[2]))
        if text in AnchoredTableParser.TESTS:
            result["Test"] = AnchoredTableParser.TESTS[text]
        density = re.fullmatch(r"DENSITY\s+(\d+)\s*A\s*/\s*SQ\.?\s*INCH", text)
        if density:
            result["Test"] = f"Density {density[1]}A/SQ INCH"
        if text == "HEAT TESTED":
            result["Test"] = "Heat Test"
        phases = re.findall(r"\b(THREE|SINGLE|3|1)\s*PH(?:ASE)?\b\.?", text)
        if phases:
            result["Phase"] = "/".join(dict.fromkeys(
                "3 Phase" if phase in {"THREE", "3"} else "1 Phase" for phase in phases
            ))
        wires = re.search(r"\b([34])\s*WIRE\b", text)
        if wires:
            result["Wires"] = wires[1]
        hinge = re.search(r"\bHINGED\s+(LEFT|RIGHT)", text)
        if hinge:
            result["Hinge Side"] = hinge[1]
        if "STANDARD" in text and "TOP" in text and "FEED" in text:
            result["Feed"] = "STANDARD"
        elif re.search(r"\b(?:REVERSE|REV)\s+FEED\b", text):
            result["Feed"] = "REVERSE"
        application = re.search(r'\bAPPLICATION\s*"?([ABC])"?\b', text)
        if application:
            result["Application"] = application[1]
        if "NO CROSS BUS" in text:
            result["Lower Cross Bus"] = "NONE"
        elif "CROSS BUS IN LOWER" in text:
            result["Lower Cross Bus"] = "PRESENT"
        if "DBL DOOR" in text or "DOUBLE DOOR" in text:
            result["Door Type"] = "DOUBLE"
        elif "SGL" in text and "DOOR" in text:
            result["Door Type"] = "SINGLE"
            size = re.search(r'SGL\s+(\d+)', text)
            if size:
                result["Door Width"] = size[1]
        elif "BLANK" in text and "DOOR" in text:
            result["Door Type"] = "BLANK"
        if "Door Type" in result:
            result["Assembly Type"] = "DOOR"
        sockets = re.fullmatch(r"\((\d+)\)\s*(BLANK|SOCKETS?)", text)
        if sockets:
            result["Meter Socket Cover"] = "BLANK" if sockets[2] == "BLANK" else f"{sockets[1]} SOCKETS"
        meters = re.search(r"\b([12])\s+(?:METER|MTR)\b", text)
        if meters:
            result["Meter Count"] = meters[1]
        if "UGPS COVER ASSEMBLY" in text:
            result["Assembly Type"] = "UGPS COVER"
        drillings = re.fullmatch(r"\(?\s*(\d+)\s+DRILLINGS\s*\)?", text)
        if drillings:
            result["Number of Drillings"] = drillings[1]
        return result

    def parse(self, root):
        positions = [(r, c) for (r, c), value in self.original.items() if root_reference(value) == root]
        configurations = []
        for anchor_row, anchor_col in positions:
            get = lambda r, c: self.grid.get((r, c), "")
            # A vertical blank column separates independent tables.
            first = next((
                r for r in range(anchor_row, min(self.sheet.max_row, anchor_row + 16) + 1)
                if any(self.GROUP.fullmatch(get(r, c)) for c in range(anchor_col, self.sheet.max_column + 1))
            ), None)
            if first is None:
                continue
            left = anchor_col
            if anchor_col > 1 and any(get(r, anchor_col - 1) for r in range(anchor_row, first + 1)):
                left -= 1
            right = self.sheet.max_column
            for c in range(anchor_col + 1, right + 1):
                if not any(get(r, c) for r in range(anchor_row, first + 2)):
                    right = c - 1
                    break
                if root_reference(get(anchor_row, c)) and root_reference(get(anchor_row, c)) != root:
                    right = c - 1
                    break
            if not any(self.GROUP.fullmatch(get(first, c)) for c in range(left, right + 1)):
                continue
            inline = first == anchor_row
            header_end = first
            if inline:
                while header_end > 1 and any(
                    self.GROUP.fullmatch(get(header_end - 1, c)) for c in range(left, right + 1)
                ):
                    header_end -= 1
            top = header_end
            while top > 1 and anchor_row - top < 12:
                previous = [get(top - 1, c) for c in range(left, right + 1)]
                if any(self.GROUP.fullmatch(v) for v in previous):
                    break
                top -= 1
            local_top = top
            if inline and not any(
                re.sub(r"[^A-Z0-9]", "", get(r, c).upper()) in self.LABELS
                for r in range(top, header_end) for c in range(left, min(right, anchor_col + 2) + 1)
            ):
                # Contiguous assembly blocks can share an explicit column header above earlier roots.
                for r in range(top - 1, max(0, header_end - 64), -1):
                    if not any(get(r, c) for c in range(left, right + 1)):
                        break
                    title = re.sub(r"[^A-Z0-9]", "", self.original.get((r, anchor_col), "").upper())
                    if title not in {"ASSYDWG", "ASSEMBLYDRAWING", "ASSEMBLYDWG", "DRAWING"}:
                        continue
                    axes = [
                        c for c in range(left, min(right, anchor_col + 2) + 1) if c != anchor_col
                        and self.LABELS.get(re.sub(r"[^A-Z0-9]", "", get(r, c).upper())) == "Amperage"
                        and re.fullmatch(r"\d+\s*A", get(first, c), re.I)
                    ]
                    if axes:
                        top = r
                    break
            header_rows = list(range(top, header_end))
            if top < local_top:
                shared_end = next((
                    r for r in range(top, local_top)
                    if any(self.GROUP.fullmatch(get(r, c)) for c in range(left, right + 1))
                ), local_top)
                header_rows = list(range(top, shared_end)) + list(range(local_top, header_end))
            bottom = first
            while bottom < self.sheet.max_row:
                following = [get(bottom + 1, c) for c in range(left, right + 1)]
                if not any(following) or any(
                    root_reference(self.original.get((bottom + 1, c), ""))
                    for c in range(left, right + 1)
                ):
                    break
                if inline and (
                    get(bottom + 1, anchor_col) not in ("", get(anchor_row, anchor_col))
                    or not any(self.GROUP.fullmatch(value) for value in following)
                ):
                    break
                bottom += 1
            constants = self.attributes(self.sheet.title)
            # Titles in the root/row-label column apply to the whole local table.
            for r in range(local_top, header_end):
                constants.update(self.attributes(get(r, anchor_col)))
            labels = {}
            for c in range(left, right + 1):
                for r in header_rows:
                    value = re.sub(r"[^A-Z0-9]", "", get(r, c).upper())
                    if value in self.LABELS:
                        labels[c] = self.LABELS[value]
            group_columns = sorted({
                c for r in range(first, bottom + 1) for c in range(left, right + 1)
                if self.GROUP.fullmatch(get(r, c))
            })
            column_attributes = {c: {} for c in group_columns}
            for r in header_rows:
                by_name = {}
                for c in range(left, right + 1):
                    for name, value in self.attributes(get(r, c)).items():
                        by_name.setdefault(name, []).append((c, value))
                for name, cells in by_name.items():
                    distinct = list(dict.fromkeys(value for _, value in cells))
                    if len(distinct) == 1:
                        for attributes in column_attributes.values():
                            attributes[name] = distinct[0]
                    elif name == "Structure Width" and len(cells) == len(distinct) and len(group_columns) % len(cells) == 0:
                        # Unmerged centred width headings, e.g. E/I/M above D:G/H:K/L:O.
                        size = len(group_columns) // len(cells)
                        if all(group_columns[i * size] <= c <= group_columns[(i + 1) * size - 1]
                               for i, (c, _) in enumerate(cells)):
                            for i, c in enumerate(group_columns):
                                column_attributes[c][name] = cells[i // size][1]
                    else:
                        for c in group_columns:
                            preceding = [value for column, value in cells if column <= c]
                            if preceding:
                                column_attributes[c][name] = preceding[-1]
            for row in range(first, bottom + 1):
                for col in range(left, right + 1):
                    group = get(row, col)
                    if not self.GROUP.fullmatch(group):
                        continue
                    config = dict(constants)
                    secondary_headers = [
                        get(r, c) for r in range(first, row) for c in range(left, right + 1)
                        if not any(self.GROUP.fullmatch(get(r, k)) for k in range(left, right + 1))
                        and get(r, c)
                    ]
                    for c in range(left, min(group_columns)):
                        config.update(self.attributes(get(row, c)))
                    for axis_col, name in labels.items():
                        value = get(row, axis_col)
                        if not value:
                            candidates = {
                                get(r, axis_col) for r in range(first, bottom + 1)
                                if get(r, axis_col) and not self.GROUP.fullmatch(get(r, axis_col))
                            }
                            if len(candidates) == 1:
                                value = next(iter(candidates))
                        if value and not self.GROUP.fullmatch(value) and not root_reference(value):
                            if name == "Amperage":
                                rating = re.match(r"(\d+(?:\s*-\s*\d+)?)\s*A?\b", value, re.I)
                                value = rating[1] if rating else value
                            elif name == "Phase":
                                value = re.sub(r"\bSINGLE\b", "1", value, flags=re.I)
                                value = re.sub(r"\bTHREE\b", "3", value, flags=re.I)
                                value = re.sub(r"\s+OR\s+", "/", value, flags=re.I)
                            config[name] = value
                    config.update(column_attributes[col])
                    for r in header_rows:
                        value = get(r, col)
                        if re.sub(r"[^A-Z]", "", value.upper()) == "CIRCUITNUMBERS":
                            circuit = next((get(k, col) for k in range(r + 1, header_end) if get(k, col)), "")
                            if circuit:
                                config["Circuit Numbers"] = circuit
                        if re.fullmatch(r'\d+(?:\.\d+)?', value):
                            label = next((labels[c] for c in range(col - 1, left - 1, -1) if c in labels), None)
                            if label:
                                config[label] = value
                    config["Group No."] = group.upper()
                    if secondary_headers:
                        # A new header band invalidates the earlier row and column axes.
                        config = dict(constants)
                        for text in secondary_headers:
                            config.update(self.attributes(text))
                        config["Unparsed Table Context"] = " / ".join(dict.fromkeys(secondary_headers))
                        config["Group No."] = group.upper()
                    if config not in configurations:
                        configurations.append(config)
        return {"Assembly": root, "Configurations": configurations}

# ==========================================================
# MAIN Generalized Hierarchical Parser
# ==========================================================
class GeneralizedHierarchicalParser:

    def __init__(
        self,
        excel_file,
        metadata_file,
        semantic_file,
        workbook=None,
    ):
        self.source_file_name = Path(
        excel_file).stem.upper()
        self.wb = workbook if workbook is not None else load_workbook(excel_file, data_only=True)
        self.ws = self.wb.active

        self.metadata = metadata_rows(metadata_file)
        self._metadata_matches = {}
        self.semantic = semantic_data(semantic_file)
        self._semantic_aliases = {}
        for parent, info in self.semantic.get("parameters", {}).items():
            for label in (parent, *info.get("aliases", ())):
                self._semantic_aliases.setdefault(label.lower(), parent)

        self.expanded_grid = self.expand_sheet()
    
    def get_metadata_sheet_name(self):

        for row in self.metadata:

            if row[0] is None:
                continue

            if str(row[0]).strip().upper() == self.source_file_name:

                return str(row[1]).strip()

        return None
    # --------------------------------------------------
    # EXPAND MERGED CELLS
    # --------------------------------------------------

    def expand_sheet(self):

        grid = {}

        for row in self.ws.iter_rows():
            for cell in row:

                if cell.value is not None:
                    grid[(cell.row, cell.column)] = str(cell.value).strip()

        for merged in self.ws.merged_cells.ranges:

            min_col, min_row, max_col, max_row = merged.bounds

            value = self.ws.cell(min_row, min_col).value

            if value is None:
                continue

            value = str(value).strip()

            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    grid[(r, c)] = value

        return grid

    # --------------------------------------------------
    # ASSEMBLY FORMAT
    # --------------------------------------------------

    def is_assembly(self, text):

        if text is None:
            return False

        text = str(text).strip()

        return bool(
            re.match(
                r'^[0-9A-Z]{6,}$',
                text
            )
        )

    # --------------------------------------------------
    # FIND ASSEMBLY CELL
    # --------------------------------------------------

    def locate_assembly(self, assembly):

        positions = []

        for (r, c), value in self.expanded_grid.items():

            if str(value).strip() == str(assembly).strip():

                positions.append(
                    (r, c)
                )

        return positions

    # --------------------------------------------------
    # WORD COUNT
    # --------------------------------------------------

    def word_count(self, text):

        return len(
            str(text).split()
        )
    
    # --------------------------------------------------
    # LOOKUP IN METADATA
    # --------------------------------------------------

    def metadata_match(self, text, sheet_name):

        target = str(text).strip().lower()

        key = sheet_name.upper()
        if key not in self._metadata_matches:
            parameters = {}
            for row in self.metadata:
                if str(row[0] or "").strip().upper() in {key, self.source_file_name}:
                    for cell in row[1:]:
                        if cell is not None:
                            parameter = str(cell).strip()
                            parameters.setdefault(parameter.lower(), parameter)
                    break
            self._metadata_matches[key] = parameters
        parameter = self._metadata_matches[key].get(target)
        return parameter is not None, parameter

    # --------------------------------------------------
    # LOOKUP IN SEMANTICS
    # --------------------------------------------------

    def semantic_match(self, text):

        # target = str(text).strip()

        # ---------------------------
        # Pattern Based Matching
        # ---------------------------
       #if re.search(
       #    r"\bWIDE\s+STRUCTURE\b",
       #    target,
       #    re.IGNORECASE
        #):
        #    return True, "Wide Structure" 
        
       
        target = str(text).strip()

        match_target = re.sub(
            r"^[0-9\-\.\s]+",
            "",
            target
        ).strip()

        target_lower = match_target.lower()
        parent = self._semantic_aliases.get(target_lower)
        return parent is not None, parent

    # --------------------------------------------------
    # HEADER SEARCH
    # --------------------------------------------------

    def find_header(self, assembly_row, assembly_col):

        sheet_name = self.ws.title

        header_bottom = None

        current_row = assembly_row

        while current_row > 1:

            current_row -= 1

            row_has_header_signal = False

            # Check assembly column + nearby hierarchy columns
            for check_col in range(
                assembly_col,
                min(
                    assembly_col + 5,
                    self.ws.max_column + 1
                )
            ):

                value = self.expanded_grid.get(
                    (current_row, check_col)
                )

                if not value:
                    continue

                value = str(value).strip()

                if self.is_assembly(value):
                    continue

                found, parent = self.metadata_match(
                    value,
                    sheet_name
                )

                if found:
                    row_has_header_signal = True
                    break

                found, parent = self.semantic_match(
                    value
                )

                if found:
                    row_has_header_signal = True
                    break

                # Check merged value with row above
                upper = self.expanded_grid.get(
                    (current_row - 1, check_col)
                )

                if upper:

                    combined = (
                        str(upper).strip()
                        + " "
                        + value
                    )

                    found, parent = self.metadata_match(
                        combined,
                        sheet_name
                    )

                    if found:
                        row_has_header_signal = True
                        break

                    found, parent = self.semantic_match(
                        combined
                    )

                    if found:
                        row_has_header_signal = True
                        break

            if row_has_header_signal:

                header_bottom = current_row

                logging.debug(
                    f"HEADER SIGNAL FOUND AT ROW "
                    f"{current_row}"
                )

                break

        if header_bottom is None:
            return None, None

        header_top = header_bottom

        while header_top > 1:

            previous_row = header_top - 1

            row_values = []

            for c in range(
                1,
                self.ws.max_column + 1
            ):

                value = self.expanded_grid.get(
                    (previous_row, c)
                )

                if value:
                    row_values.append(
                        str(value).strip()
                    )

            # Completely blank row
            if len(row_values) == 0:
                break

            # Long title / note row
            long_text_found = any(
                self.word_count(v) > 5
                for v in row_values
            )

            if long_text_found:
                break

            header_top -= 1

        logging.debug(
            f"HEADER BAND: "
            f"{header_top} -> {header_bottom}"
        )

        return header_top, None

    # --------------------------------------------------
    # UPPER BOUNDARY
    # --------------------------------------------------

    def find_upper_boundary(
        self,
        assembly_row,
        assembly_col
    ):

        header_row, _ = self.find_header(
            assembly_row,
            assembly_col
        )

        if header_row is None:
            return 1

        return header_row

    # --------------------------------------------------
    # LOWER BOUNDARY
    # --------------------------------------------------

    def find_lower_boundary(
        self,
        assembly_row,
        assembly_col
    ):

        row = assembly_row + 1

        while row <= self.ws.max_row:

            value = self.expanded_grid.get(
                (row, assembly_col)
            )

            # -----------------------
            # Blank row found
            # -----------------------
            if value is None or str(value).strip() == "":
                return row - 1

            value = str(value).strip()

            # -----------------------
            # Long note found
            # -----------------------
            if self.word_count(value) > 5:
                return row - 1

            row += 1

        return self.ws.max_row
    def find_horizontal_boundaries(
        self,
        top_row,
        assembly_row,
        assembly_col
    ):

        # -------------------------
        # LEFT BOUNDARY
        # -------------------------
        left_col = assembly_col

        while left_col > 1:

            previous_col = left_col - 1

            column_has_content = False

            for r in range(
                top_row,
                assembly_row + 1
            ):

                value = self.expanded_grid.get(
                    (r, previous_col)
                )

                if value is not None and str(value).strip():

                    column_has_content = True
                    break

            if not column_has_content:
                break

            left_col -= 1

        # -------------------------
        # RIGHT BOUNDARY
        # -------------------------
        right_col = assembly_col

        empty_streak = 0

        while right_col < self.ws.max_column:

            next_col = right_col + 1

            column_has_content = False

            for r in range(
                top_row,
                assembly_row + 1
            ):

                value = self.expanded_grid.get(
                    (r, next_col)
                )

                if value is not None and str(value).strip():

                    column_has_content = True
                    break

            if not column_has_content:

                empty_streak += 1

                if empty_streak >= 2:
                    break

            else:

                empty_streak = 0

            right_col += 1

        logging.debug(
            f"TABLE WIDTH => "
            f"LEFT={left_col}, "
            f"RIGHT={right_col}"
        )

        return left_col, right_col

    def get_nearest_value(
        self,
        row,
        col,
        top_row,
        bottom_row,
        expected_attribute
    ):

        current = self.expanded_grid.get(
            (row, col)
        )

        if (
            current is not None
            and str(current).strip()
        ):
            return str(current).strip()

        sheet_name = self.ws.title

        # -----------------------------------
        # Helper: Ignore Notes / Headers
        # -----------------------------------
        def is_invalid_candidate(candidate):

            if candidate is None:
                return True

            candidate = str(candidate).strip()

            if not candidate:
                return True

            # Ignore Notes row
            if candidate.lower() in ["notes", "notes:"]:
                return True

            # Ignore metadata headers
            found, parent = self.metadata_match(
                candidate,
                sheet_name
            )

            if found:
                return True

            # Ignore semantic headers
            found, parent = self.semantic_match(
                candidate
            )

            if found:
                return True

            return False

        up_distance = None
        up_value = None

        # -----------------------------------
        # SEARCH UP
        # -----------------------------------
        r = row - 1

        while r >= top_row:

            candidate = self.expanded_grid.get(
                (r, col)
            )

            if (
                candidate is not None
                and str(candidate).strip()
            ):

                candidate = str(candidate).strip()

                if is_invalid_candidate(
                    candidate
                ):
                    r -= 1
                    continue

                up_distance = row - r
                up_value = candidate
                break

            r -= 1

        down_distance = None
        down_value = None

        # -----------------------------------
        # SEARCH DOWN
        # -----------------------------------
        r = row + 1

        while r <= bottom_row:

            candidate = self.expanded_grid.get(
                (r, col)
            )

            if (
                candidate is not None
                and str(candidate).strip()
            ):

                candidate = str(candidate).strip()

                if is_invalid_candidate(
                    candidate
                ):
                    r += 1
                    continue

                down_distance = r - row
                down_value = candidate
                break

            r += 1

        # -----------------------------------
        # RESOLVE RESULT
        # -----------------------------------

        if up_value is None and down_value is None:
            return None

        if up_value is None:
            return down_value

        if down_value is None:
            return up_value

        if up_distance <= down_distance:
            return up_value

        return down_value
 
    def repair_unmerged_hierarchy(
        self,
        header_matrix
    ):

        if not header_matrix:
            return header_matrix

        levels = sorted(
            header_matrix.keys()
        )

        if len(levels) < 2:
            return header_matrix

        # -----------------------------------
        # Bottom-up repair
        # -----------------------------------

        for level_idx in range(
            len(levels) - 2,
            -1,
            -1
        ):

            parent_level = levels[level_idx]
            child_level = levels[level_idx + 1]

            parent_row = header_matrix[parent_level]
            child_row = header_matrix[child_level]

            all_cols = sorted(
                child_row.keys()
            )

            if not all_cols:
                continue

            min_col = min(all_cols)
            max_col = max(all_cols)

            current_parent = None

            for col in range(
                min_col,
                max_col + 1
            ):

                parent_value = parent_row.get(
                    col
                )

                child_value = child_row.get(
                    col
                )

                if not child_value:
                    continue

                # Explicit parent found
                if (
                    parent_value is not None
                    and
                    str(parent_value).strip()
                ):

                    current_parent = str(
                        parent_value
                    ).strip()

                    continue

                # Parent missing
                if current_parent:

                    parent_row[col] = (
                        current_parent
                    )

        return header_matrix

    def clean_assembly_number(
        self,
        text
    ):

        if text is None:
            return None

        text = str(text).strip().upper()

        match = re.search(
            r'\d+[A-Z]+\d+',
            text
        )

        if match:
            return match.group(0)

        return text

    # --------------------------------------------------
    # EXTRACT TABLE
    # --------------------------------------------------

    def extract_configurations(
        self,
        assembly_number,
        assembly_row,
        top_row,
        bottom_row,
        left_col,
        right_col,
        external_layout
    ):

        configurations = []

               # -----------------------------------------------------
        # Build header hierarchy for every column
        # -----------------------------------------------------

        column_headers = {}

        header_end_row = assembly_row

        if external_layout:

            header_end_row = top_row

            while header_end_row <= bottom_row:

                found_group = False

                for c in range(
                    left_col,
                    right_col + 1
                ):

                    value = self.expanded_grid.get(
                        (header_end_row, c)
                    )

                    if value is None:
                        continue

                    value = str(value).strip()

                    if re.match(
                        r"^G\d+$",
                        value,
                        re.IGNORECASE
                    ):
                        found_group = True
                        break

                if found_group:
                    break

                header_end_row += 1

            logging.debug(
                f"EXTERNAL HEADER BAND => "
                f"{top_row} - {header_end_row - 1}"
            )

        # -----------------------------------------------------
        # PHASE A
        # Build raw hierarchy matrix
        # -----------------------------------------------------

        header_matrix = {}

        header_start = top_row

        header_stop = (
            header_end_row
            if external_layout
            else assembly_row
        )

        level = 0

        for r in range(
            header_start,
            header_stop
        ):

            header_matrix[level] = {}

            for c in range(
                left_col,
                right_col + 1
            ):

                value = self.expanded_grid.get(
                    (r, c)
                )

                if value is None:
                    continue

                value = str(value).strip()

                if value:
                    header_matrix[level][c] = value

            level += 1

        # -----------------------------------------------------
        # PHASE B
        # Repair non-merged hierarchy gaps
        # -----------------------------------------------------

        header_matrix = (
            self.repair_unmerged_hierarchy(
                header_matrix
            )
        )

        # -----------------------------------------------------
        # PHASE C
        # Rebuild column headers
        # -----------------------------------------------------

        column_headers = {}

        for c in range(
            left_col,
            right_col + 1
        ):

            headers = []

            for level in sorted(
                header_matrix.keys()
            ):

                value = header_matrix[
                    level
                ].get(c)

                if not value:
                    continue

                if value not in headers:
                    headers.append(value)

            column_headers[c] = headers

            logging.debug(
                f"COL {c} HEADERS => {headers}"
            )
        # -----------------------------------------------------
        # Resolve column attributes dynamically
        # -----------------------------------------------------

        column_attributes = {}

        sheet_name = self.ws.title

        for col, headers in column_headers.items():

            attribute_name = None

            for i in range(
                len(headers) - 1,
                -1,
                -1
            ):

                candidate = headers[i]

                found, parent = self.metadata_match(
                    candidate,
                    sheet_name
                )

                if found:
                    attribute_name = parent
                    break

                found, parent = self.semantic_match(
                    candidate
                )

                if found:
                    attribute_name = parent
                    break

                if i > 0:

                    merged = (
                        headers[i - 1]
                        + " "
                        + headers[i]
                    )

                    found, parent = self.metadata_match(
                        merged,
                        sheet_name
                    )

                    if found:
                        attribute_name = parent
                        break

                    found, parent = self.semantic_match(
                        merged
                    )

                    if found:
                        attribute_name = parent
                        break

            column_attributes[col] = attribute_name

        logging.debug("COLUMN ATTRIBUTES")

        for col, attr in column_attributes.items():

            logging.debug(
                f"COL {col} => {attr}"
            )

        # -----------------------------------------------------
        # Determine rows to process
        # -----------------------------------------------------

        if external_layout:

            assembly_rows = []

            for r in range(
                header_end_row,
                bottom_row + 1
            ):

                row_contains_group = False

                for c in range(
                    left_col,
                    right_col + 1
                ):

                    value = self.expanded_grid.get(
                        (r, c)
                    )

                    if value is None:
                        continue

                    value = str(value).strip()

                    if re.match(
                        r"^G\d+$",
                        value,
                        re.IGNORECASE
                    ):
                        row_contains_group = True
                        break

                if row_contains_group:
                    assembly_rows.append(r)

        else:

            assembly_rows = []

            for r in range(
                top_row,
                bottom_row + 1
            ):

                for c in range(
                    left_col,
                    right_col + 1
                ):

                    value = self.expanded_grid.get(
                        (r, c)
                    )

                    if (
                        value is not None
                        and str(value).strip()
                        == str(assembly_number).strip()
                    ):

                        assembly_rows.append(r)
                        break

        logging.debug(
            f"ASSEMBLY ROWS => {assembly_rows}"
        )

        # -----------------------------------------------------
        # Process rows
        # -----------------------------------------------------

        for row in assembly_rows:

            row_attributes = {}

            for c in range(
                left_col,
                right_col + 1
            ):

                attribute = column_attributes.get(c)

                search_top = top_row

                if external_layout:
                    search_top = header_end_row

                value = self.get_nearest_value(
                    row,
                    c,
                    search_top,
                    bottom_row,
                    attribute
                )

                if value is None:
                    continue

                if re.fullmatch(r"G\d+", str(value).strip(), re.IGNORECASE):
                    continue

                if attribute:
                    row_attributes[attribute] = str(value).strip()
                    continue

                if attribute:

                    row_attributes[
                        attribute
                    ] = str(value).strip()

            for c in range(
                left_col,
                right_col + 1
            ):

                value = self.expanded_grid.get(
                    (row, c)
                )

                if value is None:
                    continue

                value = str(value).strip()

                if not re.match(
                    r"^G\d+$",
                    value,
                    re.IGNORECASE
                ):
                    continue

                headers = column_headers.get(
                    c,
                    []
                )

                config = dict(
                    row_attributes
                )

                config["Group No."] = value

                for header_value in headers:

                    found, parent = self.semantic_match(
                        header_value
                    )

                    if found and header_value.strip().casefold() != parent.strip().casefold():

                        config[parent] = header_value

                for level in sorted(header_matrix):
                    header_value = header_matrix[level].get(c)
                    if not header_value:
                        continue
                    found, _parent = self.semantic_match(header_value)
                    if found:
                        continue
                    for label_col in range(c - 1, left_col - 1, -1):
                        label = header_matrix[level].get(label_col)
                        if not label:
                            continue
                        parameter = self.resolve_parameter(label, self.ws.title)
                        metadata_header, _name = self.metadata_match(label, self.ws.title)
                        if parameter and (metadata_header or label.strip().casefold() == parameter.strip().casefold()):
                            config[parameter] = header_value
                            break

                configurations.append(
                    config
                )

        return configurations

    def resolve_parameter(
            self,
            text,
            sheet_name
        ):

            found, name = self.metadata_match(
                text,
                sheet_name
            )

            if found:
                return name

            found, name = self.semantic_match(
                text
            )

            if found:
                return name

            return None
    def resolve_parameter_with_merge(
        self,
        column,
        row,
        sheet_name
    ):

        value = self.expanded_grid.get(
            (row, column)
        )

        if not value:
            return None

        value = str(value).strip()

        param = self.resolve_parameter(
            value,
            sheet_name
        )

        if param:
            return param

        upper = self.expanded_grid.get(
            (row - 1, column)
        )

        if upper:

            merged = (
                str(upper).strip()
                + " "
                + value
            )

            param = self.resolve_parameter(
                merged,
                sheet_name
            )

            if param:
                return param

        return None

#############################################
###  External Assembly Layout
#############################################
    def is_external_assembly(
        self,
        assembly_row,
        assembly_col
    ):
        """
        Internal assembly:
            Assembly | Data | Data

        External assembly:
            Assembly | Blank | ...
        """

        check_1 = self.expanded_grid.get(
            (assembly_row, assembly_col + 1)
        )

        check_2 = self.expanded_grid.get(
            (assembly_row, assembly_col + 2)
        )

        if any(
            re.fullmatch(r"G\d+", str(value).strip(), re.IGNORECASE)
            for (row, _column), value in self.expanded_grid.items()
            if row == assembly_row
        ):
            return False

        check_1_filled = (
            check_1 is not None
            and str(check_1).strip() != ""
        )

        check_2_filled = (
            check_2 is not None
            and str(check_2).strip() != ""
        )

        # Both populated means
        # assembly is embedded in table

        if check_1_filled and check_2_filled:
            return False

        # At least one empty means
        # assembly likely outside table

        return True

    
    def find_external_horizontal_boundaries(
        self,
        assembly_row,
        assembly_col
    ):
        """
        Used when Assembly Number is outside the table.

        Boundary condition:
        If a column has blanks in
        current row and next 2 rows,
        treat it as outside table.

        left boundary  = previous column
        right boundary = previous column
        """

        # ----------------------------------
        # LEFT BOUNDARY
        # ----------------------------------
        left_col = assembly_col

        while left_col > 1:

            test_col = left_col - 1

            blank_count = 0

            for r in range(
                assembly_row,
                min(
                    assembly_row + 3,
                    self.ws.max_row + 1
                )
            ):

                value = self.expanded_grid.get(
                    (r, test_col)
                )

                if (
                    value is None
                    or str(value).strip() == ""
                ):
                    blank_count += 1

            if blank_count == 3:
                break

            left_col -= 1

        # ----------------------------------
        # RIGHT BOUNDARY
        # ----------------------------------
        right_col = assembly_col

        while right_col < self.ws.max_column:

            test_col = right_col + 1

            blank_count = 0

            for r in range(
                assembly_row,
                min(
                    assembly_row + 3,
                    self.ws.max_row + 1
                )
            ):

                value = self.expanded_grid.get(
                    (r, test_col)
                )

                if (
                    value is None
                    or str(value).strip() == ""
                ):
                    blank_count += 1

            if blank_count == 3:
                break

            right_col += 1

        logging.debug(
            f"EXTERNAL TABLE WIDTH => "
            f"LEFT={left_col}, "
            f"RIGHT={right_col}"
        )

        return left_col, right_col

    def find_external_lower_boundary(
        self,
        assembly_row,
        assembly_col
    ):

        row = assembly_row + 1

        while row <= self.ws.max_row:

            value = self.expanded_grid.get(
                (row, assembly_col)
            )

            if value:

                value = str(value).strip()

                # Existing note logic
                if self.word_count(value) > 5:
                    return row - 1

                # New assembly detection
                if self.is_assembly(value):
                    return row - 1

                # New semantic alias detection
                found, parent = self.semantic_match(
                    value
                )

                if found and parent.lower() == "assembly":
                    return row - 1

            row += 1

        return self.ws.max_row

    def apply_keyword_mapping(self, obj):

        mapping = self.semantic.get(
            "keyword_mapping",
            {}
        )

        if isinstance(obj, dict):

            result = {}

            for key, value in obj.items():

                result[key] = self.apply_keyword_mapping(
                    value
                )

            return result

        elif isinstance(obj, list):

            return [
                self.apply_keyword_mapping(item)
                for item in obj
            ]

        elif isinstance(obj, str):

            return mapping.get(
                obj.strip(),
                obj
            )

        return obj
    # --------------------------------------------------
    # MAIN
    # --------------------------------------------------

    def extract_for_assembly(
        self,
        assembly_number
    ):

        positions = self.locate_assembly(
            assembly_number
        )

        if not positions:
            return None
 
        all_configurations = []
        table_details = []

        processed_tables = set()

        for position in positions:

            assembly_row, assembly_col = position

            external_layout = (
                self.is_external_assembly(
                    assembly_row,
                    assembly_col
                )
            )

            if external_layout:

                top_row = assembly_row + 1

                left_col, right_col = (
                    self.find_external_horizontal_boundaries(
                        assembly_row,
                        assembly_col
                    )
                )

                bottom_row = (
                    self.find_external_lower_boundary(
                        assembly_row,
                        assembly_col
                    )
                )

            else:

                top_row = self.find_upper_boundary(
                    assembly_row,
                    assembly_col
                )

                left_col, right_col = (
                    self.find_horizontal_boundaries(
                        top_row,
                        assembly_row,
                        assembly_col
                    )
                )

                bottom_row = self.find_lower_boundary(
                    assembly_row,
                    assembly_col
                )

            table_key = (
            top_row,
            bottom_row,
            left_col,
            right_col
        )

            if table_key in processed_tables:
                continue

            processed_tables.add(
                table_key
            )

            configurations = self.extract_configurations(
                assembly_number,
                assembly_row,
                top_row,
                bottom_row,
                left_col,
                right_col,
                external_layout
            )

            all_configurations.extend(
                configurations
            )

            table_details.append(
                {
                    "row": assembly_row,
                    "column": assembly_col,
                    "top": top_row,
                    "bottom": bottom_row
                }
            )

        result = {
        "Assembly": self.clean_assembly_number(
            assembly_number
        ),
        "Occurrences": table_details,
        "Configurations": all_configurations
        }

        return self.apply_keyword_mapping(
            result
        )

# ==========================================================
# MAIN GeneralizedMatrixParser
# ==========================================================

class GeneralizedMatrixParser:

    def __init__(
        self,
        workbook_file,
        metadata_file,
        semantic_file,
        workbook=None,
    ):

        self.workbook_file = workbook_file
        self.metadata_file = metadata_file
        self.semantic_file = semantic_file

        self.workbook_name = (
            Path(workbook_file)
            .stem
            .upper()
        )

        self.metadata = (
            self.load_metadata()
        )

        self.semantic_alignment = (
            self.load_semantic_alignment()
        )

        self.alias_dictionary = (
            self.build_alias_dictionary()
        )

        self.wb = workbook if workbook is not None else load_workbook(
            workbook_file,
            data_only=True
        )

    # ======================================================
    # LOAD METADATA
    # ======================================================

    def load_metadata(self):
        for row in metadata_rows(self.metadata_file)[1:]:
            if str(row[0] or "").strip().upper() == self.workbook_name:
                return [str(value).strip() for value in row[1:] if value is not None]
        return []

    # ======================================================
    # SEMANTIC
    # ======================================================

    def load_semantic_alignment(self):
        return semantic_data(self.semantic_file)

    def build_alias_dictionary(self):

        aliases = {}

        parameters = (
            self.semantic_alignment
            .get("parameters", {})
        )

        for parameter, info in (
            parameters.items()
        ):

            aliases[
                parameter.upper()
            ] = parameter

            for alias in (
                info.get(
                    "aliases",
                    []
                )
            ):

                aliases[
                    alias.upper()
                ] = parameter

        return aliases

    # ======================================================
    # MERGED CELLS
    # ======================================================

    def get_value(
        self,
        ws,
        row,
        col
    ):

        cell = ws.cell(
            row=row,
            column=col
        )

        if cell.value is not None:
            return cell.value

        for merged in (
            ws.merged_cells.ranges
        ):

            if (
                cell.coordinate
                in merged
            ):

                return ws.cell(
                    merged.min_row,
                    merged.min_col
                ).value

        return None

    # ======================================================
    # BUILD MATRIX
    # ======================================================

    def build_matrix(
        self,
        ws
    ):

        matrix = []

        for r in range(
            1,
            ws.max_row + 1
        ):

            row = []

            for c in range(
                1,
                ws.max_column + 1
            ):

                row.append(
                    self.get_value(
                        ws,
                        r,
                        c
                    )
                )

            matrix.append(row)

        return matrix

    # ======================================================
    # NORMALIZATION
    # ======================================================

    def normalize(
        self,
        value
    ):

        if value is None:
            return None

        text = (
            str(value)
            .strip()
            .upper()
        )

        return self.alias_dictionary.get(
            text,
            str(value).strip()
        )

    # ======================================================
    # GROUP DETECTION
    # ======================================================

    def is_group(
        self,
        value
    ):

        if value is None:
            return False

        return bool(
            re.fullmatch(
                r"G\d+",
                str(value)
                .strip()
                .upper()
            )
        )

    # ======================================================
    # ASSEMBLY DETECTION
    # ======================================================

    def is_assembly(
        self,
        value
    ):

        if value is None:
            return False

        value = (
            str(value)
            .strip()
            .upper()
        )

        patterns = [

            r"\d+C\d+",

            r"\d+[A-Z]+\d+"
        ]

        return any(
            re.fullmatch(
                p,
                value
            )
            for p in patterns
        )

    # ======================================================
    # HEADER DETECTION
    # ======================================================

    def detect_matrix_headers(
        self,
        ws
    ):

        candidates = []

        for r in range(
            1,
            ws.max_row + 1
        ):

            group_count = 0

            for c in range(
                1,
                ws.max_column + 1
            ):

                value = self.get_value(
                    ws,
                    r,
                    c
                )

                if self.is_group(
                    value
                ):

                    group_count += 1

            if group_count >= 2:

                candidates.append(r)

        if not candidates:
            return None

        return candidates[0] - 1

    # ======================================================
    # FIND DIMENSIONS
    # ======================================================

    def identify_dimensions(
        self,
        ws,
        header_row
    ):

        column_dimension_values = {}

        for c in range(
            2,
            ws.max_column + 1
        ):

            value = self.get_value(
                ws,
                header_row,
                c
            )

            if value is not None:

                column_dimension_values[
                    c
                ] = value

        return column_dimension_values

    # ======================================================
    # PARAMETER CLASSIFICATION
    # ======================================================

    def classify_table(
        self,
        ws
    ):

        header_row = (
            self.detect_matrix_headers(
                ws
            )
        )

        if not header_row:

            return {
                "table_type":
                    "UNKNOWN"
            }

        return {
            "table_type":
                "MATRIX",
            "header_row":
                header_row
        }



    # ======================================================
    # GENERALIZED MATRIX PARSER
    # ======================================================

    def parse_matrix(
        self,
        ws,
        header_row
    ):

        results = {}

        column_dimensions = (
            self.identify_dimensions(
                ws,
                header_row
            )
        )

        current_assembly = None

        row = header_row + 1

        while row <= ws.max_row:

            row_values = []

            for c in range(
                1,
                ws.max_column + 1
            ):

                row_values.append(
                    self.get_value(
                        ws,
                        row,
                        c
                    )
                )

            assembly_found = False

            for value in row_values:

                if self.is_assembly(
                    value
                ):

                    current_assembly = (
                        str(value)
                        .strip()
                    )

                    assembly_found = True
                    break

            if (
                current_assembly
                is None
            ):

                row += 1
                continue

            row_dimension = None

            for value in row_values:

                try:

                    numeric = int(
                        float(value)
                    )

                    row_dimension = (
                        numeric
                    )

                    break

                except:
                    pass

            if (
                row_dimension
                is None
            ):

                row += 1
                continue

            if (
                current_assembly
                not in results
            ):

                results[
                    current_assembly
                ] = {

                    "Assembly":
                        current_assembly,

                    "ParentAttributes":
                        {},

                    "Configurations":
                        []
                }

            for (
                col,
                col_dimension
            ) in (
                column_dimensions.items()
            ):

                group = (
                    self.get_value(
                        ws,
                        row,
                        col
                    )
                )

                if not self.is_group(
                    group
                ):
                    continue

                config = {

                    "Dimension_1":
                        row_dimension,

                    "Dimension_2":
                        col_dimension,

                    "Group":
                        str(group)
                }

                results[
                    current_assembly
                ][
                    "Configurations"
                ].append(
                    config
                )

            row += 1

        return results
    def normalize_numeric_text(self, value):

        if value is None:
            return None

        text = str(value).strip()

        # 30" -> 30
        # 36" -> 36
        # 48" -> 48

        match = re.match(
            r"^(\d+(?:\.\d+)?)\s*\"$",
            text
        )

        if match:
            return match.group(1)

        return text

    # ======================================================
    # DIMENSION NAME RESOLUTION
    # ======================================================

    def resolve_dimension_names(
        self,
        ws,
        assembly_data
    ):

        header_row = self.detect_matrix_headers(ws)

        if not header_row:
            return assembly_data

        row_dimension_name = "Dimension_1"
        column_dimension_name = "Dimension_2"

        # --------------------------------------------------
        # Get first configuration from parsed data
        # --------------------------------------------------

        first_config = None

        for assembly in assembly_data.values():

            if assembly["Configurations"]:

                first_config = (
                    assembly["Configurations"][0]
                )

                break

        if not first_config:
            return assembly_data

        first_dim1 = first_config["Dimension_1"]
        first_dim2 = first_config["Dimension_2"]

        # --------------------------------------------------
        # MATERIAL CHECK
        # --------------------------------------------------

        material_keywords = {
            "ALUMINUM",
            "ALUM",
            "COPPER",
            "STEEL",
            "STAINLESS",
            "BRASS"
        }

        if (
            isinstance(first_dim1, str)
            and
            first_dim1.strip().upper()
            in material_keywords
        ):

            row_dimension_name = "Material"

        if (
            isinstance(first_dim2, str)
            and
            first_dim2.strip().upper()
            in material_keywords
        ):

            column_dimension_name = "Material"

        # --------------------------------------------------
        # FIND DIMENSION_1 HEADER
        # --------------------------------------------------

        if row_dimension_name == "Dimension_1":

            found = False

            for r in range(
                header_row + 1,
                ws.max_row + 1
            ):

                for c in range(
                    1,
                    ws.max_column + 1
                ):

                    value = self.get_value(
                        ws,
                        r,
                        c
                    )

                    try:

                        if (
                            int(float(value))
                            ==
                            int(first_dim1)
                        ):

                            # 1 row above
                            for offset in [1, 2]:

                                header = self.get_value(
                                    ws,
                                    r - offset,
                                    c
                                )

                                if (
                                    header
                                    and
                                    str(header).strip().upper()
                                    != "ASSEMBLY DRAWING"
                                ):

                                    row_dimension_name = (
                                        str(header).strip()
                                    )

                                    found = True
                                    break

                            # if nothing above check left
                            if not found:

                                for offset in [1, 2]:

                                    if c - offset < 1:
                                        continue

                                    header = self.get_value(
                                        ws,
                                        r,
                                        c - offset
                                    )

                                    if (
                                        header
                                        and
                                        str(header).strip().upper()
                                        != "ASSEMBLY DRAWING"
                                    ):

                                        row_dimension_name = (
                                            str(header).strip()
                                        )

                                        found = True
                                        break

                            break

                    except:
                        pass

                if found:
                    break

        # --------------------------------------------------
        # FIND DIMENSION_2 HEADER
        # --------------------------------------------------

        if column_dimension_name == "Dimension_2":

            found = False

            for r in range(
                1,
                header_row + 1
            ):

                for c in range(
                    1,
                    ws.max_column + 1
                ):

                    value = self.get_value(
                        ws,
                        r,
                        c
                    )

                    try:

                        matched = False

                        # -----------------------------------
                        # Original numeric logic
                        # -----------------------------------
                        try:

                            if (
                                int(float(value))
                                ==
                                int(float(first_dim2))
                            ):
                                matched = True

                        except:
                            pass

                        # -----------------------------------
                        # Only for values like 30"
                        # -----------------------------------
                        if not matched:

                            value_text = str(value).strip()
                            dim2_text = str(first_dim2).strip()

                            if (
                                '"' in value_text
                                and
                                '"' in dim2_text
                            ):

                                if (
                                    value_text.replace('"', '').strip()
                                    ==
                                    dim2_text.replace('"', '').strip()
                                ):
                                    matched = True

                        if matched:
                                                    
                            # 1-2 rows above first

                            for offset in [1, 2]:

                                header = self.get_value(
                                    ws,
                                    r - offset,
                                    c
                                )

                                if (
                                    header
                                    and
                                    str(header).strip().upper()
                                    != "ASSEMBLY DRAWING"
                                ):

                                    column_dimension_name = (
                                        str(header).strip()
                                    )

                                    found = True
                                    break

                            # if not found check left

                            if not found:

                                for offset in [1, 2]:

                                    if c - offset < 1:
                                        continue

                                    header = self.get_value(
                                        ws,
                                        r,
                                        c - offset
                                    )

                                    if (
                                        header
                                        and
                                        str(header).strip().upper()
                                        != "ASSEMBLY DRAWING"
                                    ):

                                        column_dimension_name = (
                                            str(header).strip()
                                        )

                                        found = True
                                        break

                            break

                    except:
                        pass

                if found:
                    break

        # --------------------------------------------------
        # RENAME ONLY
        # --------------------------------------------------

        for assembly in assembly_data.values():

            new_configs = []

            for config in assembly["Configurations"]:

                new_configs.append({

                    row_dimension_name:
                        config["Dimension_1"],

                    column_dimension_name:
                        config["Dimension_2"],

                    "Group":
                        config["Group"]
                })

            assembly["Configurations"] = new_configs

        return assembly_data


    # ======================================================
    # BUILD
    # ======================================================

    def build(self):

        assembly_index = {}

        for ws in (
            self.wb.worksheets
        ):

            classification = (
                self.classify_table(
                    ws
                )
            )

            if (
                classification[
                    "table_type"
                ]
                !=
                "MATRIX"
            ):

                logging.debug(
                    f"Skipping {ws.title}"
                )

                continue

            parsed = (
                self.parse_matrix(
                    ws,
                    classification[
                        "header_row"
                    ]
                )
            )

            parsed = self.resolve_dimension_names(
                ws,
                parsed
            )

            assembly_index.update(
                parsed
            )

        return assembly_index


# ==========================================================
# LOOKUP
# ==========================================================

def lookup(
    assembly_no,
    assembly_data,
):

    result = (
        assembly_data.get(
            assembly_no.upper()
        )
    )

    if result:

        pprint(result)

    else:

        print(
            "Assembly not found"
        )


class TableClassifier:

    def __init__(
        self,
        excel_file,
        assembly_no,
        workbook=None,
    ):

        self.excel_file = excel_file
        self.assembly_no = assembly_no

        self.wb = workbook if workbook is not None else load_workbook(
            excel_file,
            data_only=True
        )

    # --------------------------------------------------
    # GROUP DETECTION
    # --------------------------------------------------

    def is_group(self, value):

        if value is None:
            return False

        return bool(
            re.fullmatch(
                r"G\d+",
                str(value).strip(),
                re.IGNORECASE
            )
        )

    # --------------------------------------------------
    # CLASSIFICATION
    # --------------------------------------------------

    def classify(self, worksheets=None):

        ignore_headers = {

            "ASSEMBLY",
            "ASSEMBLY NO",
            "ASSEMBLY NUMBER",
            "ASSEMBLY DWG",
            "ASSEMBLY DRAWING"

        }

        max_depth = 0

        for ws in self.wb.worksheets if worksheets is None else worksheets:

            data_row = None

            # ----------------------------------
            # Locate first row containing groups
            # ----------------------------------

            for r in range(
                1,
                ws.max_row + 1
            ):

                found_group = False

                for c in range(
                    1,
                    ws.max_column + 1
                ):

                    value = ws.cell(
                        r,
                        c
                    ).value

                    if self.is_group(
                        value
                    ):

                        found_group = True
                        break

                if found_group:

                    data_row = r
                    break

            if not data_row:
                continue

            group_columns = [
                c for c in range(1, ws.max_column + 1)
                if self.is_group(ws.cell(data_row, c).value)
            ]
            has_row_dimension = any(
                re.fullmatch(r"\d+(?:\.\d+)?", str(ws.cell(data_row, c).value or "").strip())
                for c in range(1, ws.max_column + 1) if c not in group_columns
            )
            if data_row > 1 and len(group_columns) >= 2 and has_row_dimension:
                axis_values = [ws.cell(data_row - 1, c).value for c in group_columns]
                if all(re.fullmatch(r'\d+(?:\.\d+)?\s*"?', str(value or "").strip()) for value in axis_values):
                    return "MATRIX"

            # ----------------------------------
            # Determine header depth
            # ----------------------------------

            for c in range(
                1,
                ws.max_column + 1
            ):

                contains_group = False

                for r in range(
                    data_row,
                    ws.max_row + 1
                ):

                    value = ws.cell(
                        r,
                        c
                    ).value

                    if self.is_group(
                        value
                    ):

                        contains_group = True
                        break

                if not contains_group:
                    continue

                headers = []

                for r in range(
                    1,
                    data_row
                ):

                    value = ws.cell(
                        r,
                        c
                    ).value

                    if value is None:
                        continue

                    value = str(value).strip()

                    if not value:
                        continue

                    if (
                        value.upper()
                        in ignore_headers
                    ):
                        continue

                    if value not in headers:

                        headers.append(
                            value
                        )

                max_depth = max(
                    max_depth,
                    len(headers)
                )

        # ----------------------------------
        # Decision
        # ----------------------------------

        if max_depth < 2:

            return "MATRIX"

        return "HIERARCHICAL"


class WorkbookAssemblyParser:
    """Prepare each sheet once and reuse its matrix/grid for every assembly."""

    def __init__(self, excel_file, metadata_file, semantic_file, workbook):
        self.excel_file = excel_file
        self.metadata_file = metadata_file
        self.semantic_file = semantic_file
        self.workbook = workbook
        self.classifier = TableClassifier(excel_file, "", workbook=workbook)
        self.hierarchical = None
        self.matrix = None
        self._prepared = {}
        self._anchored = {}
        self._sheet_values = {
            sheet.title: {root_reference(value) for row in sheet.values for value in row if value is not None}
            for sheet in workbook.worksheets
            if not sheet.title.upper().startswith("OBSOLETE")
        }

    def parse(self, assembly_no):
        assembly_no = assembly_no.upper()
        results = []
        for sheet in self.workbook.worksheets:
            if assembly_no not in self._sheet_values.get(sheet.title, ()):
                continue
            if sheet.title not in self._prepared:
                parser_type = self.classifier.classify((sheet,))
                if parser_type == "MATRIX":
                    if self.matrix is None:
                        self.matrix = GeneralizedMatrixParser(
                            self.excel_file, self.metadata_file, self.semantic_file, workbook=self.workbook,
                        )
                    header = self.matrix.detect_matrix_headers(sheet)
                    data = self.matrix.resolve_dimension_names(
                        sheet, self.matrix.parse_matrix(sheet, header),
                    ) if header else {}
                else:
                    if self.hierarchical is None:
                        self.hierarchical = GeneralizedHierarchicalParser(
                            self.excel_file, self.metadata_file, self.semantic_file, workbook=self.workbook,
                        )
                    self.hierarchical.ws = sheet
                    data = self.hierarchical.expand_sheet()
                self._prepared[sheet.title] = (parser_type, data)
            parser_type, data = self._prepared[sheet.title]
            if sheet.title not in self._anchored:
                self._anchored[sheet.title] = AnchoredTableParser(sheet)
            anchored = self._anchored[sheet.title].parse(assembly_no)
            if parser_type == "MATRIX":
                result = data.get(assembly_no) or anchored
            else:
                if anchored["Configurations"]:
                    result = anchored
                else:
                    self.hierarchical.ws = sheet
                    self.hierarchical.expanded_grid = data
                    result = self.hierarchical.extract_for_assembly(assembly_no)
            if result and result["Configurations"]:
                results.append({
                    "sheet_name": sheet.title,
                    "parser_type": parser_type,
                    "result": result,
                })
        return results


def parse_assembly(excel_file, assembly_no, metadata_file, semantic_file, workbook=None):
    """Return per-sheet configurations, accepting a merged-cell-preserving live snapshot.

    The caller retains ownership of a supplied openpyxl workbook. No output files
    are created; both standalone and live workbooks use the same parser selection.
    """
    owns_workbook = workbook is None
    source = load_workbook(excel_file, data_only=True) if owns_workbook else workbook
    try:
        return WorkbookAssemblyParser(excel_file, metadata_file, semantic_file, source).parse(assembly_no)
    finally:
        if owns_workbook:
            source.close()


if __name__ == "__main__":
    import argparse

    arguments = argparse.ArgumentParser(description="Extract assembly configurations from Excel.")
    arguments.add_argument("workbook")
    arguments.add_argument("assembly")
    arguments.add_argument("--metadata", default=str(Path(__file__).with_name("Metadata.xlsx")))
    arguments.add_argument("--semantic", default=str(Path(__file__).with_name("semantic_aliases.example.json")))
    arguments.add_argument("--output", type=Path)
    args = arguments.parse_args()
    result = parse_assembly(args.workbook, args.assembly, args.metadata, args.semantic)
    if args.output:
        args.output.write_text(json.dumps(result, indent=4), encoding="utf-8")
    else:
        pprint(result)